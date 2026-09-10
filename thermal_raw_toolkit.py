import os
import sys
import csv
import glob
import json
import re
import time
import threading
from collections import deque
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np

try:
    import cv2
except ImportError:
    cv2 = None

try:
    from PyQt5.QtCore import (
        QPointF, QRect, QRectF, QSize, Qt, QThread, QTimer, QUrl, pyqtSignal
    )
    from PyQt5.QtGui import (
        QColor, QDesktopServices, QFont, QImage, QPainter, QPainterPath,
        QPen, QPixmap
    )
    from PyQt5.QtWidgets import (
        QApplication, QMainWindow, QWidget, QFileDialog, QMessageBox,
        QVBoxLayout, QHBoxLayout, QGridLayout, QFormLayout, QLabel,
        QPushButton, QComboBox, QSpinBox, QDoubleSpinBox, QLineEdit,
        QTextEdit, QGroupBox, QStackedWidget, QSplitter, QCheckBox,
        QListWidget, QListWidgetItem, QScrollArea, QFrame, QSizePolicy,
        QDialog, QTextBrowser, QDialogButtonBox, QStyle, QStyleOptionGroupBox,
        QTabWidget, QMenu, QAction, QActionGroup
    )
except ImportError as exc:
    raise SystemExit("PyQt5 is required. Install with: pip install PyQt5") from exc

try:
    from PyQt5.QtMultimedia import QCameraInfo
except (ImportError, OSError):
    QCameraInfo = None


APP_VERSION = "v1.4.0"
APP_TITLE = f"Thermal RAW Toolkit {APP_VERSION}"


def bundled_resource_path(filename):
    source_path = Path(__file__).resolve().parent / filename
    if not getattr(sys, "frozen", False):
        return source_path

    external_path = Path(sys.executable).resolve().parent / filename
    if external_path.is_file():
        return external_path

    return Path(getattr(sys, "_MEIPASS", source_path.parent)) / filename


def toolkit_root_path():
    return (
        Path(sys.executable).resolve().parent
        if getattr(sys, "frozen", False)
        else Path(__file__).resolve().parent
    )


def toolkit_settings_path():
    return toolkit_root_path() / "toolkit_settings.json"


def default_output_root():
    return str(toolkit_root_path() / "Thermal_RAW_Results")


def default_toolkit_settings():
    return {
        "log_font_size": 17,
        "raw_input": {
            "resolution": [DEFAULT_WIDTH, DEFAULT_HEIGHT],
            "bit_depth": DEFAULT_BIT_DEPTH,
            "output_root": default_output_root(),
        },
        "timeline": {
            "metric": "overall_score",
            "history_seconds": 60,
        },
        "camera_mirror": {
            "camera_a": True,
            "camera_b": False,
        },
        "quality_weights": {
            "spatial": dict(DEFAULT_SPATIAL_QUALITY_WEIGHTS),
            "uvc_temporal": dict(DEFAULT_UVC_TEMPORAL_WEIGHTS),
        },
    }


def validated_toolkit_settings(settings):
    defaults = default_toolkit_settings()
    if not isinstance(settings, dict):
        return defaults
    result = default_toolkit_settings()
    result["log_font_size"] = int(
        np.clip(settings.get("log_font_size", 17), 10, 28)
    )
    raw_input = settings.get("raw_input", {})
    try:
        resolution = tuple(
            int(value)
            for value in raw_input.get(
                "resolution", (DEFAULT_WIDTH, DEFAULT_HEIGHT)
            )
        )
    except (TypeError, ValueError):
        resolution = (DEFAULT_WIDTH, DEFAULT_HEIGHT)
    if resolution not in RESOLUTION_OPTIONS:
        resolution = (DEFAULT_WIDTH, DEFAULT_HEIGHT)
    try:
        bit_depth = int(raw_input.get("bit_depth", DEFAULT_BIT_DEPTH))
    except (TypeError, ValueError):
        bit_depth = DEFAULT_BIT_DEPTH
    if bit_depth not in BIT_DEPTH_OPTIONS:
        bit_depth = DEFAULT_BIT_DEPTH
    output_root = str(raw_input.get("output_root", "")).strip()
    if not output_root:
        output_root = default_output_root()
    result["raw_input"] = {
        "resolution": list(resolution),
        "bit_depth": bit_depth,
        "output_root": output_root,
    }
    timeline = settings.get("timeline", {})
    valid_metrics = {metric for _label, metric in QUALITY_TIMELINE_METRICS}
    metric = timeline.get("metric", "overall_score")
    result["timeline"]["metric"] = (
        metric if metric in valid_metrics else "overall_score"
    )
    result["timeline"]["history_seconds"] = int(
        np.clip(timeline.get("history_seconds", 60), 10, 600)
    )
    camera_mirror = settings.get("camera_mirror", {})
    if not isinstance(camera_mirror, dict):
        camera_mirror = {}
    result["camera_mirror"] = {
        "camera_a": bool(camera_mirror.get("camera_a", True)),
        "camera_b": bool(camera_mirror.get("camera_b", False)),
    }
    source_weights = settings.get("quality_weights", {})
    for section, defaults_for_section in (
        ("spatial", DEFAULT_SPATIAL_QUALITY_WEIGHTS),
        ("uvc_temporal", DEFAULT_UVC_TEMPORAL_WEIGHTS),
    ):
        candidate = source_weights.get(section, {})
        values = {
            key: int(np.clip(candidate.get(key, value), 0, 100))
            for key, value in defaults_for_section.items()
        }
        if sum(values.values()) != 100:
            values = dict(defaults_for_section)
        result["quality_weights"][section] = values
    return result


def load_toolkit_settings():
    path = toolkit_settings_path()
    if not path.is_file():
        return default_toolkit_settings()
    try:
        return validated_toolkit_settings(
            json.loads(path.read_text(encoding="utf-8"))
        )
    except (OSError, ValueError, TypeError):
        return default_toolkit_settings()


def save_toolkit_settings(settings):
    path = toolkit_settings_path()
    temporary_path = path.with_suffix(path.suffix + ".new")
    temporary_path.write_text(
        json.dumps(settings, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary_path.replace(path)


def apply_quality_weights(settings):
    weights = settings["quality_weights"]
    ACTIVE_SPATIAL_QUALITY_WEIGHTS.clear()
    ACTIVE_SPATIAL_QUALITY_WEIGHTS.update(weights["spatial"])
    ACTIVE_UVC_TEMPORAL_WEIGHTS.clear()
    ACTIVE_UVC_TEMPORAL_WEIGHTS.update(weights["uvc_temporal"])


DEFAULT_WIDTH = 640
DEFAULT_HEIGHT = 512
RESOLUTION_OPTIONS = [
    (256, 192),
    (320, 240),
    (384, 288),
    (640, 512),
    (1024, 768),
]
RAW_DTYPE = "<u2"       # RAW16 Little Endian
DEFAULT_BIT_DEPTH = 14
BIT_DEPTH_OPTIONS = [8, 10, 12, 14, 16]
NOVATEK_NUC_GAIN_SHIFT = 8
NOVATEK_NUC_GAIN_ONE = 1 << NOVATEK_NUC_GAIN_SHIFT
DEFAULT_CAMERA_INDEX = 1
CAMERA_INDEX_OPTIONS = list(range(10))
CAMERA_DEVICE_ID_ROLE = Qt.UserRole + 1
CAMERA_FRIENDLY_NAME_ROLE = Qt.UserRole + 2
CAMERA_REFRESH_INTERVAL_MS = 1000
DEFAULT_UVC_FPS = 30
UVC_CAPTURE_FOURCC = "MJPG"
DEFAULT_UVC_CAPTURE_MODE = "auto_detect"
UVC_CAPTURE_MODE_OPTIONS = [
    ("Auto Detect (MJPG -> Generic)", "auto_detect"),
    ("Novatek Thermal (MJPG Strict)", "novatek_mjpg"),
    ("Generic Webcam (Auto Format)", "generic_auto"),
]
UVC_AUTO_DETECT_FRAME_TIMEOUT_SECONDS = 2.0
UVC_FRAME_WAIT_SECONDS = 0.5
UVC_RECONNECT_SECONDS = 5.0
UVC_CLOSE_TIMEOUT_MS = 3000
UVC_RECORDING_FOURCC = "MJPG"
UVC_RECORDING_QUEUE_SIZE = 8
UVC_VIEWER_ANALYSIS_INTERVAL_SECONDS = 0.5
DUAL_UVC_QUALITY_INTERVAL_SECONDS = 0.5
UVC_QUALITY_ANALYSIS_MAX_PIXELS = 192 * 144
UVC_TEMPORAL_ANALYSIS_MAX_PIXELS = 160 * 120
UVC_TEMPORAL_SAMPLE_INTERVAL_SECONDS = 0.10
UVC_TEMPORAL_WINDOW_SECONDS = 2.5
UVC_TEMPORAL_MIN_DURATION_SECONDS = 1.5
UVC_TEMPORAL_MIN_SAMPLES = 8
DEFAULT_SPATIAL_QUALITY_WEIGHTS = {
    "detail_clarity_score": 22,
    "contrast_quality_score": 18,
    "flat_cleanliness_score": 18,
    "thermal_uniformity_score": 16,
    "clipping_control_score": 10,
    "artifact_control_score": 16,
}
DEFAULT_UVC_TEMPORAL_WEIGHTS = {
    "spatial_quality_score": 70,
    "flicker_control_score": 15,
    "temporal_stability_score": 15,
}
QUALITY_TIMELINE_METRICS = (
    ("Total Score", "overall_score"),
    ("Spatial Quality", "spatial_quality_score"),
    ("Detail Clarity", "detail_clarity_score"),
    ("Contrast Quality", "contrast_quality_score"),
    ("Flat Cleanliness", "flat_cleanliness_score"),
    ("Thermal Uniformity", "thermal_uniformity_score"),
    ("Clipping Control", "clipping_control_score"),
    ("Artifact Control", "artifact_control_score"),
    ("Flicker Control", "flicker_control_score"),
    ("Temporal Stability", "temporal_stability_score"),
)
ACTIVE_SPATIAL_QUALITY_WEIGHTS = dict(DEFAULT_SPATIAL_QUALITY_WEIGHTS)
ACTIVE_UVC_TEMPORAL_WEIGHTS = dict(DEFAULT_UVC_TEMPORAL_WEIGHTS)
DUAL_UVC_IMAGE_NAME_OPTIONS = [
    ("PEGA", "PEGA"),
    ("TianYuan", "TianYuan"),
    ("Viewsec", "Viewsec")
]


def camera_device_id_text(device_name):
    if isinstance(device_name, str):
        return device_name
    try:
        return bytes(device_name).decode("utf-8", errors="replace")
    except (TypeError, ValueError):
        return str(device_name)


def enumerate_camera_devices():
    """Return camera index/name/stable ID in Qt multimedia enumeration order."""
    if QCameraInfo is not None:
        try:
            devices = []
            for camera_index, camera_info in enumerate(
                QCameraInfo.availableCameras()
            ):
                friendly_name = str(camera_info.description()).strip()
                device_id = camera_device_id_text(camera_info.deviceName())
                devices.append(
                    {
                        "index": camera_index,
                        "name": friendly_name or f"Camera {camera_index}",
                        "device_id": device_id or f"camera-index:{camera_index}",
                    }
                )
            return devices
        except Exception:
            pass

    return [
        {
            "index": camera_index,
            "name": f"Camera {camera_index}",
            "device_id": f"camera-index:{camera_index}",
        }
        for camera_index in CAMERA_INDEX_OPTIONS
    ]


BPR_MARK_RADIUS = 5
BPR_HIGHLIGHT_RADIUS = 2
IMAGE_ALGORITHM_OPTIONS = [
    ("Percentile Auto Contrast", "auto_contrast"),
    ("Fieldscale Local Adaptive AGC (Firmware Core)", "fieldscale"),
    ("Histogram Equalization", "hist_equalization"),
    ("CLAHE", "clahe"),
    ("Gamma", "gamma"),
    ("Bilateral Denoise", "bilateral"),
    ("Unsharp Mask", "unsharp"),
    ("DDE Multi-scale Detail Enhancement", "dde"),
    ("Auto Contrast + CLAHE", "auto_contrast_clahe"),
    ("P1/P99 Percentile Preview (Original Values)", "raw_percentile_preview"),
]
PIPELINE_ALGORITHM_OPTIONS = [
    ("Percentile Auto Contrast", "auto_contrast"),
    ("Fieldscale Local Adaptive AGC (Firmware Core)", "fieldscale"),
    ("Histogram Equalization", "hist_equalization"),
    ("CLAHE", "clahe"),
    ("Gamma", "gamma"),
    ("Bilateral Denoise", "bilateral"),
    ("Unsharp Mask", "unsharp"),
    ("DDE Multi-scale Detail Enhancement", "dde"),
]
RAW_INPUT_MAPPING_OPTIONS = [
    ("Fixed Bit Depth Scale", "fixed_scale"),
    ("P1/P99 from Original RAW DN", "raw_percentile_preview"),
]
IMAGE_ALGORITHM_DESCRIPTIONS = {
    "auto_contrast": (
        "在 8-bit 工作影像上計算 P1/P99，將 P1 映射為黑色、P99 映射為白色，"
        "範圍外像素裁切；不包含 CLAHE。"
    ),
    "fieldscale": (
        "Firmware-matching Fieldscale core: builds 8x8 local Min/Max fields, "
        "applies configurable LES and message passing, and locally maps the input "
        "range to the full display range. Gamma and CLAHE are not applied. RAW "
        "Compare uses original DN; UVC Compare also provides Temporal Smoothing."
    ),
    "hist_equalization": (
        "對整張 8-bit 影像執行全域直方圖均衡化，重新分配灰階以提升整體對比。"
    ),
    "clahe": (
        "將 8-bit 影像分區做限制對比的局部直方圖均衡化；Clip Limit 限制放大量，"
        "Grid Size 控制分區數量。"
    ),
    "gamma": (
        "以 Gamma 曲線重新映射 8-bit 灰階；目前公式中 Gamma 小於 1 會提亮中間灰階，"
        "大於 1 會壓暗。"
    ),
    "bilateral": (
        "在 8-bit 影像上做保邊雙邊濾波，降低雜訊並盡量保留物體邊緣。"
    ),
    "unsharp": (
        "以原圖減去模糊影像取得高頻細節，再依 Amount 加回原圖以增強邊緣。"
    ),
    "dde": (
        "四層 Gaussian Pyramid 多尺度細節增強：Beta 控制基礎影像保留量，"
        "Alpha 放大正向細節，Sigma 控制 NLM 降噪強度，最後執行 P2/P98 對比拉伸。"
    ),
    "auto_contrast_clahe": (
        "先在 8-bit 工作影像做 P1/P99 Auto Contrast，再執行 CLAHE 局部對比增強。"
    ),
    "raw_percentile_preview": (
        "直接在輸入的原始數值上計算 P1/P99。P1 是排除最低 1% 像素後的灰階下限，"
        "P99 是排除最高 1% 像素前的灰階上限；將 P1 映射為黑色 0、P99 映射為白色 255，"
        "範圍外的像素裁切，最後輸出成 8-bit 顯示影像。"
    ),
}

DEFAULT_ALGORITHM_PARAMETERS = {
    "auto_contrast": {"low_percentile": 1.0, "high_percentile": 99.0},
    "fieldscale": {
        "grid_count": 8,
        "les_distance": 2,
        "les_threshold_raw14": 400,
        "message_passing_distance": 1,
        "message_passing_iterations": 7,
    },
    "hist_equalization": {},
    "clahe": {"clahe_clip": 2.0, "clahe_grid": 8},
    "gamma": {"gamma": 0.7},
    "bilateral": {"bilateral_strength": 2},
    "unsharp": {"unsharp_amount": 1.0},
    "dde": {"dde_sigma": 1.0, "dde_alpha": 1.5, "dde_beta": 0.5},
}


def algorithm_display_name(algorithm_key):
    for display_name, key in PIPELINE_ALGORITHM_OPTIONS:
        if key == algorithm_key:
            return display_name
    return algorithm_key


def default_algorithm_parameters(algorithm_key):
    if algorithm_key not in DEFAULT_ALGORITHM_PARAMETERS:
        raise ValueError(f"Unsupported pipeline algorithm: {algorithm_key}")
    return dict(DEFAULT_ALGORITHM_PARAMETERS[algorithm_key])


@dataclass
class AlgorithmStage:
    algorithm_key: str
    enabled: bool = True
    parameters: dict = field(default_factory=dict)

    def __post_init__(self):
        defaults = default_algorithm_parameters(self.algorithm_key)
        defaults.update(self.parameters)
        self.parameters = defaults

    def clone(self):
        return AlgorithmStage(self.algorithm_key, self.enabled, dict(self.parameters))

    def to_dict(self):
        return {
            "algorithm": self.algorithm_key,
            "name": algorithm_display_name(self.algorithm_key),
            "enabled": bool(self.enabled),
            "parameters": dict(self.parameters),
        }


def timestamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]


def unique_result_dir(base_dir, function_name):
    root = Path(base_dir).expanduser().resolve() if base_dir else Path.cwd() / "Thermal_RAW_Results"
    function_root = root / function_name
    function_root.mkdir(parents=True, exist_ok=True)
    out = function_root / timestamp()
    counter = 1
    original = out
    while out.exists():
        out = Path(f"{original}_{counter:02d}")
        counter += 1
    out.mkdir(parents=True, exist_ok=False)
    return out


def load_raw(path, width, height):
    expected_pixels = width * height
    expected_bytes = expected_pixels * 2
    size = os.path.getsize(path)
    if size != expected_bytes:
        raise ValueError(
            f"RAW size error:\n{path}\nActual: {size:,} bytes\nExpected: {expected_bytes:,} bytes "
            f"({width}x{height}x2)"
        )
    raw = np.fromfile(path, dtype=RAW_DTYPE)
    return raw.reshape((height, width))


IMAGE_INPUT_FILTER = "RAW / Images (*.raw *.RAW *.bin *.BIN *.jpg *.JPG *.jpeg *.JPEG *.png *.PNG *.bmp *.BMP *.tif *.TIF *.tiff *.TIFF)"
IMAGE_INPUT_NOTE = ("JPG / PNG / BMP / TIFF: automatic dimensions and bit depth; color is converted to grayscale. "
                    "RAW resolution / bit-depth settings apply only to RAW / BIN files.")
IMAGE_INPUT_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}


def load_analysis_image(path, width, height, bit_depth):
    """Decode image files to native-depth grayscale; preserve RAW loading rules."""
    is_image = Path(path).suffix.lower() in IMAGE_INPUT_EXTENSIONS
    if is_image:
        pixels = cv2.imdecode(np.fromfile(path, dtype=np.uint8), cv2.IMREAD_UNCHANGED)
        if pixels is None or pixels.size == 0:
            raise ValueError(f"Cannot decode image: {path}")
        if pixels.dtype not in (np.uint8, np.uint16):
            raise ValueError("Only unsigned 8-bit and 16-bit images are supported.")
        if pixels.ndim == 3:
            channels = pixels.shape[2]
            if channels not in (3, 4):
                raise ValueError("Unsupported image channel count.")
            pixels = cv2.cvtColor(pixels, cv2.COLOR_BGRA2GRAY if channels == 4 else cv2.COLOR_BGR2GRAY)
        if pixels.ndim != 2:
            raise ValueError("Expected a single grayscale image.")
        bit_depth = pixels.dtype.itemsize * 8
    else:
        pixels = load_raw(path, width, height)
    return pixels, {
        "width": int(pixels.shape[1]), "height": int(pixels.shape[0]),
        "bit_depth": int(bit_depth), "value_max": (1 << bit_depth) - 1,
        "value_name": "Gray level" if is_image else "RAW DN",
        "source_type": "image" if is_image else "raw",
    }


def find_raw_files(folder):
    files = []
    for ext in ("*.raw", "*.RAW", "*.bin", "*.BIN"):
        files.extend(glob.glob(os.path.join(folder, ext)))
    return sorted(set(files))


def average_raw_folder(folder, width, height):
    files = find_raw_files(folder)
    if not files:
        raise FileNotFoundError(f"No RAW/BIN files found:\n{folder}")
    total = np.zeros((height, width), dtype=np.float64)
    valid = []
    skipped = []
    for path in files:
        try:
            total += load_raw(path, width, height)
            valid.append(path)
        except Exception as exc:
            skipped.append((path, str(exc)))
    if not valid:
        raise ValueError("No valid RAW files in selected folder.")
    avg = np.rint(total / len(valid)).astype(np.uint16)
    return avg, valid, skipped


def raw_to_display_u8(image, raw_max):
    scaled = np.clip(image, 0, raw_max).astype(np.float64) * 255.0 / raw_max
    return np.rint(scaled).astype(np.uint8)


def signed_delta_to_display_u8(image, raw_max):
    scaled = (
        np.clip(image, -raw_max, raw_max).astype(np.float64) + raw_max
    ) * 255.0 / (2.0 * raw_max)
    return np.rint(scaled).astype(np.uint8)


def percentile_stretch_gray(gray, low_p=1.0, high_p=99.0):
    gray_f = gray.astype(np.float32)
    vmin = float(np.percentile(gray_f, low_p))
    vmax = float(np.percentile(gray_f, high_p))
    if vmax <= vmin:
        return gray.copy(), vmin, vmax
    stretched = (gray_f - vmin) * 255.0 / (vmax - vmin)
    return np.clip(stretched, 0, 255).astype(np.uint8), vmin, vmax


def percentile_preview_u8(image, p_low=1.0, p_high=99.0):
    arr = image.astype(np.float64)
    vmin = float(np.percentile(arr, p_low))
    vmax = float(np.percentile(arr, p_high))
    if vmax <= vmin:
        vmin = float(arr.min())
        vmax = float(arr.max())
    if vmax <= vmin:
        return np.zeros(arr.shape, dtype=np.uint8), vmin, vmax
    stretched = (arr - vmin) * 255.0 / (vmax - vmin)
    return np.clip(stretched, 0, 255).astype(np.uint8), vmin, vmax


def dde_multiscale_u8(image, sigma=1.0, alpha=1.5, beta=0.5, levels=4):
    """Executable OpenCV/NumPy form of the supplied multi-scale DDE reference."""
    if image.ndim != 2:
        raise ValueError("DDE input must be a 2D grayscale image.")
    minimum_size = 1 << (levels - 1)
    if min(image.shape) < minimum_size:
        raise ValueError(
            f"DDE requires image width and height >= {minimum_size}."
        )

    sigma = max(0.0, float(sigma))
    alpha = max(0.0, float(alpha))
    beta = max(0.0, float(beta))
    source_u8 = np.ascontiguousarray(
        np.clip(image, 0, 255).astype(np.uint8)
    )
    source = source_u8.astype(np.float32) / 255.0
    pyramid = [source]
    for _ in range(1, levels):
        pyramid.append(cv2.pyrDown(pyramid[-1]))

    # The supplied sample creates a base value but never adds it to its output and
    # never uses beta. Here beta is the retained source/base contribution.
    enhanced = beta * source
    source_size = (source.shape[1], source.shape[0])
    for level in range(1, levels):
        previous = pyramid[level - 1]
        expanded = cv2.resize(
            pyramid[level],
            (previous.shape[1], previous.shape[0]),
            interpolation=cv2.INTER_LINEAR,
        )
        detail = np.clip((previous - expanded) * alpha, 0.0, 1.0)
        if detail.shape != source.shape:
            detail = cv2.resize(
                detail, source_size, interpolation=cv2.INTER_LINEAR
            )
        enhanced += detail

    enhanced_u8 = np.rint(
        np.clip(enhanced, 0.0, 1.0) * 255.0
    ).astype(np.uint8)
    if sigma > 0.0:
        denoised = cv2.fastNlMeansDenoising(
            enhanced_u8,
            None,
            h=10.0 * sigma,
            templateWindowSize=7,
            searchWindowSize=21,
        )
    else:
        denoised = enhanced_u8

    p2, p98 = np.percentile(denoised, (2.0, 98.0))
    if p98 <= p2:
        output = denoised.copy()
    else:
        output = np.clip(
            (denoised.astype(np.float32) - p2) * 255.0 / (p98 - p2),
            0.0,
            255.0,
        ).astype(np.uint8)
    return output, [
        f"DDE Gaussian pyramid levels={levels}",
        f"Sigma={sigma:.2f}, Alpha={alpha:.2f}, Beta={beta:.2f}",
        f"NLM h={10.0 * sigma:.2f}; contrast P2={p2:.2f}, P98={p98:.2f}",
    ]


FIELDSCALE_RAW14_MAX = 16383
FIELDSCALE_Q8_ONE = 256
FIELDSCALE_INTERP_ONE = 65536


def _fieldscale_suppress_extrema(
    grid, threshold_q8, suppress_maximum, distance=2
):
    result = grid.copy()
    height, width = grid.shape
    distance = max(1, int(distance))
    for y in range(height):
        for x in range(width):
            neighbors = []
            for dy in range(-distance, distance + 1):
                for dx in range(-distance, distance + 1):
                    if dx == 0 and dy == 0:
                        continue
                    nx = x + dx
                    ny = y + dy
                    if 0 <= nx < width and 0 <= ny < height:
                        neighbors.append(int(grid[ny, nx]))
            if not neighbors:
                continue
            current = int(grid[y, x])
            mean_value = sum(neighbors) // len(neighbors)
            if suppress_maximum:
                if (
                    current >= max(neighbors)
                    and current > mean_value
                    and current - mean_value > threshold_q8
                ):
                    result[y, x] = mean_value + threshold_q8
            elif (
                current <= min(neighbors)
                and mean_value > current
                and mean_value - current > threshold_q8
            ):
                result[y, x] = mean_value - threshold_q8
    return result


def _fieldscale_message_pass(grid, increase, distance=1):
    result = grid.copy()
    height, width = grid.shape
    distance = max(1, int(distance))
    for y in range(height):
        for x in range(width):
            values = []
            for dy in range(-distance, distance + 1):
                for dx in range(-distance, distance + 1):
                    nx = x + dx
                    ny = y + dy
                    if 0 <= nx < width and 0 <= ny < height:
                        values.append(int(grid[ny, nx]))
            current = int(grid[y, x])
            mean_value = sum(values) // len(values)
            result[y, x] = (
                max(mean_value, current) if increase else min(mean_value, current)
            )
    return result


def _fieldscale_resize_q8(grid, output_width, output_height):
    grid_height, grid_width = grid.shape
    x_start_q16 = grid_width * 32768 // output_width - 32768
    x_step_q16 = grid_width * FIELDSCALE_INTERP_ONE // output_width
    y_start_q16 = grid_height * 32768 // output_height - 32768
    y_step_q16 = grid_height * FIELDSCALE_INTERP_ONE // output_height
    source_x = x_start_q16 + np.arange(output_width, dtype=np.int64) * x_step_q16
    source_y = y_start_q16 + np.arange(output_height, dtype=np.int64) * y_step_q16
    source_x = np.clip(source_x, 0, (grid_width - 1) * FIELDSCALE_INTERP_ONE)
    source_y = np.clip(source_y, 0, (grid_height - 1) * FIELDSCALE_INTERP_ONE)
    x0 = (source_x >> 16).astype(np.intp)
    y0 = (source_y >> 16).astype(np.intp)
    x1 = np.minimum(x0 + 1, grid_width - 1)
    y1 = np.minimum(y0 + 1, grid_height - 1)
    fx = source_x & 0xFFFF
    fy = source_y & 0xFFFF

    horizontal = (
        grid[:, x0] * (FIELDSCALE_INTERP_ONE - fx)[None, :]
        + grid[:, x1] * fx[None, :]
        + FIELDSCALE_INTERP_ONE // 2
    ) // FIELDSCALE_INTERP_ONE
    return (
        horizontal[y0, :] * (FIELDSCALE_INTERP_ONE - fy)[:, None]
        + horizontal[y1, :] * fy[:, None]
        + FIELDSCALE_INTERP_ONE // 2
    ) // FIELDSCALE_INTERP_ONE


def fieldscale_firmware_preview_u8(
    image,
    input_max,
    state=None,
    temporal=False,
    grid_count=8,
    les_distance=2,
    les_threshold_raw14=400,
    message_passing_distance=1,
    message_passing_iterations=7,
    temporal_smoothing_percent=90,
):
    if image.ndim != 2:
        raise ValueError("Fieldscale input must be a 2D grayscale image.")
    height, width = image.shape
    grid_count = int(grid_count)
    if (
        grid_count < 2
        or width < grid_count
        or height < grid_count
        or width % grid_count
        or height % grid_count
    ):
        raise ValueError(
            "Fieldscale requires width and height divisible by the selected "
            f"Grid count ({grid_count}x{grid_count})."
        )
    if input_max <= 0:
        raise ValueError("Fieldscale input maximum must be positive.")

    source = np.clip(image, 0, input_max).astype(np.int64)
    patch_height = height // grid_count
    patch_width = width // grid_count
    patches = source.reshape(
        grid_count,
        patch_height,
        grid_count,
        patch_width,
    )
    min_grid = patches.min(axis=(1, 3)) * FIELDSCALE_Q8_ONE
    max_grid = patches.max(axis=(1, 3)) * FIELDSCALE_Q8_ONE
    les_distance = max(1, min(grid_count - 1, int(les_distance)))
    message_passing_distance = max(
        1,
        min(grid_count - 1, int(message_passing_distance)),
    )
    message_passing_iterations = max(0, int(message_passing_iterations))
    les_threshold_raw14 = max(0, int(les_threshold_raw14))
    temporal_smoothing_percent = max(
        0, min(99, int(temporal_smoothing_percent))
    )
    threshold_q8 = max(
        0,
        int(round(
            les_threshold_raw14 * input_max * FIELDSCALE_Q8_ONE
            / FIELDSCALE_RAW14_MAX
        )),
    )
    max_grid = _fieldscale_suppress_extrema(
        max_grid, threshold_q8, True, les_distance
    )
    max_grid = _fieldscale_suppress_extrema(
        max_grid, threshold_q8, False, les_distance
    )
    for _ in range(message_passing_iterations):
        min_grid = _fieldscale_message_pass(
            min_grid, False, message_passing_distance
        )
        max_grid = _fieldscale_message_pass(
            max_grid, True, message_passing_distance
        )

    temporal_applied = False
    if temporal and state is not None:
        previous_min = state.get("min_grid")
        previous_max = state.get("max_grid")
        previous_input_max = state.get("input_max")
        if (
            previous_min is not None
            and previous_max is not None
            and previous_min.shape == min_grid.shape
            and previous_max.shape == max_grid.shape
            and previous_input_max == input_max
        ):
            previous_weight = temporal_smoothing_percent
            current_weight = 100 - previous_weight
            min_grid = (
                current_weight * min_grid
                + previous_weight * previous_min
                + 50
            ) // 100
            max_grid = (
                current_weight * max_grid
                + previous_weight * previous_max
                + 50
            ) // 100
            temporal_applied = True
        state["min_grid"] = min_grid.copy()
        state["max_grid"] = max_grid.copy()
        state["input_max"] = input_max

    min_field = _fieldscale_resize_q8(min_grid, width, height)
    max_field = _fieldscale_resize_q8(max_grid, width, height)
    low = np.minimum(min_field, max_field)
    high = np.maximum(min_field, max_field)
    value_q8 = source * FIELDSCALE_Q8_ONE
    output = np.zeros(source.shape, dtype=np.int64)
    output[value_q8 >= high] = input_max
    valid = (value_q8 > low) & (value_q8 < high) & (high > low)
    output[valid] = (
        (value_q8[valid] - low[valid]) * input_max
        // (high[valid] - low[valid])
    )
    preview = np.rint(output.astype(np.float64) * 255.0 / input_max).astype(np.uint8)
    threshold_levels = threshold_q8 / FIELDSCALE_Q8_ONE
    details = [
        f"Fieldscale firmware core: grid={grid_count}x{grid_count}, "
        f"LES distance={les_distance}",
        f"LES threshold={les_threshold_raw14} RAW14 DN "
        f"({threshold_levels:.2f} input levels)",
        f"message passing distance={message_passing_distance}, "
        f"iterations={message_passing_iterations}",
        "Gamma=off, CLAHE=off",
    ]
    if temporal:
        details.append(
            f"temporal=current {100 - temporal_smoothing_percent}% / "
            f"previous {temporal_smoothing_percent}%"
            if temporal_applied
            else "temporal=first frame"
        )
    return preview, details


def apply_image_algorithm(
    gray,
    algorithm_key,
    low_percentile=1.0,
    high_percentile=99.0,
    clahe_clip=2.0,
    clahe_grid=8,
    gamma=0.7,
    bilateral_strength=2,
    unsharp_amount=1.0,
    clahe=None,
    input_max=255,
    fieldscale_state=None,
    fieldscale_temporal=False,
    grid_count=8,
    les_distance=2,
    les_threshold_raw14=400,
    message_passing_distance=1,
    message_passing_iterations=7,
    temporal_smoothing_percent=90,
    dde_sigma=1.0,
    dde_alpha=1.5,
    dde_beta=0.5,
):
    if cv2 is None:
        raise RuntimeError("OpenCV not installed. pip install opencv-python")
    if gray.ndim != 2:
        raise ValueError("Image algorithm input must be a 2D grayscale image.")

    details = []

    if algorithm_key == "raw_percentile_preview":
        processed, vmin, vmax = percentile_preview_u8(gray)
        details.append(f"Original-value P1.0={vmin:.1f}, P99.0={vmax:.1f}")
        return processed, details

    if algorithm_key == "fieldscale":
        return fieldscale_firmware_preview_u8(
            gray,
            input_max,
            state=fieldscale_state,
            temporal=fieldscale_temporal,
            grid_count=grid_count,
            les_distance=les_distance,
            les_threshold_raw14=les_threshold_raw14,
            message_passing_distance=message_passing_distance,
            message_passing_iterations=message_passing_iterations,
            temporal_smoothing_percent=temporal_smoothing_percent,
        )

    if algorithm_key == "dde":
        return dde_multiscale_u8(
            gray,
            sigma=dde_sigma,
            alpha=dde_alpha,
            beta=dde_beta,
        )

    processed = np.ascontiguousarray(np.clip(gray, 0, 255).astype(np.uint8))

    if algorithm_key in ("auto_contrast", "auto_contrast_clahe"):
        if input_max > 255:
            processed, vmin, vmax = percentile_preview_u8(
                gray, low_percentile, high_percentile
            )
        else:
            processed, vmin, vmax = percentile_stretch_gray(
                processed, low_percentile, high_percentile
            )
        details.append(
            f"P{low_percentile:.1f}={vmin:.1f}, P{high_percentile:.1f}={vmax:.1f}"
        )

    if algorithm_key in ("clahe", "auto_contrast_clahe"):
        if clahe is None:
            clahe = cv2.createCLAHE(
                clipLimit=clahe_clip,
                tileGridSize=(clahe_grid, clahe_grid),
            )
        processed = clahe.apply(processed)
        details.append(f"CLAHE clip={clahe_clip:.1f}, grid={clahe_grid}x{clahe_grid}")

    if algorithm_key == "hist_equalization":
        processed = cv2.equalizeHist(processed)
        details.append("Global histogram equalization")
    elif algorithm_key == "gamma":
        lut = np.array(
            [((value / 255.0) ** gamma) * 255.0 for value in range(256)],
            dtype=np.uint8,
        )
        processed = cv2.LUT(processed, lut)
        details.append(f"Gamma={gamma:.1f}")
    elif algorithm_key == "bilateral":
        diameter = bilateral_strength * 2 + 1
        sigma = bilateral_strength * 20
        processed = cv2.bilateralFilter(processed, diameter, sigma, sigma)
        details.append(f"Bilateral strength={bilateral_strength}")
    elif algorithm_key == "unsharp":
        blurred = cv2.GaussianBlur(processed, (0, 0), 1.0)
        processed = cv2.addWeighted(
            processed, 1.0 + unsharp_amount, blurred, -unsharp_amount, 0
        )
        details.append(f"Unsharp amount={unsharp_amount:.1f}")

    if not details:
        raise ValueError(f"Unknown image algorithm: {algorithm_key}")
    return processed, details


def apply_algorithm_pipeline(
    gray,
    stages,
    input_max=255,
    clahe_cache=None,
    fieldscale_states=None,
    fieldscale_temporal=False,
    allow_empty=False,
    pad_fieldscale_input=False,
):
    """Apply enabled stages in top-to-bottom order and keep every step result."""
    if cv2 is None:
        raise RuntimeError("OpenCV not installed. pip install opencv-python")
    if gray.ndim != 2:
        raise ValueError("Algorithm pipeline input must be a 2D grayscale image.")
    if not any(stage.enabled for stage in stages):
        if allow_empty:
            return raw_to_display_u8(gray, input_max), []
        raise ValueError("Add and enable at least one algorithm stage.")

    processed = np.ascontiguousarray(gray.copy())
    current_max = int(input_max)
    records = []
    execution_index = 0
    for pipeline_index, stage in enumerate(stages, start=1):
        if not stage.enabled:
            continue
        execution_index += 1
        params = default_algorithm_parameters(stage.algorithm_key)
        params.update(stage.parameters)
        clahe = None
        if stage.algorithm_key == "clahe" and clahe_cache is not None:
            cache_key = (
                pipeline_index,
                float(params["clahe_clip"]),
                int(params["clahe_grid"]),
            )
            clahe = clahe_cache.get(cache_key)
            if clahe is None:
                clahe = cv2.createCLAHE(
                    clipLimit=cache_key[1],
                    tileGridSize=(cache_key[2], cache_key[2]),
                )
                clahe_cache[cache_key] = clahe

        fieldscale_state = None
        if stage.algorithm_key == "fieldscale" and fieldscale_states is not None:
            fieldscale_state = fieldscale_states.setdefault(pipeline_index, {})

        started = time.perf_counter()
        original_shape = processed.shape
        padding = (0, 0)
        if pad_fieldscale_input and stage.algorithm_key == "fieldscale":
            grid = int(params.get("grid_count", 8))
            if grid >= 2:
                padding = ((-original_shape[0]) % grid, (-original_shape[1]) % grid)
                if any(padding):
                    processed = np.pad(processed, ((0, padding[0]), (0, padding[1])), mode="edge")
        processed, details = apply_image_algorithm(
            processed,
            stage.algorithm_key,
            low_percentile=params.get("low_percentile", 1.0),
            high_percentile=params.get("high_percentile", 99.0),
            clahe_clip=params.get("clahe_clip", 2.0),
            clahe_grid=params.get("clahe_grid", 8),
            gamma=params.get("gamma", 0.7),
            bilateral_strength=params.get("bilateral_strength", 2),
            unsharp_amount=params.get("unsharp_amount", 1.0),
            clahe=clahe,
            input_max=current_max,
            fieldscale_state=fieldscale_state,
            fieldscale_temporal=(
                fieldscale_temporal and stage.algorithm_key == "fieldscale"
            ),
            grid_count=params.get("grid_count", 8),
            les_distance=params.get("les_distance", 2),
            les_threshold_raw14=params.get("les_threshold_raw14", 400),
            message_passing_distance=params.get(
                "message_passing_distance", 1
            ),
            message_passing_iterations=params.get(
                "message_passing_iterations", 7
            ),
            temporal_smoothing_percent=params.get(
                "temporal_smoothing_percent", 90
            ),
            dde_sigma=params.get("dde_sigma", 1.0),
            dde_alpha=params.get("dde_alpha", 1.5),
            dde_beta=params.get("dde_beta", 0.5),
        )
        if any(padding):
            processed = processed[:original_shape[0], :original_shape[1]].copy()
            details.append(
                f"Image edge padding: bottom={padding[0]}, right={padding[1]}; cropped back to original size"
            )
        elapsed_ms = (time.perf_counter() - started) * 1000.0
        current_max = 255
        records.append(
            {
                "pipeline_index": pipeline_index,
                "execution_index": execution_index,
                "algorithm": stage.algorithm_key,
                "name": algorithm_display_name(stage.algorithm_key),
                "parameters": dict(params),
                "details": list(details),
                "elapsed_ms": elapsed_ms,
                "image": processed,
            }
        )
    return processed, records


def pipeline_execution_json(records):
    return [
        {
            "pipeline_index": record["pipeline_index"],
            "execution_index": record["execution_index"],
            "algorithm": record["algorithm"],
            "name": record["name"],
            "parameters": dict(record["parameters"]),
            "details": list(record["details"]),
            "elapsed_ms": round(float(record["elapsed_ms"]), 3),
        }
        for record in records
    ]


def pipeline_summary_text(stages):
    names = [algorithm_display_name(stage.algorithm_key) for stage in stages if stage.enabled]
    return " -> ".join(names) if names else "No enabled algorithm"


def pipeline_record_image(records, pipeline_index, final_image):
    if pipeline_index is None:
        return final_image
    for record in records:
        if record["pipeline_index"] == pipeline_index:
            return record["image"]
    return final_image


def uvc_to_gray(frame):
    if frame.ndim == 2:
        return frame
    if frame.shape[2] == 1:
        return frame[:, :, 0]
    if frame.shape[2] == 4:
        return cv2.cvtColor(frame, cv2.COLOR_BGRA2GRAY)
    return cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)


def image_stats(image):
    arr = np.asarray(image, dtype=np.float64)
    return {
        "count": int(arr.size),
        "min": float(arr.min()),
        "max": float(arr.max()),
        "mean": float(arr.mean()),
        "median": float(np.median(arr)),
        "std": float(arr.std()),
        "p1": float(np.percentile(arr, 1.0)),
        "p99": float(np.percentile(arr, 99.0)),
    }


def gray_image_analysis(image):
    gray = np.clip(np.asarray(image), 0, 255).astype(np.uint8)
    flat = gray.reshape(-1)
    histogram = np.bincount(flat, minlength=256).astype(np.int64)
    percentiles = np.percentile(flat, [1.0, 5.0, 50.0, 95.0, 99.0])
    probabilities = histogram[histogram > 0].astype(np.float64) / flat.size
    entropy = float(-np.sum(probabilities * np.log2(probabilities)))
    zero_count = int(histogram[0])
    full_count = int(histogram[255])
    limited_low_count = int(histogram[:17].sum())
    limited_high_count = int(histogram[235:].sum())
    stats = {
        "count": int(flat.size),
        "min": int(flat.min()),
        "max": int(flat.max()),
        "mean": float(flat.mean()),
        "median": float(percentiles[2]),
        "std": float(flat.std()),
        "p1": float(percentiles[0]),
        "p5": float(percentiles[1]),
        "p95": float(percentiles[3]),
        "p99": float(percentiles[4]),
        "p99_p1": float(percentiles[4] - percentiles[0]),
        "unique_levels": int(np.count_nonzero(histogram)),
        "entropy_bits": entropy,
        "zero_count": zero_count,
        "zero_percent": float(zero_count * 100.0 / flat.size),
        "full_count": full_count,
        "full_percent": float(full_count * 100.0 / flat.size),
        "limited_low_count": limited_low_count,
        "limited_low_percent": float(limited_low_count * 100.0 / flat.size),
        "limited_high_count": limited_high_count,
        "limited_high_percent": float(limited_high_count * 100.0 / flat.size),
    }
    return stats, histogram


def raw_dn_image_analysis(image, raw_max):
    raw = np.asarray(image).astype(np.uint16, copy=False)
    flat = raw.reshape(-1)
    percentiles = np.percentile(flat, [1.0, 5.0, 50.0, 95.0, 99.0])
    _levels, level_counts = np.unique(flat, return_counts=True)
    probabilities = level_counts.astype(np.float64) / flat.size
    entropy = float(-np.sum(probabilities * np.log2(probabilities)))
    histogram, _edges = np.histogram(
        np.clip(flat, 0, raw_max),
        bins=256,
        range=(0, int(raw_max) + 1),
    )
    zero_count = int(np.count_nonzero(flat == 0))
    max_count = int(np.count_nonzero(flat == raw_max))
    over_max_count = int(np.count_nonzero(flat > raw_max))
    p99_p1 = float(percentiles[4] - percentiles[0])
    stats = {
        "count": int(flat.size),
        "min": int(flat.min()),
        "max": int(flat.max()),
        "mean": float(flat.mean()),
        "median": float(percentiles[2]),
        "std": float(flat.std()),
        "p1": float(percentiles[0]),
        "p5": float(percentiles[1]),
        "p95": float(percentiles[3]),
        "p99": float(percentiles[4]),
        "p99_p1": p99_p1,
        "p99_p1_percent": float(p99_p1 * 100.0 / raw_max),
        "unique_levels": int(level_counts.size),
        "entropy_bits": entropy,
        "zero_count": zero_count,
        "zero_percent": float(zero_count * 100.0 / flat.size),
        "max_count": max_count,
        "max_percent": float(max_count * 100.0 / flat.size),
        "over_max_count": over_max_count,
        "over_max_percent": float(over_max_count * 100.0 / flat.size),
        "raw_max": int(raw_max),
    }
    return stats, histogram.astype(np.int64)


def render_gray_histogram_rgb(
    histogram,
    stats,
    width=800,
    height=360,
    value_max=255,
    title="Gray Histogram (256 bins)",
):
    canvas = np.full((height, width, 3), (23, 25, 29), dtype=np.uint8)
    left, right, top, bottom = 58, width - 22, 48, height - 48
    plot_width = right - left
    plot_height = bottom - top

    cv2.rectangle(canvas, (left, top), (right, bottom), (67, 73, 82), 1)
    for fraction in (0.25, 0.5, 0.75):
        y = int(round(bottom - fraction * plot_height))
        cv2.line(canvas, (left, y), (right, y), (45, 49, 56), 1)

    peak_count = max(1, int(np.max(histogram)))
    for bin_index, count in enumerate(histogram):
        x0 = left + int(bin_index * plot_width / 256.0)
        x1 = left + int((bin_index + 1) * plot_width / 256.0)
        y = bottom - int(round(float(count) * plot_height / peak_count))
        cv2.rectangle(
            canvas,
            (x0, max(top, y)),
            (max(x0, x1), bottom),
            (184, 126, 67),
            -1,
        )

    cumulative = np.cumsum(histogram).astype(np.float64)
    if cumulative[-1] > 0:
        cumulative /= cumulative[-1]
        cdf_points = []
        for bin_index, value in enumerate(cumulative):
            x = left + int(round(bin_index * plot_width / 255.0))
            y = bottom - int(round(value * plot_height))
            cdf_points.append((x, y))
        cv2.polylines(
            canvas,
            [np.asarray(cdf_points, dtype=np.int32)],
            False,
            (72, 205, 235),
            2,
            cv2.LINE_AA,
        )

    marker_specs = (
        (stats["p1"], "P1", (105, 190, 105)),
        (stats["median"], "P50", (225, 225, 225)),
        (stats["p99"], "P99", (90, 145, 245)),
    )
    for marker_index, (value, label, color) in enumerate(marker_specs):
        x = max(
            left,
            min(
                left + int(round(float(value) * plot_width / value_max)),
                right,
            ),
        )
        cv2.line(canvas, (x, top), (x, bottom), color, 1, cv2.LINE_AA)
        cv2.putText(
            canvas,
            f"{label} {value:.1f}",
            (
                max(left, min(x - 28, right - 72)),
                top + 16 + marker_index * 16,
            ),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.42,
            color,
            1,
            cv2.LINE_AA,
        )

    tick_values = [
        int(round(value_max * fraction))
        for fraction in (0.0, 0.25, 0.5, 0.75, 1.0)
    ]
    for value in tick_values:
        x = left + int(round(value * plot_width / value_max))
        cv2.line(canvas, (x, bottom), (x, bottom + 5), (150, 157, 166), 1)
        cv2.putText(
            canvas,
            str(value),
            (x - max(7, len(str(value)) * 4), bottom + 22),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.43,
            (190, 196, 204),
            1,
            cv2.LINE_AA,
        )

    cv2.putText(
        canvas,
        title,
        (left, 24),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.62,
        (235, 238, 242),
        1,
        cv2.LINE_AA,
    )
    cv2.putText(
        canvas,
        f"Peak count {peak_count:,}",
        (left, 42),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.42,
        (170, 177, 186),
        1,
        cv2.LINE_AA,
    )
    cv2.line(canvas, (right - 145, 22), (right - 115, 22), (72, 205, 235), 2)
    cv2.putText(
        canvas,
        "CDF",
        (right - 107, 27),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.45,
        (205, 211, 218),
        1,
        cv2.LINE_AA,
    )
    return cv2.cvtColor(canvas, cv2.COLOR_BGR2RGB)


def _excel_cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Path):
        return str(value)
    return json.dumps(value, ensure_ascii=False)


def _flatten_excel_fields(values, prefix=""):
    for key, value in values.items():
        field_name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            yield from _flatten_excel_fields(value, field_name)
        else:
            yield field_name, _excel_cell_value(value)


def save_histogram_excel(
    path,
    report_title,
    metadata,
    analyses,
    value_max=255,
    value_name="Gray",
):
    """Save statistics, 256-bin Histogram/CDF data, and native Excel charts."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel output requires openpyxl. Install dependencies with: "
            "pip install -r requirements_thermal_raw_toolkit.txt"
        ) from exc

    normalized_analyses = []
    used_sheet_names = set()
    for analysis_index, analysis in enumerate(analyses, start=1):
        label = str(analysis["label"])
        histogram = np.asarray(analysis["histogram"], dtype=np.int64)
        if histogram.shape != (256,):
            raise ValueError(
                f"Excel Histogram must contain 256 bins, got {histogram.shape}."
            )
        base_name = re.sub(r"[\\/*?:\[\]]", "_", f"Histogram {label}")[:31]
        sheet_name = base_name or f"Histogram {analysis_index}"
        suffix = 2
        while sheet_name in used_sheet_names:
            suffix_text = f" {suffix}"
            sheet_name = f"{base_name[:31 - len(suffix_text)]}{suffix_text}"
            suffix += 1
        used_sheet_names.add(sheet_name)
        normalized_analyses.append(
            {
                "label": label,
                "stats": dict(analysis["stats"]),
                "histogram": histogram,
                "sheet_name": sheet_name,
            }
        )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append([report_title])
    summary_sheet.merge_cells("A1:C1")
    summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor="2F75B5")
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    summary_sheet.append(["Section", "Field", "Value"])
    for cell in summary_sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    for field_name, value in _flatten_excel_fields(metadata):
        summary_sheet.append(["Report", field_name, value])
    for analysis in normalized_analyses:
        for field_name, value in _flatten_excel_fields(analysis["stats"]):
            summary_sheet.append([analysis["label"], field_name, value])
    summary_sheet.freeze_panes = "A3"
    summary_sheet.auto_filter.ref = f"A2:C{summary_sheet.max_row}"
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 32
    summary_sheet.column_dimensions["C"].width = 48

    bin_width = (int(value_max) + 1) / 256.0
    histogram_sheets = []
    for analysis in normalized_analyses:
        worksheet = workbook.create_sheet(analysis["sheet_name"])
        histogram_sheets.append((analysis, worksheet))
        worksheet.append(
            [
                "Bin",
                f"{value_name} Start",
                f"{value_name} End",
                f"{value_name} Center",
                f"{analysis['label']} Count",
                f"{analysis['label']} Percentage",
                f"{analysis['label']} Cumulative Count",
                f"{analysis['label']} CDF Percentage",
            ]
        )
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        histogram = analysis["histogram"]
        total_count = int(histogram.sum())
        cumulative_count = 0
        for bin_index, count_value in enumerate(histogram):
            count = int(count_value)
            cumulative_count += count
            range_start = int(np.ceil(bin_index * bin_width))
            range_end = int(np.ceil((bin_index + 1) * bin_width) - 1)
            range_start = max(0, min(range_start, int(value_max)))
            range_end = max(range_start, min(range_end, int(value_max)))
            center = (range_start + range_end) / 2.0
            percentage = count / total_count if total_count else 0.0
            cdf_percentage = (
                cumulative_count / total_count if total_count else 0.0
            )
            worksheet.append(
                [
                    bin_index,
                    range_start,
                    range_end,
                    center,
                    count,
                    percentage,
                    cumulative_count,
                    cdf_percentage,
                ]
            )
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.column_dimensions["A"].width = 10
        worksheet.column_dimensions["B"].width = 16
        worksheet.column_dimensions["C"].width = 16
        worksheet.column_dimensions["D"].width = 17
        worksheet.column_dimensions["E"].width = 16
        worksheet.column_dimensions["F"].width = 16
        worksheet.column_dimensions["G"].width = 20
        worksheet.column_dimensions["H"].width = 18
        for row_index in range(2, 258):
            worksheet.cell(row_index, 6).number_format = "0.000000%"
            worksheet.cell(row_index, 8).number_format = "0.000000%"

    charts_sheet = workbook.create_sheet("Charts")
    histogram_chart = BarChart()
    histogram_chart.type = "col"
    histogram_chart.style = 10
    histogram_chart.title = "Histogram Comparison"
    histogram_chart.y_axis.title = "Pixel Count"
    histogram_chart.x_axis.title = value_name
    histogram_chart.gapWidth = 0
    histogram_chart.height = 10
    histogram_chart.width = 24
    cdf_chart = LineChart()
    cdf_chart.style = 13
    cdf_chart.title = "CDF Comparison"
    cdf_chart.y_axis.title = "Cumulative Percentage"
    cdf_chart.x_axis.title = value_name
    cdf_chart.height = 10
    cdf_chart.width = 24
    for _analysis, worksheet in histogram_sheets:
        categories = Reference(
            worksheet, min_col=4, min_row=2, max_row=257
        )
        histogram_chart.add_data(
            Reference(worksheet, min_col=5, min_row=1, max_row=257),
            titles_from_data=True,
        )
        histogram_chart.set_categories(categories)
        cdf_chart.add_data(
            Reference(worksheet, min_col=8, min_row=1, max_row=257),
            titles_from_data=True,
        )
        cdf_chart.set_categories(categories)
    cdf_chart.y_axis.numFmt = "0%"
    charts_sheet.add_chart(histogram_chart, "A1")
    charts_sheet.add_chart(cdf_chart, "A21")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _robust_sigma(values):
    samples = np.asarray(values, dtype=np.float32).reshape(-1)
    if samples.size == 0:
        return 0.0
    median = float(np.median(samples))
    return float(1.4826 * np.median(np.abs(samples - median)))


def _weighted_quantile(values, weights, quantile):
    values = np.asarray(values, dtype=np.float64).reshape(-1)
    weights = np.asarray(weights, dtype=np.float64).reshape(-1)
    if values.size == 0 or weights.size != values.size:
        return 0.0
    valid = np.isfinite(values) & np.isfinite(weights) & (weights > 0)
    if not np.any(valid):
        return float(np.median(values[np.isfinite(values)])) if np.any(np.isfinite(values)) else 0.0
    values = values[valid]
    weights = weights[valid]
    order = np.argsort(values)
    values = values[order]
    weights = weights[order]
    cumulative = np.cumsum(weights)
    target = float(np.clip(quantile, 0.0, 1.0)) * cumulative[-1]
    index = int(np.searchsorted(cumulative, target, side="left"))
    return float(values[min(index, values.size - 1)])


def _quality_input_to_normalized(image, input_max):
    if input_max <= 0:
        raise ValueError("Quality analysis input_max must be positive.")
    source = np.asarray(image)
    if source.ndim != 2 or source.size == 0:
        raise ValueError("Quality analysis requires a non-empty gray image.")
    return np.clip(source.astype(np.float32), 0.0, float(input_max)) / float(input_max)


def _entropy_bits_from_normalized(image_f):
    gray_u8 = np.rint(np.clip(image_f, 0.0, 1.0) * 255.0).astype(np.uint8)
    histogram = np.bincount(gray_u8.reshape(-1), minlength=256).astype(np.float64)
    probabilities = histogram[histogram > 0] / gray_u8.size
    if probabilities.size == 0:
        return 0.0
    return float(-np.sum(probabilities * np.log2(probabilities)))


def _line_profile_seam_metrics(profile):
    """Measure repeated near-periodic discontinuity lines in one orientation."""
    profile = np.asarray(profile, dtype=np.float32).reshape(-1)
    if profile.size == 0:
        return {
            "ratio": 1.0,
            "strongest": 0.0,
            "baseline": 0.0,
            "peak_percent": 0.0,
            "peak_group_count": 0,
            "periodicity_score": 0.0,
            "penalty": 0.0,
        }

    baseline = float(np.median(profile))
    spread = _robust_sigma(profile)
    top_count = max(1, int(np.ceil(profile.size * 0.05)))
    strongest = np.partition(profile, profile.size - top_count)[-top_count:]
    strongest_mean = float(np.mean(strongest))
    denominator = max(baseline + spread, 1.0e-6)
    ratio = strongest_mean / denominator

    # Path A: very strong repeated line outliers.
    peak_threshold = baseline + 3.0 * max(spread, 1.0 / 255.0)
    peak_indices = np.flatnonzero(profile > peak_threshold)
    peak_percent = float(peak_indices.size * 100.0 / profile.size)
    peak_groups = []
    if peak_indices.size:
        start = previous = int(peak_indices[0])
        for index in peak_indices[1:]:
            index = int(index)
            if index <= previous + 1:
                previous = index
            else:
                peak_groups.append((start, previous))
                start = previous = index
        peak_groups.append((start, previous))

    peak_periodicity = 0.0
    if len(peak_groups) >= 4:
        centers = np.asarray(
            [(start + end) * 0.5 for start, end in peak_groups],
            dtype=np.float64,
        )
        gaps = np.diff(centers)
        mean_gap = float(np.mean(gaps)) if gaps.size else 0.0
        if mean_gap > 0.0:
            gap_cv = float(np.std(gaps) / mean_gap)
            regularity = float(np.exp(-((gap_cv / 0.30) ** 2)))
            repetition_support = float(
                np.clip((len(peak_groups) - 3) / 4.0, 0.0, 1.0)
            )
            peak_periodicity = 100.0 * regularity * repetition_support
    peak_strength_gate = float(np.clip((ratio - 1.10) / 1.80, 0.0, 1.0))
    peak_penalty = peak_periodicity * peak_strength_gate

    # Path B: a weaker but repeated comb pattern. This catches visible grid/block lines
    # that are individually weaker than natural object boundaries. The P25 sequence
    # strength requirement prevents a few accidental peaks from creating a false grid.
    smooth = cv2.GaussianBlur(
        profile.reshape(1, -1),
        (0, 0),
        1.0,
    ).reshape(-1)
    smooth_baseline = float(np.median(smooth))
    smooth_spread = _robust_sigma(smooth)
    residual = np.maximum(smooth - smooth_baseline, 0.0)
    positive = residual[residual > 0.0]
    comb_penalty = 0.0
    comb_periodicity = 0.0
    if positive.size >= 4:
        cap = float(np.percentile(positive, 85.0))
        if cap > 1.0e-10:
            normalized_residual = np.minimum(residual, cap) / cap
            sample_count = smooth.size
            comb_denominator = max(smooth_baseline + smooth_spread, 1.0e-6)
            # Search a compact set of plausible grid densities (4..20 cells) and
            # evaluate all phases at once. This is much faster than scanning every
            # possible period/phase pair in Python.
            tested_periods = set()
            for cell_count in range(4, 21):
                period = int(round(sample_count / float(cell_count)))
                period = max(4, period)
                if period in tested_periods:
                    continue
                tested_periods.add(period)
                repetitions = sample_count // period
                if repetitions < 4:
                    continue
                usable = repetitions * period
                normalized_matrix = normalized_residual[:usable].reshape(
                    repetitions, period
                )
                actual_matrix = smooth[:usable].reshape(repetitions, period)
                repetition_support = float(
                    np.clip((repetitions - 3) / 5.0, 0.0, 1.0)
                )
                pattern_support = (
                    0.60 * np.mean(normalized_matrix, axis=0)
                    + 0.40 * np.percentile(normalized_matrix, 25.0, axis=0)
                ) * repetition_support
                sequence_strength_ratio = (
                    np.percentile(actual_matrix, 25.0, axis=0)
                    / comb_denominator
                )
                strength_gate = np.clip(
                    (sequence_strength_ratio - 1.50) / 4.0,
                    0.0,
                    1.0,
                )
                candidate_penalty = 100.0 * pattern_support * strength_gate
                best_phase = int(np.argmax(candidate_penalty))
                best_penalty = float(candidate_penalty[best_phase])
                if best_penalty > comb_penalty:
                    comb_penalty = best_penalty
                    comb_periodicity = float(
                        100.0 * pattern_support[best_phase]
                    )

    periodicity_score = max(peak_periodicity, comb_periodicity)
    penalty = float(np.clip(max(peak_penalty, comb_penalty), 0.0, 100.0))
    return {
        "ratio": float(ratio),
        "strongest": strongest_mean,
        "baseline": baseline,
        "peak_percent": peak_percent,
        "peak_group_count": int(len(peak_groups)),
        "periodicity_score": float(periodicity_score),
        "penalty": penalty,
    }

def _whole_image_seam_ratio(image):
    """Detect repeated horizontal/vertical block seams while tolerating natural edges."""
    image_f = np.asarray(image, dtype=np.float32)
    candidates = []
    if image_f.shape[1] > 1:
        dx = np.abs(np.diff(image_f, axis=1))
        candidates.append(_line_profile_seam_metrics(np.mean(dx, axis=0)))
    if image_f.shape[0] > 1:
        dy = np.abs(np.diff(image_f, axis=0))
        candidates.append(_line_profile_seam_metrics(np.mean(dy, axis=1)))
    if not candidates:
        return 1.0, 0.0, 0.0, 0.0, 0.0, 0, 0.0

    # Report the orientation with the strongest artifact evidence. If none is periodic,
    # use the strongest ratio only for diagnostics, not for the penalty.
    selected = max(candidates, key=lambda item: (item["penalty"], item["ratio"]))
    max_periodicity = max(item["periodicity_score"] for item in candidates)
    max_group_count = max(item["peak_group_count"] for item in candidates)
    max_penalty = max(item["penalty"] for item in candidates)
    return (
        float(selected["ratio"]),
        float(selected["strongest"]),
        float(selected["baseline"]),
        float(selected["peak_percent"]),
        float(max_periodicity),
        int(max_group_count),
        float(max_penalty),
    )


def _edge_halo_overshoot_ratio(
    source_f,
    gradient_x,
    gradient_y,
    edge_mask,
    max_samples=4096,
):
    """Estimate true edge halo/ringing by sampling across the edge normal.

    A normal edge should transition between its two plateaus without leaving their range.
    Sharpening halos create overshoot/undershoot outside that local plateau range.
    """
    source_f = np.asarray(source_f, dtype=np.float32)
    edge_mask = np.asarray(edge_mask, dtype=bool)
    ys, xs = np.nonzero(edge_mask)
    if ys.size == 0:
        return 0.0, 0

    gx = gradient_x[ys, xs].astype(np.float32)
    gy = gradient_y[ys, xs].astype(np.float32)
    magnitude = np.sqrt(gx * gx + gy * gy)
    valid_direction = magnitude > 1e-6
    ys = ys[valid_direction]
    xs = xs[valid_direction]
    gx = gx[valid_direction]
    gy = gy[valid_direction]
    magnitude = magnitude[valid_direction]
    if ys.size == 0:
        return 0.0, 0

    if ys.size > int(max_samples):
        keep = np.argpartition(magnitude, -int(max_samples))[-int(max_samples):]
        ys = ys[keep]
        xs = xs[keep]
        gx = gx[keep]
        gy = gy[keep]
        magnitude = magnitude[keep]

    nx = gx / np.maximum(magnitude, 1e-6)
    ny = gy / np.maximum(magnitude, 1e-6)
    profiles = []
    for distance in (-4, -3, -2, -1, 0, 1, 2, 3, 4):
        map_x = (xs.astype(np.float32) + nx * distance).reshape(-1, 1)
        map_y = (ys.astype(np.float32) + ny * distance).reshape(-1, 1)
        samples = cv2.remap(
            source_f,
            map_x,
            map_y,
            cv2.INTER_LINEAR,
            borderMode=cv2.BORDER_REFLECT_101,
        ).reshape(-1)
        profiles.append(samples)
    profile = np.stack(profiles, axis=1)

    left_plateau = np.mean(profile[:, 0:2], axis=1)
    right_plateau = np.mean(profile[:, -2:], axis=1)
    edge_contrast = np.abs(right_plateau - left_plateau)
    valid = edge_contrast >= 0.03
    if np.count_nonzero(valid) < 16:
        return 0.0, int(np.count_nonzero(valid))

    center = profile[valid, 2:7]
    left_plateau = left_plateau[valid]
    right_plateau = right_plateau[valid]
    edge_contrast = edge_contrast[valid]
    lower = np.minimum(left_plateau, right_plateau)
    upper = np.maximum(left_plateau, right_plateau)
    overshoot_high = np.maximum(np.max(center, axis=1) - upper, 0.0)
    overshoot_low = np.maximum(lower - np.min(center, axis=1), 0.0)
    severity = np.maximum(overshoot_high, overshoot_low) / np.maximum(
        edge_contrast, 0.03
    )

    # A P75/P90 blend is robust to a few textured edges while still reacting when
    # overshoot is repeated across many real edges.
    severity_p75 = float(np.percentile(severity, 75.0))
    severity_p90 = float(np.percentile(severity, 90.0))
    ratio = 0.70 * severity_p75 + 0.30 * severity_p90
    return float(ratio), int(severity.size)

def _isolated_impulse_analysis(image_f, max_component_area=4):
    """Detect isolated hot/cold impulse pixels without treating edge halos as impulses."""
    image_f = np.asarray(image_f, dtype=np.float32)
    median = cv2.medianBlur(image_f, 3)
    residual = image_f - median
    sigma = _robust_sigma(residual)
    threshold = max(6.0 * sigma, 3.0 / 255.0)
    candidate = (np.abs(residual) > threshold).astype(np.uint8)
    component_count, _labels, stats, _centroids = cv2.connectedComponentsWithStats(
        candidate,
        connectivity=8,
    )
    isolated_pixels = 0
    isolated_components = 0
    for component_id in range(1, component_count):
        area = int(stats[component_id, cv2.CC_STAT_AREA])
        if area <= int(max_component_area):
            isolated_pixels += area
            isolated_components += 1
    isolated_percent = float(
        isolated_pixels * 100.0 / image_f.size
    )
    score = float(100.0 * np.exp(-isolated_percent / 0.25))
    return {
        "score": score,
        "percent": isolated_percent,
        "pixel_count": int(isolated_pixels),
        "component_count": int(isolated_components),
        "threshold": float(threshold),
        "sigma": float(sigma),
    }


def _row_column_banding_sigma(residual, flat_mask):
    residual = np.asarray(residual, dtype=np.float32)
    flat_mask = np.asarray(flat_mask, dtype=bool)
    values = []
    minimum_row_pixels = max(8, residual.shape[1] // 12)
    minimum_col_pixels = max(8, residual.shape[0] // 12)
    for row in range(residual.shape[0]):
        mask = flat_mask[row]
        if np.count_nonzero(mask) >= minimum_row_pixels:
            values.append(float(np.median(residual[row][mask])))
    for column in range(residual.shape[1]):
        mask = flat_mask[:, column]
        if np.count_nonzero(mask) >= minimum_col_pixels:
            values.append(float(np.median(residual[:, column][mask])))
    return _robust_sigma(values)


def _full_flat_region_analysis(base, detail, uniform_residual, flat_mask, gray_bands=8):
    """Split flat pixels by gray level and connected region before measuring noise.

    This prevents large flat areas at different brightness levels (e.g. sky and wall) from
    being mixed into one noise population. It also keeps spatially separate regions apart.
    """
    flat_values = base[flat_mask]
    if flat_values.size == 0:
        return [], 0
    low, high = np.percentile(flat_values, [1.0, 99.0])
    if high <= low:
        low = float(np.min(flat_values))
        high = float(np.max(flat_values))
    if high <= low:
        low, high = 0.0, 1.0

    minimum_area = max(96, int(np.ceil(base.size * 0.0015)))
    regions = []
    for band_index in range(gray_bands):
        band_low = low + (high - low) * band_index / gray_bands
        band_high = low + (high - low) * (band_index + 1) / gray_bands
        if band_index == gray_bands - 1:
            band_mask = flat_mask & (base >= band_low) & (base <= band_high)
        else:
            band_mask = flat_mask & (base >= band_low) & (base < band_high)
        if np.count_nonzero(band_mask) < minimum_area:
            continue
        count, labels, stats, _centroids = cv2.connectedComponentsWithStats(
            band_mask.astype(np.uint8), connectivity=8
        )
        for label_id in range(1, count):
            area = int(stats[label_id, cv2.CC_STAT_AREA])
            if area < minimum_area:
                continue
            x = int(stats[label_id, cv2.CC_STAT_LEFT])
            y = int(stats[label_id, cv2.CC_STAT_TOP])
            width = int(stats[label_id, cv2.CC_STAT_WIDTH])
            height = int(stats[label_id, cv2.CC_STAT_HEIGHT])
            local_labels = labels[y:y + height, x:x + width]
            local_mask = local_labels == label_id
            local_detail = detail[y:y + height, x:x + width][local_mask]
            local_uniform = uniform_residual[y:y + height, x:x + width][local_mask]
            local_base = base[y:y + height, x:x + width][local_mask]
            regions.append(
                {
                    "band": int(band_index),
                    "area": area,
                    "mean_level": float(np.mean(local_base)),
                    "noise_sigma": _robust_sigma(local_detail),
                    "uniform_sigma": _robust_sigma(local_uniform),
                }
            )

    if not regions and np.count_nonzero(flat_mask) >= 16:
        regions.append(
            {
                "band": 0,
                "area": int(np.count_nonzero(flat_mask)),
                "mean_level": float(np.mean(base[flat_mask])),
                "noise_sigma": _robust_sigma(detail[flat_mask]),
                "uniform_sigma": _robust_sigma(uniform_residual[flat_mask]),
            }
        )

    regions.sort(key=lambda item: item["area"], reverse=True)
    regions = regions[:48]
    return regions, len({item["band"] for item in regions})


def _spatial_overall_score(detail_score, contrast_score, noise_score, uniformity_score, clipping_score, artifact_score):
    scores = {
        "detail_clarity_score": detail_score,
        "contrast_quality_score": contrast_score,
        "flat_cleanliness_score": noise_score,
        "thermal_uniformity_score": uniformity_score,
        "clipping_control_score": clipping_score,
        "artifact_control_score": artifact_score,
    }
    return float(
        np.clip(
            sum(
                ACTIVE_SPATIAL_QUALITY_WEIGHTS[key] * scores[key]
                for key in scores
            )
            / 100.0,
            0.0,
            100.0,
        )
    )


def thermal_raw_quality_analysis(image, input_max):
    """Full no-reference spatial analysis for RAW/still-image comparison.

    The full profile keeps native resolution, separates base/detail layers, segments flat
    areas by both brightness and connectivity, mixes typical and worst-region statistics,
    and adds row/column banding diagnostics. Scores are relative same-scene indicators.
    """
    source_f = _quality_input_to_normalized(image, input_max)
    height, width = source_f.shape

    # Edge-preserving base/detail decomposition. Working in normalized 0..1 makes the
    # thresholds independent of RAW bit depth.
    base = cv2.bilateralFilter(source_f, 7, 0.08, 5.0)
    detail = source_f - base
    detail_abs = np.abs(detail)

    gradient_x = cv2.Sobel(base, cv2.CV_32F, 1, 0, ksize=3)
    gradient_y = cv2.Sobel(base, cv2.CV_32F, 0, 1, ksize=3)
    gradient = cv2.magnitude(gradient_x, gradient_y) / 4.0
    local_mean = cv2.blur(base, (9, 9))
    local_mean_square = cv2.blur(base * base, (9, 9))
    local_std = np.sqrt(np.maximum(local_mean_square - local_mean * local_mean, 0.0))

    edge_threshold = max(float(np.percentile(gradient, 75.0)), 0.015)
    edge_mask = gradient >= edge_threshold
    if np.count_nonzero(edge_mask) < max(16, source_f.size // 500):
        edge_mask = gradient > 0.005

    flat_gradient_threshold = float(
        np.clip(np.percentile(gradient, 35.0), 0.004, 0.020)
    )
    flat_std_threshold = float(
        np.clip(np.percentile(local_std, 50.0), 0.008, 0.035)
    )
    flat_mask = (
        (gradient <= flat_gradient_threshold)
        & (local_std <= flat_std_threshold)
    )
    # Keep high-gradient edge neighborhoods out of flat-noise statistics.
    edge_guard = cv2.dilate(edge_mask.astype(np.uint8), np.ones((3, 3), np.uint8)) > 0
    guarded_flat_mask = flat_mask & ~edge_guard
    if np.count_nonzero(guarded_flat_mask) >= max(32, source_f.size // 200):
        flat_mask = guarded_flat_mask
    elif np.count_nonzero(flat_mask) < max(16, source_f.size // 100):
        flat_mask = gradient <= flat_gradient_threshold

    broad_base = cv2.GaussianBlur(base, (0, 0), 5.0)
    uniform_residual = base - broad_base
    flat_regions, flat_gray_band_count = _full_flat_region_analysis(
        base, detail, uniform_residual, flat_mask, gray_bands=8
    )

    if flat_regions:
        region_areas = np.asarray([region["area"] for region in flat_regions], dtype=np.float64)
        region_noise = np.asarray([region["noise_sigma"] for region in flat_regions], dtype=np.float64)
        region_uniformity = np.asarray([region["uniform_sigma"] for region in flat_regions], dtype=np.float64)
        flat_noise_typical = _weighted_quantile(region_noise, region_areas, 0.50)
        # Worst-region terms are intentionally not area-weighted. Otherwise a large clean
        # sky can hide a smaller noisy wall/object. Very small fragments were already
        # rejected by the connected-component minimum-area rule above.
        flat_noise_worst = float(np.percentile(region_noise, 80.0))
        uniform_typical = _weighted_quantile(region_uniformity, region_areas, 0.50)
        uniform_worst = float(np.percentile(region_uniformity, 80.0))
        analyzed_flat_pixels = int(np.sum(region_areas))
    else:
        flat_noise_typical = _robust_sigma(detail[flat_mask]) if np.any(flat_mask) else 0.0
        flat_noise_worst = flat_noise_typical
        uniform_typical = _robust_sigma(uniform_residual[flat_mask]) if np.any(flat_mask) else 0.0
        uniform_worst = uniform_typical
        analyzed_flat_pixels = int(np.count_nonzero(flat_mask))

    noise_typical_score = 100.0 * np.exp(-flat_noise_typical / 0.025)
    noise_worst_score = 100.0 * np.exp(-flat_noise_worst / 0.025)
    robust_noise_score = float(
        np.clip(
            0.65 * noise_typical_score + 0.35 * noise_worst_score,
            0.0,
            100.0,
        )
    )
    impulse_result = _isolated_impulse_analysis(
        source_f,
        max_component_area=4,
    )
    impulse_outlier_score = float(impulse_result["score"])
    impulse_outlier_percent = float(impulse_result["percent"])
    impulse_threshold = float(impulse_result["threshold"])
    noise_score = float(
        np.clip(
            0.85 * robust_noise_score + 0.15 * impulse_outlier_score,
            0.0,
            100.0,
        )
    )

    uniform_typical_score = 100.0 * np.exp(-uniform_typical / 0.020)
    uniform_worst_score = 100.0 * np.exp(-uniform_worst / 0.020)
    banding_sigma = _row_column_banding_sigma(uniform_residual, flat_mask)
    banding_score = 100.0 * np.exp(-banding_sigma / 0.012)
    uniformity_score = float(
        np.clip(
            0.50 * uniform_typical_score
            + 0.30 * uniform_worst_score
            + 0.20 * banding_score,
            0.0,
            100.0,
        )
    )

    edge_strength = float(np.mean(gradient[edge_mask])) if np.any(edge_mask) else 0.0
    edge_detail = float(np.mean(detail_abs[edge_mask])) if np.any(edge_mask) else 0.0
    reference_noise = max(flat_noise_typical, 1.0 / 255.0)
    detail_snr = edge_detail / reference_noise
    sharpness_score = float(np.clip(edge_strength / 0.16 * 100.0, 0.0, 100.0))
    detail_snr_score = float(np.clip(detail_snr / 5.0 * 100.0, 0.0, 100.0))

    halo_ratio, halo_edge_sample_count = _edge_halo_overshoot_ratio(
        source_f,
        gradient_x,
        gradient_y,
        edge_mask,
        max_samples=4096,
    )
    halo_penalty = float(
        np.clip((halo_ratio - 0.10) / 0.50 * 100.0, 0.0, 100.0)
    )
    detail_clarity = float(
        np.clip(
            0.50 * sharpness_score
            + 0.50 * detail_snr_score
            - 0.08 * halo_penalty,
            0.0,
            100.0,
        )
    )

    p1, p99 = np.percentile(source_f, [1.0, 99.0])
    contrast_range = float(p99 - p1)
    entropy_bits = _entropy_bits_from_normalized(source_f)
    contrast_span_score = float(np.clip(contrast_range / 0.60 * 100.0, 0.0, 100.0))
    entropy_score = float(np.clip(entropy_bits / 7.0 * 100.0, 0.0, 100.0))
    contrast_score = float(np.clip(0.70 * contrast_span_score + 0.30 * entropy_score, 0.0, 100.0))

    clipping_percent = float(
        np.count_nonzero((source_f <= 1.0 / 255.0) | (source_f >= 254.0 / 255.0))
        * 100.0
        / source_f.size
    )
    clipping_score = float(100.0 * np.exp(-clipping_percent / 3.0))

    (
        seam_ratio,
        seam_strongest_difference_norm,
        seam_baseline_difference_norm,
        seam_peak_line_percent,
        seam_periodicity_score,
        seam_peak_group_count,
        seam_penalty,
    ) = _whole_image_seam_ratio(base)
    artifact_control = float(
        np.clip(
            0.50 * (100.0 - halo_penalty)
            + 0.50 * (100.0 - seam_penalty),
            0.0,
            100.0,
        )
    )

    overall_score = _spatial_overall_score(
        detail_clarity,
        contrast_score,
        noise_score,
        uniformity_score,
        clipping_score,
        artifact_control,
    )

    edge_region_percent = float(np.count_nonzero(edge_mask) * 100.0 / source_f.size)
    flat_region_percent = float(np.count_nonzero(flat_mask) * 100.0 / source_f.size)
    analyzed_flat_percent = float(analyzed_flat_pixels * 100.0 / source_f.size)
    flat_confidence = min(flat_region_percent / 15.0, 1.0)
    edge_confidence = min(edge_region_percent / 5.0, 1.0)
    region_confidence = min(len(flat_regions) / 3.0, 1.0)
    confidence_score = float(
        np.clip(
            (
                0.45 * np.sqrt(flat_confidence * edge_confidence)
                + 0.35 * flat_confidence
                + 0.20 * region_confidence
            )
            * 100.0,
            0.0,
            100.0,
        )
    )

    base_u8 = np.rint(np.clip(base, 0.0, 1.0) * 255.0).astype(np.uint8)
    detail_display = np.clip(detail / 0.125 * 127.0 + 128.0, 0, 255).astype(np.uint8)
    metrics = {
        "quality_profile": "RAW Full",
        "overall_score": overall_score,
        "detail_clarity_score": detail_clarity,
        "contrast_quality_score": contrast_score,
        "flat_cleanliness_score": noise_score,
        "robust_noise_score": robust_noise_score,
        "impulse_outlier_control_score": impulse_outlier_score,
        "thermal_uniformity_score": uniformity_score,
        "clipping_control_score": clipping_score,
        "artifact_control_score": artifact_control,
        "quality_confidence_score": confidence_score,
        "sharpness_score": sharpness_score,
        "detail_snr_score": detail_snr_score,
        "contrast_span_score": contrast_span_score,
        "entropy_score": entropy_score,
        "banding_control_score": float(banding_score),
        "halo_penalty_score": halo_penalty,
        "seam_artifact_penalty_score": seam_penalty,
        "detail_snr": float(detail_snr),
        "flat_noise_dn": float(flat_noise_typical * input_max),
        "flat_noise_typical_dn": float(flat_noise_typical * input_max),
        "flat_noise_worst_dn": float(flat_noise_worst * input_max),
        "flat_noise_typical_percent": float(flat_noise_typical * 100.0),
        "flat_noise_worst_percent": float(flat_noise_worst * 100.0),
        "impulse_outlier_percent": float(impulse_outlier_percent),
        "impulse_pixel_count": int(impulse_result["pixel_count"]),
        "impulse_component_count": int(impulse_result["component_count"]),
        "impulse_threshold_dn": float(impulse_threshold * input_max),
        "impulse_residual_sigma_dn": float(impulse_result["sigma"] * input_max),
        "edge_strength_dn": float(edge_strength * input_max),
        "base_residual_noise_dn": float(uniform_typical * input_max),
        "uniformity_typical_dn": float(uniform_typical * input_max),
        "uniformity_worst_dn": float(uniform_worst * input_max),
        "uniformity_typical_percent": float(uniform_typical * 100.0),
        "uniformity_worst_percent": float(uniform_worst * 100.0),
        "row_column_banding_dn": float(banding_sigma * input_max),
        "row_column_banding_percent": float(banding_sigma * 100.0),
        "halo_ratio": float(halo_ratio),
        "halo_edge_sample_count": int(halo_edge_sample_count),
        "seam_artifact_ratio": float(seam_ratio),
        "seam_periodicity_score": float(seam_periodicity_score),
        "seam_peak_group_count": int(seam_peak_group_count),
        "seam_strongest_difference_dn": float(seam_strongest_difference_norm * input_max),
        "seam_baseline_difference_dn": float(seam_baseline_difference_norm * input_max),
        "seam_peak_line_percent": float(seam_peak_line_percent),
        "p1": float(p1 * input_max),
        "p99": float(p99 * input_max),
        "p99_p1": float((p99 - p1) * input_max),
        "p99_p1_percent": float((p99 - p1) * 100.0),
        "entropy_bits": float(entropy_bits),
        "clipping_percent": clipping_percent,
        "edge_region_percent": edge_region_percent,
        "flat_region_percent": flat_region_percent,
        "flat_analyzed_region_percent": analyzed_flat_percent,
        "flat_region_group_count": int(len(flat_regions)),
        "flat_gray_band_count": int(flat_gray_band_count),
        "analysis_width": int(width),
        "analysis_height": int(height),
        "input_max": float(input_max),
    }
    return {
        "metrics": metrics,
        "base_u8": base_u8,
        "detail_u8": detail_display,
        "flat_regions": [
            {
                "band": region["band"],
                "area": region["area"],
                "mean_level_percent": round(region["mean_level"] * 100.0, 3),
                "noise_dn": round(region["noise_sigma"] * input_max, 4),
                "uniformity_dn": round(region["uniform_sigma"] * input_max, 4),
            }
            for region in flat_regions
        ],
    }


def thermal_uvc_quality_analysis(image, max_analysis_pixels=None):
    """Lightweight spatial quality analysis for live UVC streams.

    It intentionally avoids RAW Full's connected-component/gray-band segmentation and
    row/column banding pass. A downsampled single-scale Gaussian/Sobel analysis is used so
    live video output is not burdened by the quality monitor.
    """
    source_gray = np.clip(np.asarray(image), 0, 255).astype(np.uint8)
    if source_gray.ndim != 2 or source_gray.size == 0:
        raise ValueError("UVC quality analysis requires a non-empty gray image.")
    gray = source_gray
    if max_analysis_pixels and source_gray.size > int(max_analysis_pixels):
        scale = np.sqrt(float(max_analysis_pixels) / source_gray.size)
        analysis_width = max(32, int(round(source_gray.shape[1] * scale)))
        analysis_height = max(32, int(round(source_gray.shape[0] * scale)))
        gray = cv2.resize(
            source_gray,
            (analysis_width, analysis_height),
            interpolation=cv2.INTER_AREA,
        )

    image_f = gray.astype(np.float32) / 255.0
    base = cv2.GaussianBlur(image_f, (5, 5), 0.9)
    detail = image_f - base
    detail_abs = np.abs(detail)
    gradient_x = cv2.Sobel(base, cv2.CV_32F, 1, 0, ksize=3)
    gradient_y = cv2.Sobel(base, cv2.CV_32F, 0, 1, ksize=3)
    gradient = cv2.magnitude(gradient_x, gradient_y) / 4.0
    local_mean = cv2.blur(base, (7, 7))
    local_mean_square = cv2.blur(base * base, (7, 7))
    local_std = np.sqrt(np.maximum(local_mean_square - local_mean * local_mean, 0.0))

    edge_threshold = max(float(np.percentile(gradient, 75.0)), 0.015)
    edge_mask = gradient >= edge_threshold
    flat_gradient_threshold = float(
        np.clip(np.percentile(gradient, 40.0), 0.004, 0.025)
    )
    flat_std_threshold = float(
        np.clip(np.percentile(local_std, 55.0), 0.008, 0.040)
    )
    flat_mask = (gradient <= flat_gradient_threshold) & (local_std <= flat_std_threshold)
    if np.count_nonzero(flat_mask) < max(16, gray.size // 100):
        flat_mask = gradient <= flat_gradient_threshold

    edge_strength = float(np.mean(gradient[edge_mask])) if np.any(edge_mask) else 0.0
    edge_detail = float(np.mean(detail_abs[edge_mask])) if np.any(edge_mask) else 0.0
    flat_noise = _robust_sigma(detail[flat_mask]) if np.any(flat_mask) else 0.0
    detail_snr = edge_detail / max(flat_noise, 1.0 / 255.0)
    sharpness_score = float(np.clip(edge_strength / 0.16 * 100.0, 0.0, 100.0))
    detail_snr_score = float(np.clip(detail_snr / 5.0 * 100.0, 0.0, 100.0))

    # Lightweight halo proxy: high second derivative relative to the blurred edge
    # gradient rises with oversharpening/ringing, but costs only one Laplacian pass.
    laplacian = np.abs(cv2.Laplacian(image_f, cv2.CV_32F, ksize=3)) / 4.0
    halo_ratio = (
        float(np.mean(laplacian[edge_mask]))
        / max(edge_strength, 1.0 / 255.0)
        if np.any(edge_mask)
        else 0.0
    )
    halo_penalty = float(
        np.clip((halo_ratio - 0.75) / 0.80 * 100.0, 0.0, 100.0)
    )
    detail_clarity = float(
        np.clip(0.50 * sharpness_score + 0.50 * detail_snr_score - 0.08 * halo_penalty, 0.0, 100.0)
    )

    broad_base = cv2.GaussianBlur(base, (0, 0), 3.0)
    base_residual_noise = _robust_sigma((base - broad_base)[flat_mask]) if np.any(flat_mask) else 0.0
    robust_noise_score = float(100.0 * np.exp(-flat_noise / 0.025))
    impulse_result = _isolated_impulse_analysis(
        image_f,
        max_component_area=3,
    )
    impulse_outlier_score = float(impulse_result["score"])
    impulse_outlier_percent = float(impulse_result["percent"])
    impulse_threshold = float(impulse_result["threshold"])
    noise_score = float(
        np.clip(
            0.85 * robust_noise_score + 0.15 * impulse_outlier_score,
            0.0,
            100.0,
        )
    )
    uniformity_score = float(100.0 * np.exp(-base_residual_noise / 0.020))

    p1, p99 = np.percentile(image_f, [1.0, 99.0])
    contrast_range = float(p99 - p1)
    entropy_bits = _entropy_bits_from_normalized(image_f)
    contrast_span_score = float(np.clip(contrast_range / 0.60 * 100.0, 0.0, 100.0))
    entropy_score = float(np.clip(entropy_bits / 7.0 * 100.0, 0.0, 100.0))
    contrast_score = float(np.clip(0.75 * contrast_span_score + 0.25 * entropy_score, 0.0, 100.0))
    clipping_percent = float(
        np.count_nonzero((gray <= 1) | (gray >= 254)) * 100.0 / gray.size
    )
    clipping_score = float(100.0 * np.exp(-clipping_percent / 3.0))

    (
        seam_ratio,
        seam_strongest_difference_norm,
        seam_baseline_difference_norm,
        seam_peak_line_percent,
        seam_periodicity_score,
        seam_peak_group_count,
        seam_penalty,
    ) = _whole_image_seam_ratio(base)
    artifact_control = float(
        np.clip(
            0.50 * (100.0 - halo_penalty)
            + 0.50 * (100.0 - seam_penalty),
            0.0,
            100.0,
        )
    )
    overall_score = _spatial_overall_score(
        detail_clarity,
        contrast_score,
        noise_score,
        uniformity_score,
        clipping_score,
        artifact_control,
    )

    base_u8 = np.rint(np.clip(base, 0.0, 1.0) * 255.0).astype(np.uint8)
    detail_display = np.clip(detail / 0.125 * 127.0 + 128.0, 0, 255).astype(np.uint8)
    output_base = base_u8
    output_detail = detail_display
    if gray.shape != source_gray.shape:
        output_base = cv2.resize(base_u8, (source_gray.shape[1], source_gray.shape[0]), interpolation=cv2.INTER_LINEAR)
        output_detail = cv2.resize(detail_display, (source_gray.shape[1], source_gray.shape[0]), interpolation=cv2.INTER_LINEAR)

    metrics = {
        "quality_profile": "UVC Lightweight",
        "overall_score": overall_score,
        "detail_clarity_score": detail_clarity,
        "contrast_quality_score": contrast_score,
        "flat_cleanliness_score": noise_score,
        "robust_noise_score": robust_noise_score,
        "impulse_outlier_control_score": impulse_outlier_score,
        "thermal_uniformity_score": uniformity_score,
        "clipping_control_score": clipping_score,
        "artifact_control_score": artifact_control,
        "sharpness_score": sharpness_score,
        "detail_snr_score": detail_snr_score,
        "contrast_span_score": contrast_span_score,
        "entropy_score": entropy_score,
        "halo_penalty_score": halo_penalty,
        "seam_artifact_penalty_score": seam_penalty,
        "detail_snr": float(detail_snr),
        "flat_noise_dn": float(flat_noise * 255.0),
        "impulse_outlier_percent": float(impulse_outlier_percent),
        "impulse_pixel_count": int(impulse_result["pixel_count"]),
        "impulse_component_count": int(impulse_result["component_count"]),
        "impulse_threshold_dn": float(impulse_threshold * 255.0),
        "impulse_residual_sigma_dn": float(impulse_result["sigma"] * 255.0),
        "edge_strength_dn": float(edge_strength * 255.0),
        "base_residual_noise_dn": float(base_residual_noise * 255.0),
        "halo_ratio": float(halo_ratio),
        "seam_artifact_ratio": float(seam_ratio),
        "seam_periodicity_score": float(seam_periodicity_score),
        "seam_peak_group_count": int(seam_peak_group_count),
        "seam_strongest_difference_dn": float(seam_strongest_difference_norm * 255.0),
        "seam_baseline_difference_dn": float(seam_baseline_difference_norm * 255.0),
        "seam_peak_line_percent": float(seam_peak_line_percent),
        "p1": float(p1 * 255.0),
        "p99": float(p99 * 255.0),
        "p99_p1": float((p99 - p1) * 255.0),
        "entropy_bits": float(entropy_bits),
        "clipping_percent": clipping_percent,
        "edge_region_percent": float(np.count_nonzero(edge_mask) * 100.0 / gray.size),
        "flat_region_percent": float(np.count_nonzero(flat_mask) * 100.0 / gray.size),
        "analysis_width": int(gray.shape[1]),
        "analysis_height": int(gray.shape[0]),
    }
    return {
        "metrics": metrics,
        "base_u8": output_base,
        "detail_u8": output_detail,
    }


def thermal_image_quality_analysis(image, max_analysis_pixels=None):
    """Backward-compatible alias for the lightweight 8-bit UVC profile."""
    return thermal_uvc_quality_analysis(image, max_analysis_pixels=max_analysis_pixels)


def quality_analysis_manifest(result):
    manifest = {"metrics": dict(result["metrics"])}
    if "flat_regions" in result:
        manifest["flat_regions"] = list(result["flat_regions"])
    return manifest


class UvcTemporalQualityTracker:
    """Low-cost rolling temporal analysis for one UVC image stream."""

    def __init__(self):
        self.reset()

    def reset(self):
        self.brightness_samples = deque()
        self.temporal_samples = deque()
        self.previous_frame = None
        self.previous_time = None
        self.frame_shape = None

    @staticmethod
    def _analysis_frame(image):
        gray = np.clip(np.asarray(image), 0, 255).astype(np.uint8)
        if gray.ndim != 2 or gray.size == 0:
            raise ValueError("Temporal analysis requires a non-empty gray image.")
        if gray.size <= UVC_TEMPORAL_ANALYSIS_MAX_PIXELS:
            return gray
        scale = np.sqrt(UVC_TEMPORAL_ANALYSIS_MAX_PIXELS / gray.size)
        width = max(32, int(round(gray.shape[1] * scale)))
        height = max(32, int(round(gray.shape[0] * scale)))
        return cv2.resize(gray, (width, height), interpolation=cv2.INTER_AREA)

    def add_frame(self, image, frame_time=None):
        now = time.perf_counter() if frame_time is None else float(frame_time)
        if self.previous_time is not None:
            if now <= self.previous_time:
                self.reset()
            elif now - self.previous_time < UVC_TEMPORAL_SAMPLE_INTERVAL_SECONDS:
                return False
        frame = self._analysis_frame(image)
        if self.frame_shape != frame.shape:
            self.reset()
        self.frame_shape = frame.shape

        p5, p95 = np.percentile(frame, [5.0, 95.0])
        trimmed = frame[(frame >= p5) & (frame <= p95)]
        brightness = float(np.mean(trimmed)) if trimmed.size else float(np.mean(frame))
        self.brightness_samples.append((now, brightness))

        if self.previous_frame is not None:
            current_f = frame.astype(np.float32)
            previous_f = self.previous_frame.astype(np.float32)
            difference = current_f - previous_f
            global_offset = float(np.median(difference))
            residual = difference - global_offset
            midpoint = (current_f + previous_f) * 0.5
            gradient_x = cv2.Sobel(midpoint, cv2.CV_32F, 1, 0, ksize=3)
            gradient_y = cv2.Sobel(midpoint, cv2.CV_32F, 0, 1, ksize=3)
            gradient = cv2.magnitude(gradient_x, gradient_y) / 4.0
            flat_threshold = max(float(np.percentile(gradient, 60.0)), 2.0)
            flat_mask = gradient <= flat_threshold
            residual_abs = np.abs(residual)
            if np.count_nonzero(flat_mask) >= 64:
                motion_threshold = max(
                    float(np.percentile(residual_abs[flat_mask], 85.0)), 2.0
                )
                stable_mask = flat_mask & (residual_abs <= motion_threshold)
            else:
                stable_mask = residual_abs <= max(
                    float(np.percentile(residual_abs, 70.0)), 2.0
                )
            stable_count = int(np.count_nonzero(stable_mask))
            if stable_count >= 64:
                temporal_noise_dn = _robust_sigma(residual[stable_mask]) / np.sqrt(2.0)
                stable_percent = stable_count * 100.0 / frame.size
                self.temporal_samples.append(
                    (now, float(temporal_noise_dn), float(stable_percent))
                )

        self.previous_frame = frame.copy()
        self.previous_time = now
        cutoff = now - UVC_TEMPORAL_WINDOW_SECONDS
        while self.brightness_samples and self.brightness_samples[0][0] < cutoff:
            self.brightness_samples.popleft()
        while self.temporal_samples and self.temporal_samples[0][0] < cutoff:
            self.temporal_samples.popleft()
        return True

    def result(self):
        brightness = list(self.brightness_samples)
        temporal = list(self.temporal_samples)
        duration = (
            float(brightness[-1][0] - brightness[0][0])
            if len(brightness) >= 2
            else 0.0
        )
        ready = (
            duration >= UVC_TEMPORAL_MIN_DURATION_SECONDS
            and len(brightness) >= UVC_TEMPORAL_MIN_SAMPLES
            and len(temporal) >= UVC_TEMPORAL_MIN_SAMPLES - 1
        )
        result = {
            "temporal_ready": bool(ready),
            "temporal_frame_count": len(brightness),
            "temporal_window_seconds": duration,
            "flicker_dn": None,
            "temporal_noise_dn": None,
            "stable_region_percent": None,
            "flicker_control_score": None,
            "temporal_stability_score": None,
        }
        if not ready:
            return result

        times = np.asarray([item[0] for item in brightness], dtype=np.float64)
        levels = np.asarray([item[1] for item in brightness], dtype=np.float64)
        relative_times = times - times[0]
        if np.ptp(relative_times) > 0.0:
            trend = np.polyval(np.polyfit(relative_times, levels, 1), relative_times)
        else:
            trend = np.full_like(levels, np.median(levels))
        flicker_dn = float(1.4826 * np.median(np.abs(levels - trend)))
        temporal_noise_dn = float(np.median([item[1] for item in temporal]))
        stable_percent = float(np.median([item[2] for item in temporal]))
        result.update(
            {
                "flicker_dn": float(flicker_dn),
                "temporal_noise_dn": temporal_noise_dn,
                "stable_region_percent": stable_percent,
                "flicker_control_score": float(100.0 * np.exp(-flicker_dn / 2.0)),
                "temporal_stability_score": float(
                    100.0 * np.exp(-temporal_noise_dn / 3.0)
                ),
            }
        )
        return result


def add_uvc_temporal_quality(spatial_result, temporal_result):
    result = dict(spatial_result)
    result["metrics"] = dict(spatial_result["metrics"])
    metrics = result["metrics"]
    spatial_score = float(metrics["overall_score"])
    metrics["spatial_quality_score"] = spatial_score
    metrics.update(temporal_result)
    if temporal_result.get("temporal_ready"):
        metrics["overall_score"] = float(
            (
                ACTIVE_UVC_TEMPORAL_WEIGHTS["spatial_quality_score"]
                * spatial_score
                + ACTIVE_UVC_TEMPORAL_WEIGHTS["flicker_control_score"]
                * temporal_result["flicker_control_score"]
                + ACTIVE_UVC_TEMPORAL_WEIGHTS["temporal_stability_score"]
                * temporal_result["temporal_stability_score"]
            )
            / 100.0
        )
    else:
        metrics["overall_score"] = None
    return result


def write_quality_summary(summary, label, result):
    metrics = result["metrics"] if "metrics" in result else result

    def score_text(key):
        value = metrics.get(key)
        return "Collecting..." if value is None else f"{float(value):.3f}"

    profile = metrics.get("quality_profile", "Quality")
    summary.write(f"{label} quality [{profile}] (0-100, higher is better):\n")
    summary.write(
        f"  Total Score / detail clarity / contrast: "
        f"{score_text('overall_score')} / "
        f"{score_text('detail_clarity_score')} / "
        f"{score_text('contrast_quality_score')}\n"
    )
    summary.write(
        f"  Flat cleanliness / thermal uniformity / clipping control / "
        f"artifact control: {score_text('flat_cleanliness_score')} / "
        f"{score_text('thermal_uniformity_score')} / "
        f"{score_text('clipping_control_score')} / "
        f"{score_text('artifact_control_score')}\n"
    )
    if metrics.get("quality_confidence_score") is not None:
        summary.write(
            f"  Quality confidence / flat groups / gray bands: "
            f"{score_text('quality_confidence_score')} / "
            f"{int(metrics.get('flat_region_group_count', 0))} / "
            f"{int(metrics.get('flat_gray_band_count', 0))}\n"
        )
    if "flicker_control_score" in metrics:
        summary.write(
            f"  Flicker control / temporal stability: "
            f"{score_text('flicker_control_score')} / "
            f"{score_text('temporal_stability_score')}\n"
        )
        stable_text = score_text("stable_region_percent")
        summary.write(
            f"  Temporal window / frames / stable region: "
            f"{float(metrics.get('temporal_window_seconds', 0.0)):.3f} s / "
            f"{int(metrics.get('temporal_frame_count', 0))} / "
            f"{stable_text}%\n"
        )
    summary.write(
        f"  Flat noise / base residual noise: "
        f"{float(metrics.get('flat_noise_dn', 0.0)):.3f} / "
        f"{float(metrics.get('base_residual_noise_dn', 0.0)):.3f} DN\n"
    )
    if metrics.get("flat_noise_worst_dn") is not None:
        summary.write(
            f"  Flat noise typical / worst: "
            f"{float(metrics.get('flat_noise_typical_dn', 0.0)):.3f} / "
            f"{float(metrics.get('flat_noise_worst_dn', 0.0)):.3f} DN\n"
        )
        summary.write(
            f"  Uniformity typical / worst / row-column banding: "
            f"{float(metrics.get('uniformity_typical_dn', 0.0)):.3f} / "
            f"{float(metrics.get('uniformity_worst_dn', 0.0)):.3f} / "
            f"{float(metrics.get('row_column_banding_dn', 0.0)):.3f} DN\n"
        )
    if metrics.get("impulse_outlier_control_score") is not None:
        summary.write(
            f"  Robust noise / impulse control / impulse pixels: "
            f"{score_text('robust_noise_score')} / "
            f"{score_text('impulse_outlier_control_score')} / "
            f"{float(metrics.get('impulse_outlier_percent', 0.0)):.4f}%\n"
        )
    summary.write(
        f"  Detail SNR / halo ratio / seam periodicity: "
        f"{float(metrics.get('detail_snr', 0.0)):.3f} / "
        f"{float(metrics.get('halo_ratio', 0.0)):.3f} / "
        f"{float(metrics.get('seam_periodicity_score', 0.0)):.3f}\n"
    )


def save_quality_excel(
    path,
    report_title,
    metadata,
    analyses,
    temporal_timelines=None,
    timeline_metric_definitions=None,
):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel output requires openpyxl. Install dependencies with: "
            "pip install -r requirements_thermal_raw_toolkit.txt"
        ) from exc

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append([report_title])
    summary_sheet.merge_cells("A1:C1")
    summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor="2F75B5")
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    summary_sheet.append(["Section", "Field", "Value"])
    for field_name, value in _flatten_excel_fields(metadata):
        summary_sheet.append(["Report", field_name, value])
    for metric_name, value in ACTIVE_SPATIAL_QUALITY_WEIGHTS.items():
        summary_sheet.append(
            ["Spatial Weight", metric_name, f"{value}%"]
        )
    for metric_name, value in ACTIVE_UVC_TEMPORAL_WEIGHTS.items():
        summary_sheet.append(
            ["UVC Temporal Weight", metric_name, f"{value}%"]
        )
    score_metric_names = [
        "overall_score",
        "detail_clarity_score",
        "contrast_quality_score",
        "flat_cleanliness_score",
        "thermal_uniformity_score",
        "clipping_control_score",
        "artifact_control_score",
    ]
    if analyses and all(
        "quality_confidence_score" in analysis["metrics"] for analysis in analyses
    ):
        score_metric_names.append("quality_confidence_score")
    if analyses and all(
        "flicker_control_score" in analysis["metrics"] for analysis in analyses
    ):
        score_metric_names.extend(
            ["flicker_control_score", "temporal_stability_score"]
        )

    def score_display_name(metric_name):
        if metric_name == "overall_score":
            return "Total Score"
        return metric_name

    def excel_value(value):
        if value is None:
            return None
        if isinstance(value, (bool, np.bool_)):
            return bool(value)
        if isinstance(value, (int, float, np.integer, np.floating)):
            return float(value)
        return str(value)

    for analysis in analyses:
        for metric_name in score_metric_names:
            summary_sheet.append(
                [
                    analysis["label"],
                    score_display_name(metric_name),
                    excel_value(analysis["metrics"][metric_name]),
                ]
            )
    for cell in summary_sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    summary_sheet.freeze_panes = "A3"
    summary_sheet.auto_filter.ref = f"A2:C{summary_sheet.max_row}"
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 36
    summary_sheet.column_dimensions["C"].width = 48

    metrics_sheet = workbook.create_sheet("Score Comparison")
    metrics_sheet.append(["Score"] + [analysis["label"] for analysis in analyses])
    for metric_name in score_metric_names:
        metrics_sheet.append(
            [score_display_name(metric_name)]
            + [excel_value(analysis["metrics"][metric_name]) for analysis in analyses]
        )
    diagnostic_metric_names = [
        metric_name
        for metric_name in analyses[0]["metrics"]
        if metric_name not in score_metric_names
    ]
    diagnostics_sheet = workbook.create_sheet("Measurements")
    diagnostics_sheet.append(
        ["Measurement"] + [analysis["label"] for analysis in analyses]
    )
    for metric_name in diagnostic_metric_names:
        diagnostics_sheet.append(
            [metric_name]
            + [excel_value(analysis["metrics"][metric_name]) for analysis in analyses]
        )
    for worksheet in (metrics_sheet, diagnostics_sheet):
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            worksheet.column_dimensions[column_cells[0].column_letter].width = 20

    temporal_sheet = None
    timeline_chart_columns = []
    if temporal_timelines:
        temporal_sheet = workbook.create_sheet("Temporal Timeline")
        timeline_metrics = timeline_metric_definitions or [
            ("Total Score", "overall_score"),
            ("Spatial", "spatial_quality_score"),
            ("Flicker Control", "flicker_control_score"),
            ("Temporal Stability", "temporal_stability_score"),
            ("Flicker DN", "flicker_dn"),
            ("Temporal Noise DN", "temporal_noise_dn"),
            ("Stable Region %", "stable_region_percent"),
        ]
        headers = ["Elapsed Seconds"]
        for stream_label in temporal_timelines:
            for display_name, _metric_name in timeline_metrics:
                headers.append(f"{stream_label} {display_name}")
        temporal_sheet.append(headers)
        max_samples = max(len(items) for items in temporal_timelines.values())
        first_timeline = next(iter(temporal_timelines.values()))
        for sample_index in range(max_samples):
            elapsed = (
                first_timeline[sample_index]["elapsed_seconds"]
                if sample_index < len(first_timeline)
                else None
            )
            row = [elapsed]
            for items in temporal_timelines.values():
                metrics = (
                    items[sample_index]["metrics"]
                    if sample_index < len(items)
                    else {}
                )
                row.extend(
                    excel_value(metrics.get(metric_name))
                    for _display_name, metric_name in timeline_metrics
                )
            temporal_sheet.append(row)
        for cell in temporal_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        temporal_sheet.freeze_panes = "A2"
        temporal_sheet.auto_filter.ref = temporal_sheet.dimensions
        temporal_sheet.column_dimensions["A"].width = 18
        for column_index in range(2, temporal_sheet.max_column + 1):
            temporal_sheet.column_dimensions[
                temporal_sheet.cell(1, column_index).column_letter
            ].width = 24
        metrics_per_stream = len(timeline_metrics)
        for stream_index in range(len(temporal_timelines)):
            first_column = 2 + stream_index * metrics_per_stream
            timeline_chart_columns.extend(
                first_column + metric_index
                for metric_index, (_display_name, metric_name) in enumerate(
                    timeline_metrics
                )
                if metric_name.endswith("_score")
            )
        temporal_summary_sheet = workbook.create_sheet("Temporal Summary")
        temporal_summary_sheet.append(
            ["Stream", "Score", "Average", "P10", "Worst"]
        )
        for stream_label, items in temporal_timelines.items():
            for display_name, metric_name in timeline_metrics:
                values = np.asarray(
                    [
                        item["metrics"].get(metric_name)
                        for item in items
                        if item["metrics"].get(metric_name) is not None
                    ],
                    dtype=np.float64,
                )
                temporal_summary_sheet.append(
                    [
                        stream_label,
                        display_name,
                        None if values.size == 0 else float(np.mean(values)),
                        None
                        if values.size == 0
                        else float(np.percentile(values, 10.0)),
                        None if values.size == 0 else float(np.min(values)),
                    ]
                )
        for cell in temporal_summary_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        temporal_summary_sheet.freeze_panes = "A2"
        temporal_summary_sheet.auto_filter.ref = temporal_summary_sheet.dimensions
        for column_cells in temporal_summary_sheet.columns:
            temporal_summary_sheet.column_dimensions[
                column_cells[0].column_letter
            ].width = 22

    charts_sheet = workbook.create_sheet("Charts")
    score_chart = BarChart()
    score_chart.type = "col"
    score_chart.style = 10
    score_chart.title = "Quality Score Comparison"
    score_chart.y_axis.title = "Score (0-100)"
    score_chart.height = 10
    score_chart.width = 24
    score_chart.add_data(
        Reference(
            metrics_sheet,
            min_col=2,
            max_col=1 + len(analyses),
            min_row=1,
            max_row=1 + len(score_metric_names),
        ),
        titles_from_data=True,
    )
    score_chart.set_categories(
        Reference(
            metrics_sheet,
            min_col=1,
            min_row=2,
            max_row=1 + len(score_metric_names),
        )
    )
    charts_sheet.add_chart(score_chart, "A1")
    if temporal_sheet is not None and temporal_sheet.max_row >= 2:
        temporal_chart = LineChart()
        temporal_chart.style = 13
        temporal_chart.title = "UVC Temporal Quality Timeline"
        temporal_chart.y_axis.title = "Score (0-100)"
        temporal_chart.x_axis.title = "Elapsed Seconds"
        temporal_chart.height = 11
        temporal_chart.width = 24
        for column_index in timeline_chart_columns:
            temporal_chart.add_data(
                Reference(
                    temporal_sheet,
                    min_col=column_index,
                    min_row=1,
                    max_row=temporal_sheet.max_row,
                ),
                titles_from_data=True,
            )
        temporal_chart.set_categories(
            Reference(
                temporal_sheet,
                min_col=1,
                min_row=2,
                max_row=temporal_sheet.max_row,
            )
        )
        charts_sheet.add_chart(temporal_chart, "A22")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def calculate_nuc(low, high, min_delta):
    low_f = low.astype(np.float64)
    high_f = high.astype(np.float64)
    low_avg = low_f.mean()
    high_avg = high_f.mean()
    target_delta = high_avg - low_avg
    pixel_delta = high_f - low_f

    valid = pixel_delta > min_delta
    gain = np.ones_like(low_f, dtype=np.float64)
    offset = np.zeros_like(low_f, dtype=np.float64)
    gain[valid] = target_delta / pixel_delta[valid]
    offset[valid] = low_avg - gain[valid] * low_f[valid]
    return gain, offset, valid, pixel_delta, low_avg, high_avg


def convert_additive_b_to_gfx_b(old_b, stored_k):
    old_b_i64 = np.asarray(old_b, dtype=np.int64)
    stored_k_i64 = np.asarray(stored_k, dtype=np.int64)
    gfx_b = np.zeros(old_b_i64.shape, dtype=np.int64)
    nonzero = stored_k_i64 != 0
    numerator = -(old_b_i64[nonzero] * NOVATEK_NUC_GAIN_ONE)
    denominator = stored_k_i64[nonzero]
    rounded = (np.abs(numerator) + denominator // 2) // denominator
    gfx_b[nonzero] = np.where(numerator < 0, -rounded, rounded)
    return np.clip(gfx_b, -32768, 32767).astype(np.int16)


def build_novatek_nuc_tables(low, high, min_delta):
    low_i64 = np.asarray(low, dtype=np.int64)
    high_i64 = np.asarray(high, dtype=np.int64)
    pixel_count = low_i64.size
    low_avg = int(low_i64.sum() // pixel_count)
    high_avg = int(high_i64.sum() // pixel_count)
    target_delta = high_avg - low_avg
    if target_delta <= 0:
        raise ValueError(
            f"Novatek NUC requires HIGH average > LOW average; got LOW={low_avg}, HIGH={high_avg}."
        )

    pixel_delta = high_i64 - low_i64
    valid = pixel_delta > min_delta
    stored_k_i64 = np.full(low_i64.shape, NOVATEK_NUC_GAIN_ONE, dtype=np.int64)
    raw_k = np.zeros(low_i64.shape, dtype=np.int64)
    raw_k[valid] = (target_delta << NOVATEK_NUC_GAIN_SHIFT) // pixel_delta[valid]
    table_valid = valid & (raw_k > 0)
    stored_k_i64[table_valid] = np.clip(raw_k[table_valid], 1, 65535)

    old_b = np.zeros(low_i64.shape, dtype=np.int64)
    old_b[table_valid] = low_avg - (
        (stored_k_i64[table_valid] * low_i64[table_valid]) >> NOVATEK_NUC_GAIN_SHIFT
    )
    gfx_b = np.zeros(low_i64.shape, dtype=np.int16)
    gfx_b[table_valid] = convert_additive_b_to_gfx_b(
        old_b[table_valid], stored_k_i64[table_valid]
    )
    stored_k = stored_k_i64.astype(np.uint16)
    return stored_k, gfx_b, table_valid, low_avg, high_avg


def apply_novatek_nuc(raw, stored_k, gfx_b, raw_max):
    corrected = (
        stored_k.astype(np.int64) * (raw.astype(np.int64) - gfx_b.astype(np.int64))
    ) >> NOVATEK_NUC_GAIN_SHIFT
    return np.clip(corrected, 0, raw_max).astype(np.uint16)


def save_csv_matrix(path, data, fmt):
    np.savetxt(path, data, delimiter=",", fmt=fmt)


def load_novatek_nuc_bin(path, width, height, dtype, table_name):
    if not path:
        raise ValueError(f"Select the Novatek {table_name} BIN file.")
    expected_bytes = width * height * np.dtype(dtype).itemsize
    actual_bytes = os.path.getsize(path)
    if actual_bytes != expected_bytes:
        raise ValueError(
            f"{table_name} BIN size error:\n{path}\nActual: {actual_bytes:,} bytes\n"
            f"Expected: {expected_bytes:,} bytes ({width}x{height}x2)"
        )
    return np.fromfile(path, dtype=dtype).reshape((height, width))


def deduplicate_coordinates(coordinates):
    unique = []
    seen = set()
    for x, y in coordinates:
        coordinate = (int(x), int(y))
        if coordinate in seen:
            continue
        seen.add(coordinate)
        unique.append(coordinate)
    return unique


def load_bpr_sie_coordinates(path):
    if not path:
        raise ValueError("Select a BPR dpc_table.bin file.")
    actual_bytes = os.path.getsize(path)
    if actual_bytes < 4 or actual_bytes % 2 != 0:
        raise ValueError(
            f"Invalid BPR BIN size: {actual_bytes:,} bytes. Expected an even-sized SIE table."
        )

    words = np.fromfile(path, dtype="<u2")
    coordinates = []
    current_y = None
    terminator_word = None
    for word_index, raw_value in enumerate(words):
        value = int(raw_value)
        if value == 0xFFFF:
            terminator_word = word_index
            break
        if value & 0x8000:
            current_y = value & 0x7FFF
            continue
        if current_y is None:
            raise ValueError(
                f"Invalid BPR SIE table: X={value} appears before a Y row marker at word {word_index}."
            )
        coordinates.append((value, current_y))

    if terminator_word is None:
        raise ValueError("Invalid BPR SIE table: 0xFFFF terminator was not found.")
    if terminator_word + 1 >= words.size or int(words[terminator_word + 1]) != 0:
        raise ValueError("Invalid BPR SIE table: expected 0x0000 after the 0xFFFF terminator.")

    unique = deduplicate_coordinates(coordinates)
    return unique, {
        "format": "Novatek SIE dpc_table.bin",
        "entry_count": len(coordinates),
        "unique_count": len(unique),
        "duplicate_count": len(coordinates) - len(unique),
        "terminator_word": terminator_word,
    }


def _parse_coordinate_integer(token):
    token = token.strip().strip("()[]{}")
    if token.lower().startswith(("0x", "+0x", "-0x")):
        return int(token, 16)
    return int(token, 10)


def load_bpr_text_coordinates(path):
    if not path:
        raise ValueError("Select a BPR coordinate CSV/TXT file.")
    coordinates = []
    with open(path, "r", encoding="utf-8-sig") as source:
        for line_number, original_line in enumerate(source, 1):
            line = original_line.split("#", 1)[0].split("//", 1)[0].strip()
            if not line:
                continue
            fields = [field for field in re.split(r"[,;\s]+", line) if field]
            if len(fields) < 2:
                raise ValueError(
                    f"Invalid coordinate at line {line_number}: expected x,y."
                )
            try:
                x = _parse_coordinate_integer(fields[0])
                y = _parse_coordinate_integer(fields[1])
            except ValueError as exc:
                if (
                    not coordinates
                    and fields[0].lower() == "x"
                    and fields[1].lower() == "y"
                ):
                    continue
                raise ValueError(
                    f"Invalid coordinate at line {line_number}: {original_line.strip()}"
                ) from exc
            coordinates.append((x, y))

    if not coordinates:
        raise ValueError("The BPR coordinate file does not contain any x,y coordinates.")
    unique = deduplicate_coordinates(coordinates)
    return unique, {
        "format": "Text coordinates (x,y)",
        "entry_count": len(coordinates),
        "unique_count": len(unique),
        "duplicate_count": len(coordinates) - len(unique),
        "terminator_word": None,
    }


def load_bpr_coordinates(path):
    suffix = Path(path).suffix.lower() if path else ""
    if suffix == ".bin":
        return load_bpr_sie_coordinates(path)
    if suffix in (".csv", ".txt"):
        return load_bpr_text_coordinates(path)
    raise ValueError("BPR coordinates must be a Novatek .bin table or an x,y .csv/.txt file.")


def split_coordinates_by_bounds(coordinates, width, height):
    valid = []
    outside = []
    for coordinate in coordinates:
        x, y = coordinate
        if 0 <= x < width and 0 <= y < height:
            valid.append(coordinate)
        else:
            outside.append(coordinate)
    return valid, outside


def format_bpr_coordinate_report(table_info, coordinates, outside_coordinates, width, height):
    outside_set = set(outside_coordinates)
    marked_count = len(coordinates) - len(outside_coordinates)
    lines = [
        f"Table entries           : {table_info['entry_count']}",
        f"Unique bad pixels       : {table_info['unique_count']}",
        f"Duplicate entries       : {table_info['duplicate_count']}",
        f"Marked in {width}x{height}      : {marked_count}",
        f"Out-of-bounds           : {len(outside_coordinates)}",
        "",
        "All unique coordinates (top-left origin; OUT = not marked):",
    ]
    coordinate_tokens = []
    index_width = max(3, len(str(len(coordinates))))
    for index, (x, y) in enumerate(coordinates, 1):
        suffix = " [OUT]" if (x, y) in outside_set else ""
        coordinate_tokens.append(
            f"{index:0{index_width}d}:({x},{y}){suffix}"
        )
    for start in range(0, len(coordinate_tokens), 4):
        lines.append("    ".join(coordinate_tokens[start:start + 4]))
    return "\n".join(lines)


def draw_bpr_markers(
    base_gray,
    coordinates,
    radius=BPR_MARK_RADIUS,
    highlight_bright_points=False,
    highlight_radius=BPR_HIGHLIGHT_RADIUS,
):
    if cv2 is None:
        raise RuntimeError("OpenCV is required for BPR marking. Install: pip install opencv-python")
    if base_gray.ndim != 2:
        raise ValueError("BPR marker base image must be a 2D grayscale image.")
    marked_rgb = cv2.cvtColor(
        np.ascontiguousarray(np.clip(base_gray, 0, 255).astype(np.uint8)),
        cv2.COLOR_GRAY2RGB,
    )
    for x, y in coordinates:
        cv2.circle(
            marked_rgb,
            (int(x), int(y)),
            int(radius),
            (255, 0, 0),
            1,
            cv2.LINE_AA,
        )
        if highlight_bright_points:
            cv2.circle(
                marked_rgb,
                (int(x), int(y)),
                int(highlight_radius),
                (255, 255, 255),
                -1,
                cv2.LINE_AA,
            )
    return marked_rgb


def save_png(path, image):
    if cv2 is None:
        raise RuntimeError("OpenCV is required to save PNG. Install: pip install opencv-python")
    ok = cv2.imwrite(str(path), image)
    if not ok:
        raise RuntimeError(f"Failed to save image: {path}")


def save_rgb_png(path, image_rgb):
    if cv2 is None:
        raise RuntimeError("OpenCV is required to save PNG. Install: pip install opencv-python")
    if image_rgb.ndim != 3 or image_rgb.shape[2] < 3:
        raise ValueError("RGB PNG output requires a three-channel image.")
    bgr = cv2.cvtColor(np.ascontiguousarray(image_rgb[:, :, :3]), cv2.COLOR_RGB2BGR)
    ok = cv2.imwrite(str(path), bgr)
    if not ok:
        raise RuntimeError(f"Failed to save image: {path}")


def save_pipeline_overview_png(path, items, max_columns=3):
    """Save numbered pipeline images as one labeled contact sheet."""
    if cv2 is None:
        raise RuntimeError("OpenCV is required to save PNG. Install: pip install opencv-python")
    if not items:
        raise ValueError("Pipeline overview requires at least one image.")

    prepared = []
    for label, image, display_max in items:
        if image.ndim != 2:
            raise ValueError("Pipeline overview supports grayscale images only.")
        if display_max <= 0:
            raise ValueError("Pipeline overview display maximum must be positive.")
        image_u8 = np.rint(
            np.clip(image, 0, display_max).astype(np.float64)
            * 255.0
            / display_max
        ).astype(np.uint8)
        prepared.append((str(label), image_u8))

    source_height, source_width = prepared[0][1].shape
    if any(image.shape != (source_height, source_width) for _, image in prepared):
        raise ValueError("All pipeline overview images must have the same dimensions.")

    scale = min(1.0, 480.0 / source_width, 360.0 / source_height)
    tile_width = max(1, int(round(source_width * scale)))
    tile_height = max(1, int(round(source_height * scale)))
    header_height = 42
    gap = 8
    margin = 8
    columns = min(max(1, int(max_columns)), len(prepared))
    rows = (len(prepared) + columns - 1) // columns
    panel_width = tile_width
    panel_height = header_height + tile_height
    canvas_width = margin * 2 + columns * panel_width + (columns - 1) * gap
    canvas_height = margin * 2 + rows * panel_height + (rows - 1) * gap
    canvas = np.full((canvas_height, canvas_width), 24, dtype=np.uint8)

    for item_index, (label, image_u8) in enumerate(prepared):
        row = item_index // columns
        column = item_index % columns
        x = margin + column * (panel_width + gap)
        y = margin + row * (panel_height + gap)
        resized = (
            image_u8
            if image_u8.shape == (tile_height, tile_width)
            else cv2.resize(
                image_u8,
                (tile_width, tile_height),
                interpolation=cv2.INTER_AREA,
            )
        )
        canvas[y + header_height:y + panel_height, x:x + tile_width] = resized
        cv2.rectangle(
            canvas,
            (x, y),
            (x + panel_width - 1, y + panel_height - 1),
            72,
            1,
        )
        font_scale = min(
            0.62,
            max(0.36, panel_width / max(500.0, len(label) * 17.0)),
        )
        text_size, _ = cv2.getTextSize(
            label, cv2.FONT_HERSHEY_SIMPLEX, font_scale, 1
        )
        text_x = x + max(6, (panel_width - text_size[0]) // 2)
        cv2.putText(
            canvas,
            label,
            (text_x, y + 28),
            cv2.FONT_HERSHEY_SIMPLEX,
            font_scale,
            255,
            1,
            cv2.LINE_AA,
        )

    save_png(path, canvas)


def save_uvc_comparison_png(path, original, processed, algorithm_title):
    if cv2 is None:
        raise RuntimeError("OpenCV is required to save PNG. Install: pip install opencv-python")
    if original.ndim != 2 or processed.ndim != 2:
        raise ValueError("UVC comparison PNG requires two grayscale images.")
    if original.shape != processed.shape:
        raise ValueError(
            f"UVC comparison image size mismatch: original={original.shape}, processed={processed.shape}"
        )

    original_u8 = np.ascontiguousarray(np.clip(original, 0, 255).astype(np.uint8))
    processed_u8 = np.ascontiguousarray(np.clip(processed, 0, 255).astype(np.uint8))
    height, width = original_u8.shape
    header_height = 42
    gap = 8
    canvas = np.full((height + header_height, width * 2 + gap), 24, dtype=np.uint8)
    canvas[header_height:, :width] = original_u8
    canvas[header_height:, width + gap:] = processed_u8

    labels = ("Original UVC", f"Processed: {algorithm_title}")
    panel_x = (0, width + gap)
    for label, x_offset in zip(labels, panel_x):
        font_scale = min(0.7, max(0.42, width / max(420.0, len(label) * 18.0)))
        text_size, _ = cv2.getTextSize(
            label, cv2.FONT_HERSHEY_SIMPLEX, font_scale, 1
        )
        text_x = x_offset + max(6, (width - text_size[0]) // 2)
        cv2.putText(
            canvas,
            label,
            (text_x, 28),
            cv2.FONT_HERSHEY_SIMPLEX,
            font_scale,
            255,
            1,
            cv2.LINE_AA,
        )

    save_png(path, canvas)


class PreviewLabel(QLabel):
    double_clicked = pyqtSignal()
    image_changed = pyqtSignal()

    def __init__(self, title):
        super().__init__()
        self.title = title
        self._pixmap = None
        self.setAlignment(Qt.AlignCenter)
        self.setMinimumSize(360, 200)
        self.setStyleSheet("background:#121417; color:#aeb7c2; border:1px solid #424850; border-radius:8px; font-size:15px; font-weight:600;")
        self.setText("Waiting for image...")

    def set_title(self, title):
        self.title = title
        if self._pixmap is None:
            self.setText("Waiting for image...")

    def set_array(self, array):
        if array is None:
            self._pixmap = None
            self.setText("Waiting for image...")
            self.image_changed.emit()
            return
        if array.ndim == 2:
            img8 = np.ascontiguousarray(np.clip(array, 0, 255).astype(np.uint8))
            qimg = QImage(img8.data, img8.shape[1], img8.shape[0], img8.strides[0], QImage.Format_Grayscale8).copy()
        else:
            rgb = np.ascontiguousarray(array[:, :, :3])
            qimg = QImage(rgb.data, rgb.shape[1], rgb.shape[0], rgb.strides[0], QImage.Format_RGB888).copy()
        self._pixmap = QPixmap.fromImage(qimg)
        self._rescale()
        self.image_changed.emit()

    def _rescale(self):
        if self._pixmap:
            self.setPixmap(self._pixmap.scaled(self.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation))

    def resizeEvent(self, event):
        self._rescale()
        super().resizeEvent(event)

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.double_clicked.emit()
            event.accept()
            return
        super().mouseDoubleClickEvent(event)


class SwipeCompareWidget(QWidget):
    double_clicked = pyqtSignal()
    position_changed = pyqtSignal(float)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.frames = (None, None)
        self.names = ("Image A", "Image B")
        self.position = 0.5
        self.dragging = False
        self.setMinimumSize(360, 200)
        self.setMouseTracking(True)

    def set_frames(self, frames, names):
        self.frames = tuple(frames)
        self.names = tuple(names)
        self.update()

    def set_position(self, position):
        self.position = max(0.0, min(1.0, float(position)))
        self.update()

    def image_rect(self):
        if any(frame is None for frame in self.frames):
            return QRectF()
        width = max(frame.width() for frame in self.frames)
        height = max(frame.height() for frame in self.frames)
        scale = min(self.width() / width, self.height() / height)
        return QRectF((self.width() - width * scale) / 2,
                      (self.height() - height * scale) / 2,
                      width * scale, height * scale)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.fillRect(self.rect(), QColor("#000000"))
        area = self.image_rect()
        if area.isEmpty():
            painter.setPen(QColor("#aeb7c2"))
            painter.drawText(self.rect(), Qt.AlignCenter, "Waiting for both images...")
            return
        painter.setRenderHint(QPainter.SmoothPixmapTransform)
        split = area.left() + area.width() * self.position
        clips = (QRectF(area.left(), area.top(), split - area.left(), area.height()),
                 QRectF(split, area.top(), area.right() - split, area.height()))
        for frame, clip in zip(self.frames, clips):
            scale = min(area.width() / frame.width(), area.height() / frame.height())
            target = QRectF(0, 0, frame.width() * scale, frame.height() * scale)
            target.moveCenter(area.center())
            painter.save()
            painter.setClipRect(clip)
            painter.drawPixmap(target, frame, QRectF(frame.rect()))
            painter.restore()
        painter.setPen(QPen(QColor("#ffffff"), 2))
        painter.drawLine(QPointF(split, area.top()), QPointF(split, area.bottom()))
        font = QFont("Segoe UI", 11)
        font.setBold(False)
        painter.setFont(font)
        for name, alignment in zip(self.names, (Qt.AlignLeft, Qt.AlignRight)):
            label = QRectF(area.left() + 10, area.top() + 8, max(0, area.width() - 20), 26)
            painter.setPen(QColor("#000000"))
            painter.drawText(label.translated(1, 1), alignment | Qt.AlignVCenter, name)
            painter.setPen(QColor("#ffffff"))
            painter.drawText(label, alignment | Qt.AlignVCenter, name)

    def mousePressEvent(self, event):
        area = self.image_rect()
        split = area.left() + area.width() * self.position
        if (event.button() == Qt.LeftButton and not area.isEmpty()
                and area.top() <= event.y() <= area.bottom()
                and abs(event.x() - split) <= 16):
            self.dragging = True
            self.setCursor(Qt.SplitHCursor)
            event.accept()
            return
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        area = self.image_rect()
        if not area.isEmpty():
            if self.dragging:
                self.set_position((event.x() - area.left()) / area.width())
                self.position_changed.emit(self.position)
            near_bar = abs(event.x() - area.left() - area.width() * self.position) <= 16
            self.setCursor(Qt.SplitHCursor if self.dragging or near_bar else Qt.ArrowCursor)
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.dragging = False
        super().mouseReleaseEvent(event)

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.dragging = False
            self.double_clicked.emit()
            event.accept()
            return
        super().mouseDoubleClickEvent(event)


class QualityHelpTitle(QLabel):
    help_requested = pyqtSignal()

    def __init__(self, title, parent=None):
        super().__init__(f"{title}  ⓘ", parent)
        self.setCursor(Qt.PointingHandCursor)
        self.setToolTip("Click to view score definitions and formulas")
        self.setAccessibleName(f"{title} help")
        self._set_hovered(False)

    def _set_hovered(self, hovered):
        if hovered:
            self.setStyleSheet(
                "color:#86c8f4; font-size:18px; font-weight:800; "
                "padding:2px 4px;"
            )
        else:
            self.setStyleSheet(
                "color:#c9e2f2; font-size:18px; font-weight:600; "
                "padding:2px 4px;"
            )

    def enterEvent(self, event):
        self._set_hovered(True)
        super().enterEvent(event)

    def leaveEvent(self, event):
        self._set_hovered(False)
        super().leaveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.LeftButton and self.rect().contains(event.pos()):
            self.help_requested.emit()
            event.accept()
            return
        super().mouseReleaseEvent(event)


class CollapsibleGroupBox(QGroupBox):
    """QGroupBox whose title toggles the visibility of its body."""

    def __init__(self, title, parent=None):
        super().__init__(parent)
        self._section_title = str(title)
        self._expanded = True
        self._stored_visibility = {}
        self.setProperty("collapsibleGroup", True)
        self.setMouseTracking(True)
        self.setFocusPolicy(Qt.StrongFocus)
        self._update_title()

    def _update_title(self):
        marker = "-" if self._expanded else "+"
        super().setTitle(f"{marker}  {self._section_title}")
        self.setToolTip(
            "Click the title to collapse" if self._expanded
            else "Click the title to expand"
        )

    def setTitle(self, title):
        if not hasattr(self, "_section_title"):
            super().setTitle(title)
            return
        self._section_title = str(title)
        self._update_title()

    def _title_rect(self):
        option = QStyleOptionGroupBox()
        self.initStyleOption(option)
        return self.style().subControlRect(
            QStyle.CC_GroupBox,
            option,
            QStyle.SC_GroupBoxLabel,
            self,
        )

    @staticmethod
    def _layout_widgets(layout):
        widgets = []
        for index in range(layout.count()):
            item = layout.itemAt(index)
            widget = item.widget()
            if widget is not None:
                widgets.append(widget)
                continue
            child_layout = item.layout()
            if child_layout is not None:
                widgets.extend(CollapsibleGroupBox._layout_widgets(child_layout))
        return widgets

    def set_expanded(self, expanded):
        expanded = bool(expanded)
        if expanded == self._expanded:
            return
        layout = self.layout()
        if layout is None:
            return

        if not expanded:
            self._stored_visibility = {}
            for widget in self._layout_widgets(layout):
                self._stored_visibility[widget] = not widget.isHidden()
                widget.hide()
        else:
            for widget, was_visible in self._stored_visibility.items():
                widget.setVisible(was_visible)
            self._stored_visibility.clear()

        self._expanded = expanded
        self._update_title()
        self.setMaximumHeight(
            16777215 if expanded else self._title_rect().bottom() + 10
        )
        self.updateGeometry()

    def toggle_expanded(self):
        self.set_expanded(not self._expanded)

    def is_expanded(self):
        return self._expanded

    def mouseMoveEvent(self, event):
        self.setCursor(
            Qt.PointingHandCursor
            if self._title_rect().contains(event.pos())
            else Qt.ArrowCursor
        )
        super().mouseMoveEvent(event)

    def leaveEvent(self, event):
        self.setCursor(Qt.ArrowCursor)
        super().leaveEvent(event)

    def mouseReleaseEvent(self, event):
        if (
            event.button() == Qt.LeftButton
            and self._title_rect().contains(event.pos())
        ):
            self.toggle_expanded()
            event.accept()
            return
        super().mouseReleaseEvent(event)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key_Space, Qt.Key_Return, Qt.Key_Enter):
            self.toggle_expanded()
            event.accept()
            return
        super().keyPressEvent(event)


class QualityScoreCheckBox(QCheckBox):
    """Large dark-theme checkbox with an explicit painted check mark."""

    def __init__(self, text="Calculate Quality Score", parent=None):
        super().__init__(text, parent)
        self.setObjectName("qualityScoreToggle")
        self._hovered = False
        font = self.font()
        font.setFamily("Segoe UI")
        font.setPixelSize(20)
        font.setWeight(QFont.DemiBold)
        font.setStretch(QFont.Unstretched)
        self.setFont(font)
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(42)
        self.setMinimumWidth(280)

    def sizeHint(self):
        return QSize(300, 42)

    def enterEvent(self, event):
        self._hovered = True
        self.update()
        super().enterEvent(event)

    def leaveEvent(self, event):
        self._hovered = False
        self.update()
        super().leaveEvent(event)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing, True)

        box_size = 22
        box_x = 2
        box_y = (self.height() - box_size) // 2
        box_rect = QRect(box_x, box_y, box_size, box_size)

        if self.isChecked():
            fill_color = QColor("#4f9dcc" if not self._hovered else "#62b2df")
            border_color = QColor("#a9ddf8")
        else:
            fill_color = QColor("#1b1f24" if not self._hovered else "#252b31")
            border_color = QColor("#7e8995" if not self._hovered else "#a9b4bf")
        if not self.isEnabled():
            fill_color = QColor("#2a2d31")
            border_color = QColor("#555b62")

        painter.setPen(QPen(border_color, 1.5))
        painter.setBrush(fill_color)
        painter.drawRoundedRect(box_rect, 5, 5)

        if self.isChecked():
            check_pen = QPen(QColor("#ffffff"), 2.6)
            check_pen.setCapStyle(Qt.RoundCap)
            check_pen.setJoinStyle(Qt.RoundJoin)
            painter.setPen(check_pen)
            painter.setBrush(Qt.NoBrush)
            painter.drawLine(box_x + 5, box_y + 11, box_x + 9, box_y + 15)
            painter.drawLine(box_x + 9, box_y + 15, box_x + 17, box_y + 7)

        text_color = QColor("#edf4f8" if self.isEnabled() else "#7d858d")
        if self._hovered and self.isEnabled():
            text_color = QColor("#ffffff")
        painter.setPen(text_color)
        painter.setFont(self.font())
        metrics = painter.fontMetrics()
        text_y = (self.height() + metrics.ascent() - metrics.descent()) // 2
        painter.drawText(34, text_y, self.text())


class FocusWheelGuard:
    """Allow wheel changes only after explicit mouse or keyboard focus."""

    def _initialize_wheel_guard(self):
        self.setFocusPolicy(Qt.StrongFocus)
        self._wheel_input_enabled = False

    def mousePressEvent(self, event):
        self._wheel_input_enabled = True
        super().mousePressEvent(event)

    def focusInEvent(self, event):
        if event.reason() in (
            Qt.MouseFocusReason,
            Qt.TabFocusReason,
            Qt.BacktabFocusReason,
        ):
            self._wheel_input_enabled = True
        elif event.reason() != Qt.PopupFocusReason:
            self._wheel_input_enabled = False
        super().focusInEvent(event)

    def focusOutEvent(self, event):
        # Keep the marker while a popup or child editor temporarily owns focus.
        # wheelEvent() still rejects input whenever hasFocus() is false;
        # the next focusInEvent decides whether wheel input is enabled again.
        super().focusOutEvent(event)

    def wheelEvent(self, event):
        if not self.hasFocus() or not self._wheel_input_enabled:
            event.ignore()
            return
        super().wheelEvent(event)


class FocusWheelComboBox(FocusWheelGuard, QComboBox):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._initialize_wheel_guard()


class FocusWheelSpinBox(FocusWheelGuard, QSpinBox):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._initialize_wheel_guard()


class FocusWheelDoubleSpinBox(FocusWheelGuard, QDoubleSpinBox):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._initialize_wheel_guard()


class QualityScoreTimelineChart(QWidget):
    open_requested = pyqtSignal()
    data_changed = pyqtSignal()

    def __init__(self, series, history_seconds=60.0, parent=None):
        super().__init__(parent)
        self.series = tuple(series)
        self.history_seconds = float(history_seconds)
        self.metric_name = "overall_score"
        self.metric_display_name = "Total Score"
        self.samples = deque(maxlen=1200)
        self.paused = False
        self.setMinimumSize(360, 210)
        self.setCursor(Qt.PointingHandCursor)
        self.setToolTip("Double-click to open a resizable quality-score timeline")

    def set_display_metric(self, metric_name):
        metric_map = {
            metric: display for display, metric in QUALITY_TIMELINE_METRICS
        }
        if metric_name not in metric_map:
            metric_name = "overall_score"
        self.metric_name = metric_name
        self.metric_display_name = metric_map[metric_name]
        self.update()
        self.data_changed.emit()

    def set_history_seconds(self, seconds):
        self.history_seconds = float(max(10, min(600, int(seconds))))
        self.update()
        self.data_changed.emit()

    def set_series_labels(self, **labels):
        self.series = tuple(
            (key, str(labels.get(key, label)), color)
            for key, label, color in self.series
        )
        self.update()
        self.data_changed.emit()

    def append_metrics(self, timestamp=None, **stream_metrics):
        if self.paused:
            return
        timestamp = time.perf_counter() if timestamp is None else float(timestamp)
        values = {}
        for key, _label, _color in self.series:
            metrics = stream_metrics.get(key)
            values[key] = None if metrics is None else dict(metrics)
        self.samples.append((timestamp, values))
        cutoff = timestamp - self.history_seconds * 1.25
        while self.samples and self.samples[0][0] < cutoff:
            self.samples.popleft()
        self.update()
        self.data_changed.emit()

    def append_scores(self, timestamp=None, **scores):
        self.append_metrics(
            timestamp,
            **{
                key: {self.metric_name: value}
                for key, value in scores.items()
            },
        )

    def set_paused(self, paused):
        paused = bool(paused)
        if paused != self.paused and self.samples:
            self.samples.append((time.perf_counter(), {}))
        self.paused = paused
        self.update()
        self.data_changed.emit()

    def clear(self):
        self.samples.clear()
        self.update()
        self.data_changed.emit()

    def copy_state_from(self, source):
        self.series = tuple(source.series)
        self.samples = deque(source.samples, maxlen=1200)
        self.paused = source.paused
        self.history_seconds = source.history_seconds
        self.metric_name = source.metric_name
        self.metric_display_name = source.metric_display_name
        self.update()

    def excel_metric_definition(self):
        return [(self.metric_display_name, self.metric_name)]

    def _metric_value(self, values, key):
        stream = values.get(key)
        if isinstance(stream, dict):
            return stream.get(self.metric_name)
        return stream

    def excel_timelines(self):
        if not self.samples:
            return {}
        end_time = self.samples[-1][0]
        start_time = end_time - self.history_seconds
        visible_samples = [
            item for item in self.samples if item[0] >= start_time
        ]
        if not visible_samples:
            return {}
        first_time = visible_samples[0][0]
        timelines = {}
        for key, label, _color in self.series:
            timelines[label] = [
                {
                    "elapsed_seconds": round(timestamp - first_time, 3),
                    "metrics": {
                        self.metric_name: self._metric_value(values, key)
                    },
                }
                for timestamp, values in visible_samples
            ]
        return timelines

    def _series_values(self, key, visible_samples):
        return [
            self._metric_value(values, key)
            for _timestamp, values in visible_samples
            if self._metric_value(values, key) is not None
        ]

    @staticmethod
    def _trend_text(values):
        if len(values) < 6:
            return "Collecting"
        recent = values[-10:]
        half = max(1, len(recent) // 2)
        change = float(np.mean(recent[half:]) - np.mean(recent[:half]))
        if change > 1.0:
            return "Improving"
        if change < -1.0:
            return "Declining"
        return "Stable"

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing, True)
        painter.fillRect(self.rect(), QColor("#1b1e23"))

        plot = QRectF(48, 42, max(40, self.width() - 66), max(40, self.height() - 76))
        painter.setPen(QPen(QColor("#39414a"), 1))
        painter.setBrush(QColor("#20252b"))
        painter.drawRoundedRect(plot, 5, 5)

        label_font = QFont("Segoe UI")
        label_font.setPixelSize(11)
        painter.setFont(label_font)
        for score in range(0, 101, 20):
            y = plot.bottom() - (score / 100.0) * plot.height()
            painter.setPen(QPen(QColor("#3b444d"), 1, Qt.DotLine))
            painter.drawLine(QPointF(plot.left(), y), QPointF(plot.right(), y))
            painter.setPen(QColor("#aeb8c2"))
            painter.drawText(QRectF(4, y - 9, 39, 18), Qt.AlignRight | Qt.AlignVCenter, str(score))

        visible_samples = list(self.samples)
        if visible_samples:
            end_time = visible_samples[-1][0]
        else:
            end_time = time.perf_counter()
        start_time = end_time - self.history_seconds
        visible_samples = [item for item in visible_samples if item[0] >= start_time]

        for seconds_ago in (int(self.history_seconds), int(self.history_seconds / 2), 0):
            x = plot.right() - (seconds_ago / self.history_seconds) * plot.width()
            painter.setPen(QColor("#8f9aa5"))
            text = "Now" if seconds_ago == 0 else f"-{seconds_ago}s"
            painter.drawText(QRectF(x - 24, plot.bottom() + 5, 48, 18), Qt.AlignCenter, text)

        legend_x = plot.left()
        for key, label, color_name in self.series:
            color = QColor(color_name)
            values = self._series_values(key, visible_samples)
            latest = "--" if not values else f"{values[-1]:.1f}"
            average = "--" if not values else f"{float(np.mean(values)):.1f}"
            legend = f"{label}  {latest}  Avg {average}  {self._trend_text(values)}"
            painter.setPen(QPen(color, 3))
            painter.drawLine(QPointF(legend_x, 23), QPointF(legend_x + 18, 23))
            painter.setPen(QColor("#e5edf3"))
            painter.drawText(QRectF(legend_x + 24, 10, 250, 26), Qt.AlignLeft | Qt.AlignVCenter, legend)
            legend_x += min(285, max(190, self.width() // max(1, len(self.series))))

            path = QPainterPath()
            path_active = False
            for timestamp, score_values in visible_samples:
                value = self._metric_value(score_values, key)
                if value is None:
                    path_active = False
                    continue
                x = plot.left() + ((timestamp - start_time) / self.history_seconds) * plot.width()
                y = plot.bottom() - (max(0.0, min(100.0, value)) / 100.0) * plot.height()
                if not path_active:
                    path.moveTo(x, y)
                    path_active = True
                else:
                    path.lineTo(x, y)
            painter.setPen(QPen(color, 2.3))
            painter.setBrush(Qt.NoBrush)
            painter.drawPath(path)

        if not visible_samples:
            painter.setPen(QColor("#84909b"))
            painter.drawText(
                plot,
                Qt.AlignCenter,
                f"Waiting for {self.metric_display_name}...",
            )
        if self.paused:
            painter.fillRect(plot, QColor(20, 23, 27, 145))
            paused_font = QFont("Segoe UI")
            paused_font.setPixelSize(17)
            paused_font.setBold(True)
            painter.setFont(paused_font)
            painter.setPen(QColor("#d5dde4"))
            painter.drawText(plot, Qt.AlignCenter, "Quality scoring paused")

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.open_requested.emit()
            event.accept()
            return
        super().mouseDoubleClickEvent(event)


class QualityTimelineDialog(QDialog):
    def __init__(self, source_chart, title, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setWindowFlags(
            self.windowFlags()
            | Qt.WindowMinimizeButtonHint
            | Qt.WindowMaximizeButtonHint
        )
        self.resize(1050, 620)
        self.setMinimumSize(680, 420)
        layout = QVBoxLayout(self)
        self.chart = QualityScoreTimelineChart(
            source_chart.series, source_chart.history_seconds, self
        )
        layout.addWidget(self.chart)
        self.source_chart = source_chart
        source_chart.data_changed.connect(self.sync_from_source)
        self.sync_from_source()

    def sync_from_source(self):
        self.chart.copy_state_from(self.source_chart)


class ToolkitSettingsDialog(QDialog):
    settings_applied = pyqtSignal(object)

    SPATIAL_LABELS = {
        "detail_clarity_score": "Detail Clarity",
        "contrast_quality_score": "Contrast Quality",
        "flat_cleanliness_score": "Flat Cleanliness",
        "thermal_uniformity_score": "Thermal Uniformity",
        "clipping_control_score": "Clipping Control",
        "artifact_control_score": "Artifact Control",
    }
    TEMPORAL_LABELS = {
        "spatial_quality_score": "Spatial Quality",
        "flicker_control_score": "Flicker Control",
        "temporal_stability_score": "Temporal Stability",
    }

    def __init__(self, settings, parent=None):
        super().__init__(parent)
        self.setObjectName("toolkitSettingsDialog")
        self.setWindowTitle("Toolkit Setting")
        self.setWindowFlags(
            self.windowFlags()
            | Qt.WindowMinimizeButtonHint
            | Qt.WindowMaximizeButtonHint
        )
        self.resize(720, 760)
        self.setMinimumSize(640, 680)
        self.setSizeGripEnabled(True)
        root = QVBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(10)
        self.tabs = QTabWidget()
        root.addWidget(self.tabs, 1)

        general_page = QWidget()
        general_form = QFormLayout(general_page)
        general_form.setContentsMargins(18, 18, 18, 18)
        general_form.setHorizontalSpacing(16)
        general_form.setVerticalSpacing(14)
        self.resolution_combo = FocusWheelComboBox()
        for width, height in RESOLUTION_OPTIONS:
            self.resolution_combo.addItem(
                f"{width} × {height}", (width, height)
            )
        self.bit_depth_combo = FocusWheelComboBox()
        for bits in BIT_DEPTH_OPTIONS:
            self.bit_depth_combo.addItem(
                f"{bits}-bit (Max {(1 << bits) - 1})", bits
            )
        self.output_root_row = PathRow(
            "folder", dialog_title="Select Output Root Folder"
        )
        general_form.addRow("Resolution", self.resolution_combo)
        general_form.addRow("Bit Depth", self.bit_depth_combo)
        general_form.addRow("Output Root", self.output_root_row)
        self.camera_a_mirror_check = QCheckBox("Mirror Camera A horizontally")
        self.camera_b_mirror_check = QCheckBox("Mirror Camera B horizontally")
        mirror_options = QWidget()
        mirror_layout = QVBoxLayout(mirror_options)
        mirror_layout.setContentsMargins(0, 0, 0, 0)
        mirror_layout.setSpacing(6)
        mirror_layout.addWidget(self.camera_a_mirror_check)
        mirror_layout.addWidget(self.camera_b_mirror_check)
        general_form.addRow("Camera Mirror", mirror_options)
        self.tabs.addTab(general_page, "General")

        appearance_page = QWidget()
        appearance_form = QFormLayout(appearance_page)
        appearance_form.setContentsMargins(18, 18, 18, 18)
        appearance_form.setHorizontalSpacing(16)
        appearance_form.setVerticalSpacing(14)
        self.log_font_spin = FocusWheelSpinBox()
        self.log_font_spin.setRange(10, 28)
        self.log_font_spin.setSuffix(" px")
        appearance_form.addRow("Log Font Size", self.log_font_spin)
        self.tabs.addTab(appearance_page, "Appearance")

        timeline_page = QWidget()
        timeline_form = QFormLayout(timeline_page)
        self.timeline_metric_combo = FocusWheelComboBox()
        for display_name, metric_name in QUALITY_TIMELINE_METRICS:
            self.timeline_metric_combo.addItem(display_name, metric_name)
        self.timeline_history_spin = FocusWheelSpinBox()
        self.timeline_history_spin.setRange(10, 600)
        self.timeline_history_spin.setSuffix(" s")
        timeline_note = QLabel(
            "Single shows one stream; Dual shows Camera A/B; Algorithm "
            "shows Original/Final. Excel exports the selected metric."
        )
        timeline_note.setWordWrap(True)
        timeline_note.setStyleSheet("color:#aeb8c2; padding-top:6px;")
        timeline_form.addRow("Displayed Score", self.timeline_metric_combo)
        timeline_form.addRow("History", self.timeline_history_spin)
        timeline_form.addRow(timeline_note)
        self.tabs.addTab(timeline_page, "Timeline")

        weights_page = QWidget()
        weights_layout = QVBoxLayout(weights_page)
        spatial_group = QGroupBox("Spatial Total Score (RAW and UVC)")
        spatial_group.setObjectName("settingSection")
        spatial_form = QFormLayout(spatial_group)
        self.spatial_weight_spins = self._create_weight_spins(
            spatial_form, self.SPATIAL_LABELS
        )
        self.spatial_total_label = QLabel()
        spatial_form.addRow("Total", self.spatial_total_label)
        weights_layout.addWidget(spatial_group)

        temporal_group = QGroupBox("UVC Final Total Score")
        temporal_group.setObjectName("settingSection")
        temporal_form = QFormLayout(temporal_group)
        self.temporal_weight_spins = self._create_weight_spins(
            temporal_form, self.TEMPORAL_LABELS
        )
        self.temporal_total_label = QLabel()
        temporal_form.addRow("Total", self.temporal_total_label)
        weights_layout.addWidget(temporal_group)
        weights_note = QLabel(
            "Each section must total exactly 100%. Applying new weights "
            "clears UVC timelines so different scoring definitions are not mixed."
        )
        weights_note.setWordWrap(True)
        weights_note.setStyleSheet("color:#aeb8c2; padding:6px;")
        weights_layout.addWidget(weights_note)
        weights_layout.addStretch()
        self.tabs.addTab(weights_page, "Quality Weights")

        self.apply_status = QLabel()
        self.apply_status.setAlignment(Qt.AlignCenter)
        self.apply_status.setMinimumHeight(28)
        root.addWidget(self.apply_status)

        button_row = QHBoxLayout()
        self.defaults_button = QPushButton("Reset to Default")
        self.apply_button = QPushButton("Apply")
        self.apply_button.setObjectName("primaryButton")
        close_button = QPushButton("Close")
        button_row.addWidget(self.defaults_button)
        button_row.addStretch()
        button_row.addWidget(self.apply_button)
        button_row.addWidget(close_button)
        root.addLayout(button_row)
        self.defaults_button.clicked.connect(self.reset_defaults)
        self.apply_button.clicked.connect(self.apply_settings)
        close_button.clicked.connect(self.close)
        self.resolution_combo.currentIndexChanged.connect(
            self.clear_apply_status
        )
        self.bit_depth_combo.currentIndexChanged.connect(
            self.clear_apply_status
        )
        self.output_root_row.edit.textChanged.connect(
            self.clear_apply_status
        )
        self.camera_a_mirror_check.toggled.connect(self.clear_apply_status)
        self.camera_b_mirror_check.toggled.connect(self.clear_apply_status)
        self.log_font_spin.valueChanged.connect(self.clear_apply_status)
        self.timeline_metric_combo.currentIndexChanged.connect(
            self.clear_apply_status
        )
        self.timeline_history_spin.valueChanged.connect(
            self.clear_apply_status
        )
        self.set_settings(settings)

    def _create_weight_spins(self, form, labels):
        spins = {}
        for metric_name, display_name in labels.items():
            spin = FocusWheelSpinBox()
            spin.setRange(0, 100)
            spin.setSuffix(" %")
            spin.valueChanged.connect(self.update_weight_validation)
            form.addRow(display_name, spin)
            spins[metric_name] = spin
        return spins

    def set_settings(self, settings):
        settings = validated_toolkit_settings(settings)
        raw_input = settings["raw_input"]
        selected_resolution = tuple(raw_input["resolution"])
        resolution_index = next(
            (
                index
                for index in range(self.resolution_combo.count())
                if tuple(self.resolution_combo.itemData(index))
                == selected_resolution
            ),
            -1,
        )
        self.resolution_combo.setCurrentIndex(max(resolution_index, 0))
        bit_depth_index = self.bit_depth_combo.findData(
            raw_input["bit_depth"]
        )
        self.bit_depth_combo.setCurrentIndex(max(bit_depth_index, 0))
        self.output_root_row.setText(raw_input["output_root"])
        camera_mirror = settings["camera_mirror"]
        self.camera_a_mirror_check.setChecked(camera_mirror["camera_a"])
        self.camera_b_mirror_check.setChecked(camera_mirror["camera_b"])
        self.log_font_spin.setValue(settings["log_font_size"])
        timeline = settings["timeline"]
        metric_index = self.timeline_metric_combo.findData(timeline["metric"])
        self.timeline_metric_combo.setCurrentIndex(max(metric_index, 0))
        self.timeline_history_spin.setValue(timeline["history_seconds"])
        weights = settings["quality_weights"]
        for key, spin in self.spatial_weight_spins.items():
            spin.setValue(weights["spatial"][key])
        for key, spin in self.temporal_weight_spins.items():
            spin.setValue(weights["uvc_temporal"][key])
        self.update_weight_validation()

    def current_settings(self):
        return {
            "log_font_size": self.log_font_spin.value(),
            "raw_input": {
                "resolution": list(self.resolution_combo.currentData()),
                "bit_depth": self.bit_depth_combo.currentData(),
                "output_root": self.output_root_row.text(),
            },
            "timeline": {
                "metric": self.timeline_metric_combo.currentData(),
                "history_seconds": self.timeline_history_spin.value(),
            },
            "camera_mirror": {
                "camera_a": self.camera_a_mirror_check.isChecked(),
                "camera_b": self.camera_b_mirror_check.isChecked(),
            },
            "quality_weights": {
                "spatial": {
                    key: spin.value()
                    for key, spin in self.spatial_weight_spins.items()
                },
                "uvc_temporal": {
                    key: spin.value()
                    for key, spin in self.temporal_weight_spins.items()
                },
            },
        }

    def update_weight_validation(self, _value=None):
        self.clear_apply_status()
        spatial_total = sum(
            spin.value() for spin in self.spatial_weight_spins.values()
        )
        temporal_total = sum(
            spin.value() for spin in self.temporal_weight_spins.values()
        )
        valid = spatial_total == 100 and temporal_total == 100
        for label, total in (
            (self.spatial_total_label, spatial_total),
            (self.temporal_total_label, temporal_total),
        ):
            color = "#8ed0a2" if total == 100 else "#f08080"
            label.setText(f"{total}%")
            label.setStyleSheet(
                f"color:{color}; font-size:16px; font-weight:700;"
            )
        self.apply_button.setEnabled(valid)

    def reset_defaults(self):
        self.set_settings(default_toolkit_settings())

    def clear_apply_status(self, _value=None):
        self.apply_status.clear()

    def show_apply_status(self, message, success=True):
        color = "#8ed0a2" if success else "#f0c078"
        self.apply_status.setText(message)
        self.apply_status.setStyleSheet(
            f"color:{color}; background:#202a25; border:1px solid {color}; "
            "border-radius:5px; padding:5px; font-size:14px; font-weight:700;"
        )

    def apply_settings(self):
        if not self.apply_button.isEnabled():
            return
        self.settings_applied.emit(self.current_settings())


class QualityScoreHelpDialog(QDialog):
    def __init__(self, include_temporal=False, parent=None):
        super().__init__(parent)
        self.setWindowFlags(
            self.windowFlags()
            | Qt.WindowMinimizeButtonHint
            | Qt.WindowMaximizeButtonHint
        )
        source_name = "UVC" if include_temporal else "RAW"
        self.setWindowTitle(f"{source_name} Quality Score Guide")
        self.setModal(False)
        self.resize(820, 720)
        self.setMinimumSize(620, 480)
        self.include_temporal = bool(include_temporal)

        layout = QVBoxLayout(self)
        self.browser = QTextBrowser()
        self.browser.setOpenExternalLinks(False)
        self.reload_guide()
        layout.addWidget(self.browser, 1)

        buttons = QDialogButtonBox(QDialogButtonBox.Close)
        buttons.rejected.connect(self.close)
        layout.addWidget(buttons)

    def reload_guide(self):
        guide_path = bundled_resource_path("quality_score_guide.html")
        if guide_path.is_file():
            self.browser.document().setBaseUrl(
                QUrl.fromLocalFile(f"{guide_path.parent}{os.sep}")
            )
            self.browser.setHtml(guide_path.read_text(encoding="utf-8"))
            self.browser.scrollToAnchor(
                "uvc" if self.include_temporal else "raw"
            )
        else:
            self.browser.setHtml(self._help_html(self.include_temporal))

    @staticmethod
    def _help_html(include_temporal):
        profile_title = "UVC Lightweight Quality" if include_temporal else "Image Quality"
        profile_description = (
            "UVC 即時評分採降採樣與單尺度 Gaussian/Sobel 分析，避免品質監測影響影像輸出；"
            "另以約 10 Hz 取樣計算 Flicker 與 Temporal Stability。"
            if include_temporal
            else
            "RAW 靜態評分使用原始位深與完整解析度。Flat 區先依灰階分成 8 個 level，"
            "再依連通區域拆開，因此天空、牆面等不同灰階或不同位置的平坦區不會直接混成同一組統計。"
        )
        total_formula = (
            "0.70 × Spatial Score + 0.15 × Flicker Control + 0.15 × Temporal Stability"
            if include_temporal
            else
            "0.22 × Detail Clarity + 0.18 × Contrast Quality + 0.18 × Flat Cleanliness + "
            "0.16 × Thermal Uniformity + 0.10 × Clipping Control + 0.16 × Artifact Control"
        )
        temporal_rows = ""
        if include_temporal:
            temporal_rows = """
            <tr><td><b>閃爍控制<br>Flicker Control</b></td>
                <td>去除亮度線性趨勢後的 frame brightness 波動。</td>
                <td><code>100 × exp(-flicker_DN / 2)</code></td>
                <td>越高代表可見亮度閃爍越少。</td></tr>
            <tr><td><b>時間穩定度<br>Temporal Stability</b></td>
                <td>穩定平坦區內，去除全域 offset 後的 frame-to-frame robust residual noise。</td>
                <td><code>100 × exp(-temporal_noise_DN / 3)</code></td>
                <td>越高代表連續影像越穩定。</td></tr>
            """

        raw_extra = "" if include_temporal else """
          <h2>RAW Flat 區域如何拆分</h2>
          <p>1. Bilateral Filter 建立 Base Layer，Detail = Original - Base。</p>
          <p>2. Sobel gradient + 9×9 Local Std 找出 Flat candidate，並避開 Edge 周圍。</p>
          <p>3. Flat pixel 依 Base gray level 的 P1~P99 範圍切成 8 個灰階帶。</p>
          <p>4. 每個灰階帶再用 8-connectivity Connected Components 分成獨立區域。</p>
          <p>5. 每個區域各自算 Robust Noise / Uniformity；Typical 使用面積加權 P50，Worst 使用各有效區域的 P80（不按面積加權）。
             Score 不是只看最大面積區域，因此大面積乾淨天空不會完全掩蓋較小但較差的牆面。</p>
          <p><b>Quality Confidence</b> 只表示本張影像有多少 Flat / Edge 可供分析，不直接加入 Total Score。</p>
        """

        return f"""
        <html><head><style>
          body {{ color:#e6edf2; background:#202328; font-family:'Segoe UI','Microsoft JhengHei'; font-size:14px; line-height:1.55; }}
          h1 {{ color:#86c8f4; font-size:24px; margin-bottom:6px; }}
          h2 {{ color:#c9e2f2; font-size:18px; margin-top:20px; }}
          table {{ border-collapse:collapse; width:100%; margin-top:10px; }}
          th {{ background:#31546c; color:#ffffff; padding:8px; text-align:left; }}
          td {{ border:1px solid #53606c; padding:8px; vertical-align:top; }}
          code {{ color:#ffd38a; font-family:'Cascadia Mono','Consolas'; }}
          .notice {{ background:#2b323a; border-left:4px solid #d6a85f; padding:10px; }}
        </style></head><body>
          <h1>{profile_title}</h1>
          <p class="notice"><b>判讀原則：</b>0~100 分，越高越好。這仍是 No-Reference、同場景相對比較指標，
          不是絕對 PASS/FAIL 規格；沒有 Ground Truth 時，任何單一總分都不能完全代表主觀畫質。</p>
          <p>{profile_description}</p>

          <h2>Total Score</h2><p><code>{total_formula}</code></p>
          {"<p>Spatial Score 使用相同六個空間維度，但 UVC 使用 Lightweight 計算。時序資料至少累積 1.5 秒才顯示完整 Total Score。</p>" if include_temporal else ""}

          <h2>主要分數</h2>
          <table>
            <tr><th>分數</th><th>評估內容</th><th>核心方法</th><th>高分代表</th></tr>
            <tr><td><b>細節清晰度<br>Detail Clarity</b></td>
                <td>真正邊緣強度與 Detail SNR，同時扣除 Halo。</td>
                <td>Sobel Edge + Base/Detail + Robust Flat Noise。RAW：<code>50% Sharpness + 50% Detail-SNR - 8% Halo Penalty</code></td>
                <td>細節強、相對 noise 高，而且不是單純 oversharpen。</td></tr>
            <tr><td><b>對比品質<br>Contrast Quality</b></td>
                <td>Robust P1~P99 動態範圍 + Histogram Entropy。</td>
                <td>RAW：70% Span + 30% Entropy；UVC：75% Span + 25% Entropy。</td>
                <td>灰階利用充分，且不是只靠少數極端值拉大範圍。</td></tr>
            <tr><td><b>平坦區潔淨度<br>Flat Cleanliness</b></td>
                <td>Flat 區的高頻 residual noise。</td>
                <td>RAW 對每個灰階/連通區域分開算 MAD robust sigma；Typical 使用面積加權 P50，Worst 使用各有效區域 P80（不按面積加權）。另外以 3×3 Median residual + 小型 Connected Component 抓 isolated hot/cold impulse。</td>
                <td>不同亮度平坦區都較乾淨，且少量異常 pixel 不會被 robust MAD 完全忽略。</td></tr>
            <tr><td><b>熱影像均勻度<br>Thermal Uniformity</b></td>
                <td>低頻不均、局部 residual，以及 RAW 額外的 row/column banding。</td>
                <td>RAW：50% Typical + 30% Worst + 20% Banding Control；UVC 使用簡化 residual。</td>
                <td>平坦熱區更均勻，條紋/帶狀 fixed-pattern 較少。</td></tr>
            <tr><td><b>裁切控制<br>Clipping Control</b></td>
                <td>靠近黑端/白端的 clipping pixel 比例。</td>
                <td><code>100 × exp(-clipping_percent / 3)</code></td>
                <td>黑白細節較少因 AGC/contrast 拉伸而消失。</td></tr>
            <tr><td><b>假影控制<br>Artifact Control</b></td>
                <td>Edge Halo + 水平/垂直 Seam/Block discontinuity；重複 seam 比單一自然邊界懲罰更重。</td>
                <td>50% Halo Control + 50% Seam Control；RAW Halo 以 edge-normal overshoot/undershoot 判斷，Seam 主要抓重複且近似週期性的 block/grid line。</td>
                <td>較少 halo、ringing、block/grid boundary。</td></tr>
            {temporal_rows}
          </table>
          {raw_extra}
        </body></html>
        """


class UvcVideoWindow(QWidget):
    """Resizable image-only window for UVC views and recording layouts."""

    def __init__(self, parent=None):
        super().__init__(parent, Qt.Window)
        self._pixmap = None
        self.setWindowTitle("UVC Recording View")
        self.resize(1280, 720)
        self.setMinimumSize(480, 320)
        self.setStyleSheet("background:#000000;")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        self.image_label = QLabel()
        self.image_label.setAlignment(Qt.AlignCenter)
        self.image_label.setStyleSheet("background:#000000; border:none;")
        layout.addWidget(self.image_label)

    def set_bgr_frame(self, frame):
        array = np.asarray(frame)
        if array.ndim != 3 or array.shape[2] < 3:
            return
        bgr = np.ascontiguousarray(array[:, :, :3].astype(np.uint8))
        bgr_format = getattr(QImage, "Format_BGR888", None)
        if bgr_format is not None:
            image = QImage(
                bgr.data,
                bgr.shape[1],
                bgr.shape[0],
                bgr.strides[0],
                bgr_format,
            ).copy()
        else:
            rgb = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)
            image = QImage(
                rgb.data,
                rgb.shape[1],
                rgb.shape[0],
                rgb.strides[0],
                QImage.Format_RGB888,
            ).copy()
        self._pixmap = QPixmap.fromImage(image)
        self._rescale()

    def clear_frame(self):
        self._pixmap = None
        self.image_label.clear()

    def _rescale(self):
        if self._pixmap is not None and not self._pixmap.isNull():
            self.image_label.setPixmap(
                self._pixmap.scaled(
                    self.image_label.size(),
                    Qt.KeepAspectRatio,
                    Qt.SmoothTransformation,
                )
            )

    def resizeEvent(self, event):
        self._rescale()
        super().resizeEvent(event)


class PathRow(QWidget):
    path_selected = pyqtSignal(str)

    def __init__(self, mode="file", filter_text="All Files (*)", dialog_title=None):
        super().__init__()
        self.mode = mode
        self.filter_text = filter_text
        self.dialog_title = dialog_title or ("Select Folder" if mode == "folder" else "Select File")
        lay = QHBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        self.edit = QLineEdit()
        self.edit.setPlaceholderText(self.dialog_title)
        self.btn = QPushButton("Browse")
        self.btn.setProperty("buttonRole", "browse")
        self.btn.setToolTip(self.dialog_title)
        self.btn.clicked.connect(self.browse)
        lay.addWidget(self.edit, 1)
        lay.addWidget(self.btn)

    def browse(self):
        if self.mode == "folder":
            p = QFileDialog.getExistingDirectory(self, self.dialog_title)
        else:
            p, _ = QFileDialog.getOpenFileName(self, self.dialog_title, "", self.filter_text)
        if p:
            self.edit.setText(p)
            self.path_selected.emit(p)

    def text(self):
        return self.edit.text().strip()

    def setText(self, text):
        self.edit.setText(text)


class AlgorithmPipelineEditor(QWidget):
    pipeline_changed = pyqtSignal()

    def __init__(
        self,
        parent=None,
        allow_step_algorithm_change=True,
        fieldscale_temporal=False,
    ):
        super().__init__(parent)
        self._stages = []
        self._updating = False
        self._allow_step_algorithm_change = bool(
            allow_step_algorithm_change
        )
        self._fieldscale_temporal = bool(fieldscale_temporal)

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(7)
        title = QLabel("Algorithm Pipeline (top to bottom)")
        title.setWordWrap(True)
        title.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        title.setStyleSheet("font-size:15px; font-weight:700; color:#c9e2f2;")
        root.addWidget(title)

        self.add_algorithm_combo = FocusWheelComboBox()
        for display_name, algorithm_key in PIPELINE_ALGORITHM_OPTIONS:
            self.add_algorithm_combo.addItem(display_name, algorithm_key)
        self.add_algorithm_combo.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.add_button = QPushButton("Add Step")
        self.add_button.setObjectName("primaryButton")
        root.addWidget(self.add_algorithm_combo)
        root.addWidget(self.add_button)

        self.stage_list = QListWidget()
        self.stage_list.setMinimumHeight(132)
        self.stage_list.setMaximumHeight(190)
        self.stage_list.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.stage_list.setTextElideMode(Qt.ElideRight)
        self.stage_list.setToolTip(
            "Checked steps run from top to bottom. Duplicate algorithms are allowed."
        )
        root.addWidget(self.stage_list)

        order_grid = QGridLayout()
        order_grid.setHorizontalSpacing(7)
        order_grid.setVerticalSpacing(7)
        self.move_up_button = QPushButton("Move Up")
        self.move_down_button = QPushButton("Move Down")
        self.remove_button = QPushButton("Remove")
        self.reset_all_parameters_button = QPushButton("Reset All Parameters")
        self.reset_all_parameters_button.setToolTip(
            "Reset parameters for all listed steps, including unchecked steps."
        )
        self.reset_pipeline_button = QPushButton("Reset Pipeline")
        self.reset_pipeline_button.setObjectName("resetButton")
        order_grid.addWidget(self.move_up_button, 0, 0)
        order_grid.addWidget(self.move_down_button, 0, 1)
        order_grid.addWidget(self.remove_button, 1, 0)
        order_grid.addWidget(self.reset_all_parameters_button, 1, 1)
        order_grid.addWidget(self.reset_pipeline_button, 2, 0, 1, 2)
        for button in (
            self.move_up_button,
            self.move_down_button,
            self.remove_button,
            self.reset_all_parameters_button,
            self.reset_pipeline_button,
        ):
            button.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
        root.addLayout(order_grid)

        self.parameter_group = CollapsibleGroupBox("Step Parameters")
        self.parameter_group.setObjectName("pipelineParameterCard")
        self.parameter_form = QVBoxLayout(self.parameter_group)
        self.parameter_form.setContentsMargins(10, 10, 10, 10)
        self.parameter_form.setSpacing(8)
        self.reset_step_button = QPushButton("Reset Step")
        self.reset_step_button.setToolTip(
            "Reset parameters for the currently selected step."
        )
        self.parameter_rows = {}
        self.stage_algorithm_combo = FocusWheelComboBox()
        for display_name, algorithm_key in PIPELINE_ALGORITHM_OPTIONS:
            self.stage_algorithm_combo.addItem(display_name, algorithm_key)
        self.stage_algorithm_combo.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Fixed
        )
        self.stage_description = QLabel()
        self.stage_description.setWordWrap(True)
        self.stage_description.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        self.stage_description.setStyleSheet(
            "color:#c8d0d9; background:#292d33; border-left:3px solid #79afd3; "
            "border-radius:5px; padding:7px;"
        )
        self.stage_low_percentile = FocusWheelDoubleSpinBox()
        self.stage_low_percentile.setRange(0.0, 99.5)
        self.stage_low_percentile.setDecimals(1)
        self.stage_low_percentile.setSingleStep(0.5)
        self.stage_high_percentile = FocusWheelDoubleSpinBox()
        self.stage_high_percentile.setRange(0.5, 100.0)
        self.stage_high_percentile.setDecimals(1)
        self.stage_high_percentile.setSingleStep(0.5)
        self.stage_clahe_clip = FocusWheelDoubleSpinBox()
        self.stage_clahe_clip.setRange(0.1, 100.0)
        self.stage_clahe_clip.setDecimals(1)
        self.stage_clahe_clip.setSingleStep(0.1)
        self.stage_clahe_grid = FocusWheelSpinBox()
        self.stage_clahe_grid.setRange(1, 64)
        self.stage_gamma = FocusWheelDoubleSpinBox()
        self.stage_gamma.setRange(0.1, 3.0)
        self.stage_gamma.setDecimals(1)
        self.stage_gamma.setSingleStep(0.1)
        self.stage_bilateral_strength = FocusWheelSpinBox()
        self.stage_bilateral_strength.setRange(1, 7)
        self.stage_unsharp_amount = FocusWheelDoubleSpinBox()
        self.stage_unsharp_amount.setRange(0.1, 3.0)
        self.stage_unsharp_amount.setDecimals(1)
        self.stage_unsharp_amount.setSingleStep(0.1)
        self.stage_dde_sigma = FocusWheelDoubleSpinBox()
        self.stage_dde_sigma.setRange(0.0, 5.0)
        self.stage_dde_sigma.setDecimals(1)
        self.stage_dde_sigma.setSingleStep(0.1)
        self.stage_dde_sigma.setValue(1.0)
        self.stage_dde_sigma.setToolTip(
            "Noise-reduction strength; OpenCV NLM uses h = 10 x Sigma."
        )
        self.stage_dde_alpha = FocusWheelDoubleSpinBox()
        self.stage_dde_alpha.setRange(0.0, 5.0)
        self.stage_dde_alpha.setDecimals(1)
        self.stage_dde_alpha.setSingleStep(0.1)
        self.stage_dde_alpha.setValue(1.5)
        self.stage_dde_alpha.setToolTip(
            "Gain applied to positive multi-scale detail layers."
        )
        self.stage_dde_beta = FocusWheelDoubleSpinBox()
        self.stage_dde_beta.setRange(0.0, 2.0)
        self.stage_dde_beta.setDecimals(1)
        self.stage_dde_beta.setSingleStep(0.1)
        self.stage_dde_beta.setValue(0.5)
        self.stage_dde_beta.setToolTip(
            "Amount of the source/base image retained before detail enhancement."
        )
        self.stage_fieldscale_grid_count = FocusWheelComboBox()
        for grid_count in (2, 4, 8, 16, 32):
            self.stage_fieldscale_grid_count.addItem(
                f"{grid_count} x {grid_count}", grid_count
            )
        self.stage_fieldscale_grid_count.setCurrentIndex(
            self.stage_fieldscale_grid_count.findData(8)
        )
        self.stage_fieldscale_grid_count.setToolTip(
            "Number of local Min/Max regions on each axis; firmware default is 8x8."
        )
        self.stage_fieldscale_les_distance = FocusWheelSpinBox()
        self.stage_fieldscale_les_distance.setRange(1, 7)
        self.stage_fieldscale_les_distance.setValue(2)
        self.stage_fieldscale_les_distance.setToolTip(
            "Neighborhood radius on the 8x8 Min/Max grid; firmware default is 2."
        )
        self.stage_fieldscale_les_threshold = FocusWheelSpinBox()
        self.stage_fieldscale_les_threshold.setRange(0, FIELDSCALE_RAW14_MAX)
        self.stage_fieldscale_les_threshold.setValue(400)
        self.stage_fieldscale_les_threshold.setSuffix(" RAW14 DN")
        self.stage_fieldscale_les_threshold.setToolTip(
            "Local-extrema suppression threshold in RAW14 DN."
        )
        self.stage_fieldscale_message_distance = FocusWheelSpinBox()
        self.stage_fieldscale_message_distance.setRange(1, 7)
        self.stage_fieldscale_message_distance.setValue(1)
        self.stage_fieldscale_message_distance.setToolTip(
            "Neighborhood radius for each message-passing iteration."
        )
        self.stage_fieldscale_message_iterations = FocusWheelSpinBox()
        self.stage_fieldscale_message_iterations.setRange(0, 32)
        self.stage_fieldscale_message_iterations.setValue(7)
        self.stage_fieldscale_message_iterations.setToolTip(
            "Number of message-passing iterations; 0 disables message passing."
        )
        self.stage_fieldscale_temporal_smoothing = FocusWheelSpinBox()
        self.stage_fieldscale_temporal_smoothing.setRange(0, 99)
        self.stage_fieldscale_temporal_smoothing.setValue(90)
        self.stage_fieldscale_temporal_smoothing.setSuffix(" % previous frame")
        self.stage_fieldscale_temporal_smoothing.setToolTip(
            "UVC only: previous-frame weight. 0% disables temporal smoothing."
        )
        if self._allow_step_algorithm_change:
            self._add_parameter_row("Algorithm", self.stage_algorithm_combo)
        self.parameter_form.addWidget(self.stage_description)
        self._add_parameter_row(
            "Grid 數量", self.stage_fieldscale_grid_count
        )
        self._add_parameter_row(
            "LES Distance", self.stage_fieldscale_les_distance
        )
        self._add_parameter_row(
            "LES Threshold", self.stage_fieldscale_les_threshold
        )
        self._add_parameter_row(
            "Message Passing Distance", self.stage_fieldscale_message_distance
        )
        self._add_parameter_row(
            "Message Passing 次數 N",
            self.stage_fieldscale_message_iterations,
        )
        self._add_parameter_row(
            "Temporal Smoothing",
            self.stage_fieldscale_temporal_smoothing,
        )
        self._add_parameter_row("Low Percentile", self.stage_low_percentile)
        self._add_parameter_row("High Percentile", self.stage_high_percentile)
        self._add_parameter_row("CLAHE Clip Limit", self.stage_clahe_clip)
        self._add_parameter_row("CLAHE Grid Size", self.stage_clahe_grid)
        self._add_parameter_row("Gamma", self.stage_gamma)
        self._add_parameter_row("Sigma", self.stage_dde_sigma)
        self._add_parameter_row("Alpha", self.stage_dde_alpha)
        self._add_parameter_row("Beta", self.stage_dde_beta)
        self._add_parameter_row(
            "Bilateral Strength", self.stage_bilateral_strength
        )
        self._add_parameter_row("Unsharp Amount", self.stage_unsharp_amount)
        self.parameter_form.addWidget(self.reset_step_button)
        self.parameter_form.addStretch()
        root.addWidget(self.parameter_group)

        self.add_button.clicked.connect(self.add_stage)
        self.move_up_button.clicked.connect(lambda: self.move_selected(-1))
        self.move_down_button.clicked.connect(lambda: self.move_selected(1))
        self.remove_button.clicked.connect(self.remove_selected)
        self.reset_step_button.clicked.connect(self.reset_selected_step)
        self.reset_all_parameters_button.clicked.connect(self.reset_all_parameters)
        self.reset_pipeline_button.clicked.connect(self.reset_pipeline)
        self.stage_list.currentRowChanged.connect(self.load_selected_step)
        self.stage_list.itemChanged.connect(self.stage_check_changed)
        if self._allow_step_algorithm_change:
            self.stage_algorithm_combo.currentIndexChanged.connect(
                self.selected_algorithm_changed
            )
        for widget in (
            self.stage_low_percentile,
            self.stage_high_percentile,
            self.stage_clahe_clip,
            self.stage_clahe_grid,
            self.stage_gamma,
            self.stage_bilateral_strength,
            self.stage_unsharp_amount,
            self.stage_dde_sigma,
            self.stage_dde_alpha,
            self.stage_dde_beta,
            self.stage_fieldscale_les_distance,
            self.stage_fieldscale_les_threshold,
            self.stage_fieldscale_message_distance,
            self.stage_fieldscale_message_iterations,
            self.stage_fieldscale_temporal_smoothing,
        ):
            widget.valueChanged.connect(self.selected_parameter_changed)
        self.stage_fieldscale_grid_count.currentIndexChanged.connect(
            self.selected_fieldscale_grid_changed
        )
        self.set_stages([])

    def _add_parameter_row(self, label_text, widget):
        row = QWidget()
        row_layout = QVBoxLayout(row)
        row_layout.setContentsMargins(0, 0, 0, 0)
        row_layout.setSpacing(4)
        row_layout.addWidget(QLabel(label_text))
        row_layout.addWidget(widget)
        self.parameter_form.addWidget(row)
        self.parameter_rows[widget] = row

    def stages(self):
        stages = [stage.clone() for stage in self._stages]
        for stage in stages:
            if stage.algorithm_key != "fieldscale":
                continue
            if self._fieldscale_temporal:
                stage.parameters.setdefault("temporal_smoothing_percent", 90)
            else:
                stage.parameters.pop("temporal_smoothing_percent", None)
        return stages

    def enabled_stages(self):
        return [stage for stage in self.stages() if stage.enabled]

    def configuration(self):
        return [stage.to_dict() for stage in self.stages()]

    def summary(self):
        return pipeline_summary_text(self._stages)

    def set_stages(self, stages, selected_row=0):
        self._stages = [
            stage.clone()
            if isinstance(stage, AlgorithmStage)
            else AlgorithmStage(
                stage["algorithm"],
                stage.get("enabled", True),
                stage.get("parameters", {}),
            )
            for stage in stages
        ]
        self.refresh_stage_list(selected_row)

    def parameter_summary(self, stage):
        params = stage.parameters
        if stage.algorithm_key == "auto_contrast":
            return f"P{params['low_percentile']:.1f}/P{params['high_percentile']:.1f}"
        if stage.algorithm_key == "fieldscale":
            summary = (
                f"Grid={params.get('grid_count', 8)}x"
                f"{params.get('grid_count', 8)}; "
                f"LES d={params.get('les_distance', 2)}, "
                f"th={params.get('les_threshold_raw14', 400)}; "
                f"MP d={params.get('message_passing_distance', 1)}, "
                f"N={params.get('message_passing_iterations', 7)}"
            )
            if self._fieldscale_temporal:
                summary += (
                    "; Temporal="
                    f"{params.get('temporal_smoothing_percent', 90)}%"
                )
            return summary
        if stage.algorithm_key == "clahe":
            return f"clip={params['clahe_clip']:.1f}, grid={params['clahe_grid']}"
        if stage.algorithm_key == "gamma":
            return f"gamma={params['gamma']:.1f}"
        if stage.algorithm_key == "dde":
            return (
                f"sigma={params.get('dde_sigma', 1.0):.1f}, "
                f"alpha={params.get('dde_alpha', 1.5):.1f}, "
                f"beta={params.get('dde_beta', 0.5):.1f}"
            )
        if stage.algorithm_key == "bilateral":
            return f"strength={params['bilateral_strength']}"
        if stage.algorithm_key == "unsharp":
            return f"amount={params['unsharp_amount']:.1f}"
        return "no parameters"

    def refresh_stage_list(self, selected_row=None):
        if selected_row is None:
            selected_row = self.stage_list.currentRow()
        self._updating = True
        self.stage_list.blockSignals(True)
        self.stage_list.clear()
        for index, stage in enumerate(self._stages, start=1):
            text = (
                f"{index}. {algorithm_display_name(stage.algorithm_key)}"
                f"    [{self.parameter_summary(stage)}]"
            )
            item = QListWidgetItem(text)
            item.setToolTip(text)
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Checked if stage.enabled else Qt.Unchecked)
            self.stage_list.addItem(item)
        self.stage_list.blockSignals(False)
        self._updating = False
        if self._stages:
            self.stage_list.setCurrentRow(
                min(max(selected_row, 0), len(self._stages) - 1)
            )
        else:
            self.load_selected_step(-1)
        self.update_button_states()

    def add_stage(self, _checked=False):
        self._stages.append(AlgorithmStage(self.add_algorithm_combo.currentData()))
        self.refresh_stage_list(len(self._stages) - 1)
        self.pipeline_changed.emit()

    def remove_selected(self, _checked=False):
        row = self.stage_list.currentRow()
        if row < 0:
            return
        del self._stages[row]
        self.refresh_stage_list(min(row, len(self._stages) - 1))
        self.pipeline_changed.emit()

    def move_selected(self, offset):
        row = self.stage_list.currentRow()
        target = row + offset
        if row < 0 or target < 0 or target >= len(self._stages):
            return
        self._stages[row], self._stages[target] = self._stages[target], self._stages[row]
        self.refresh_stage_list(target)
        self.pipeline_changed.emit()

    def reset_selected_step(self, _checked=False):
        row = self.stage_list.currentRow()
        if row < 0:
            return
        self._stages[row].parameters = default_algorithm_parameters(
            self._stages[row].algorithm_key
        )
        self.refresh_stage_list(row)
        self.pipeline_changed.emit()

    def reset_all_parameters(self, _checked=False):
        if not self._stages:
            return
        for stage in self._stages:
            stage.parameters = default_algorithm_parameters(stage.algorithm_key)
        self.refresh_stage_list()
        self.pipeline_changed.emit()

    def reset_pipeline(self, _checked=False):
        self.set_stages([])
        self.pipeline_changed.emit()

    def stage_check_changed(self, item):
        if self._updating:
            return
        row = self.stage_list.row(item)
        if 0 <= row < len(self._stages):
            self._stages[row].enabled = item.checkState() == Qt.Checked
            self.pipeline_changed.emit()

    def selected_algorithm_changed(self, _combo_index):
        if self._updating:
            return
        row = self.stage_list.currentRow()
        if row < 0:
            return
        self._stages[row] = AlgorithmStage(
            self.stage_algorithm_combo.currentData(),
            enabled=self._stages[row].enabled,
        )
        self.refresh_stage_list(row)
        self.pipeline_changed.emit()

    def selected_parameter_changed(self, _value=None):
        if self._updating:
            return
        row = self.stage_list.currentRow()
        if row < 0:
            return
        self.update_percentile_limits()
        stage = self._stages[row]
        if stage.algorithm_key == "auto_contrast":
            stage.parameters["low_percentile"] = self.stage_low_percentile.value()
            stage.parameters["high_percentile"] = self.stage_high_percentile.value()
        elif stage.algorithm_key == "fieldscale":
            stage.parameters["grid_count"] = (
                self.stage_fieldscale_grid_count.currentData()
            )
            stage.parameters["les_distance"] = (
                self.stage_fieldscale_les_distance.value()
            )
            stage.parameters["les_threshold_raw14"] = (
                self.stage_fieldscale_les_threshold.value()
            )
            stage.parameters["message_passing_distance"] = (
                self.stage_fieldscale_message_distance.value()
            )
            stage.parameters["message_passing_iterations"] = (
                self.stage_fieldscale_message_iterations.value()
            )
            if self._fieldscale_temporal:
                stage.parameters["temporal_smoothing_percent"] = (
                    self.stage_fieldscale_temporal_smoothing.value()
                )
            else:
                stage.parameters.pop("temporal_smoothing_percent", None)
        elif stage.algorithm_key == "clahe":
            stage.parameters["clahe_clip"] = self.stage_clahe_clip.value()
            stage.parameters["clahe_grid"] = self.stage_clahe_grid.value()
        elif stage.algorithm_key == "gamma":
            stage.parameters["gamma"] = self.stage_gamma.value()
        elif stage.algorithm_key == "dde":
            stage.parameters["dde_sigma"] = self.stage_dde_sigma.value()
            stage.parameters["dde_alpha"] = self.stage_dde_alpha.value()
            stage.parameters["dde_beta"] = self.stage_dde_beta.value()
        elif stage.algorithm_key == "bilateral":
            stage.parameters["bilateral_strength"] = self.stage_bilateral_strength.value()
        elif stage.algorithm_key == "unsharp":
            stage.parameters["unsharp_amount"] = self.stage_unsharp_amount.value()
        self.refresh_stage_list(row)
        self.pipeline_changed.emit()

    def selected_fieldscale_grid_changed(self, _index=None):
        if self._updating:
            return
        self._updating = True
        self.update_fieldscale_distance_limits()
        self._updating = False
        self.selected_parameter_changed()

    def update_fieldscale_distance_limits(self):
        grid_count = self.stage_fieldscale_grid_count.currentData()
        maximum_distance = max(1, int(grid_count or 8) - 1)
        self.stage_fieldscale_les_distance.setMaximum(maximum_distance)
        self.stage_fieldscale_message_distance.setMaximum(maximum_distance)

    def set_parameter_visible(self, widget, visible):
        row = self.parameter_rows.get(widget)
        if row is not None:
            row.setVisible(visible)
        else:
            widget.setVisible(visible)
        self.parameter_group.updateGeometry()
        self.updateGeometry()

    def load_selected_step(self, row):
        valid = 0 <= row < len(self._stages)
        self.reset_step_button.setVisible(valid)
        self.set_parameter_visible(
            self.stage_algorithm_combo, valid and self._allow_step_algorithm_change
        )
        self._updating = True
        if valid:
            stage = self._stages[row]
            index = self.stage_algorithm_combo.findData(stage.algorithm_key)
            self.stage_algorithm_combo.setCurrentIndex(max(index, 0))
            params = stage.parameters
            self.stage_low_percentile.setValue(params.get("low_percentile", 1.0))
            self.stage_high_percentile.setValue(params.get("high_percentile", 99.0))
            self.stage_clahe_clip.setValue(params.get("clahe_clip", 2.0))
            self.stage_clahe_grid.setValue(params.get("clahe_grid", 8))
            self.stage_gamma.setValue(params.get("gamma", 0.7))
            self.stage_bilateral_strength.setValue(params.get("bilateral_strength", 2))
            self.stage_unsharp_amount.setValue(params.get("unsharp_amount", 1.0))
            self.stage_dde_sigma.setValue(params.get("dde_sigma", 1.0))
            self.stage_dde_alpha.setValue(params.get("dde_alpha", 1.5))
            self.stage_dde_beta.setValue(params.get("dde_beta", 0.5))
            grid_index = self.stage_fieldscale_grid_count.findData(
                params.get("grid_count", 8)
            )
            self.stage_fieldscale_grid_count.setCurrentIndex(
                grid_index if grid_index >= 0 else 2
            )
            self.update_fieldscale_distance_limits()
            self.stage_fieldscale_les_distance.setValue(
                params.get("les_distance", 2)
            )
            self.stage_fieldscale_les_threshold.setValue(
                params.get("les_threshold_raw14", 400)
            )
            self.stage_fieldscale_message_distance.setValue(
                params.get("message_passing_distance", 1)
            )
            self.stage_fieldscale_message_iterations.setValue(
                params.get("message_passing_iterations", 7)
            )
            self.stage_fieldscale_temporal_smoothing.setValue(
                params.get("temporal_smoothing_percent", 90)
            )
            self.stage_description.setText(
                IMAGE_ALGORITHM_DESCRIPTIONS.get(stage.algorithm_key, "")
            )
            is_fieldscale = stage.algorithm_key == "fieldscale"
            for widget in (
                self.stage_fieldscale_grid_count,
                self.stage_fieldscale_les_distance,
                self.stage_fieldscale_les_threshold,
                self.stage_fieldscale_message_distance,
                self.stage_fieldscale_message_iterations,
            ):
                self.set_parameter_visible(widget, is_fieldscale)
            self.set_parameter_visible(
                self.stage_fieldscale_temporal_smoothing,
                is_fieldscale and self._fieldscale_temporal,
            )
            self.set_parameter_visible(
                self.stage_low_percentile, stage.algorithm_key == "auto_contrast"
            )
            self.set_parameter_visible(
                self.stage_high_percentile, stage.algorithm_key == "auto_contrast"
            )
            self.set_parameter_visible(
                self.stage_clahe_clip, stage.algorithm_key == "clahe"
            )
            self.set_parameter_visible(
                self.stage_clahe_grid, stage.algorithm_key == "clahe"
            )
            self.set_parameter_visible(self.stage_gamma, stage.algorithm_key == "gamma")
            for widget in (
                self.stage_dde_sigma,
                self.stage_dde_alpha,
                self.stage_dde_beta,
            ):
                self.set_parameter_visible(
                    widget, stage.algorithm_key == "dde"
                )
            self.set_parameter_visible(
                self.stage_bilateral_strength, stage.algorithm_key == "bilateral"
            )
            self.set_parameter_visible(
                self.stage_unsharp_amount, stage.algorithm_key == "unsharp"
            )
            self.update_percentile_limits()
        else:
            for parameter_row in self.parameter_rows.values():
                parameter_row.hide()
            self.stage_description.setText(
                "Waiting for an algorithm. Select an algorithm above and click Add Step."
            )
        self.parameter_group.updateGeometry()
        self.updateGeometry()
        self._updating = False
        self.update_button_states()

    def update_percentile_limits(self):
        self.stage_high_percentile.setMinimum(
            min(100.0, self.stage_low_percentile.value() + 0.5)
        )
        self.stage_low_percentile.setMaximum(
            max(0.0, self.stage_high_percentile.value() - 0.5)
        )

    def update_button_states(self):
        row = self.stage_list.currentRow()
        valid = 0 <= row < len(self._stages)
        self.move_up_button.setEnabled(valid and row > 0)
        self.move_down_button.setEnabled(valid and row < len(self._stages) - 1)
        self.remove_button.setEnabled(valid)
        self.reset_step_button.setEnabled(valid)
        self.reset_all_parameters_button.setEnabled(bool(self._stages))


class LatestFrameBuffer:
    """Thread-safe single-frame mailbox; newer frames replace pending old ones."""

    def __init__(self):
        self._lock = threading.Lock()
        self._frame = None
        self._frame_time = 0.0
        self._sequence = 0

    def submit(self, frame):
        if frame is None or not getattr(frame, "size", 0):
            return
        with self._lock:
            self._frame = frame.copy()
            self._frame_time = time.perf_counter()
            self._sequence += 1

    def take_latest(self, previous_sequence):
        with self._lock:
            if self._frame is None or self._sequence == previous_sequence:
                return None, previous_sequence, 0.0
            return self._frame.copy(), self._sequence, self._frame_time

    def clear(self):
        with self._lock:
            self._frame = None
            self._frame_time = 0.0
            self._sequence = 0


class UvcFpsCsvLogger:
    """Count received frames; write interval rates independently of capture/UI stalls."""

    def __init__(self, path, status_callback):
        self.path = Path(path)
        self._status = status_callback
        self._lock = threading.Lock()
        self._stop = threading.Event()
        self._count = 0
        self._last_count = 0
        self._started = self._last_sample = time.perf_counter()
        self._file = self.path.open("w", newline="", encoding="utf-8")
        self._writer = csv.writer(self._file)
        self._writer.writerow(["timestamp", "elapsed_s", "interval_s", "received_frames", "fps", "total_frames"])
        self._file.flush()
        self._thread = threading.Thread(target=self._run, name="UVC FPS CSV", daemon=True)
        self._thread.start()

    def frame_received(self):
        with self._lock:
            if not self._stop.is_set():
                self._count += 1

    def _sample(self):
        with self._lock:
            now = time.perf_counter()
            count = self._count
        interval = now - self._last_sample
        if interval <= 0:
            return
        received = count - self._last_count
        self._writer.writerow([
            datetime.now().astimezone().isoformat(timespec="milliseconds"),
            f"{now - self._started:.6f}", f"{interval:.6f}", received,
            f"{received / interval:.6f}", count,
        ])
        self._last_sample = now
        self._last_count = count

    def _run(self):
        try:
            samples = 0
            while not self._stop.wait(1.0):
                self._sample()
                samples += 1
                if samples % 5 == 0:
                    self._file.flush()
            self._sample()  # Include the final partial interval.
            self._file.flush()
        except Exception as exc:
            self._status(f"FPS CSV write failed: {exc}; file: {self.path}")
        finally:
            self._file.close()
            self._status(f"FPS CSV closed: {self.path}")

    def close(self):
        self._stop.set()
        self._thread.join(timeout=2.0)


class UvcCaptureThread(QThread):
    frame_ready = pyqtSignal(object)
    frame_seen = pyqtSignal(float)
    camera_opened = pyqtSignal(int, int, int, float)
    status_changed = pyqtSignal(str)
    stream_interrupted = pyqtSignal(str)

    def __init__(
        self,
        camera_index,
        width,
        height,
        fps,
        reconnect_seconds,
        parent=None,
        frame_consumer=None,
        capture_mode=DEFAULT_UVC_CAPTURE_MODE,
        camera_device_id=None,
        camera_name=None,
    ):
        super().__init__(parent)
        self.camera_index = camera_index
        self.width = width
        self.height = height
        self.fps = fps
        self.reconnect_seconds = reconnect_seconds
        self.frame_consumer = frame_consumer
        self.capture_mode = capture_mode
        self.camera_device_id = camera_device_id
        self.camera_name = camera_name or f"Camera {camera_index}"
        self.selected_capture_mode = None
        self.fps_logger = None
        self._camera_info_lock = threading.Lock()
        self._stop_event = threading.Event()
        self._reconnect_event = threading.Event()
        self._capture_lock = threading.Lock()
        self._capture = None

    def stop(self):
        self._stop_event.set()
        if self.fps_logger is not None:
            self.fps_logger.close()

    def force_reconnect(self):
        self._reconnect_event.set()

    def update_camera_device(self, camera_index, camera_name=None):
        with self._camera_info_lock:
            changed = self.camera_index != camera_index
            self.camera_index = camera_index
            if camera_name:
                self.camera_name = camera_name
        if changed:
            self._reconnect_event.set()

    def _camera_device_snapshot(self):
        with self._camera_info_lock:
            return self.camera_index, self.camera_name

    def _set_capture(self, capture):
        with self._capture_lock:
            self._capture = capture

    def _clear_capture(self, capture):
        with self._capture_lock:
            if self._capture is capture:
                self._capture = None
        try:
            capture.release()
        except Exception:
            pass

    def _open_capture(self, capture_mode=None, camera_index=None):
        capture_mode = capture_mode or self.capture_mode
        if camera_index is None:
            camera_index, _camera_name = self._camera_device_snapshot()
        if camera_index is None:
            return None
        if capture_mode == "generic_auto":
            if sys.platform.startswith("win"):
                backends = [cv2.CAP_DSHOW]
                if hasattr(cv2, "CAP_MSMF"):
                    backends.append(cv2.CAP_MSMF)
                for backend in backends:
                    try:
                        capture = cv2.VideoCapture(camera_index, backend)
                    except (TypeError, cv2.error):
                        continue
                    if capture.isOpened():
                        return capture
                    capture.release()
            return cv2.VideoCapture(camera_index)

        if sys.platform.startswith("win"):
            mjpg_fourcc = cv2.VideoWriter_fourcc(*UVC_CAPTURE_FOURCC)
            open_params = [
                cv2.CAP_PROP_FRAME_WIDTH,
                self.width,
                cv2.CAP_PROP_FRAME_HEIGHT,
                self.height,
                cv2.CAP_PROP_FOURCC,
                mjpg_fourcc,
            ]
            try:
                # DirectShow applies these values before starting the device.
                # Setting CAP_PROP_FPS afterward can reopen it without FOURCC
                # and make the Novatek host request format=0 instead of MJPEG.
                capture = cv2.VideoCapture(
                    camera_index,
                    cv2.CAP_DSHOW,
                    open_params,
                )
            except (TypeError, cv2.error):
                capture = cv2.VideoCapture(camera_index, cv2.CAP_DSHOW)

            if capture.isOpened():
                return capture

            capture.release()
            capture = cv2.VideoCapture(camera_index, cv2.CAP_DSHOW)
            if capture.isOpened():
                capture.set(cv2.CAP_PROP_FOURCC, mjpg_fourcc)
                capture.set(cv2.CAP_PROP_FRAME_WIDTH, self.width)
                capture.set(cv2.CAP_PROP_FRAME_HEIGHT, self.height)
                capture.set(cv2.CAP_PROP_FOURCC, mjpg_fourcc)
            return capture

        capture = cv2.VideoCapture(camera_index)
        if capture.isOpened():
            capture.set(
                cv2.CAP_PROP_FOURCC,
                cv2.VideoWriter_fourcc(*UVC_CAPTURE_FOURCC),
            )
            capture.set(cv2.CAP_PROP_FRAME_WIDTH, self.width)
            capture.set(cv2.CAP_PROP_FRAME_HEIGHT, self.height)
            capture.set(cv2.CAP_PROP_FPS, self.fps)
        return capture

    @staticmethod
    def _capture_mode_name(capture_mode):
        for mode_name, mode_key in UVC_CAPTURE_MODE_OPTIONS:
            if mode_key == capture_mode:
                return mode_name
        return capture_mode

    @staticmethod
    def _alternate_auto_mode(capture_mode):
        if capture_mode == "novatek_mjpg":
            return "generic_auto"
        return "novatek_mjpg"

    @staticmethod
    def _fourcc_text(value):
        fourcc = int(value)
        text = "".join(chr((fourcc >> (8 * index)) & 0xFF) for index in range(4))
        return text.strip("\x00 ") or "UNKNOWN"

    def run(self):
        capture = None
        next_auto_mode = "novatek_mjpg"
        stream_started = False
        try:
            while not self._stop_event.is_set():
                self._reconnect_event.clear()
                camera_index, camera_name = self._camera_device_snapshot()
                if camera_index is None:
                    if stream_started:
                        message = (
                            f"{camera_name} disconnected after streaming "
                            "started; stopping camera."
                        )
                        self.status_changed.emit(message)
                        self.stream_interrupted.emit(message)
                        return
                    self.status_changed.emit(
                        f"{camera_name} disconnected; waiting for hot-plug..."
                    )
                    if self._stop_event.wait(1.0):
                        break
                    continue
                active_capture_mode = self.capture_mode
                if self.capture_mode == "auto_detect":
                    active_capture_mode = next_auto_mode
                    self.status_changed.emit(
                        f"Auto Detect: trying {camera_name} "
                        f"(index {camera_index}) with "
                        f"{self._capture_mode_name(active_capture_mode)}..."
                    )
                else:
                    self.status_changed.emit(
                        f"Opening {camera_name} (index {camera_index}) "
                        f"({self._capture_mode_name(active_capture_mode)})..."
                    )
                try:
                    capture = self._open_capture(
                        active_capture_mode, camera_index
                    )
                except Exception as exc:
                    capture = None
                    self.status_changed.emit(f"Camera open error: {exc}; retrying...")

                if capture is None or not capture.isOpened():
                    if capture is not None:
                        self._clear_capture(capture)
                        capture = None
                    if stream_started:
                        message = (
                            f"Cannot reopen {camera_name} after streaming "
                            "was interrupted; stopping camera."
                        )
                        self.status_changed.emit(message)
                        self.stream_interrupted.emit(message)
                        return
                    if self.capture_mode == "auto_detect":
                        next_auto_mode = self._alternate_auto_mode(
                            active_capture_mode
                        )
                    self.status_changed.emit(
                        f"Cannot open {camera_name} (index {camera_index}); "
                        "retrying..."
                    )
                    if self._stop_event.wait(1.0):
                        break
                    continue

                self._set_capture(capture)
                if not sys.platform.startswith("win") and hasattr(cv2, "CAP_PROP_BUFFERSIZE"):
                    capture.set(cv2.CAP_PROP_BUFFERSIZE, 1)
                if not sys.platform.startswith("win") and hasattr(cv2, "CAP_PROP_READ_TIMEOUT_MSEC"):
                    capture.set(cv2.CAP_PROP_READ_TIMEOUT_MSEC, 1000)
                actual_width = int(capture.get(cv2.CAP_PROP_FRAME_WIDTH))
                actual_height = int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT))
                actual_fps = float(capture.get(cv2.CAP_PROP_FPS))
                actual_fourcc_value = int(capture.get(cv2.CAP_PROP_FOURCC))
                actual_fourcc = self._fourcc_text(actual_fourcc_value)
                if (
                    sys.platform.startswith("win")
                    and active_capture_mode == "novatek_mjpg"
                    and actual_fourcc != UVC_CAPTURE_FOURCC
                ):
                    self.status_changed.emit(
                        f"{camera_name} negotiated {actual_fourcc!r}, "
                        f"expected {UVC_CAPTURE_FOURCC}; retrying..."
                    )
                    self._clear_capture(capture)
                    capture = None
                    if self.capture_mode == "auto_detect":
                        next_auto_mode = "generic_auto"
                    if self._stop_event.wait(1.0):
                        break
                    continue
                self.status_changed.emit(
                    f"{camera_name} (index {camera_index}): {actual_fourcc} "
                    f"{actual_width}x{actual_height} {actual_fps:.2f} FPS"
                )
                self.camera_opened.emit(
                    camera_index, actual_width, actual_height, actual_fps
                )

                failed_since = None
                received_frame = False
                last_frame_seen_emit = 0.0
                while not self._stop_event.is_set():
                    try:
                        ok, frame = capture.read()
                    except Exception:
                        ok, frame = False, None

                    if self._stop_event.is_set():
                        break
                    if self._reconnect_event.is_set():
                        if stream_started:
                            message = (
                                f"{camera_name} stream interrupted; "
                                "stopping camera."
                            )
                            self.status_changed.emit(message)
                            self.stream_interrupted.emit(message)
                            return
                        self.status_changed.emit(
                            f"Reconnecting {camera_name} (index {camera_index})..."
                        )
                        break
                    if ok and frame is not None and frame.size:
                        if self.fps_logger is not None:
                            self.fps_logger.frame_received()
                        failed_since = None
                        stream_started = True
                        if not received_frame:
                            received_frame = True
                            self.selected_capture_mode = active_capture_mode
                            if self.capture_mode == "auto_detect":
                                next_auto_mode = active_capture_mode
                                self.status_changed.emit(
                                    f"Auto Detect selected: "
                                    f"{self._capture_mode_name(active_capture_mode)}; "
                                    f"{actual_fourcc} {actual_width}x{actual_height} "
                                    f"{actual_fps:.2f} FPS"
                                )
                        now = time.perf_counter()
                        if now - last_frame_seen_emit >= 0.1:
                            self.frame_seen.emit(now)
                            last_frame_seen_emit = now
                        if self.frame_consumer is None:
                            self.frame_ready.emit(frame)
                        else:
                            self.frame_consumer(frame)
                        continue

                    now = time.perf_counter()
                    if failed_since is None:
                        failed_since = now
                    no_frame_timeout = self.reconnect_seconds
                    if self.capture_mode == "auto_detect" and not received_frame:
                        no_frame_timeout = UVC_AUTO_DETECT_FRAME_TIMEOUT_SECONDS
                    if now - failed_since >= no_frame_timeout:
                        if stream_started:
                            message = (
                                f"{camera_name} stream interrupted for "
                                f"{no_frame_timeout:.1f}s; stopping camera."
                            )
                            self.status_changed.emit(message)
                            self.stream_interrupted.emit(message)
                            return
                        if self.capture_mode == "auto_detect" and not received_frame:
                            next_auto_mode = self._alternate_auto_mode(
                                active_capture_mode
                            )
                        self.status_changed.emit(
                            f"No frame for {no_frame_timeout:.1f}s; "
                            f"reconnecting {camera_name} "
                            f"(index {camera_index})..."
                        )
                        break
                    if self._stop_event.wait(0.03):
                        break

                self._clear_capture(capture)
                capture = None
                if not self._stop_event.is_set():
                    self._stop_event.wait(0.25)
        finally:
            if self.fps_logger is not None:
                self.fps_logger.close()
            if capture is not None:
                self._clear_capture(capture)


def process_uvc_pipeline_frame(
    frame,
    config,
    clahe_cache=None,
    fieldscale_states=None,
):
    started = time.perf_counter()
    original_gray = uvc_to_gray(frame)
    preprocessing_details = []
    pipeline_input_max = 255

    first_enabled = next(
        (stage for stage in config["stages"] if stage.enabled), None
    )
    use_native_fieldscale = (
        first_enabled is not None
        and first_enabled.algorithm_key == "fieldscale"
        and original_gray.dtype.itemsize > 1
    )
    if use_native_fieldscale:
        pipeline_input = original_gray.copy()
        pipeline_input_max = config["raw_max"]
        preprocessing_details.append("Native DN input for first Fieldscale step")
    elif original_gray.dtype == np.uint8:
        pipeline_input = original_gray.copy()
    else:
        pipeline_input = raw_to_display_u8(original_gray, config["raw_max"])

    final_image, records = apply_algorithm_pipeline(
        pipeline_input,
        config["stages"],
        input_max=pipeline_input_max,
        clahe_cache=clahe_cache,
        fieldscale_states=fieldscale_states,
        fieldscale_temporal=True,
    )
    total_ms = (time.perf_counter() - started) * 1000.0
    original_display = (
        original_gray.copy()
        if original_gray.dtype == np.uint8
        else raw_to_display_u8(original_gray, config["raw_max"])
    )
    return {
        "revision": config["revision"],
        "original_source": original_gray.copy(),
        "original_display": original_display,
        "pipeline_input": pipeline_input,
        "pipeline_input_max": pipeline_input_max,
        "final_image": final_image,
        "records": records,
        "preprocessing_details": preprocessing_details,
        "total_ms": total_ms,
    }


class UvcProcessingThread(QThread):
    frame_processed = pyqtSignal(object)
    frame_failed = pyqtSignal(object)

    def __init__(self, parent=None):
        super().__init__(parent)
        self._condition = threading.Condition()
        self._pending = None
        self._config = None
        self._config_revision = None
        self._stop_requested = False
        self._clahe_cache = {}
        self._fieldscale_states = {}

    def set_config(self, config):
        with self._condition:
            revision = None if config is None else config.get("revision")
            if revision != self._config_revision:
                self._clahe_cache.clear()
                self._fieldscale_states.clear()
                self._config_revision = revision
            self._config = config
            self._pending = None

    def submit_frame(self, frame):
        with self._condition:
            if self._stop_requested or self._config is None:
                return
            self._pending = (frame.copy(), self._config)
            self._condition.notify()

    def submit(self, frame, config):
        with self._condition:
            if self._stop_requested:
                return
            self._pending = (frame.copy(), config)
            self._condition.notify()

    def stop(self):
        with self._condition:
            self._stop_requested = True
            self._pending = None
            self._config = None
            self._condition.notify_all()

    def run(self):
        while True:
            with self._condition:
                while self._pending is None and not self._stop_requested:
                    self._condition.wait()
                if self._stop_requested:
                    return
                frame, config = self._pending
                self._pending = None
            try:
                payload = process_uvc_pipeline_frame(
                    frame,
                    config,
                    clahe_cache=self._clahe_cache,
                    fieldscale_states=self._fieldscale_states,
                )
            except Exception as exc:
                self.frame_failed.emit(
                    {"revision": config.get("revision"), "message": str(exc)}
                )
            else:
                self.frame_processed.emit(payload)


class UvcRecordingThread(QThread):
    """Encode synchronized Original/Final comparison frames off the GUI thread."""

    def __init__(
        self,
        video_path,
        fps,
        parent=None,
        left_label="Original UVC",
        right_label="Final Pipeline Output",
    ):
        super().__init__(parent)
        self.video_path = Path(video_path)
        self.fps = max(1.0, float(fps))
        self.left_label = str(left_label)
        self.right_label = str(right_label)
        self._condition = threading.Condition()
        self._queue = deque()
        self._stop_requested = False
        self._stop_frame_time = None
        self._submitted_frames = 0
        self._queue_dropped_frames = 0
        self.result_summary = None
        self.error_message = None

    def submit(self, original, processed, frame_time=None):
        if original is None or processed is None:
            return False
        timestamp_value = (
            time.perf_counter() if frame_time is None else float(frame_time)
        )
        with self._condition:
            if self._stop_requested:
                return False
            if len(self._queue) >= UVC_RECORDING_QUEUE_SIZE:
                self._queue.popleft()
                self._queue_dropped_frames += 1
            self._queue.append(
                (original.copy(), processed.copy(), timestamp_value)
            )
            self._submitted_frames += 1
            self._condition.notify()
        return True

    def stop(self, frame_time=None):
        with self._condition:
            if not self._stop_requested:
                self._stop_frame_time = (
                    time.perf_counter()
                    if frame_time is None
                    else float(frame_time)
                )
            self._stop_requested = True
            self._condition.notify_all()

    @staticmethod
    def _display_bgr(image):
        array = np.asarray(image)
        if array.ndim != 2:
            raise ValueError("UVC recording requires 2D grayscale images.")
        if array.dtype == np.uint8:
            gray = np.ascontiguousarray(array)
        else:
            gray = np.clip(array, 0, 255).astype(np.uint8)
        return cv2.cvtColor(gray, cv2.COLOR_GRAY2BGR)

    @classmethod
    def _compose_frame(
        cls,
        original,
        processed,
        left_label="Original UVC",
        right_label="Final Pipeline Output",
    ):
        left = cls._display_bgr(original)
        right = cls._display_bgr(processed)
        if right.shape[:2] != left.shape[:2]:
            right = cv2.resize(
                right,
                (left.shape[1], left.shape[0]),
                interpolation=cv2.INTER_AREA,
            )
        canvas = np.ascontiguousarray(np.hstack((left, right)))
        if canvas.shape[0] % 2:
            canvas = cv2.copyMakeBorder(
                canvas, 0, 1, 0, 0, cv2.BORDER_CONSTANT, value=(0, 0, 0)
            )
        if canvas.shape[1] % 2:
            canvas = cv2.copyMakeBorder(
                canvas, 0, 0, 0, 1, cv2.BORDER_CONSTANT, value=(0, 0, 0)
            )

        panel_width = left.shape[1]
        cv2.rectangle(canvas, (0, 0), (panel_width - 1, 34), (0, 0, 0), -1)
        cv2.rectangle(
            canvas,
            (panel_width, 0),
            (canvas.shape[1] - 1, 34),
            (0, 0, 0),
            -1,
        )
        return cls._draw_comparison_labels(
            canvas, panel_width, left_label, right_label
        )

    @staticmethod
    def _draw_comparison_labels(canvas, panel_width, left_label, right_label):
        """Draw Unicode camera names into a BGR frame using Qt's system fonts."""
        # Paint on RGB32: Qt does not support painting on every BGR format.
        rgb = cv2.cvtColor(canvas, cv2.COLOR_BGR2RGB)
        image = QImage(
            rgb.data, rgb.shape[1], rgb.shape[0], rgb.strides[0], QImage.Format_RGB888
        ).convertToFormat(QImage.Format_RGB32)

        painter = QPainter(image)
        painter.setRenderHint(QPainter.TextAntialiasing, True)
        font = QFont()
        font.setPixelSize(18)
        font.setBold(True)
        painter.setFont(font)
        painter.setPen(QColor(255, 255, 255))
        painter.drawText(12, 24, str(left_label))
        painter.drawText(panel_width + 12, 24, str(right_label))
        painter.end()

        image = image.convertToFormat(QImage.Format_RGB888)
        bits = image.bits()
        bits.setsize(image.sizeInBytes())
        painted = np.frombuffer(bits, dtype=np.uint8).reshape(
            image.height(), image.bytesPerLine()
        )[:, :image.width() * 3].reshape(image.height(), image.width(), 3).copy()
        return cv2.cvtColor(painted, cv2.COLOR_RGB2BGR)


    def run(self):
        writer = None
        first_frame_time = None
        last_frame_time = None
        last_canvas = None
        frame_size = None
        written_frames = 0
        duplicate_frames = 0
        timestamp_dropped_frames = 0

        try:
            if cv2 is None:
                raise RuntimeError(
                    "OpenCV not installed. pip install opencv-python"
                )
            while True:
                with self._condition:
                    while not self._queue and not self._stop_requested:
                        self._condition.wait()
                    if self._queue:
                        original, processed, frame_time = self._queue.popleft()
                    elif self._stop_requested:
                        break
                    else:
                        continue

                if first_frame_time is None:
                    first_frame_time = frame_time
                elapsed_seconds = max(0.0, frame_time - first_frame_time)
                canvas = self._compose_frame(
                    original,
                    processed,
                    self.left_label,
                    self.right_label,
                )
                current_size = (canvas.shape[1], canvas.shape[0])
                if writer is None:
                    frame_size = current_size
                    fourcc = cv2.VideoWriter_fourcc(*UVC_RECORDING_FOURCC)
                    writer = cv2.VideoWriter(
                        str(self.video_path), fourcc, self.fps, frame_size, True
                    )
                    if not writer.isOpened():
                        raise RuntimeError(
                            f"Cannot open UVC recording writer: {self.video_path}"
                        )
                elif current_size != frame_size:
                    raise RuntimeError(
                        "UVC recording frame size changed from "
                        f"{frame_size[0]}x{frame_size[1]} to "
                        f"{current_size[0]}x{current_size[1]}"
                    )

                target_index = int(round(elapsed_seconds * self.fps))
                if target_index < written_frames:
                    timestamp_dropped_frames += 1
                    last_frame_time = frame_time
                    continue
                while written_frames < target_index and last_canvas is not None:
                    writer.write(last_canvas)
                    written_frames += 1
                    duplicate_frames += 1
                writer.write(canvas)
                written_frames += 1
                last_canvas = canvas
                last_frame_time = frame_time

            if (
                writer is not None
                and last_canvas is not None
                and first_frame_time is not None
                and self._stop_frame_time is not None
            ):
                stop_elapsed = max(
                    0.0, self._stop_frame_time - first_frame_time
                )
                target_count = int(round(stop_elapsed * self.fps)) + 1
                while written_frames < target_count:
                    writer.write(last_canvas)
                    written_frames += 1
                    duplicate_frames += 1
        except Exception as exc:
            self.error_message = str(exc)
        finally:
            if writer is not None:
                writer.release()
            duration_seconds = 0.0
            if first_frame_time is not None:
                end_time = self._stop_frame_time
                if end_time is None:
                    end_time = last_frame_time
                if end_time is not None:
                    duration_seconds = max(0.0, end_time - first_frame_time)
            file_size = (
                self.video_path.stat().st_size
                if self.video_path.is_file()
                else 0
            )
            self.result_summary = {
                "codec": UVC_RECORDING_FOURCC,
                "fps": round(self.fps, 3),
                "frame_size": (
                    None
                    if frame_size is None
                    else {"width": frame_size[0], "height": frame_size[1]}
                ),
                "duration_seconds": round(duration_seconds, 3),
                "submitted_frames": self._submitted_frames,
                "written_frames": written_frames,
                "duplicate_frames": duplicate_frames,
                "queue_dropped_frames": self._queue_dropped_frames,
                "timestamp_dropped_frames": timestamp_dropped_frames,
                "file_size_bytes": int(file_size),
            }


class UvcViewerRecordingThread(UvcRecordingThread):
    """Encode one grayscale frame from Single UVC Statistic without text overlays."""

    def submit(self, frame, frame_time=None):
        if frame is None:
            return False
        timestamp_value = (
            time.perf_counter() if frame_time is None else float(frame_time)
        )
        with self._condition:
            if self._stop_requested:
                return False
            if len(self._queue) >= UVC_RECORDING_QUEUE_SIZE:
                self._queue.popleft()
                self._queue_dropped_frames += 1
            self._queue.append((frame.copy(), None, timestamp_value))
            self._submitted_frames += 1
            self._condition.notify()
        return True

    @classmethod
    def _compose_frame(
        cls,
        original,
        _processed=None,
        _left_label=None,
        _right_label=None,
    ):
        canvas = cls._display_bgr(original)
        if canvas.shape[0] % 2:
            canvas = cv2.copyMakeBorder(
                canvas, 0, 1, 0, 0, cv2.BORDER_CONSTANT, value=(0, 0, 0)
            )
        if canvas.shape[1] % 2:
            canvas = cv2.copyMakeBorder(
                canvas, 0, 0, 0, 1, cv2.BORDER_CONSTANT, value=(0, 0, 0)
            )
        return canvas


class ThermalRawToolkit(QMainWindow):
    FUNCTIONS = [
        "Single Image Statistics",
        "Average RAW Folder",
        "NUC Calculate",
        "NUC Apply",
        "High / Low Analyzer",
        "Temporal Drift Analyzer",
        "BPR Bad Pixel Marking",
        "Dual Image Compare",
        "Single UVC Statistic",
        "Dual UVC Compare",
        "Image Algorithm Compare",
        "UVC Algorithm Compare",
    ]
    FUNCTION_GROUPS = [
        (
            "Calibration",
            (
                "Average RAW Folder",
                "NUC Calculate",
            ),
        ),
        (
            "Analyzer",
            (
                "NUC Apply",
                "BPR Bad Pixel Marking",
                "High / Low Analyzer",
                "Temporal Drift Analyzer",
            ),
        ),
        (
            "Image",
            (
                "Single Image Statistics",
                "Dual Image Compare",
                "Image Algorithm Compare",
            ),
        ),
        (
            "UVC",
            (
                "Single UVC Statistic",
                "Dual UVC Compare",
                "UVC Algorithm Compare",
            ),
        ),
    ]
    FUNCTION_DESCRIPTIONS = [
        "選擇 RAW 檔後立即顯示影像、RAW DN 統計、Histogram／CDF 與單張品質評分；Run／Save Analysis 會保存完整分析結果。",
        "將資料夾內的 RAW 影像逐像素平均，並依 LOW／HIGH 後綴儲存平均結果。",
        "使用 LOW 與 HIGH RAW 計算 NUC 參數，產生 Novatek pointN_K.bin 與 pointN_B.bin；右側以固定 Bit Depth 比例驗證校正後 HIGH。",
        "讀取 RAW、pointN_K.bin 與 pointN_B.bin 執行 NUC 計算，顯示並儲存原始與 NUC 後的 RAW 和 PNG。",
        "逐像素比較 LOW 與 HIGH RAW，影像區顯示 HIGH - LOW Difference，並分析 DN 差異及異常像素。",
        "分析資料夾內多張 RAW 的時間變化，計算標準差與 Peak-to-Peak。",
        "讀取單張 RAW，選擇輸入映射，再依自訂數量與順序執行影像演算法 Pipeline。",
        "開啟 UVC Camera：預設 Auto Detect 先嘗試 Novatek MJPG，失敗自動改用相機原生格式，再依自訂順序即時執行 Pipeline。",
    ]
    PREVIEW_LAYOUTS = [
        (False, "", "RAW Preview (Fixed Scale)"),
        (False, "", "Average RAW (Fixed Scale)"),
        (True, "HIGH Input (Fixed Scale)", "NUC Corrected HIGH (Fixed Scale)"),
        (True, "Original RAW (Fixed Scale)", "NUC Corrected RAW (Fixed Scale)"),
        (False, "", "HIGH - LOW Difference (Fixed Scale)"),
        (False, "", "Temporal Standard Deviation (Fixed Scale)"),
        (True, "RAW Fixed-Scale / White Canvas", "BPR Bad Pixels Marked"),
        (True, "RAW Image A", "RAW Image B"),
        (False, "", "Single UVC Statistic"),
        (True, "UVC Camera A", "UVC Camera B"),
        (True, "Original RAW (Fixed Scale)", "Final Pipeline Output"),
        (True, "Original UVC", "Final Pipeline Output"),
    ]
    FUNCTION_DESCRIPTIONS.insert(
        6,
        "Load a Novatek BPR dpc_table.bin or x,y coordinate file. Optionally overlay bad-pixel circles on a RAW image; otherwise use a white canvas.",
    )
    FUNCTION_DESCRIPTIONS.insert(
        7,
        "Compare two same-scene images with Image Quality: native-bit-depth Base/Detail analysis, gray-level/connected flat-region separation, robust typical/worst noise, uniformity, clipping, and artifact scores.",
    )
    FUNCTION_DESCRIPTIONS.insert(
        8,
        "Open one UVC camera with Auto Detect for a clean live view. Double-click the image to open a resizable and maximizable image-only window.",
    )
    FUNCTION_DESCRIPTIONS.insert(
        9,
        "Open two independent UVC cameras with per-camera Auto Detect for side-by-side monitoring, paired capture, and synchronized comparison recording.",
    )
    FUNCTION_DESCRIPTIONS[FUNCTIONS.index("Image Algorithm Compare")] = (
        "Select a RAW file to immediately compare the original image with the "
        "selected pipeline output. Algorithm and mapping changes refresh the "
        "preview automatically; Save Analysis Results writes the current analysis."
    )

    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(1450, 900)
        self.active_function_index = 0
        self.app_settings = load_toolkit_settings()
        apply_quality_weights(self.app_settings)
        self.settings_dialog = None
        self.quality_help_dialogs = {}
        self.quality_timeline_dialogs = {}
        self.current_original = None
        self.current_result = None
        self.last_output_dir = None
        self.uvc_video_window = None
        self.viewer_uvc_histogram_window = None
        self.raw_viewer_histogram_window = None
        self.raw_viewer_stats = None
        self.raw_viewer_histogram = None
        self.raw_viewer_histogram_rgb = None
        self.raw_viewer_quality = None
        self.raw_quality_result_a = None
        self.raw_quality_result_b = None
        self.camera_devices = []
        self._camera_device_signature = None
        self.viewer_uvc_worker = None
        self.viewer_uvc_frame = None
        self.viewer_uvc_actual_fps = 0.0
        self.viewer_uvc_output_fps = 0.0
        self.viewer_uvc_frame_count = 0
        self.viewer_uvc_fps_started = 0.0
        self.viewer_uvc_last_analysis_time = 0.0
        self.viewer_uvc_stats = None
        self.viewer_uvc_histogram = None
        self.viewer_uvc_histogram_rgb = None
        self.viewer_uvc_quality = None
        self.viewer_uvc_temporal_tracker = UvcTemporalQualityTracker()
        self.viewer_uvc_restart_when_stopped = False
        self.viewer_uvc_recording_worker = None
        self.viewer_uvc_recording_output_dir = None
        self.viewer_uvc_recording_manifest = None
        self.viewer_uvc_recording_start_histogram = None
        self.viewer_uvc_recording_quality_start = None
        self.viewer_uvc_recording_quality_timeline = []
        self.viewer_uvc_recording_stop_reason = None
        self.viewer_uvc_recording_started_perf = 0.0
        self.viewer_uvc_buffer = LatestFrameBuffer()
        self.viewer_uvc_buffer_sequence = 0
        self.viewer_uvc_display_timer = QTimer(self)
        self.viewer_uvc_display_timer.setInterval(10)
        self.viewer_uvc_display_timer.timeout.connect(
            self.consume_uvc_viewer_latest_frame
        )
        self.raw_algorithm_source = None
        self.raw_pipeline_input = None
        self.raw_pipeline_input_max = 255
        self.raw_pipeline_final = None
        self.raw_pipeline_records = []
        self.raw_algorithm_quality_original = None
        self.raw_algorithm_quality_final = None
        self.dual_uvc_worker_a = None
        self.dual_uvc_worker_b = None
        self.dual_uvc_frame_a = None
        self.dual_uvc_frame_b = None
        self.dual_uvc_frame_time_a = 0.0
        self.dual_uvc_frame_time_b = 0.0
        self.dual_uvc_actual_fps_a = 0.0
        self.dual_uvc_actual_fps_b = 0.0
        self.dual_uvc_output_fps_a = 0.0
        self.dual_uvc_output_fps_b = 0.0
        self.dual_uvc_frame_count_a = 0
        self.dual_uvc_frame_count_b = 0
        self.dual_uvc_quality_a = None
        self.dual_uvc_quality_b = None
        self.dual_uvc_quality_last_update = 0.0
        self.dual_uvc_temporal_tracker_a = UvcTemporalQualityTracker()
        self.dual_uvc_temporal_tracker_b = UvcTemporalQualityTracker()
        self.dual_uvc_recording_quality_timeline = []
        self.dual_uvc_fps_started_a = 0.0
        self.dual_uvc_fps_started_b = 0.0
        self.dual_uvc_recording_worker = None
        self.dual_uvc_recording_output_dir = None
        self.dual_uvc_recording_manifest = None
        self.dual_uvc_recording_quality_start_a = None
        self.dual_uvc_recording_quality_start_b = None
        self.dual_uvc_recording_stop_reason = None
        self.dual_uvc_recording_started_perf = 0.0
        self.dual_uvc_restart_when_stopped = False
        self.dual_uvc_close_pending = False
        self.dual_uvc_buffer_a = LatestFrameBuffer()
        self.dual_uvc_buffer_b = LatestFrameBuffer()
        self.dual_uvc_buffer_sequence_a = 0
        self.dual_uvc_buffer_sequence_b = 0
        self.dual_uvc_display_timer = QTimer(self)
        self.dual_uvc_display_timer.setInterval(10)
        self.dual_uvc_display_timer.timeout.connect(
            self.consume_dual_uvc_latest_frames
        )
        self.uvc_worker = None
        self.uvc_processing_worker = None
        self.uvc_recording_worker = None
        self.uvc_recording_output_dir = None
        self.uvc_recording_manifest = None
        self.uvc_recording_quality_start_original = None
        self.uvc_recording_quality_start_processed = None
        self.uvc_recording_stop_reason = None
        self.uvc_recording_started_perf = 0.0
        self.uvc_actual_fps = 0.0
        self.uvc_original_source = None
        self.uvc_original_frame = None
        self.uvc_processed_frame = None
        self.uvc_quality_original = None
        self.uvc_quality_processed = None
        self.uvc_quality_last_update = 0.0
        self.uvc_temporal_tracker_original = UvcTemporalQualityTracker()
        self.uvc_temporal_tracker_processed = UvcTemporalQualityTracker()
        self.uvc_recording_quality_timeline = []
        self.uvc_pipeline_input = None
        self.uvc_pipeline_input_max = 255
        self.uvc_pipeline_records = []
        self.uvc_pipeline_total_ms = 0.0
        self.uvc_processing_revision = 0
        self.uvc_last_processing_error = None
        self.uvc_frame_count = 0
        self.uvc_fps_display = 0.0
        self.uvc_last_time = 0.0
        self.uvc_last_frame_time = 0.0
        self.uvc_waiting_for_frame = False
        self.uvc_reconnect_requested = False
        self.uvc_restart_when_stopped = False
        self.uvc_close_pending = False
        self.uvc_watchdog_timer = QTimer(self)
        self.uvc_watchdog_timer.setInterval(200)
        self.uvc_watchdog_timer.timeout.connect(self.check_uvc_stream)
        self._build_ui()
        self._apply_ui_style()
        self._apply_category_badge_font()
        self.apply_toolkit_settings(
            self.app_settings, persist=False, announce=False
        )
        self.refresh_camera_devices(log_changes=False)
        self.camera_refresh_timer = QTimer(self)
        self.camera_refresh_timer.setInterval(CAMERA_REFRESH_INTERVAL_MS)
        self.camera_refresh_timer.timeout.connect(self.refresh_camera_devices)
        self.camera_refresh_timer.start()

    def _apply_ui_style(self):
        self.setStyleSheet("""
            QMainWindow {
                background-color: #181a1e;
            }
            QMessageBox {
                background-color: #22252a;
                color: #e7e9ed;
                font-family: "Segoe UI";
                font-size: 13px;
            }
            QMessageBox QLabel#qt_msgbox_label {
                color: #e7e9ed;
                background: transparent;
                padding: 8px 10px;
                min-width: 300px;
            }
            QMessageBox QLabel#qt_msgboxex_icon_label {
                background: transparent;
                padding: 8px 4px 8px 10px;
            }
            QMessageBox QDialogButtonBox {
                background: transparent;
            }
            QMessageBox QPushButton {
                background-color: #3a4048;
                color: #f5f6f7;
                border: 1px solid #59636e;
                border-radius: 7px;
                min-width: 84px;
                min-height: 28px;
                padding: 5px 14px;
                font-weight: 600;
            }
            QMessageBox QPushButton:hover {
                background-color: #4a515b;
                border-color: #74808c;
            }
            QMessageBox QPushButton:pressed {
                background-color: #2d3239;
            }
            QMessageBox QPushButton:default {
                background-color: #3f7eae;
                border-color: #69a5d0;
                color: #ffffff;
            }
            QMessageBox QPushButton:default:hover {
                background-color: #4d90c1;
            }
            QDialog#toolkitSettingsDialog {
                background-color: #181a1e;
                color: #e7e9ed;
                font-family: "Segoe UI";
                font-size: 13px;
            }
            QDialog#toolkitSettingsDialog QTabWidget::pane {
                background-color: #22252a;
                border: 1px solid #424850;
                border-radius: 7px;
                top: -1px;
            }
            QDialog#toolkitSettingsDialog QTabBar::tab {
                background-color: #292d33;
                color: #bdc5ce;
                border: 1px solid #424850;
                padding: 8px 15px;
                min-width: 110px;
            }
            QDialog#toolkitSettingsDialog QTabBar::tab:selected {
                background-color: #31546c;
                color: #ffffff;
                border-color: #5685a3;
            }
            QDialog#toolkitSettingsDialog QGroupBox#settingSection {
                background-color: #20252b;
                border: 1px solid #46515c;
                border-radius: 8px;
                margin-top: 14px;
                padding: 12px 8px 8px 8px;
                color: #c9e2f2;
                font-weight: 700;
            }
            QDialog#toolkitSettingsDialog QGroupBox#settingSection::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 5px;
                color: #c9e2f2;
            }
            QWidget#centralWidget {
                background-color: #181a1e;
                color: #e7e9ed;
                font-family: "Segoe UI";
                font-size: 13px;
            }
            QWidget#controlPanel, QWidget#previewPanel {
                background-color: #22252a;
                border: 1px solid #3a3f47;
                border-radius: 10px;
            }
            QWidget#infoPanel {
                background-color: #202328;
                border: 1px solid #424850;
                border-radius: 8px;
            }
            QGroupBox#settingsCard {
                background-color: #22252a;
                border: 1px solid #3a3f47;
                border-radius: 10px;
            }
            QGroupBox#dualUvcCameraGroup {
                background-color: #22252a;
                border: 1px solid #3a3f47;
                border-radius: 10px;
            }
            QGroupBox#dualUvcCameraGroup::title {
                subcontrol-origin: padding;
                subcontrol-position: top left;
                left: 12px;
                top: 7px;
                padding: 0;
            }
            QGroupBox#viewerStatsCard {
                background-color: #20252b;
                border: 1px solid #46515c;
                border-radius: 8px;
                margin-top: 10px;
                padding: 10px 8px 8px 8px;
                font-weight: 700;
                color: #c9e2f2;
            }
            QGroupBox#viewerStatsCard::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 5px;
            }
            QGroupBox#pipelineParameterCard {
                background-color: #202328;
                border: 1px solid #424850;
                border-radius: 7px;
                margin-top: 10px;
                padding: 10px 8px 8px 8px;
                color: #c9e2f2;
                font-weight: 700;
            }
            QGroupBox#pipelineParameterCard::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 5px;
                color: #c9e2f2;
            }
            QGroupBox[collapsibleGroup="true"]::title {
                font-family: "Segoe UI";
                font-size: 16px;
                font-weight: 700;
                color: #c9e2f2;
            }
            QLabel[viewerStatValue="true"] {
                color: #ffffff;
                background-color: #191c21;
                border: 1px solid #3f4852;
                border-radius: 5px;
                padding: 5px 7px;
                font-family: "Cascadia Mono", "Consolas";
                font-weight: 700;
            }
            QLabel {
                color: #e0e3e8;
                background: transparent;
            }
            QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox, QTextEdit, QListWidget {
                background-color: #191c21;
                color: #f0f1f3;
                border: 1px solid #454b54;
                border-radius: 6px;
                padding: 6px 8px;
                selection-background-color: #4f8fbd;
            }
            QLineEdit:focus, QComboBox:focus, QSpinBox:focus, QDoubleSpinBox:focus,
            QTextEdit:focus, QListWidget:focus {
                border: 1px solid #86bce2;
            }
            QComboBox::drop-down {
                border: none;
                width: 24px;
            }
            QComboBox QAbstractItemView {
                background-color: #25292f;
                color: #f0f1f3;
                border: 1px solid #4a515b;
                selection-background-color: #4f8fbd;
                padding: 4px;
            }
            QPushButton#categoryBadge {
                background-color: #376f8d;
                color: #f4fbff;
                border: 1px solid #65a4c5;
                border-radius: 7px;
                padding: 7px 24px 7px 22px;
                font-family: "Segoe UI";
                font-size: 20px;
                font-weight: normal;
            }
            QPushButton#categoryBadge:hover,
            QPushButton#categoryBadge:focus {
                background-color: #43819f;
                border-color: #8bc8e7;
            }
            QPushButton#categoryBadge:pressed {
                background-color: #2f607a;
                border-color: #5795b5;
            }
            QPushButton#categoryBadge::menu-indicator {
                subcontrol-origin: padding;
                subcontrol-position: right center;
                right: 9px;
            }
            QMenu {
                background-color: #25292f;
                color: #f0f1f3;
                border: 1px solid #555d68;
                border-radius: 6px;
                padding: 5px;
                font-family: "Segoe UI";
                font-size: 17px;
                font-weight: 700;
            }
            QMenu::item {
                border-radius: 4px;
                padding: 7px 30px 7px 12px;
            }
            QMenu::item:selected {
                background-color: #43819f;
                color: #ffffff;
            }
            QMenu::item:checked {
                color: #55b5ff;
                font-weight: 700;
            }
            QMenu::item:checked:selected {
                background-color: #19456b;
                color: #55b5ff;
            }
            QMenu::separator {
                height: 1px;
                background-color: #454b54;
                margin: 4px 8px;
            }
            QListWidget::item {
                color: #e7e9ed;
                padding: 4px 5px;
                border-radius: 3px;
            }
            QListWidget::item:selected {
                background-color: #4f8fbd;
                color: #ffffff;
            }
            QPushButton {
                background-color: #3a4048;
                color: #f5f6f7;
                border: 1px solid #555d68;
                border-radius: 7px;
                padding: 7px 14px;
                font-weight: 600;
            }
            QPushButton:hover {
                background-color: #4a515b;
                border-color: #69727e;
            }
            QPushButton:pressed {
                background-color: #2d3239;
            }
            QPushButton:disabled {
                background-color: #292d33;
                color: #737a84;
                border-color: #383d44;
            }
            QPushButton#primaryButton {
                background-color: #3f7eae;
                border-color: #69a5d0;
            }
            QPushButton#primaryButton:hover {
                background-color: #4d90c1;
            }
            QPushButton#startButton, QPushButton#saveButton {
                background-color: #397851;
                border-color: #5ca476;
            }
            QPushButton#startButton:hover, QPushButton#saveButton:hover {
                background-color: #458c60;
            }
            QPushButton#stopButton {
                background-color: #984646;
                border-color: #c66a6a;
            }
            QPushButton#stopButton:hover {
                background-color: #ad5050;
            }
            QPushButton#captureButton {
                background-color: #70518c;
                border-color: #9874b5;
            }
            QPushButton#captureButton:hover {
                background-color: #80609d;
            }
            QPushButton#resetButton {
                background-color: #a96f2e;
                border-color: #d09a58;
                color: #fff8ef;
            }
            QPushButton#resetButton:hover {
                background-color: #bc7e35;
            }
            QPushButton#startButton:disabled, QPushButton#stopButton:disabled,
            QPushButton#captureButton:disabled, QPushButton#resetButton:disabled,
            QPushButton#saveButton:disabled {
                background-color: #292d33;
                color: #737a84;
                border-color: #383d44;
            }
            QPushButton[buttonRole="browse"] {
                background-color: #376f8d;
                border-color: #65a4c5;
                min-width: 66px;
            }
            QPushButton[buttonRole="browse"]:hover {
                background-color: #43819f;
            }
            QCheckBox {
                color: #e7e9ed;
                spacing: 8px;
                font-size: 14px;
                font-weight: 600;
            }
            QCheckBox#qualityScoreToggle {
                color: #edf4f8;
                font-family: "Segoe UI";
                font-size: 20px;
                font-weight: 600;
                spacing: 0px;
                padding: 0px;
            }
            QScrollArea#formScroll,
            QWidget#formViewport,
            QWidget#formContent {
                background-color: #22252a;
                color: #e7e9ed;
                border: none;
            }
            QScrollBar:vertical {
                background-color: #191c21;
                width: 12px;
                margin: 0;
                border: none;
            }
            QScrollBar::handle:vertical {
                background-color: #4a515b;
                min-height: 28px;
                border-radius: 5px;
            }
            QScrollBar:horizontal {
                background-color: #191c21;
                height: 12px;
                margin: 0;
                border: none;
            }
            QScrollBar::handle:horizontal {
                background-color: #4a515b;
                min-width: 28px;
                border-radius: 5px;
            }
            QScrollBar::handle:hover {
                background-color: #626c78;
            }
            QScrollBar::add-line, QScrollBar::sub-line {
                width: 0;
                height: 0;
                border: none;
                background: transparent;
            }
            QScrollBar::add-page, QScrollBar::sub-page {
                background: transparent;
            }
            QTextEdit {
                font-family: "Cascadia Mono", "Consolas";
                font-size: 17px;
            }
            QSplitter::handle {
                background: transparent;
            }
            QSplitter::handle:horizontal {
                width: 6px;
            }
            QSplitter::handle:vertical {
                height: 6px;
            }
            QToolTip {
                background-color: #30343a;
                color: #f5f6f7;
                border: 1px solid #555d68;
                padding: 5px;
            }
        """)

    def _apply_category_badge_font(self):
        """Match the adjacent title font and size while keeping regular weight."""
        self.function_title.ensurePolished()
        font = QFont(self.function_title.font())
        font.setWeight(QFont.Normal)
        font.setBold(False)
        font.setStretch(QFont.Unstretched)
        self.category_badge.setFont(font)

    def _build_ui(self):
        central = QWidget()
        central.setObjectName("centralWidget")
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(18, 16, 18, 18)
        root.setSpacing(12)

        top = QHBoxLayout()
        self.category_badge = QPushButton(self.FUNCTION_GROUPS[0][0])
        self.category_badge.setObjectName("categoryBadge")
        self.category_badge.setCursor(Qt.PointingHandCursor)
        self.category_badge.setToolTip(
            "Select a Category and Function"
        )
        self.navigation_menu = QMenu(self.category_badge)
        self.navigation_action_group = QActionGroup(self)
        self.navigation_action_group.setExclusive(True)
        self.function_actions = {}
        for category_name, function_names in self.FUNCTION_GROUPS:
            category_menu = self.navigation_menu.addMenu(category_name)
            for function_name in function_names:
                function_index = self.FUNCTIONS.index(function_name)
                action = QAction(function_name, self)
                action.setCheckable(True)
                action.triggered.connect(
                    lambda _checked=False, index=function_index: (
                        self.select_function_from_menu(index)
                    )
                )
                self.navigation_action_group.addAction(action)
                self.function_actions[function_index] = action
                category_menu.addAction(action)
        self.category_badge.setMenu(self.navigation_menu)
        top.addWidget(self.category_badge)
        self.function_title = QLabel(self.FUNCTION_GROUPS[0][1][0])
        self.function_title.setStyleSheet("font-size:26px; font-weight:700; color:#86bce2;")
        top.addWidget(self.function_title)
        top.addStretch()
        root.addLayout(top)

        settings = QGroupBox()
        settings.setObjectName("settingsCard")
        sg = QVBoxLayout(settings)
        sg.setContentsMargins(14, 14, 14, 14)
        sg.setSpacing(10)
        self.function_description = QLabel()
        self.function_description.setWordWrap(True)
        self.function_description.setStyleSheet("font-size:18px; font-weight:600; color:#e6edf2; background:#2b323a; border-left:4px solid #79afd3; border-radius:6px; padding:10px 12px;")
        sg.addWidget(self.function_description)
        self.global_input_controls = QWidget(self)
        self.global_input_controls.hide()
        self.resolution_combo = FocusWheelComboBox(
            self.global_input_controls
        )
        for width, height in RESOLUTION_OPTIONS:
            self.resolution_combo.addItem(f"{width} × {height}", (width, height))
        default_resolution = (DEFAULT_WIDTH, DEFAULT_HEIGHT)
        default_index = RESOLUTION_OPTIONS.index(default_resolution) if default_resolution in RESOLUTION_OPTIONS else 0
        self.resolution_combo.setCurrentIndex(max(default_index, 0))
        self.bit_depth_combo = FocusWheelComboBox(
            self.global_input_controls
        )
        for bits in BIT_DEPTH_OPTIONS:
            self.bit_depth_combo.addItem(f"{bits}-bit (Max {(1 << bits) - 1})", bits)
        default_bit_index = BIT_DEPTH_OPTIONS.index(DEFAULT_BIT_DEPTH) if DEFAULT_BIT_DEPTH in BIT_DEPTH_OPTIONS else 0
        self.bit_depth_combo.setCurrentIndex(default_bit_index)
        self.output_root = PathRow("folder", dialog_title="Select Output Root Folder")
        self.output_root.setParent(self.global_input_controls)
        self.output_root.setText(str(Path.cwd() / "Thermal_RAW_Results"))
        root.addWidget(settings)

        self.main_splitter = QSplitter(Qt.Horizontal)
        self.main_splitter.setHandleWidth(6)
        splitter = self.main_splitter
        root.addWidget(splitter, 1)

        left = QWidget(); left.setObjectName("controlPanel"); left_lay = QVBoxLayout(left)
        self.stack = QStackedWidget()
        left_lay.addWidget(self.stack)
        self.run_btn = QPushButton("Run / Preview")
        self.run_btn.setObjectName("primaryButton")
        self.run_btn.setMinimumHeight(44)
        self.run_btn.clicked.connect(self.run_current)
        left_lay.addWidget(self.run_btn)
        splitter.addWidget(left)

        right = QWidget(); right.setObjectName("previewPanel"); right_lay = QVBoxLayout(right)
        self.dual_compare_controls = QWidget()
        compare_controls_layout = QHBoxLayout(self.dual_compare_controls)
        compare_controls_layout.setContentsMargins(0, 0, 0, 0)
        compare_controls_layout.addWidget(QLabel("Comparison View"))
        self.dual_compare_mode = QComboBox()
        self.dual_compare_mode.addItems(["Side by Side", "Swipe Compare"])
        compare_controls_layout.addWidget(self.dual_compare_mode)
        compare_controls_layout.addStretch()
        self.dual_compare_controls.hide()
        self.right_splitter = QSplitter(Qt.Vertical)
        self.right_splitter.setHandleWidth(6)
        self.right_splitter.setChildrenCollapsible(True)
        preview_area = QWidget()
        previews = QHBoxLayout(preview_area)
        previews.setContentsMargins(0, 0, 0, 0)
        previews.setSpacing(8)
        self.preview_original = PreviewLabel("Original / Input Preview")
        self.preview_result = PreviewLabel("Result Preview")
        self.preview_original.double_clicked.connect(
            lambda: self.open_preview_window(self.preview_original)
        )
        self.preview_result.double_clicked.connect(
            lambda: self.open_preview_window(self.preview_result)
        )
        self.image_preview_windows = {}
        for preview in (self.preview_original, self.preview_result):
            preview.image_changed.connect(
                lambda preview=preview: self.update_image_preview_window(preview)
            )
        self.preview_original_container = QWidget()
        original_preview_layout = QVBoxLayout(self.preview_original_container)
        original_preview_layout.setContentsMargins(0, 0, 0, 0)
        original_preview_layout.setSpacing(6)
        self.preview_original_caption = QLabel("Original / Input Preview")
        self.preview_original_caption.setAlignment(Qt.AlignCenter)
        self.preview_original_caption.setStyleSheet("font-size:16px; font-weight:700; color:#c9e2f2; background:#2b333b; border-left:4px solid #79afd3; border-radius:6px; padding:7px;")
        original_preview_layout.addWidget(self.preview_original_caption)
        original_preview_layout.addWidget(self.preview_original, 1)
        self.preview_result_container = QWidget()
        result_preview_layout = QVBoxLayout(self.preview_result_container)
        result_preview_layout.setContentsMargins(0, 0, 0, 0)
        result_preview_layout.setSpacing(6)
        self.preview_result_caption = QLabel("Result Preview")
        self.preview_result_caption.setAlignment(Qt.AlignCenter)
        self.preview_result_caption.setStyleSheet("font-size:16px; font-weight:700; color:#d4eadc; background:#29372f; border-left:4px solid #6aaa7d; border-radius:6px; padding:7px;")
        result_preview_layout.addWidget(self.preview_result_caption)
        result_preview_layout.addWidget(self.preview_result, 1)
        previews.addWidget(self.preview_original_container, 1)
        previews.addWidget(self.preview_result_container, 1)
        self.dual_swipe_preview = SwipeCompareWidget()
        previews.addWidget(self.dual_swipe_preview, 1)
        self.dual_swipe_preview.hide()
        self.dual_swipe_preview.double_clicked.connect(
            lambda: self.open_preview_window(self.preview_original)
        )
        self.dual_swipe_preview.position_changed.connect(self.sync_dual_swipe_position)
        self.dual_compare_mode.currentIndexChanged.connect(self.update_dual_compare_mode)
        self.right_splitter.addWidget(preview_area)

        self.info_area = QWidget()
        self.info_area.setObjectName("infoPanel")
        info_layout = QVBoxLayout(self.info_area)
        info_layout.setContentsMargins(10, 8, 10, 10)
        info_layout.setSpacing(6)
        info_header = QHBoxLayout()
        self.info_title = QLabel("Information / Log")
        self.info_title.setStyleSheet("font-size:16px; font-weight:700; color:#d7dde5; padding:2px 4px;")
        self.settings_button = QPushButton("Setting")
        self.settings_button.setToolTip(
            "Configure appearance, quality timeline, and score weights"
        )
        self.settings_button.clicked.connect(self.show_toolkit_settings)
        info_header.addWidget(self.info_title)
        info_header.addStretch()
        info_layout.addLayout(info_header)
        self.log = QTextEdit(); self.log.setReadOnly(True); self.log.setMinimumHeight(120)
        self.update_log_font_size(self.app_settings["log_font_size"])
        info_layout.addWidget(self.log, 1)
        self.right_splitter.addWidget(self.info_area)
        self.right_splitter.setCollapsible(0, False)
        self.right_splitter.setCollapsible(1, True)
        self.right_splitter.setStretchFactor(0, 3)
        self.right_splitter.setStretchFactor(1, 1)
        self.right_splitter.setSizes([650, 50])
        self.right_splitter.handle(1).setToolTip(
            "Drag down to resize or fully hide Information / Log"
        )
        right_lay.addWidget(self.right_splitter, 1)
        action_row = QHBoxLayout()
        self.save_current_btn = QPushButton("Save Current Preview PNG")
        self.save_current_btn.setObjectName("saveButton")
        self.save_current_btn.clicked.connect(self.save_current_preview)
        self.open_result_btn = QPushButton("Open Last Result Folder")
        self.open_result_btn.setObjectName("openButton")
        self.open_result_btn.clicked.connect(self.open_last_result)
        action_row.addWidget(self.save_current_btn)
        action_row.addWidget(self.open_result_btn)
        action_row.addStretch()
        action_row.addWidget(self.settings_button)
        right_lay.addLayout(action_row)
        splitter.addWidget(right)
        splitter.setSizes([350, 1050])

        self._create_pages()
        self.change_page(self.FUNCTIONS.index("Single Image Statistics"))

    def update_log_font_size(self, size):
        self.log.setStyleSheet(f"font-size:{size}px;")
        if hasattr(self, "bpr_coordinate_info"):
            self.bpr_coordinate_info.setStyleSheet(
                f'font-family:"Consolas"; font-size:{size}px;'
            )

    def show_toolkit_settings(self):
        if self.settings_dialog is None:
            self.settings_dialog = ToolkitSettingsDialog(
                self.app_settings, self
            )
            self.settings_dialog.settings_applied.connect(
                self.apply_toolkit_settings
            )
        else:
            self.settings_dialog.set_settings(self.app_settings)
        self.settings_dialog.show()
        self.settings_dialog.raise_()
        self.settings_dialog.activateWindow()

    def apply_toolkit_settings(
        self, settings, persist=True, announce=True
    ):
        settings = validated_toolkit_settings(settings)
        old_settings = getattr(self, "app_settings", default_toolkit_settings())
        weights_changed = (
            settings["quality_weights"]
            != old_settings["quality_weights"]
        )
        recording_active = any(
            worker is not None
            for worker in (
                self.viewer_uvc_recording_worker,
                self.dual_uvc_recording_worker,
                self.uvc_recording_worker,
            )
        )
        if weights_changed and recording_active:
            QMessageBox.information(
                self,
                "Recording Active",
                "Stop UVC recording before changing quality-score weights.",
            )
            if self.settings_dialog is not None:
                self.settings_dialog.set_settings(self.app_settings)
            return

        self.app_settings = settings
        apply_quality_weights(settings)
        self.update_log_font_size(settings["log_font_size"])

        raw_input = settings["raw_input"]
        selected_resolution = tuple(raw_input["resolution"])
        resolution_index = next(
            (
                index
                for index in range(self.resolution_combo.count())
                if tuple(self.resolution_combo.itemData(index))
                == selected_resolution
            ),
            -1,
        )
        if resolution_index >= 0:
            self.resolution_combo.setCurrentIndex(resolution_index)
        bit_depth_index = self.bit_depth_combo.findData(
            raw_input["bit_depth"]
        )
        if bit_depth_index >= 0:
            self.bit_depth_combo.setCurrentIndex(bit_depth_index)
        self.output_root.setText(raw_input["output_root"])

        timeline = settings["timeline"]
        metric_display = next(
            display
            for display, metric in QUALITY_TIMELINE_METRICS
            if metric == timeline["metric"]
        )
        chart_groups = (
            (self.viewer_uvc_quality_chart, self.viewer_uvc_quality_chart_group),
            (self.dual_uvc_quality_chart, self.dual_uvc_quality_chart_group),
            (self.uvc_quality_chart, self.uvc_quality_chart_group),
        )
        for chart, group in chart_groups:
            chart.set_display_metric(timeline["metric"])
            chart.set_history_seconds(timeline["history_seconds"])
            group.setTitle(f"{metric_display} Timeline")
        timeline_dialog_titles = {
            "single_uvc": "Single UVC",
            "dual_uvc": "Dual UVC",
            "uvc_algorithm": "UVC Algorithm",
        }
        for key, dialog in self.quality_timeline_dialogs.items():
            dialog.setWindowTitle(
                f"{timeline_dialog_titles[key]} {metric_display} Timeline"
            )

        if weights_changed:
            for chart, _group in chart_groups:
                chart.clear()
            self.raw_viewer_quality = None
            self.raw_quality_result_a = None
            self.raw_quality_result_b = None
            self.raw_algorithm_quality_original = None
            self.raw_algorithm_quality_final = None
            self.viewer_uvc_quality = None
            self.dual_uvc_quality_a = None
            self.dual_uvc_quality_b = None
            self.uvc_quality_original = None
            self.uvc_quality_processed = None
            if hasattr(self, "raw_viewer_quality_labels"):
                self.update_single_quality_score_labels(
                    self.raw_viewer_quality_labels, None
                )
                self.update_quality_score_labels(
                    self.raw_quality_score_labels, None, None
                )
                self.update_quality_score_labels(
                    self.raw_algorithm_quality_labels, None, None
                )
                self.update_single_quality_score_labels(
                    self.viewer_uvc_quality_labels, None
                )
                self.update_quality_score_labels(
                    self.dual_uvc_quality_labels, None, None
                )
                self.update_quality_score_labels(
                    self.uvc_quality_labels, None, None
                )

        settings_saved = True
        if persist:
            try:
                save_toolkit_settings(settings)
            except OSError as exc:
                settings_saved = False
                QMessageBox.warning(
                    self,
                    "Setting Save Failed",
                    f"Setting was applied but could not be saved:\n{exc}",
                )
        if announce:
            self.log.append(
                f"Setting applied: log {settings['log_font_size']} px; "
                f"RAW {self.width} × {self.height}, "
                f"{self.bit_depth}-bit; "
                f"timeline {metric_display}, "
                f"{timeline['history_seconds']} s"
            )
            if weights_changed:
                self.log.append(
                    "Quality weights updated; score displays and UVC timelines were reset."
                )
            if self.settings_dialog is not None:
                self.settings_dialog.show_apply_status(
                    "Setting applied successfully."
                    if settings_saved
                    else "Setting applied, but the JSON file was not saved.",
                    success=settings_saved,
                )

    def _form_page(self):
        page = QWidget()
        page_layout = QVBoxLayout(page)
        page_layout.setContentsMargins(0, 0, 0, 0)
        scroll = QScrollArea()
        scroll.setObjectName("formScroll")
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.viewport().setObjectName("formViewport")
        scroll.viewport().setAutoFillBackground(True)
        content = QWidget()
        content.setObjectName("formContent")
        content.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        content_layout = QVBoxLayout(content)
        form = QFormLayout()
        form.setRowWrapPolicy(QFormLayout.WrapLongRows)
        content_layout.addLayout(form)
        content_layout.addStretch()
        scroll.setWidget(content)
        page_layout.addWidget(scroll)
        return page, form

    def create_quality_score_container(
        self,
        title,
        include_temporal=False,
        collapsible=True,
    ):
        group = (
            CollapsibleGroupBox(title)
            if collapsible
            else QGroupBox(title)
        )
        group.setObjectName("viewerStatsCard")
        outer_layout = QVBoxLayout(group)
        outer_layout.setContentsMargins(
            8 if collapsible else 2,
            6 if collapsible else 0,
            8 if collapsible else 2,
            8 if collapsible else 2,
        )
        outer_layout.setSpacing(6)
        layout = QGridLayout()
        outer_layout.addLayout(layout)
        return group, layout

    def create_quality_help_note(self, note_text, include_temporal=False):
        note = QLabel(
            f'{note_text} <a href="quality-help" '
            'style="color:#c9e2f2; text-decoration:none;">'
            '<b>說明 ⓘ</b></a>'
        )
        note.setTextFormat(Qt.RichText)
        note.setWordWrap(True)
        note.setTextInteractionFlags(
            Qt.LinksAccessibleByMouse | Qt.LinksAccessibleByKeyboard
        )
        note.setOpenExternalLinks(False)
        note.linkActivated.connect(
            lambda _href, include_temporal=include_temporal:
            self.show_quality_score_help(include_temporal)
        )
        note.setStyleSheet(
            "color:#aeb8c2; font-weight:400; padding-top:4px;"
        )
        return note

    def show_quality_score_help(self, include_temporal=False):
        guide_path = bundled_resource_path("quality_score_guide.html")
        if guide_path.is_file():
            guide_url = QUrl.fromLocalFile(str(guide_path))
            guide_url.setFragment("uvc" if include_temporal else "raw")
            if QDesktopServices.openUrl(guide_url):
                return

        dialog_key = "uvc" if include_temporal else "raw"
        dialog = self.quality_help_dialogs.get(dialog_key)
        if dialog is None:
            dialog = QualityScoreHelpDialog(include_temporal, self)
            self.quality_help_dialogs[dialog_key] = dialog
        else:
            dialog.reload_guide()
        dialog.show()
        dialog.raise_()
        dialog.activateWindow()

    def open_quality_timeline(self, key, source_chart, title):
        title = title.replace(
            "Total Score", source_chart.metric_display_name
        )
        dialog = self.quality_timeline_dialogs.get(key)
        if dialog is None:
            dialog = QualityTimelineDialog(source_chart, title, self)
            self.quality_timeline_dialogs[key] = dialog
        else:
            dialog.sync_from_source()
            dialog.setWindowTitle(title)
        dialog.show()
        dialog.raise_()
        dialog.activateWindow()

    @staticmethod
    def create_quality_timeline_group(chart):
        group = CollapsibleGroupBox("Total Score Timeline")
        group.setObjectName("viewerStatsCard")
        layout = QVBoxLayout(group)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.addWidget(chart)
        return group

    def create_quality_score_group(self, title, include_temporal=False):
        group, layout = self.create_quality_score_container(
            title,
            include_temporal,
        )
        layout.setHorizontalSpacing(10)
        layout.setVerticalSpacing(7)
        comparison_header_labels = []
        for column, header in enumerate(("Metric", "A", "B")):
            header_label = QLabel(header)
            header_label.setStyleSheet(
                "font-weight:700; color:#c9e2f2;"
            )
            header_label.setAlignment(
                Qt.AlignLeft if column == 0 else Qt.AlignRight
            )
            layout.addWidget(header_label, 0, column)
            if column > 0:
                comparison_header_labels.append(header_label)
        group.comparison_header_labels = comparison_header_labels
        fields = [
            ("Total Score", "overall_score"),
            ("Detail Clarity", "detail_clarity_score"),
            ("Contrast Quality", "contrast_quality_score"),
            ("Flat Cleanliness", "flat_cleanliness_score"),
            ("Thermal Uniformity", "thermal_uniformity_score"),
            ("Clipping Control", "clipping_control_score"),
            ("Artifact Control", "artifact_control_score"),
        ]
        if include_temporal:
            fields.extend(
                [
                    ("Flicker Control", "flicker_control_score"),
                    ("Temporal Stability", "temporal_stability_score"),
                ]
            )
        else:
            fields.append(("Quality Confidence", "quality_confidence_score"))
        value_labels = {}
        for row, (display_name, metric_name) in enumerate(fields, start=1):
            metric_label = QLabel(display_name)
            layout.addWidget(metric_label, row, 0)
            metric_labels = []
            for column in range(1, 3):
                value_label = QLabel("--")
                value_label.setProperty("viewerStatValue", True)
                value_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
                value_label.setMinimumWidth(66)
                layout.addWidget(value_label, row, column)
                metric_labels.append(value_label)
            value_labels[metric_name] = metric_labels
        if include_temporal:
            note_text = (
                "UVC Lightweight Quality：使用降採樣、Gaussian/Sobel 與簡化空間評分，"
                "降低即時影像負載。Total Score 權重由 Setting 設定；"
                "時序分數暖機期間顯示 Collecting...。"
            )
        else:
            note_text = (
                "Image Quality：以原始位深完整解析，Flat 區會依灰階層級與連通區域分開統計，"
                "並兼顧典型區域與較差區域；Quality Confidence 表示本張影像可供評分的 Flat/Edge 覆蓋程度。"
            )
        note = self.create_quality_help_note(note_text, include_temporal)
        layout.addWidget(note, len(fields) + 1, 0, 1, 3)
        return group, value_labels

    @staticmethod
    def update_quality_score_labels(value_labels, result_a, result_b):
        metrics_a = None if result_a is None else result_a["metrics"]
        metrics_b = None if result_b is None else result_b["metrics"]
        for metric_name, labels in value_labels.items():
            for label in labels:
                label.setStyleSheet("color:#c9d1d9; font-weight:400;")
                label.setToolTip("")
            if metrics_a is None or metrics_b is None:
                for label in labels:
                    label.setText("--")
                continue
            value_a = metrics_a.get(metric_name)
            value_b = metrics_b.get(metric_name)
            labels[0].setText(
                "Collecting..." if value_a is None else f"{float(value_a):.2f}"
            )
            labels[1].setText(
                "Collecting..." if value_b is None else f"{float(value_b):.2f}"
            )
            if value_a is not None and value_b is not None:
                # Compare displayed precision so visually equal scores remain neutral.
                score_a = float(f"{float(value_a):.2f}")
                score_b = float(f"{float(value_b):.2f}")
                if np.isfinite(score_a) and np.isfinite(score_b) and score_a != score_b:
                    winner = labels[0 if score_a > score_b else 1]
                    winner.setStyleSheet("color:#55b5ff; font-weight:700;")
                    winner.setToolTip("Higher score for this metric")

    def create_single_quality_score_group(
        self,
        title,
        include_temporal=False,
        collapsible=True,
    ):
        group, layout = self.create_quality_score_container(
            title,
            include_temporal,
            collapsible,
        )
        layout.setHorizontalSpacing(10)
        layout.setVerticalSpacing(7)
        fields = [
            ("Total Score", "overall_score"),
            ("Detail Clarity", "detail_clarity_score"),
            ("Contrast Quality", "contrast_quality_score"),
            ("Flat Cleanliness", "flat_cleanliness_score"),
            ("Thermal Uniformity", "thermal_uniformity_score"),
            ("Clipping Control", "clipping_control_score"),
            ("Artifact Control", "artifact_control_score"),
        ]
        if include_temporal:
            fields.extend(
                [
                    ("Flicker Control", "flicker_control_score"),
                    ("Temporal Stability", "temporal_stability_score"),
                ]
            )
        else:
            fields.append(("Quality Confidence", "quality_confidence_score"))
        value_labels = {}
        for row, (display_name, metric_name) in enumerate(fields):
            metric_label = QLabel(display_name)
            layout.addWidget(metric_label, row, 0)
            value_label = QLabel("--")
            value_label.setProperty("viewerStatValue", True)
            value_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            value_label.setMinimumWidth(90)
            layout.addWidget(value_label, row, 1)
            value_labels[metric_name] = value_label
        if include_temporal:
            note_text = (
                "UVC Lightweight Quality：以低負載空間評分加上 Flicker / Temporal Stability。"
                "Total Score 權重由 Setting 設定。"
            )
        else:
            note_text = (
                "Image Quality：使用原始位深完整分析，Flat 區依灰階與連通區域分開統計；"
                "分數適合相同場景相對比較，不作為絕對 PASS/FAIL 規格。"
            )
        note = self.create_quality_help_note(note_text, include_temporal)
        layout.addWidget(note, len(fields), 0, 1, 2)
        return group, value_labels

    @staticmethod
    def update_single_quality_score_labels(value_labels, result):
        metrics = None if result is None else result["metrics"]
        for metric_name, label in value_labels.items():
            if metrics is None:
                label.setText("--")
                continue
            value = metrics.get(metric_name)
            label.setText(
                "Collecting..." if value is None else f"{float(value):.2f}"
            )

    def _create_pages(self):
        # Viewer
        p, f = self._form_page()
        self.viewer_file = PathRow(
            "file",
            IMAGE_INPUT_FILTER,
            "Select RAW / Image File to View",
        )
        image_input_note = QLabel(IMAGE_INPUT_NOTE)
        image_input_note.setWordWrap(True)
        image_input_note.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        f.addRow(image_input_note)
        self.viewer_file.path_selected.connect(self.auto_preview_raw)
        self.viewer_file.edit.returnPressed.connect(self.auto_preview_raw)
        self.viewer_file.edit.textChanged.connect(lambda: self.set_image_save_ready("single", False))
        self.resolution_combo.currentIndexChanged.connect(
            self.refresh_raw_viewer_settings
        )
        self.bit_depth_combo.currentIndexChanged.connect(
            self.refresh_raw_viewer_settings
        )
        f.addRow("RAW / Image File", self.viewer_file)

        self.viewer_stats_group = CollapsibleGroupBox("RAW Statistics (DN)")
        self.viewer_stats_group.setObjectName("viewerStatsCard")
        viewer_stats_layout = QGridLayout(self.viewer_stats_group)
        viewer_stats_layout.setHorizontalSpacing(10)
        viewer_stats_layout.setVerticalSpacing(8)
        self.viewer_stat_values = {}
        stat_fields = [
            ("Pixel Count", "count"),
            ("Unique Levels", "unique_levels"),
            ("Minimum", "min"),
            ("Maximum", "max"),
            ("Mean", "mean"),
            ("Median / P50", "median"),
            ("Std Dev", "std"),
            ("Entropy", "entropy_bits"),
            ("P1", "p1"),
            ("P5", "p5"),
            ("P95", "p95"),
            ("P99", "p99"),
            ("P99 - P1", "p99_p1"),
            ("P99-P1 Usage", "p99_p1_percent"),
            ("Zero Clip", "zero_clip"),
            ("Bit Max Clip", "max_clip"),
            ("Over Bit Max", "over_max"),
        ]
        for index, (label_text, key) in enumerate(stat_fields):
            row = index // 2
            column = (index % 2) * 2
            name_label = QLabel(label_text)
            value_label = QLabel("--")
            value_label.setProperty("viewerStatValue", True)
            value_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            value_label.setMinimumWidth(92)
            viewer_stats_layout.addWidget(name_label, row, column)
            viewer_stats_layout.addWidget(value_label, row, column + 1)
            self.viewer_stat_values[key] = value_label

        self.viewer_stats_note = QLabel(
            "Statistics use all source grayscale pixels (RAW: original DN). Std Dev uses ddof=0; "
            "P99-P1 Usage is relative to the source bit-depth maximum."
        )
        self.viewer_stats_note.setWordWrap(True)
        self.viewer_stats_note.setStyleSheet("color:#aeb8c2; font-weight:400; padding-top:4px;")
        viewer_stats_layout.addWidget(
            self.viewer_stats_note,
            (len(stat_fields) + 1) // 2,
            0,
            1,
            4,
        )
        f.addRow(self.viewer_stats_group)

        (
            self.raw_viewer_quality_group,
            self.raw_viewer_quality_labels,
        ) = self.create_single_quality_score_group(
            "Image Quality Scores",
            collapsible=True,
        )
        f.addRow(self.raw_viewer_quality_group)

        self.raw_viewer_histogram_group = CollapsibleGroupBox(
            "RAW DN Histogram + CDF"
        )
        self.raw_viewer_histogram_group.setObjectName("viewerStatsCard")
        raw_viewer_histogram_layout = QVBoxLayout(
            self.raw_viewer_histogram_group
        )
        self.raw_viewer_histogram_preview = PreviewLabel(
            "RAW DN Histogram"
        )
        self.raw_viewer_histogram_preview.setMinimumSize(320, 180)
        self.raw_viewer_histogram_preview.setToolTip(
            "Double-click to open a resizable and maximizable histogram window"
        )
        self.raw_viewer_histogram_preview.double_clicked.connect(
            self.open_raw_histogram_window
        )
        raw_viewer_histogram_layout.addWidget(
            self.raw_viewer_histogram_preview
        )
        raw_histogram_note = QLabel(
            "The chart aggregates the selected Bit Depth range into 256 "
            "display bins. Statistics and percentile markers remain in "
            "original RAW DN; over-range pixels enter the rightmost chart bin "
            "and are reported separately. Double-click the chart to enlarge."
        )
        raw_histogram_note.setWordWrap(True)
        raw_histogram_note.setStyleSheet(
            "color:#aeb8c2; font-weight:400; padding-top:3px;"
        )
        raw_viewer_histogram_layout.addWidget(raw_histogram_note)
        f.addRow(self.raw_viewer_histogram_group)
        self.stack.addWidget(p)

        # Average
        p, f = self._form_page(); self.avg_folder = PathRow("folder", dialog_title="Select RAW Folder to Average"); self.avg_suffix_combo = FocusWheelComboBox(); self.avg_suffix_combo.addItems(["LOW", "HIGH"]); f.addRow("RAW Folder", self.avg_folder); f.addRow("Filename Suffix", self.avg_suffix_combo); self.stack.addWidget(p)

        # NUC Calculate
        p, f = self._form_page()
        self.nuc_mode = FocusWheelComboBox(); self.nuc_mode.addItems(["Folder Average (Recommended)", "Single RAW"]); self.nuc_mode.currentIndexChanged.connect(self.update_nuc_mode)
        self.nuc_point_combo = FocusWheelComboBox()
        for point_number in range(1, 6):
            self.nuc_point_combo.addItem(f"Point {point_number}", point_number)
        self.low_single = PathRow("file", "RAW Files (*.raw *.RAW *.bin *.BIN)", "Select LOW RAW File for NUC")
        self.high_single = PathRow("file", "RAW Files (*.raw *.RAW *.bin *.BIN)", "Select HIGH RAW File for NUC")
        self.low_folder = PathRow("folder", dialog_title="Select LOW RAW Folder for NUC")
        self.high_folder = PathRow("folder", dialog_title="Select HIGH RAW Folder for NUC")
        self.min_delta = FocusWheelSpinBox(); self.min_delta.setRange(0, 16383); self.min_delta.setValue(50)
        f.addRow("NUC Point", self.nuc_point_combo); f.addRow("NUC Source Mode", self.nuc_mode); f.addRow("LOW RAW", self.low_single); f.addRow("HIGH RAW", self.high_single); f.addRow("LOW Folder", self.low_folder); f.addRow("HIGH Folder", self.high_folder); f.addRow("Min Delta DN", self.min_delta)
        note = QLabel("Recommended: capture multiple frames at each stable blackbody temperature, average each folder pixel-by-pixel, then calculate NUC Gain/Offset and Novatek K/B tables.")
        note.setWordWrap(True); note.setStyleSheet("color:#c98;"); f.addRow(note)
        self.stack.addWidget(p); self.update_nuc_mode()

        # NUC Apply
        p, f = self._form_page(); self.apply_raw_file = PathRow("file", "RAW Files (*.raw *.RAW *.bin *.BIN)", "Select RAW File for NUC Apply"); self.k_bin_file = PathRow("file", "Novatek K Table (*.bin *.BIN)", "Select NUC pointN_K.bin"); self.b_bin_file = PathRow("file", "Novatek B Table (*.bin *.BIN)", "Select NUC pointN_B.bin"); self.k_bin_file.path_selected.connect(self.auto_pair_nuc_b_file); f.addRow("RAW File", self.apply_raw_file); f.addRow("pointN_K.bin", self.k_bin_file); f.addRow("pointN_B.bin", self.b_bin_file); self.stack.addWidget(p)

        # High Low Analyzer
        p, f = self._form_page(); self.an_low = PathRow("file", "RAW Files (*.raw *.RAW *.bin *.BIN)", "Select LOW RAW File for High / Low Analysis"); self.an_high = PathRow("file", "RAW Files (*.raw *.RAW *.bin *.BIN)", "Select HIGH RAW File for High / Low Analysis"); self.an_min_delta = FocusWheelSpinBox(); self.an_min_delta.setRange(0, 16383); self.an_min_delta.setValue(50); f.addRow("LOW RAW", self.an_low); f.addRow("HIGH RAW", self.an_high); f.addRow("Min Delta DN", self.an_min_delta); self.stack.addWidget(p)

        # Temporal Drift
        p, f = self._form_page(); self.drift_folder = PathRow("folder", dialog_title="Select RAW Folder for Temporal Drift Analysis"); self.std_spec = FocusWheelDoubleSpinBox(); self.std_spec.setRange(0, 100000); self.std_spec.setValue(50.0); self.p2p_spec = FocusWheelDoubleSpinBox(); self.p2p_spec.setRange(0, 100000); self.p2p_spec.setValue(100.0); f.addRow("RAW Folder", self.drift_folder); f.addRow("Temporal Std Spec", self.std_spec); f.addRow("Peak-to-Peak Spec", self.p2p_spec)
        self.drift_help = QLabel(
            "Temporal Std：每個像素隨時間的整體波動程度。\n"
            "Peak-to-Peak：每個像素的最大 DN - 最小 DN；超過 Spec 會判定為異常。\n"
            "建議使用固定溫度、固定場景且沒有 FFC／Shutter 動作的 RAW。"
        )
        self.drift_help.setWordWrap(True)
        self.drift_help.setStyleSheet("color:#c8d0d9; background:#292d33; border-left:3px solid #79afd3; border-radius:5px; padding:8px;")
        f.addRow(self.drift_help)
        self.stack.addWidget(p)

        # BPR Bad Pixel Marking
        p, f = self._form_page()
        self.bpr_form = f
        self.bpr_coordinate_file = PathRow(
            "file",
            "BPR Table / Coordinates (*.bin *.BIN *.csv *.CSV *.txt *.TXT)",
            "Select Novatek dpc_table.bin or x,y coordinate file",
        )
        self.bpr_use_raw = QualityScoreCheckBox("Overlay circles on an image")
        background_font = self.bpr_use_raw.font()
        background_font.setWeight(QFont.Normal)
        self.bpr_use_raw.setFont(background_font)
        self.bpr_use_raw.setMinimumWidth(
            self.bpr_use_raw.fontMetrics().horizontalAdvance(self.bpr_use_raw.text()) + 44
        )
        self.bpr_use_raw.setChecked(False)
        self.bpr_raw_file = PathRow(
            "file",
            IMAGE_INPUT_FILTER,
            "Select optional image background",
        )
        self.bpr_circle_radius = FocusWheelSpinBox()
        self.bpr_circle_radius.setRange(1, 20)
        self.bpr_circle_radius.setValue(BPR_MARK_RADIUS)
        self.bpr_highlight_bright_points = QualityScoreCheckBox(
            "Highlight bright points"
        )
        highlight_font = self.bpr_highlight_bright_points.font()
        highlight_font.setWeight(QFont.Normal)
        self.bpr_highlight_bright_points.setFont(highlight_font)
        self.bpr_highlight_radius = FocusWheelSpinBox()
        self.bpr_highlight_radius.setRange(1, 20)
        self.bpr_highlight_radius.setValue(BPR_HIGHLIGHT_RADIUS)
        self.bpr_note = QLabel()
        self.bpr_note.setWordWrap(True)
        self.bpr_note.setStyleSheet(
            "color:#c8d0d9; background:#292d33; border-left:3px solid #79afd3; "
            "border-radius:5px; padding:8px;"
        )
        self.bpr_coordinate_info = QTextEdit()
        self.bpr_coordinate_info.setReadOnly(True)
        self.bpr_coordinate_info.setMinimumHeight(190)
        self.bpr_coordinate_info.setPlaceholderText(
            "Browse a coordinate file to preview bad-pixel locations."
        )
        self.update_log_font_size(self.app_settings["log_font_size"])
        for row in (self.bpr_coordinate_file, self.bpr_raw_file):
            row.edit.textChanged.connect(self.invalidate_bpr_preview)
            row.path_selected.connect(self.auto_preview_bpr)
            row.edit.returnPressed.connect(self.auto_preview_bpr)
        self.bpr_use_raw.toggled.connect(self.auto_preview_bpr)
        self.bpr_circle_radius.valueChanged.connect(self.auto_preview_bpr)
        self.bpr_highlight_bright_points.toggled.connect(
            self.auto_preview_bpr
        )
        self.bpr_highlight_bright_points.toggled.connect(
            self.update_bpr_marking_ui
        )
        self.bpr_highlight_radius.valueChanged.connect(self.auto_preview_bpr)
        self.bit_depth_combo.currentIndexChanged.connect(self.auto_preview_bpr)
        self.bpr_use_raw.toggled.connect(self.update_bpr_marking_ui)
        self.resolution_combo.currentIndexChanged.connect(self.auto_preview_bpr)
        self.resolution_combo.currentIndexChanged.connect(self.update_bpr_coordinate_note)
        f.addRow("BPR Table / Coordinates", self.bpr_coordinate_file)
        f.addRow("Background", self.bpr_use_raw)
        f.addRow("Image File", self.bpr_raw_file)
        f.addRow("Red Circle Radius", self.bpr_circle_radius)
        f.addRow("Bright Point", self.bpr_highlight_bright_points)
        f.addRow("Bright Point Radius", self.bpr_highlight_radius)
        f.addRow(self.bpr_note)
        f.addRow("Bad Pixel Information", self.bpr_coordinate_info)
        self.stack.addWidget(p)
        self.update_bpr_marking_ui()
        self.update_bpr_coordinate_note()

        # Dual Image Compare
        p, f = self._form_page()
        self.raw_quality_file_a = PathRow(
            "file",
            IMAGE_INPUT_FILTER,
            "Select RAW / Image Image A for Quality Compare",
        )
        self.raw_quality_file_b = PathRow(
            "file",
            IMAGE_INPUT_FILTER,
            "Select RAW / Image Image B for Quality Compare",
        )
        image_input_note = QLabel(IMAGE_INPUT_NOTE)
        image_input_note.setWordWrap(True)
        image_input_note.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        f.addRow(image_input_note)
        self.raw_quality_note = QLabel(
            "Use two images of the same scene, resolution, and Bit Depth. Image Quality "
            "uses native DN and separates flat regions by gray level + connected area before "
            "calculating typical/worst noise and uniformity, so different flat objects are not mixed."
        )
        self.raw_quality_note.setWordWrap(True)
        self.raw_quality_note.setStyleSheet(
            "color:#c8d0d9; background:#292d33; border-left:3px solid #79afd3; "
            "border-radius:5px; padding:8px;"
        )
        f.addRow(self.raw_quality_note)
        self.dual_image_name_a = FocusWheelComboBox()
        self.dual_image_name_b = FocusWheelComboBox()
        for name, tag in DUAL_UVC_IMAGE_NAME_OPTIONS:
            self.dual_image_name_a.addItem(name, tag)
            self.dual_image_name_b.addItem(name, tag)
        self.dual_image_custom_name_a = QLineEdit()
        self.dual_image_custom_name_b = QLineEdit()
        name_fields = []
        for combo, custom, fallback in (
            (self.dual_image_name_a, self.dual_image_custom_name_a, "Image A"),
            (self.dual_image_name_b, self.dual_image_custom_name_b, "Image B"),
        ):
            combo.addItem("Custom...", "custom")
            custom.setPlaceholderText(f"Enter a name (default: {fallback})")
            custom.setMaxLength(80)
            custom.hide()
            field = QWidget()
            field_layout = QVBoxLayout(field)
            field_layout.setContentsMargins(0, 0, 0, 0)
            field_layout.setSpacing(4)
            field_layout.addWidget(combo)
            field_layout.addWidget(custom)
            name_fields.append(field)
        self.dual_image_name_b.setCurrentIndex(1)
        f.addRow("Image A Name", name_fields[0])
        f.addRow("RAW / Image A", self.raw_quality_file_a)
        f.addRow("Image B Name", name_fields[1])
        f.addRow("RAW / Image B", self.raw_quality_file_b)
        f.addRow(self.dual_compare_controls)
        (
            self.raw_quality_score_group,
            self.raw_quality_score_labels,
        ) = self.create_quality_score_group("Image Quality Scores")
        f.addRow(self.raw_quality_score_group)
        self.dual_image_name_a.currentIndexChanged.connect(self.dual_image_name_changed)
        self.dual_image_name_b.currentIndexChanged.connect(self.dual_image_name_changed)
        self.dual_image_custom_name_a.textChanged.connect(self.dual_image_name_changed)
        self.dual_image_custom_name_b.textChanged.connect(self.dual_image_name_changed)
        self.dual_image_name_changed()
        self.stack.addWidget(p)
        self.raw_quality_file_a.path_selected.connect(
            self.auto_preview_raw_quality
        )
        self.raw_quality_file_b.path_selected.connect(
            self.auto_preview_raw_quality
        )
        for path_row in (self.raw_quality_file_a, self.raw_quality_file_b):
            path_row.edit.textChanged.connect(lambda: self.set_image_save_ready("dual", False))
        self.raw_quality_file_a.edit.returnPressed.connect(
            self.auto_preview_raw_quality
        )
        self.raw_quality_file_b.edit.returnPressed.connect(
            self.auto_preview_raw_quality
        )
        self.resolution_combo.currentIndexChanged.connect(
            self.refresh_raw_quality_settings
        )
        self.bit_depth_combo.currentIndexChanged.connect(
            self.refresh_raw_quality_settings
        )

        # Single UVC Statistic
        p, f = self._form_page()
        self.viewer_uvc_note = QLabel(
            "Clean UVC live view using Auto Detect. Double-click the image or "
            "Histogram to open a resizable analysis window."
        )
        self.viewer_uvc_note.setWordWrap(True)
        self.viewer_uvc_note.setSizePolicy(
            QSizePolicy.Ignored, QSizePolicy.Preferred
        )
        self.viewer_uvc_note.setMaximumHeight(72)
        self.viewer_uvc_note.setStyleSheet(
            "color:#c8d0d9; background:#292d33; border-left:3px solid #79afd3; "
            "border-radius:5px; padding:8px;"
        )
        self.viewer_uvc_camera_combo = FocusWheelComboBox()
        self.viewer_uvc_start_btn = QPushButton("Start Camera")
        self.viewer_uvc_start_btn.setObjectName("startButton")
        self.viewer_uvc_start_btn.clicked.connect(self.start_uvc_viewer)
        self.viewer_uvc_stop_btn = QPushButton("Stop Camera")
        self.viewer_uvc_stop_btn.setObjectName("stopButton")
        self.viewer_uvc_stop_btn.clicked.connect(self.stop_uvc_viewer)
        self.viewer_uvc_stop_btn.setEnabled(False)
        self.viewer_uvc_capture_btn = QPushButton("Capture")
        self.viewer_uvc_capture_btn.setObjectName("captureButton")
        self.viewer_uvc_capture_btn.clicked.connect(self.capture_uvc_viewer)
        self.viewer_uvc_capture_btn.setEnabled(False)
        self.viewer_uvc_record_start_btn = QPushButton("Start Recording")
        self.viewer_uvc_record_start_btn.setObjectName("startButton")
        self.viewer_uvc_record_start_btn.clicked.connect(
            self.start_uvc_viewer_recording
        )
        self.viewer_uvc_record_start_btn.setEnabled(False)
        self.viewer_uvc_record_stop_btn = QPushButton("Stop Recording")
        self.viewer_uvc_record_stop_btn.setObjectName("stopButton")
        self.viewer_uvc_record_stop_btn.clicked.connect(
            self.stop_uvc_viewer_recording
        )
        self.viewer_uvc_record_stop_btn.setEnabled(False)
        self.viewer_uvc_refresh_btn = QPushButton("Refresh Cameras")
        self.viewer_uvc_refresh_btn.clicked.connect(
            self.refresh_camera_devices
        )
        viewer_controls = QWidget()
        viewer_controls_layout = QGridLayout(viewer_controls)
        viewer_controls_layout.setContentsMargins(0, 0, 0, 0)
        viewer_controls_layout.addWidget(self.viewer_uvc_start_btn, 0, 0)
        viewer_controls_layout.addWidget(self.viewer_uvc_stop_btn, 0, 1)
        viewer_controls_layout.addWidget(
            self.viewer_uvc_capture_btn, 1, 0, 1, 2
        )
        viewer_controls_layout.addWidget(
            self.viewer_uvc_record_start_btn, 2, 0
        )
        viewer_controls_layout.addWidget(
            self.viewer_uvc_record_stop_btn, 2, 1
        )
        viewer_controls_layout.addWidget(
            self.viewer_uvc_refresh_btn, 3, 0, 1, 2
        )
        self.viewer_uvc_status = QLabel("Camera stopped")
        self.viewer_uvc_status.setWordWrap(True)
        self.viewer_uvc_status.setSizePolicy(
            QSizePolicy.Ignored, QSizePolicy.Preferred
        )
        f.addRow(self.viewer_uvc_note)
        f.addRow("CAMERA_INDEX", self.viewer_uvc_camera_combo)
        f.addRow(viewer_controls)
        f.addRow("Status", self.viewer_uvc_status)

        self.viewer_uvc_stats_group = CollapsibleGroupBox(
            "Live Gray Statistics (8-bit)"
        )
        self.viewer_uvc_stats_group.setObjectName("viewerStatsCard")
        viewer_uvc_stats_layout = QGridLayout(self.viewer_uvc_stats_group)
        viewer_uvc_stats_layout.setHorizontalSpacing(10)
        viewer_uvc_stats_layout.setVerticalSpacing(7)
        self.viewer_uvc_stat_values = {}
        viewer_uvc_stat_fields = [
            ("Pixel Count", "count"),
            ("Gray Levels", "unique_levels"),
            ("Minimum", "min"),
            ("Maximum", "max"),
            ("Mean", "mean"),
            ("Median / P50", "median"),
            ("Std Dev", "std"),
            ("Entropy", "entropy_bits"),
            ("P1", "p1"),
            ("P5", "p5"),
            ("P95", "p95"),
            ("P99", "p99"),
            ("P99 - P1", "p99_p1"),
            ("Zero Clip", "zero_clip"),
            ("255 Clip", "full_clip"),
            ("Y <= 16", "limited_low"),
            ("Y >= 235", "limited_high"),
        ]
        for index, (label_text, key) in enumerate(viewer_uvc_stat_fields):
            row = index // 2
            column = (index % 2) * 2
            name_label = QLabel(label_text)
            value_label = QLabel("--")
            value_label.setProperty("viewerStatValue", True)
            value_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            value_label.setMinimumWidth(90)
            viewer_uvc_stats_layout.addWidget(name_label, row, column)
            viewer_uvc_stats_layout.addWidget(
                value_label, row, column + 1
            )
            self.viewer_uvc_stat_values[key] = value_label
        viewer_uvc_stats_note = QLabel(
            "Updated twice per second from every displayed pixel. Clipping "
            "shows count and frame percentage; entropy is measured in bits."
        )
        viewer_uvc_stats_note.setWordWrap(True)
        viewer_uvc_stats_note.setStyleSheet(
            "color:#aeb8c2; font-weight:400; padding-top:4px;"
        )
        viewer_uvc_stats_layout.addWidget(
            viewer_uvc_stats_note,
            (len(viewer_uvc_stat_fields) + 1) // 2,
            0,
            1,
            4,
        )
        f.addRow(self.viewer_uvc_stats_group)

        self.viewer_uvc_quality_enabled = QualityScoreCheckBox()
        self.viewer_uvc_quality_enabled.setChecked(True)
        self.viewer_uvc_quality_enabled.setToolTip(
            "Disable to skip UVC spatial and temporal quality calculations"
        )
        self.viewer_uvc_quality_enabled.toggled.connect(
            self.update_uvc_viewer_quality_option
        )
        f.addRow(self.viewer_uvc_quality_enabled)
        (
            self.viewer_uvc_quality_group,
            self.viewer_uvc_quality_labels,
        ) = self.create_single_quality_score_group(
            "Live UVC Quality",
            include_temporal=True,
        )
        f.addRow(self.viewer_uvc_quality_group)
        self.viewer_uvc_quality_chart = QualityScoreTimelineChart(
            (("single", "Single UVC", "#67c5f5"),)
        )
        self.viewer_uvc_quality_chart.open_requested.connect(
            lambda: self.open_quality_timeline(
                "single_uvc",
                self.viewer_uvc_quality_chart,
                "Single UVC Total Score Timeline",
            )
        )
        self.viewer_uvc_quality_chart_group = self.create_quality_timeline_group(
            self.viewer_uvc_quality_chart
        )
        f.addRow(self.viewer_uvc_quality_chart_group)
        self.update_uvc_viewer_quality_option(True)

        self.viewer_uvc_histogram_group = CollapsibleGroupBox(
            "Live Gray Histogram + CDF"
        )
        self.viewer_uvc_histogram_group.setObjectName("viewerStatsCard")
        viewer_uvc_histogram_layout = QVBoxLayout(
            self.viewer_uvc_histogram_group
        )
        self.viewer_uvc_histogram_preview = PreviewLabel(
            "Live Gray Histogram"
        )
        self.viewer_uvc_histogram_preview.setMinimumSize(320, 180)
        self.viewer_uvc_histogram_preview.setToolTip(
            "Double-click to open a resizable and maximizable histogram window"
        )
        self.viewer_uvc_histogram_preview.double_clicked.connect(
            self.open_uvc_histogram_window
        )
        viewer_uvc_histogram_layout.addWidget(
            self.viewer_uvc_histogram_preview
        )
        histogram_note = QLabel(
            "Blue bars: 256-bin gray histogram. Yellow line: cumulative "
            "distribution (CDF). Markers show P1, P50, and P99."
        )
        histogram_note.setWordWrap(True)
        histogram_note.setStyleSheet(
            "color:#aeb8c2; font-weight:400; padding-top:3px;"
        )
        viewer_uvc_histogram_layout.addWidget(histogram_note)
        f.addRow(self.viewer_uvc_histogram_group)
        self.stack.addWidget(p)
        self.viewer_uvc_camera_combo.currentIndexChanged.connect(
            self.switch_uvc_viewer_setting
        )
        self.resolution_combo.currentIndexChanged.connect(
            self.switch_uvc_viewer_setting
        )

        # Dual UVC Compare
        p, f = self._form_page()
        self.dual_uvc_note = QLabel(
            "Camera A and Camera B run in independent capture threads. Capture "
            "and recording pair the latest available frames in software; this is "
            "not hardware-trigger synchronization. Camera indexes must differ."
        )
        self.dual_uvc_note.setWordWrap(True)
        self.dual_uvc_note.setStyleSheet(
            "color:#c8d0d9; background:#292d33; border-left:3px solid #79afd3; "
            "border-radius:5px; padding:8px;"
        )
        f.addRow(self.dual_uvc_note)

        self.dual_uvc_group_a = CollapsibleGroupBox("UVC Camera A")
        self.dual_uvc_group_a.setObjectName("dualUvcCameraGroup")
        dual_form_a = QFormLayout(self.dual_uvc_group_a)
        dual_form_a.setContentsMargins(12, 34, 12, 12)
        dual_form_a.setHorizontalSpacing(10)
        dual_form_a.setVerticalSpacing(8)
        self.dual_uvc_image_name_a = FocusWheelComboBox()
        self.dual_uvc_image_name_b = FocusWheelComboBox()
        for image_name, file_tag in DUAL_UVC_IMAGE_NAME_OPTIONS:
            self.dual_uvc_image_name_a.addItem(image_name, file_tag)
            self.dual_uvc_image_name_b.addItem(image_name, file_tag)
        self.dual_uvc_image_name_a.setCurrentIndex(0)
        self.dual_uvc_image_name_b.setCurrentIndex(1)
        self.dual_uvc_camera_a = FocusWheelComboBox()
        self.dual_uvc_camera_b = FocusWheelComboBox()
        self.dual_uvc_status_a = QLabel("Camera A stopped")
        self.dual_uvc_status_a.setWordWrap(True)
        dual_form_a.addRow("Image Name", self.dual_uvc_image_name_a)
        dual_form_a.addRow("CAMERA_INDEX", self.dual_uvc_camera_a)
        dual_form_a.addRow("Status", self.dual_uvc_status_a)

        self.dual_uvc_group_b = CollapsibleGroupBox("UVC Camera B")
        self.dual_uvc_group_b.setObjectName("dualUvcCameraGroup")
        dual_form_b = QFormLayout(self.dual_uvc_group_b)
        dual_form_b.setContentsMargins(12, 34, 12, 12)
        dual_form_b.setHorizontalSpacing(10)
        dual_form_b.setVerticalSpacing(8)
        self.dual_uvc_status_b = QLabel("Camera B stopped")
        self.dual_uvc_status_b.setWordWrap(True)
        dual_form_b.addRow("Image Name", self.dual_uvc_image_name_b)
        dual_form_b.addRow("CAMERA_INDEX", self.dual_uvc_camera_b)
        dual_form_b.addRow("Status", self.dual_uvc_status_b)
        f.addRow(self.dual_uvc_group_a)
        f.addRow(self.dual_uvc_group_b)
        self.dual_uvc_quality_enabled = QualityScoreCheckBox()
        self.dual_uvc_quality_enabled.setChecked(True)
        self.dual_uvc_quality_enabled.setToolTip(
            "Disable to skip both cameras' spatial and temporal quality calculations"
        )
        self.dual_uvc_quality_enabled.toggled.connect(
            self.update_dual_uvc_quality_option
        )
        f.addRow(self.dual_uvc_quality_enabled)
        (
            self.dual_uvc_quality_group,
            self.dual_uvc_quality_labels,
        ) = self.create_quality_score_group(
            "Live Relative Quality Scores", include_temporal=True
        )
        f.addRow(self.dual_uvc_quality_group)
        self.dual_uvc_quality_chart = QualityScoreTimelineChart(
            (
                ("camera_a", "Camera A", "#67c5f5"),
                ("camera_b", "Camera B", "#f1b765"),
            )
        )
        self.dual_uvc_quality_chart.open_requested.connect(
            lambda: self.open_quality_timeline(
                "dual_uvc",
                self.dual_uvc_quality_chart,
                "Dual UVC Total Score Timeline",
            )
        )
        self.dual_uvc_quality_chart_group = self.create_quality_timeline_group(
            self.dual_uvc_quality_chart
        )
        f.addRow(self.dual_uvc_quality_chart_group)
        self.update_dual_uvc_quality_option(True)

        dual_controls = QWidget()
        dual_controls_layout = QGridLayout(dual_controls)
        dual_controls_layout.setContentsMargins(0, 0, 0, 0)
        self.dual_uvc_start_btn = QPushButton("Start Cameras")
        self.dual_uvc_start_btn.setObjectName("startButton")
        self.dual_uvc_start_btn.clicked.connect(self.start_dual_uvc_cameras)
        self.dual_uvc_stop_btn = QPushButton("Stop Cameras")
        self.dual_uvc_stop_btn.setObjectName("stopButton")
        self.dual_uvc_stop_btn.clicked.connect(self.stop_dual_uvc_cameras)
        self.dual_uvc_stop_btn.setEnabled(False)
        self.dual_uvc_capture_btn = QPushButton("Capture Pair")
        self.dual_uvc_capture_btn.setObjectName("captureButton")
        self.dual_uvc_capture_btn.clicked.connect(self.capture_dual_uvc_pair)
        self.dual_uvc_capture_btn.setEnabled(False)
        self.dual_uvc_record_start_btn = QPushButton("Start Recording")
        self.dual_uvc_record_start_btn.setObjectName("startButton")
        self.dual_uvc_record_start_btn.clicked.connect(
            self.start_dual_uvc_recording
        )
        self.dual_uvc_record_start_btn.setEnabled(False)
        self.dual_uvc_record_stop_btn = QPushButton("Stop Recording")
        self.dual_uvc_record_stop_btn.setObjectName("stopButton")
        self.dual_uvc_record_stop_btn.clicked.connect(
            self.stop_dual_uvc_recording
        )
        self.dual_uvc_record_stop_btn.setEnabled(False)
        self.dual_uvc_refresh_btn = QPushButton("Refresh Cameras")
        self.dual_uvc_refresh_btn.clicked.connect(self.refresh_camera_devices)
        dual_controls_layout.addWidget(self.dual_uvc_start_btn, 0, 0)
        dual_controls_layout.addWidget(self.dual_uvc_stop_btn, 0, 1)
        dual_controls_layout.addWidget(self.dual_uvc_capture_btn, 1, 0, 1, 2)
        dual_controls_layout.addWidget(self.dual_uvc_record_start_btn, 2, 0)
        dual_controls_layout.addWidget(self.dual_uvc_record_stop_btn, 2, 1)
        dual_controls_layout.addWidget(self.dual_uvc_refresh_btn, 3, 0, 1, 2)
        f.insertRow(3, dual_controls)
        self.stack.addWidget(p)

        self.dual_uvc_camera_a.currentIndexChanged.connect(
            self.dual_uvc_setting_changed
        )
        self.dual_uvc_camera_b.currentIndexChanged.connect(
            self.dual_uvc_setting_changed
        )
        self.dual_uvc_image_name_a.currentIndexChanged.connect(
            self.dual_uvc_image_name_changed
        )
        self.dual_uvc_image_name_b.currentIndexChanged.connect(
            self.dual_uvc_image_name_changed
        )
        self.resolution_combo.currentIndexChanged.connect(
            self.dual_uvc_setting_changed
        )
        self.dual_uvc_image_name_changed()

        # Single Image Algorithm Compare
        p, f = self._form_page()
        self.raw_algorithm_file = PathRow(
            "file",
            IMAGE_INPUT_FILTER,
            "Select RAW / Image File for Algorithm Compare",
        )
        image_input_note = QLabel(IMAGE_INPUT_NOTE)
        image_input_note.setWordWrap(True)
        image_input_note.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        f.addRow(image_input_note)
        self.raw_input_mapping_combo = FocusWheelComboBox()
        for mapping_name, mapping_key in RAW_INPUT_MAPPING_OPTIONS:
            self.raw_input_mapping_combo.addItem(mapping_name, mapping_key)
        self.raw_input_mapping_combo.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Fixed
        )
        self.raw_input_mapping_note = QLabel(
            "Input Mapping runs before the display pipeline. With Fixed Bit Depth "
            "and Fieldscale as the first enabled step, Fieldscale reads native grayscale values "
            "(RAW: original DN); following steps are 8-bit. Image bit depth is detected automatically."
        )
        self.raw_input_mapping_note.setWordWrap(True)
        self.raw_input_mapping_note.setSizePolicy(
            QSizePolicy.Ignored, QSizePolicy.Preferred
        )
        self.raw_input_mapping_note.setStyleSheet("color:#c98;")
        self.raw_pipeline_editor = AlgorithmPipelineEditor(
            allow_step_algorithm_change=False
        )
        self.raw_preview_stage_combo = FocusWheelComboBox()
        self.raw_preview_stage_combo.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Fixed
        )
        f.addRow("RAW / Image File", self.raw_algorithm_file)
        f.addRow("Input Mapping", self.raw_input_mapping_combo)
        f.addRow(self.raw_input_mapping_note)
        f.addRow(self.raw_pipeline_editor)
        f.addRow("Right Preview", self.raw_preview_stage_combo)
        self.algorithm_compare_mode = FocusWheelComboBox()
        self.algorithm_compare_mode.addItems(["Side by Side", "Swipe Compare"])
        f.addRow("Comparison View", self.algorithm_compare_mode)
        self.algorithm_name_fields = []
        for label in ("Image A Name", "Image B Name"):
            combo = FocusWheelComboBox()
            combo.addItem("Automatic", "auto")
            for name, tag in DUAL_UVC_IMAGE_NAME_OPTIONS:
                combo.addItem(name, tag)
            combo.addItem("Custom...", "custom")
            custom = QLineEdit()
            custom.setMaxLength(80)
            custom.setPlaceholderText("Enter a name (blank uses automatic name)")
            custom.hide()
            field = QWidget()
            layout = QVBoxLayout(field)
            layout.setContentsMargins(0, 0, 0, 0)
            layout.addWidget(combo)
            layout.addWidget(custom)
            f.addRow(label, field)
            self.algorithm_name_fields.append((combo, custom))
            combo.currentIndexChanged.connect(self.algorithm_image_name_changed)
            custom.textChanged.connect(self.algorithm_image_name_changed)
        self.algorithm_compare_mode.currentIndexChanged.connect(self.update_dual_compare_mode)

        (
            self.raw_algorithm_quality_group,
            self.raw_algorithm_quality_labels,
        ) = self.create_quality_score_group(
            "Image Quality: Original (A) vs Final Pipeline (B)"
        )
        f.addRow(self.raw_algorithm_quality_group)
        self.stack.addWidget(p)
        self.raw_algorithm_file.edit.textChanged.connect(
            self.invalidate_raw_pipeline_source
        )
        self.raw_algorithm_file.path_selected.connect(
            self.auto_preview_raw_algorithm
        )
        self.raw_algorithm_file.edit.returnPressed.connect(
            self.auto_preview_raw_algorithm
        )
        self.raw_input_mapping_combo.currentIndexChanged.connect(
            self.update_raw_pipeline_ui
        )
        self.raw_pipeline_editor.pipeline_changed.connect(
            self.update_raw_pipeline_ui
        )
        self.raw_preview_stage_combo.currentIndexChanged.connect(
            self.update_raw_stage_preview
        )
        self.resolution_combo.currentIndexChanged.connect(
            self.refresh_raw_algorithm_source_settings
        )
        self.bit_depth_combo.currentIndexChanged.connect(
            self.update_raw_pipeline_ui
        )
        self.update_raw_pipeline_ui()

        # UVC Algorithm Compare
        p, f = self._form_page()
        self.uvc_camera_combo = FocusWheelComboBox()
        self.uvc_pipeline_editor = AlgorithmPipelineEditor(
            allow_step_algorithm_change=False,
            fieldscale_temporal=True,
        )
        self.uvc_preview_stage_combo = FocusWheelComboBox()
        self.uvc_preview_stage_combo.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Fixed
        )
        self.uvc_camera_combo.currentIndexChanged.connect(self.switch_uvc_camera)
        self.bit_depth_combo.currentIndexChanged.connect(
            self.mark_uvc_processing_config_changed
        )
        self.uvc_pipeline_editor.pipeline_changed.connect(
            self.update_uvc_pipeline_ui
        )
        self.uvc_preview_stage_combo.currentIndexChanged.connect(
            self.update_uvc_stage_preview
        )
        controls = QWidget()
        controls_layout = QGridLayout(controls)
        controls_layout.setContentsMargins(0, 0, 0, 0)
        self.uvc_start_btn = QPushButton("Start Camera")
        self.uvc_start_btn.setObjectName("startButton")
        self.uvc_start_btn.clicked.connect(self.start_uvc_camera)
        self.uvc_stop_btn = QPushButton("Stop Camera")
        self.uvc_stop_btn.setObjectName("stopButton")
        self.uvc_stop_btn.clicked.connect(self.stop_uvc_camera)
        self.uvc_stop_btn.setEnabled(False)
        self.uvc_capture_btn = QPushButton("Capture Pair")
        self.uvc_capture_btn.setObjectName("captureButton")
        self.uvc_capture_btn.clicked.connect(self.capture_uvc_pair)
        self.uvc_capture_btn.setEnabled(False)
        self.uvc_record_start_btn = QPushButton("Start Recording")
        self.uvc_record_start_btn.setObjectName("startButton")
        self.uvc_record_start_btn.setToolTip(
            "Record synchronized Original UVC and Final Pipeline Output "
            "side by side as MJPG AVI"
        )
        self.uvc_record_start_btn.clicked.connect(self.start_uvc_recording)
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn = QPushButton("Stop Recording")
        self.uvc_record_stop_btn.setObjectName("stopButton")
        self.uvc_record_stop_btn.clicked.connect(self.stop_uvc_recording)
        self.uvc_record_stop_btn.setEnabled(False)
        self.uvc_refresh_btn = QPushButton("Refresh Cameras")
        self.uvc_refresh_btn.clicked.connect(self.refresh_camera_devices)
        controls_layout.addWidget(self.uvc_start_btn, 0, 0)
        controls_layout.addWidget(self.uvc_stop_btn, 0, 1)
        controls_layout.addWidget(self.uvc_capture_btn, 1, 0, 1, 2)
        controls_layout.addWidget(self.uvc_record_start_btn, 2, 0)
        controls_layout.addWidget(self.uvc_record_stop_btn, 2, 1)
        controls_layout.addWidget(self.uvc_refresh_btn, 3, 0, 1, 2)
        self.uvc_status = QLabel("Camera stopped")
        self.uvc_status.setWordWrap(True)
        self.uvc_status.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        f.addRow("CAMERA_INDEX", self.uvc_camera_combo)
        f.addRow(controls)
        f.addRow("Status", self.uvc_status)
        f.addRow(self.uvc_pipeline_editor)
        f.addRow("Right Preview", self.uvc_preview_stage_combo)
        self.uvc_quality_enabled = QualityScoreCheckBox()
        self.uvc_quality_enabled.setChecked(True)
        self.uvc_quality_enabled.setToolTip(
            "Disable to skip Original and Final spatial and temporal quality calculations"
        )
        self.uvc_quality_enabled.toggled.connect(
            self.update_uvc_algorithm_quality_option
        )
        f.addRow(self.uvc_quality_enabled)
        (
            self.uvc_quality_group,
            self.uvc_quality_labels,
        ) = self.create_quality_score_group(
            "Live Quality: Original (A) vs Final Pipeline (B)",
            include_temporal=True,
        )
        f.addRow(self.uvc_quality_group)
        self.uvc_quality_chart = QualityScoreTimelineChart(
            (
                ("original", "Original", "#67c5f5"),
                ("final", "Final", "#76d39b"),
            )
        )
        self.uvc_quality_chart.open_requested.connect(
            lambda: self.open_quality_timeline(
                "uvc_algorithm",
                self.uvc_quality_chart,
                "UVC Algorithm Total Score Timeline",
            )
        )
        self.uvc_quality_chart_group = self.create_quality_timeline_group(
            self.uvc_quality_chart
        )
        f.addRow(self.uvc_quality_chart_group)
        self.update_uvc_algorithm_quality_option(True)
        self.stack.addWidget(p)
        self.update_uvc_pipeline_ui()
    def update_nuc_mode(self):
        folder_mode = self.nuc_mode.currentIndex() == 0
        for w in (self.low_folder, self.high_folder): w.setEnabled(folder_mode)
        for w in (self.low_single, self.high_single): w.setEnabled(not folder_mode)

    def update_bpr_marking_ui(self, _checked=None):
        self.set_form_parameter_visible(
            self.bpr_form,
            self.bpr_raw_file,
            self.bpr_use_raw.isChecked(),
        )
        self.set_form_parameter_visible(
            self.bpr_form,
            self.bpr_highlight_radius,
            self.bpr_highlight_bright_points.isChecked(),
        )

    def update_bpr_coordinate_note(self, _index=None):
        self.bpr_note.setText(
            "座標方向：左上角為原點 (0,0)；X 由左往右增加，"
            f"有效範圍 0..{self.width - 1}；Y 由上往下增加，"
            f"有效範圍 0..{self.height - 1}。座標以 (x,y) 表示，"
            "對應程式陣列 RAW[y,x]。例如 (600,79) 是從左側第 601 欄、"
            "從上方第 80 列。Novatek BIN 使用 SIE row 格式；CSV/TXT 使用 x,y；"
            "重複座標只圈選一次。"
        )

    def clear_bpr_coordinate_info(self, _value=None):
        self.bpr_coordinate_info.clear()

    def set_form_parameter_visible(self, form, widget, visible):
        widget.setVisible(visible)
        label = form.labelForField(widget)
        if label is not None:
            label.setVisible(visible)

    def rebuild_pipeline_preview_combo(self, editor, combo):
        previous_data = combo.currentData()
        combo.blockSignals(True)
        combo.clear()
        combo.addItem("Final Pipeline Output", None)
        for pipeline_index, stage in enumerate(editor.stages(), start=1):
            if stage.enabled:
                combo.addItem(
                    f"After Step {pipeline_index}: "
                    f"{algorithm_display_name(stage.algorithm_key)}",
                    pipeline_index,
                )
        target_index = combo.findData(previous_data)
        combo.setCurrentIndex(target_index if target_index >= 0 else 0)
        combo.blockSignals(False)

    def invalidate_raw_pipeline_source(self, _value=None):
        self.raw_algorithm_source = None
        self.raw_pipeline_input = None
        self.raw_pipeline_input_max = 255
        self.raw_pipeline_final = None
        self.raw_pipeline_records = []
        self.raw_algorithm_quality_original = None
        self.raw_algorithm_quality_final = None
        if hasattr(self, "raw_algorithm_quality_labels"):
            self.update_quality_score_labels(
                self.raw_algorithm_quality_labels, None, None
            )
        if self.current_function_index() == len(self.FUNCTIONS) - 2:
            self.current_original = None
            self.current_result = None
            self.preview_original.set_array(None)
            self.preview_result.set_array(None)
            self.run_btn.setEnabled(False)

    def refresh_raw_algorithm_source_settings(self, _value=None):
        self.invalidate_raw_pipeline_source()
        if (
            self.current_function_index()
            == self.FUNCTIONS.index("Image Algorithm Compare")
            and self.raw_algorithm_file.text()
        ):
            self.auto_preview_raw_algorithm()

    def load_analysis_source(self, path):
        pixels, info = load_analysis_image(path, self.width, self.height, self.bit_depth)
        self.log.append(
            f"[{info['source_type'].title()}] Loaded: {path}\n"
            f"{info['width']}x{info['height']}, "
            f"{info['bit_depth']}-bit {info['value_name']} (0..{info['value_max']}). "
            + ("Image dimensions/depth are automatic; RAW settings are ignored. Color is converted to grayscale." if info['source_type'] == "image" else "")
        )
        return pixels, info

    def sync_image_input_mapping(self):
        combo = self.raw_input_mapping_combo
        selected = combo.currentData()
        if selected in ("fixed_scale", "raw_percentile_preview"):
            self._preferred_raw_mapping = selected
        preferred = getattr(self, "_preferred_raw_mapping", "fixed_scale")
        first = next((stage for stage in self.raw_pipeline_editor.stages() if stage.enabled), None)
        first_key = first.algorithm_key if first is not None else None
        info = getattr(self, "raw_algorithm_info", {}) if self.raw_algorithm_source is not None else {}
        if info.get("source_type") == "image":
            options = [("None (Image Input)", "image_input")]
            note = "Image input uses grayscale values without percentile mapping. Native precision is kept for a first Fieldscale or Percentile Auto Contrast step; otherwise 16-bit input is scaled to 8-bit automatically."
        elif first_key == "fieldscale":
            options = [("None (Fieldscale)", "pipeline_fieldscale")]
            note = "Fieldscale receives original grayscale values. Input Mapping is bypassed."
        elif first_key == "auto_contrast":
            options = [("Percentile Auto Contrast (Pipeline)", "pipeline_auto_contrast")]
            note = "The first Percentile Auto Contrast step receives original grayscale values and applies its Step Parameters once."
        else:
            options = RAW_INPUT_MAPPING_OPTIONS
            note = "Mapping runs before the pipeline, including when the pipeline is empty. Left: fixed-scale original. Right: mapping and pipeline result."
        combo.blockSignals(True)
        combo.clear()
        for label, key in options:
            combo.addItem(label, key)
        combo.setCurrentIndex(max(0, combo.findData(preferred)))
        combo.setEnabled(options is RAW_INPUT_MAPPING_OPTIONS)
        combo.blockSignals(False)
        self.raw_input_mapping_note.setText(note)
        return first_key

    def prepare_raw_pipeline_input(self, raw):
        first_key = self.sync_image_input_mapping()
        self.update_image_preview_titles()
        info = self.raw_algorithm_info
        if first_key in ("fieldscale", "auto_contrast"):
            return (
                raw.copy(), info["value_max"],
                f"Original {info['value_name']} 0..{info['value_max']} -> first {algorithm_display_name(first_key)} step; no separate mapping",
            )
        if self.raw_input_mapping_combo.currentData() == "raw_percentile_preview":
            mapped, p1, p99 = percentile_preview_u8(raw)
            return mapped, 255, f"Original {info['value_name']} P1={p1:.1f}, P99={p99:.1f} -> 8-bit"
        return (
            raw_to_display_u8(raw, info["value_max"]), 255,
            f"Fixed 0..{info['value_max']} {info['value_name']} -> 0..255"
            + (" (automatic image format conversion; no percentile mapping)" if info["source_type"] == "image" else ""),
        )

    def update_raw_pipeline_ui(self, _value=None):
        self.sync_image_input_mapping()
        self.raw_algorithm_quality_original = None
        self.raw_algorithm_quality_final = None
        if hasattr(self, "raw_algorithm_quality_labels"):
            self.update_quality_score_labels(
                self.raw_algorithm_quality_labels, None, None
            )
        self.rebuild_pipeline_preview_combo(
            self.raw_pipeline_editor, self.raw_preview_stage_combo
        )
        if self.current_function_index() == len(self.FUNCTIONS) - 2:
            self.run_btn.setEnabled(
                self.raw_algorithm_source is not None
            )
        self.refresh_raw_pipeline_preview()
        self.update_raw_stage_preview()

    def refresh_raw_pipeline_preview(self):
        if self.raw_algorithm_source is None:
            return False
        try:
            pipeline_input, input_max, _mapping_details = (
                self.prepare_raw_pipeline_input(self.raw_algorithm_source)
            )
            final_image, records = apply_algorithm_pipeline(
                pipeline_input,
                self.raw_pipeline_editor.stages(),
                input_max=input_max, allow_empty=True,
                pad_fieldscale_input=self.raw_algorithm_info["source_type"] == "image",
            )
            quality_original = thermal_raw_quality_analysis(
                self.raw_algorithm_source, self.raw_algorithm_info["value_max"]
            )
            quality_final = thermal_raw_quality_analysis(final_image, 255)
        except Exception as exc:
            self.raw_pipeline_input = None
            self.raw_pipeline_input_max = 255
            self.raw_pipeline_final = None
            self.raw_pipeline_records = []
            self.current_result = None
            if self.current_function_index() == len(self.FUNCTIONS) - 2:
                self.preview_result.set_array(None)
                self.run_btn.setEnabled(False)
            self.log.append(f"RAW pipeline preview error: {exc}")
            return False
        self.raw_pipeline_input = pipeline_input
        self.raw_pipeline_input_max = input_max
        self.raw_pipeline_final = final_image
        self.raw_pipeline_records = records
        self.raw_algorithm_quality_original = quality_original
        self.raw_algorithm_quality_final = quality_final
        self.update_quality_score_labels(
            self.raw_algorithm_quality_labels,
            quality_original,
            quality_final,
        )
        self.current_original = raw_to_display_u8(
            self.raw_algorithm_source, self.raw_algorithm_info["value_max"]
        )
        if self.current_function_index() == len(self.FUNCTIONS) - 2:
            self.preview_original.set_array(self.current_original)
            self.run_btn.setEnabled(
                self.raw_algorithm_source is not None
            )
        self.update_raw_stage_preview()
        return True

    def update_raw_stage_preview(self, _combo_index=None):
        if self.raw_pipeline_final is None:
            return
        selected_pipeline_index = self.raw_preview_stage_combo.currentData()
        preview = pipeline_record_image(
            self.raw_pipeline_records,
            selected_pipeline_index,
            self.raw_pipeline_final,
        )
        self.current_result = preview
        if self.current_function_index() == len(self.FUNCTIONS) - 2:
            title = self.algorithm_image_names()[1]
            self.preview_result.set_title(title)
            self.preview_result_caption.setText(title)
            self.preview_result.set_array(preview)

    def category_name_for_function(self, function_name):
        return next(
            category_name
            for category_name, function_names in self.FUNCTION_GROUPS
            if function_name in function_names
        )

    def select_function_from_menu(self, function_index):
        function_index = int(function_index)
        if function_index == self.current_function_index():
            self.function_actions[function_index].setChecked(True)
            return
        self.change_page(function_index)

    def current_function_index(self):
        return int(self.active_function_index)

    def current_uvc_recording_view(self):
        function_index = self.current_function_index()
        viewer_uvc_page_index = self.FUNCTIONS.index("Single UVC Statistic")
        dual_uvc_page_index = len(self.FUNCTIONS) - 3
        uvc_page_index = len(self.FUNCTIONS) - 1
        if function_index == viewer_uvc_page_index:
            if self.viewer_uvc_frame is None:
                return None, None
            return (
                cv2.cvtColor(
                    self.viewer_uvc_frame, cv2.COLOR_GRAY2BGR
                ),
                "Single UVC Statistic",
            )
        if function_index == dual_uvc_page_index:
            if self.dual_uvc_frame_a is None or self.dual_uvc_frame_b is None:
                return None, None
            label_a, label_b = self.dual_uvc_labels()
            return (
                UvcRecordingThread._compose_frame(
                    self.dual_uvc_frame_a,
                    self.dual_uvc_frame_b,
                    label_a,
                    label_b,
                ),
                "Dual UVC Recording View",
            )
        if function_index == uvc_page_index:
            if self.uvc_original_frame is None or self.uvc_processed_frame is None:
                return None, None
            return (
                UvcRecordingThread._compose_frame(
                    self.uvc_original_frame,
                    self.uvc_processed_frame,
                    "Original UVC",
                    "Final Pipeline Output",
                ),
                "UVC Pipeline Recording View",
            )
        return None, None

    def open_preview_window(self, preview):
        function_name = self.FUNCTIONS[self.current_function_index()]
        if self.category_name_for_function(function_name) != "Image":
            self.open_uvc_video_window()
            return
        if preview._pixmap is None:
            return
        key = "dual_image" if function_name == "Dual Image Compare" else preview
        if function_name == "Image Algorithm Compare":
            key = "image_algorithm"
        window = self.image_preview_windows.get(key)
        if window is None:
            window = UvcVideoWindow(self)
            if key in ("dual_image", "image_algorithm"):
                window.swipe_preview = SwipeCompareWidget()
                window.layout().addWidget(window.swipe_preview)
                window.swipe_preview.hide()
                window.swipe_preview.position_changed.connect(
                    lambda position, key=key: self.sync_dual_swipe_position(position, key)
                )
            self.image_preview_windows[key] = window
        window.show()
        self.update_image_preview_window(preview)
        window.raise_()
        window.activateWindow()

    def update_image_preview_titles(self):
        name = self.FUNCTIONS[self.current_function_index()]
        def source_label(info):
            return "RAW" if info.get("source_type") == "raw" else "Image"
        if name == "Single Image Statistics":
            kind = source_label(getattr(self, "raw_viewer_info", {}))
            titles = (None, f"{kind} Preview (Fixed Scale)")
        elif name == "Image Algorithm Compare":
            kind = source_label(getattr(self, "raw_algorithm_info", {}))
            titles = self.algorithm_image_names()
        elif name == "Dual Image Compare":
            names = self.dual_image_names()
            infos = getattr(self, "dual_image_source_infos", ({}, {}))
            titles = tuple(f"{label} [{source_label(info)}]" for label, info in zip(names, infos))
        else:
            return
        for preview, caption, title in zip(
            (self.preview_original, self.preview_result),
            (self.preview_original_caption, self.preview_result_caption), titles
        ):
            if title is not None:
                preview.set_title(title)
                caption.setText(title)

    def dual_image_names(self):
        return tuple(
            (custom.text().strip() or fallback)
            if combo.currentData() == "custom" else combo.currentText()
            for combo, custom, fallback in (
                (self.dual_image_name_a, self.dual_image_custom_name_a, "Image A"),
                (self.dual_image_name_b, self.dual_image_custom_name_b, "Image B"),
            )
        )

    def dual_image_name_changed(self, _index=None):
        for combo, custom in (
            (self.dual_image_name_a, self.dual_image_custom_name_a),
            (self.dual_image_name_b, self.dual_image_custom_name_b),
        ):
            custom.setVisible(combo.currentData() == "custom")
        names = self.dual_image_names()
        for label, name in zip(self.raw_quality_score_group.comparison_header_labels, names):
            label.setText(name)
        if self.FUNCTIONS[self.current_function_index()] == "Dual Image Compare":
            for preview, caption, name in zip(
                (self.preview_original, self.preview_result),
                (self.preview_original_caption, self.preview_result_caption), names
            ):
                preview.set_title(name)
                caption.setText(name)
            self.update_image_preview_titles()
            self.update_image_preview_window(self.preview_original)

    def algorithm_image_names(self):
        kind = "RAW" if getattr(self, "raw_algorithm_info", {}).get("source_type") == "raw" else "Image"
        defaults = (f"Original {kind} (Fixed Scale)",
                    self.raw_preview_stage_combo.currentText() or "Final Pipeline")
        return tuple(
            (custom.text().strip() or default) if combo.currentData() == "custom"
            else default if combo.currentData() == "auto" else combo.currentText()
            for (combo, custom), default in zip(self.algorithm_name_fields, defaults)
        )

    def algorithm_image_name_changed(self, _value=None):
        for combo, custom in self.algorithm_name_fields:
            custom.setVisible(combo.currentData() == "custom")
        if self.FUNCTIONS[self.current_function_index()] == "Image Algorithm Compare":
            self.update_image_preview_titles()
            self.update_image_preview_window(self.preview_original)

    def image_comparison_key(self):
        return "image_algorithm" if self.FUNCTIONS[self.current_function_index()] == "Image Algorithm Compare" else "dual_image"

    def image_comparison_mode(self):
        return self.algorithm_compare_mode if self.image_comparison_key() == "image_algorithm" else self.dual_compare_mode

    def sync_dual_swipe_position(self, position, key=None):
        key = key or self.image_comparison_key()
        if not hasattr(self, "image_swipe_positions"):
            self.image_swipe_positions = {}
        self.image_swipe_positions[key] = position
        if key == self.image_comparison_key():
            self.dual_swipe_preview.set_position(position)
        window = self.image_preview_windows.get(key)
        if window is not None:
            window.swipe_preview.set_position(position)

    def update_dual_compare_mode(self, _index=None):
        name = self.FUNCTIONS[self.current_function_index()]
        is_comparison = name in ("Dual Image Compare", "Image Algorithm Compare")
        swipe = is_comparison and self.image_comparison_mode().currentIndex() == 1
        self.dual_compare_controls.setVisible(name == "Dual Image Compare")
        self.dual_swipe_preview.setVisible(swipe)
        self.dual_swipe_preview.set_position(
            getattr(self, "image_swipe_positions", {}).get(self.image_comparison_key(), 0.5)
        )
        show_original = self.PREVIEW_LAYOUTS[self.current_function_index()][0]
        self.preview_original_container.setVisible(show_original and not swipe)
        self.preview_result_container.setVisible(not swipe)
        if is_comparison:
            self.update_image_preview_window(self.preview_original)

    def save_dual_swipe_preview(self, path):
        # Render at source canvas resolution, independent of preview window size.
        export = SwipeCompareWidget()
        export.setMinimumSize(1, 1)
        export.resize(max(frame.width() for frame in self.dual_swipe_preview.frames),
                      max(frame.height() for frame in self.dual_swipe_preview.frames))
        export.set_frames(self.dual_swipe_preview.frames, self.dual_swipe_preview.names)
        export.set_position(self.dual_swipe_preview.position)
        pixmap = QPixmap(export.size())
        export.render(pixmap)
        if not pixmap.save(str(path), "PNG"):
            raise OSError(f"Could not save comparison: {path}")

    def compose_dual_image_preview(self, frames, names=None):
        panel_height = max(frame.shape[0] for frame in frames)
        panel_width = max(frame.shape[1] for frame in frames)
        fitted_frames = []
        for frame in frames:
            scale = min(panel_width / frame.shape[1], panel_height / frame.shape[0])
            width = max(1, min(panel_width, round(frame.shape[1] * scale)))
            height = max(1, min(panel_height, round(frame.shape[0] * scale)))
            resized = cv2.resize(frame, (width, height), interpolation=cv2.INTER_AREA)
            canvas = np.zeros((panel_height, panel_width), dtype=np.uint8)
            x = (panel_width - width) // 2
            y = (panel_height - height) // 2
            canvas[y:y + height, x:x + width] = resized
            fitted_frames.append(canvas)
        return UvcRecordingThread._compose_frame(
            *fitted_frames, *(self.dual_image_names() if names is None else names)
        )

    def update_image_preview_window(self, preview):
        function_name = self.FUNCTIONS[self.current_function_index()]
        is_dual = function_name == "Dual Image Compare"
        is_algorithm = function_name == "Image Algorithm Compare"
        if is_dual or is_algorithm:
            pixmaps = (self.preview_original._pixmap, self.preview_result._pixmap)
            names = (self.preview_original.title, self.preview_result.title)
            self.dual_swipe_preview.set_frames(pixmaps, names)
        key = "image_algorithm" if is_algorithm else "dual_image" if is_dual else preview
        window = self.image_preview_windows.get(key)
        if window is None or not window.isVisible():
            return
        if is_dual or is_algorithm:
            swipe = self.image_comparison_mode().currentIndex() == 1
            window.image_label.setVisible(not swipe)
            window.swipe_preview.setVisible(swipe)
            window.swipe_preview.set_frames(pixmaps, names)
            window.swipe_preview.set_position(self.dual_swipe_preview.position)
            window.setWindowTitle(function_name)
            if swipe:
                return
        if is_dual or is_algorithm:
            frames = []
            for panel in (self.preview_original, self.preview_result):
                if panel._pixmap is None:
                    window.clear_frame()
                    return
                image = panel._pixmap.toImage().convertToFormat(QImage.Format_Grayscale8)
                buffer = image.bits()
                buffer.setsize(image.sizeInBytes())
                frames.append(
                    np.frombuffer(buffer, dtype=np.uint8)
                    .reshape(image.height(), image.bytesPerLine())[:, :image.width()].copy()
                )
            window.setWindowTitle(function_name)
            names = (
                (self.preview_original.title, self.preview_result.title)
                if is_algorithm else self.dual_image_names()
            )
            window.set_bgr_frame(self.compose_dual_image_preview(frames, names))
            return
        if preview._pixmap is None:
            window.clear_frame()
            return
        window.setWindowTitle(f"{function_name} - {preview.title}")
        window._pixmap = QPixmap(preview._pixmap)
        window._rescale()

    def open_uvc_video_window(self):
        function_index = self.current_function_index()
        uvc_page_indexes = {
            self.FUNCTIONS.index("Single UVC Statistic"),
            len(self.FUNCTIONS) - 3,
            len(self.FUNCTIONS) - 1,
        }
        if function_index not in uvc_page_indexes:
            return
        frame, title = self.current_uvc_recording_view()
        if frame is None:
            return
        if self.uvc_video_window is None:
            self.uvc_video_window = UvcVideoWindow(self)
        self.uvc_video_window.setWindowTitle(title)
        self.uvc_video_window.set_bgr_frame(frame)
        self.uvc_video_window.show()
        self.uvc_video_window.raise_()
        self.uvc_video_window.activateWindow()

    def update_uvc_video_window(self):
        window = self.uvc_video_window
        if window is None or not window.isVisible():
            return
        frame, title = self.current_uvc_recording_view()
        if frame is None:
            return
        window.setWindowTitle(title)
        window.set_bgr_frame(frame)

    def close_uvc_video_window(self):
        if self.uvc_video_window is None:
            return
        self.uvc_video_window.hide()
        self.uvc_video_window.clear_frame()

    def clear_uvc_viewer_analysis(self, reset_timeline=True):
        self.viewer_uvc_stats = None
        self.viewer_uvc_histogram = None
        self.viewer_uvc_histogram_rgb = None
        self.viewer_uvc_quality = None
        self.viewer_uvc_temporal_tracker.reset()
        self.viewer_uvc_last_analysis_time = 0.0
        if hasattr(self, "viewer_uvc_stat_values"):
            for value_label in self.viewer_uvc_stat_values.values():
                value_label.setText("--")
        if hasattr(self, "viewer_uvc_histogram_preview"):
            self.viewer_uvc_histogram_preview.set_array(None)
        if hasattr(self, "viewer_uvc_quality_labels"):
            self.update_single_quality_score_labels(
                self.viewer_uvc_quality_labels, None
            )
        if reset_timeline and hasattr(self, "viewer_uvc_quality_chart"):
            self.viewer_uvc_quality_chart.clear()

    def update_uvc_viewer_analysis(self, image, frame_time=None):
        stats, histogram = gray_image_analysis(image)
        histogram_rgb = render_gray_histogram_rgb(histogram, stats)
        self.viewer_uvc_stats = stats
        self.viewer_uvc_histogram = histogram
        self.viewer_uvc_histogram_rgb = histogram_rgb
        quality_enabled = self.viewer_uvc_quality_enabled.isChecked()
        if quality_enabled:
            self.viewer_uvc_quality = add_uvc_temporal_quality(
                thermal_uvc_quality_analysis(
                    image,
                    max_analysis_pixels=UVC_QUALITY_ANALYSIS_MAX_PIXELS,
                ),
                self.viewer_uvc_temporal_tracker.result(),
            )
        else:
            self.viewer_uvc_quality = None
        self.viewer_uvc_last_analysis_time = (
            time.perf_counter() if frame_time is None else float(frame_time)
        )
        self.update_single_quality_score_labels(
            self.viewer_uvc_quality_labels, self.viewer_uvc_quality
        )
        if quality_enabled:
            self.viewer_uvc_quality_chart.append_metrics(
                self.viewer_uvc_last_analysis_time,
                single=self.viewer_uvc_quality["metrics"],
            )
        if quality_enabled and self.viewer_uvc_recording_worker is not None:
            elapsed = max(
                0.0,
                self.viewer_uvc_last_analysis_time
                - self.viewer_uvc_recording_started_perf,
            )
            self.viewer_uvc_recording_quality_timeline.append(
                {
                    "elapsed_seconds": round(elapsed, 3),
                    "metrics": dict(self.viewer_uvc_quality["metrics"]),
                }
            )

        values = self.viewer_uvc_stat_values
        values["count"].setText(f"{stats['count']:,}")
        values["unique_levels"].setText(str(stats["unique_levels"]))
        values["min"].setText(str(stats["min"]))
        values["max"].setText(str(stats["max"]))
        for key in ("mean", "median", "std", "p1", "p5", "p95", "p99"):
            values[key].setText(f"{stats[key]:.3f}")
        values["p99_p1"].setText(f"{stats['p99_p1']:.3f}")
        values["entropy_bits"].setText(f"{stats['entropy_bits']:.3f} bits")
        values["zero_clip"].setText(
            f"{stats['zero_count']:,} ({stats['zero_percent']:.3f}%)"
        )
        values["full_clip"].setText(
            f"{stats['full_count']:,} ({stats['full_percent']:.3f}%)"
        )
        values["limited_low"].setText(
            f"{stats['limited_low_count']:,} "
            f"({stats['limited_low_percent']:.3f}%)"
        )
        values["limited_high"].setText(
            f"{stats['limited_high_count']:,} "
            f"({stats['limited_high_percent']:.3f}%)"
        )
        self.viewer_uvc_histogram_preview.set_array(histogram_rgb)
        self.update_uvc_histogram_window()

    def open_uvc_histogram_window(self):
        if self.current_function_index() != self.FUNCTIONS.index("Single UVC Statistic"):
            return
        if self.viewer_uvc_histogram_rgb is None:
            QMessageBox.information(
                self,
                "No Histogram",
                "Start Single UVC Statistic and wait for gray statistics first.",
            )
            return
        if self.viewer_uvc_histogram_window is None:
            self.viewer_uvc_histogram_window = UvcVideoWindow(self)
        self.viewer_uvc_histogram_window.setWindowTitle(
            "Single UVC Statistic - Gray Histogram and CDF"
        )
        self.viewer_uvc_histogram_window.set_bgr_frame(
            cv2.cvtColor(
                self.viewer_uvc_histogram_rgb, cv2.COLOR_RGB2BGR
            )
        )
        self.viewer_uvc_histogram_window.show()
        self.viewer_uvc_histogram_window.raise_()
        self.viewer_uvc_histogram_window.activateWindow()

    def update_uvc_histogram_window(self):
        window = self.viewer_uvc_histogram_window
        if (
            window is None
            or not window.isVisible()
            or self.viewer_uvc_histogram_rgb is None
        ):
            return
        window.set_bgr_frame(
            cv2.cvtColor(
                self.viewer_uvc_histogram_rgb, cv2.COLOR_RGB2BGR
            )
        )

    def close_uvc_histogram_window(self):
        if self.viewer_uvc_histogram_window is None:
            return
        self.viewer_uvc_histogram_window.hide()
        self.viewer_uvc_histogram_window.clear_frame()

    def open_raw_histogram_window(self):
        if self.current_function_index() != self.FUNCTIONS.index("Single Image Statistics"):
            return
        if self.raw_viewer_histogram_rgb is None:
            QMessageBox.information(
                self,
                "No Histogram",
                "Select a RAW file and wait for the preview first.",
            )
            return
        if self.raw_viewer_histogram_window is None:
            self.raw_viewer_histogram_window = UvcVideoWindow(self)
        self.raw_viewer_histogram_window.setWindowTitle(
            "Single Image Statistics - DN Histogram and CDF"
        )
        self.raw_viewer_histogram_window.set_bgr_frame(
            cv2.cvtColor(
                self.raw_viewer_histogram_rgb, cv2.COLOR_RGB2BGR
            )
        )
        self.raw_viewer_histogram_window.show()
        self.raw_viewer_histogram_window.raise_()
        self.raw_viewer_histogram_window.activateWindow()

    def close_raw_histogram_window(self):
        if self.raw_viewer_histogram_window is None:
            return
        self.raw_viewer_histogram_window.hide()
        self.raw_viewer_histogram_window.clear_frame()

    def change_page(self, function_index):
        for window in self.image_preview_windows.values():
            window.hide()
            window.clear_frame()
        self.close_uvc_video_window()
        self.close_uvc_histogram_window()
        self.close_raw_histogram_window()
        if function_index is None:
            return
        index = int(function_index)
        if not 0 <= index < len(self.FUNCTIONS):
            return
        self.active_function_index = index
        function_name = self.FUNCTIONS[index]
        category_name = self.category_name_for_function(function_name)
        self.function_actions[index].setChecked(True)
        single_raw_page_index = self.FUNCTIONS.index("Single Image Statistics")
        viewer_uvc_page_index = self.FUNCTIONS.index("Single UVC Statistic")
        uvc_page_index = len(self.FUNCTIONS) - 1
        raw_algorithm_page_index = len(self.FUNCTIONS) - 2
        dual_uvc_page_index = len(self.FUNCTIONS) - 3
        if index != viewer_uvc_page_index and self.viewer_uvc_worker is not None:
            self.viewer_uvc_restart_when_stopped = False
            self.stop_uvc_viewer(log_message=False)
        if index != uvc_page_index and (
            self.uvc_worker is not None or self.uvc_processing_worker is not None
        ):
            self.uvc_restart_when_stopped = False
            self.stop_uvc_camera(log_message=False)
        if index != dual_uvc_page_index and (
            self.dual_uvc_worker_a is not None
            or self.dual_uvc_worker_b is not None
        ):
            self.dual_uvc_restart_when_stopped = False
            self.stop_dual_uvc_cameras(log_message=False)
        self.stack.setCurrentIndex(index)
        self.category_badge.setText(category_name)
        self.function_title.setText(function_name)
        self.function_description.setText(self.FUNCTION_DESCRIPTIONS[index])
        show_comparison, original_title, result_title = self.PREVIEW_LAYOUTS[index]
        if index == raw_algorithm_page_index:
            result_title = self.raw_preview_stage_combo.currentText()
        elif index == uvc_page_index:
            result_title = self.uvc_result_title()
        elif index == dual_uvc_page_index:
            original_title, result_title = self.dual_uvc_image_names()
        elif function_name == "Dual Image Compare":
            original_title, result_title = self.dual_image_names()
        self.preview_original_container.setVisible(show_comparison)
        self.preview_original.setVisible(show_comparison)
        self.preview_original.set_title(original_title)
        self.preview_original_caption.setText(original_title)
        self.preview_result.set_title(result_title)
        self.preview_result_caption.setText(result_title)
        self.update_image_preview_titles()
        self.preview_original.set_array(None)
        self.preview_result.set_array(None)
        self.update_dual_compare_mode()
        action_labels = {
            "Single Image Statistics": "Save Analysis Results",
            "Average RAW Folder": "Average and Save",
            "NUC Calculate": "Calculate and Save NUC",
            "NUC Apply": "Apply NUC and Save",
            "High / Low Analyzer": "Analyze and Save",
            "Temporal Drift Analyzer": "Analyze and Save",
            "BPR Bad Pixel Marking": "Save",
            "Dual Image Compare": "Save Comparison Results",
            "Single UVC Statistic": "Capture and Save Analysis",
            "Dual UVC Compare": "Start Cameras",
            "Image Algorithm Compare": "Save Analysis Results",
            "UVC Algorithm Compare": "Start Camera",
        }
        save_only = function_name in (
            "Single Image Statistics", "Dual Image Compare", "Image Algorithm Compare"
        )
        run_button_object_name = "primaryButton"
        self.run_btn.setText(action_labels[function_name])
        self.run_btn.setToolTip(
            "The preview updates automatically. Save images and analysis reports to a new result folder."
            if save_only else action_labels[function_name]
        )
        if self.run_btn.objectName() != run_button_object_name:
            self.run_btn.setObjectName(run_button_object_name)
            self.run_btn.style().unpolish(self.run_btn)
            self.run_btn.style().polish(self.run_btn)
        self.run_btn.setVisible(
            index not in (
                dual_uvc_page_index,
                uvc_page_index,
            )
        )
        if index == viewer_uvc_page_index:
            self.run_btn.setEnabled(
                self.viewer_uvc_worker is not None
                and self.viewer_uvc_frame is not None
            )
        elif function_name == "Single Image Statistics":
            self.run_btn.setEnabled(getattr(self, "_single_image_save_ready", False))
        elif function_name == "Dual Image Compare":
            self.run_btn.setEnabled(getattr(self, "_dual_image_save_ready", False))
        else:
            self.run_btn.setEnabled(
                index != raw_algorithm_page_index
                or (
                    self.raw_algorithm_source is not None
                    )
            )
        self.save_current_btn.setVisible(
            index not in (3, viewer_uvc_page_index, dual_uvc_page_index)
        )
        self.log.clear()
        self.log.append(
            f"Selected: {category_name} / {function_name}"
        )
        if (
            index == raw_algorithm_page_index
            and self.raw_algorithm_source is not None
        ):
            self.refresh_raw_pipeline_preview()

        if function_name == "BPR Bad Pixel Marking":
            self.auto_preview_bpr()

    def set_image_save_ready(self, mode, ready):
        setattr(self, f"_{mode}_image_save_ready", bool(ready))
        name = "Single Image Statistics" if mode == "single" else "Dual Image Compare"
        if self.FUNCTIONS[self.current_function_index()] == name:
            self.run_btn.setEnabled(bool(ready))

    def auto_preview_raw(self, _path=None):
        if self.current_function_index() == 0:
            QApplication.setOverrideCursor(Qt.WaitCursor)
            try:
                self.run_viewer(save_output=False)
            except Exception as exc:
                self.log.append(f"\nERROR: {exc}")
                QMessageBox.critical(self, "Error", str(exc))
            finally:
                QApplication.restoreOverrideCursor()

    def refresh_raw_viewer_settings(self, _value=None):
        if (
            self.current_function_index() == self.FUNCTIONS.index("Single Image Statistics")
            and self.viewer_file.text().strip()
        ):
            self.auto_preview_raw()

    def auto_preview_raw_quality(self, _path=None):
        if (
            self.current_function_index()
            != self.FUNCTIONS.index("Dual Image Compare")
            or not self.raw_quality_file_a.text().strip()
            or not self.raw_quality_file_b.text().strip()
        ):
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            self.run_raw_quality_compare(save_output=False)
        except Exception as exc:
            self.log.append(f"\nERROR: {exc}")
            QMessageBox.critical(self, "Error", str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def auto_preview_raw_algorithm(self, _path=None):
        if (
            self.current_function_index()
            != self.FUNCTIONS.index("Image Algorithm Compare")
            or not self.raw_algorithm_file.text()
        ):
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            raw_path = self.raw_algorithm_file.text()
            self.raw_algorithm_source, self.raw_algorithm_info = self.load_analysis_source(raw_path)
            if not self.refresh_raw_pipeline_preview():
                return
            original_score = self.raw_algorithm_quality_original["metrics"][
                "overall_score"
            ]
            final_score = self.raw_algorithm_quality_final["metrics"][
                "overall_score"
            ]
            self.log.append(f"\n[{self.raw_algorithm_info['source_type'].title()}] Preview loaded: {raw_path}")
            self.log.append(
                f"Pipeline preview: {self.raw_pipeline_editor.summary()}"
            )
            self.log.append(
                "Image Quality: "
                f"Original={original_score:.2f}, Final={final_score:.2f}"
            )
            self.log.append(
                "Preview updated. Save Analysis Results creates output files."
            )
        except Exception as exc:
            self.invalidate_raw_pipeline_source()
            self.log.append(f"\nERROR: {exc}")
            QMessageBox.critical(self, "Error", str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def refresh_raw_quality_settings(self, _value=None):
        if (
            self.current_function_index()
            == self.FUNCTIONS.index("Dual Image Compare")
            and self.raw_quality_file_a.text().strip()
            and self.raw_quality_file_b.text().strip()
        ):
            self.auto_preview_raw_quality()

    def auto_pair_nuc_b_file(self, k_path):
        path = Path(k_path)
        if path.name.lower().endswith("_k.bin"):
            candidate = path.with_name(f"{path.name[:-6]}_B.bin")
            if candidate.is_file():
                self.b_bin_file.setText(str(candidate))

    @property
    def width(self): return self.resolution_combo.currentData()[0]
    @property
    def height(self): return self.resolution_combo.currentData()[1]
    @property
    def bit_depth(self): return self.bit_depth_combo.currentData()
    @property
    def raw_max(self): return (1 << self.bit_depth) - 1

    def log_stats(self, name, image, stats=None):
        s = stats if stats is not None else image_stats(image)
        self.log.append(f"\n[{name}]\nMin={s['min']:.3f}  Max={s['max']:.3f}  Mean={s['mean']:.3f}  Median={s['median']:.3f}  Std={s['std']:.3f}")

    def clear_viewer_stats(self):
        for value_label in self.viewer_stat_values.values():
            value_label.setText("--")
        self.raw_viewer_stats = None
        self.raw_viewer_histogram = None
        self.raw_viewer_histogram_rgb = None
        self.raw_viewer_quality = None
        self.raw_viewer_histogram_preview.set_array(None)
        self.update_single_quality_score_labels(
            self.raw_viewer_quality_labels,
            None,
        )

    def update_viewer_stats(self, image):
        self.viewer_stats_group.setTitle(f"Statistics ({self.raw_viewer_info['value_name']})")
        stats, histogram = raw_dn_image_analysis(image, self.raw_viewer_info['value_max'])
        histogram_rgb = render_gray_histogram_rgb(
            histogram,
            stats,
            value_max=self.raw_viewer_info['value_max'],
            title=(
                f"{self.raw_viewer_info['bit_depth']}-bit {self.raw_viewer_info['value_name']} Histogram "
                f"(256 display bins)"
            ),
        )
        self.raw_viewer_stats = stats
        self.raw_viewer_histogram = histogram
        self.raw_viewer_histogram_rgb = histogram_rgb
        self.viewer_stat_values["count"].setText(f"{stats['count']:,}")
        self.viewer_stat_values["unique_levels"].setText(
            f"{stats['unique_levels']:,}"
        )
        self.viewer_stat_values["min"].setText(str(stats["min"]))
        self.viewer_stat_values["max"].setText(str(stats["max"]))
        for key in ("mean", "median", "std", "p1", "p5", "p95", "p99"):
            self.viewer_stat_values[key].setText(f"{stats[key]:,.3f}")
        self.viewer_stat_values["entropy_bits"].setText(
            f"{stats['entropy_bits']:.3f} bits"
        )
        self.viewer_stat_values["p99_p1"].setText(
            f"{stats['p99_p1']:,.3f}"
        )
        self.viewer_stat_values["p99_p1_percent"].setText(
            f"{stats['p99_p1_percent']:.3f}%"
        )
        self.viewer_stat_values["zero_clip"].setText(
            f"{stats['zero_count']:,} ({stats['zero_percent']:.3f}%)"
        )
        self.viewer_stat_values["max_clip"].setText(
            f"{stats['max_count']:,} ({stats['max_percent']:.3f}%)"
        )
        self.viewer_stat_values["over_max"].setText(
            f"{stats['over_max_count']:,} "
            f"({stats['over_max_percent']:.3f}%)"
        )
        self.raw_viewer_histogram_preview.set_array(histogram_rgb)
        return stats

    @staticmethod
    def write_raw_viewer_analysis_summary(summary, stats, value_name="RAW DN"):
        summary.write(f"{value_name} statistics:\n")
        summary.write(
            f"  Pixels / unique DN: {stats['count']:,} / "
            f"{stats['unique_levels']:,}\n"
        )
        summary.write(
            f"  Min / max / bit max: {stats['min']} / "
            f"{stats['max']} / {stats['raw_max']}\n"
        )
        summary.write(
            f"  Mean / median / std: {stats['mean']:.3f} / "
            f"{stats['median']:.3f} / {stats['std']:.3f}\n"
        )
        summary.write(
            f"  P1 / P5 / P95 / P99: {stats['p1']:.3f} / "
            f"{stats['p5']:.3f} / {stats['p95']:.3f} / "
            f"{stats['p99']:.3f}\n"
        )
        summary.write(
            f"  P99-P1: {stats['p99_p1']:.3f} "
            f"({stats['p99_p1_percent']:.3f}% of bit range)\n"
        )
        summary.write(
            f"  Entropy: {stats['entropy_bits']:.3f} bits\n"
        )
        summary.write(
            f"  Zero clip: {stats['zero_count']:,} "
            f"({stats['zero_percent']:.3f}%)\n"
        )
        summary.write(
            f"  Bit max clip: {stats['max_count']:,} "
            f"({stats['max_percent']:.3f}%)\n"
        )
        summary.write(
            f"  Over bit max: {stats['over_max_count']:,} "
            f"({stats['over_max_percent']:.3f}%)\n"
        )

    def make_output(self, name):
        out = unique_result_dir(self.output_root.text(), name)
        self.last_output_dir = str(out)
        self.log.append(f"\nOutput: {out}")
        return out

    def run_current(self):
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            idx = self.current_function_index()
            handlers = [
                self.run_viewer,
                self.run_average,
                self.run_nuc_calculate,
                self.run_nuc_apply,
                self.run_high_low,
                self.run_drift,
                self.run_bpr_marking,
                self.run_raw_quality_compare,
                self.capture_uvc_viewer,
                self.start_dual_uvc_cameras,
                self.run_raw_algorithm,
                self.start_uvc_camera,
            ]
            handlers[idx]()
        except Exception as exc:
            self.log.append(f"\nERROR: {exc}")
            QMessageBox.critical(self, "Error", str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def run_raw_quality_compare(self, save_output=True):
        self.set_image_save_ready("dual", False)
        name_a, name_b = self.dual_image_names()
        path_a = self.raw_quality_file_a.text().strip()
        path_b = self.raw_quality_file_b.text().strip()
        raw_a, info_a = self.load_analysis_source(path_a)
        raw_b, info_b = self.load_analysis_source(path_b)
        self.dual_image_source_infos = (info_a, info_b)
        self.update_image_preview_titles()
        display_a = raw_to_display_u8(raw_a, info_a["value_max"])
        display_b = raw_to_display_u8(raw_b, info_b["value_max"])
        result_a = thermal_raw_quality_analysis(raw_a, info_a["value_max"])
        result_b = thermal_raw_quality_analysis(raw_b, info_b["value_max"])
        self.raw_quality_result_a = result_a
        self.raw_quality_result_b = result_b
        self.set_image_save_ready("dual", True)
        self.current_original = display_a
        self.current_result = display_b
        self.preview_original.set_array(display_a)
        self.preview_result.set_array(display_b)
        self.update_quality_score_labels(
            self.raw_quality_score_labels, result_a, result_b
        )
        score_a = result_a["metrics"]["overall_score"]
        score_b = result_b["metrics"]["overall_score"]
        self.log.append(
            f"\nDual Image Compare: "
            f"[{info_a['source_type'].title()}] {name_a}={score_a:.3f}, "
            f"[{info_b['source_type'].title()}] {name_b}={score_b:.3f}"
        )
        self.log.append(
            "Image Quality no-reference score: compare only images of the same scene/settings."
        )
        if not save_output:
            return

        out = self.make_output("raw_quality_compare")
        outputs = {
            "image_a": "raw_a.png",
            "image_b": "raw_b.png",
            "comparison": "image_comparison.png",
            "base_a": "base_a.png",
            "base_b": "base_b.png",
            "detail_a": "detail_a.png",
            "detail_b": "detail_b.png",
            "excel": "quality_analysis.xlsx",
        }
        save_png(out / outputs["image_a"], display_a)
        save_png(out / outputs["image_b"], display_b)
        if info_a["source_type"] == "raw":
            outputs["raw_csv_a"] = "raw_a_dn.csv"
            save_csv_matrix(out / outputs["raw_csv_a"], raw_a, "%d")
            self.log.append(f"Saved RAW A DN CSV: {out / outputs['raw_csv_a']}")
        if info_b["source_type"] == "raw":
            outputs["raw_csv_b"] = "raw_b_dn.csv"
            save_csv_matrix(out / outputs["raw_csv_b"], raw_b, "%d")
            self.log.append(f"Saved RAW B DN CSV: {out / outputs['raw_csv_b']}")
        if self.dual_compare_mode.currentIndex() == 1:
            self.save_dual_swipe_preview(out / outputs["comparison"])
        else:
            comparison = self.compose_dual_image_preview([display_a, display_b])
            save_rgb_png(out / outputs["comparison"], cv2.cvtColor(comparison, cv2.COLOR_BGR2RGB))
        self.log.append(f"Saved comparison preview: {out / outputs['comparison']}")
        save_png(out / outputs["base_a"], result_a["base_u8"])
        save_png(out / outputs["base_b"], result_b["base_u8"])
        save_png(out / outputs["detail_a"], result_a["detail_u8"])
        save_png(out / outputs["detail_b"], result_b["detail_u8"])
        manifest = {
            "format_version": 1,
            "mode": "raw_quality_compare",
            "comparison_view": self.dual_compare_mode.currentText(),
            "swipe_position": self.dual_swipe_preview.position,
            "analyzed_at": datetime.now().astimezone().isoformat(
                timespec="seconds"
            ),
            "method": "Image Quality no-reference same-scene quality: native DN + Base/Detail + gray-level/connected flat-region robust aggregation",
            "source_a": str(Path(path_a).resolve()),
            "source_b": str(Path(path_b).resolve()),
            "image_name_a": name_a,
            "image_name_b": name_b,
            "source_info_a": info_a,
            "source_info_b": info_b,
            "quality_a": quality_analysis_manifest(result_a),
            "quality_b": quality_analysis_manifest(result_b),
            "outputs": outputs,
        }
        with open(out / "quality.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(
            out / "quality_summary.txt", "w", encoding="utf-8"
        ) as summary:
            summary.write(f"Analyzed: {manifest['analyzed_at']}\n")
            summary.write(f"RAW A: {manifest['source_a']}\n")
            summary.write(f"RAW B: {manifest['source_b']}\n")
            summary.write(
                f"Input A: {info_a}\nInput B: {info_b}\n"
            )
            summary.write(
                "Use these scores only for same-scene relative comparison.\n"
            )
            write_quality_summary(summary, name_a, result_a)
            write_quality_summary(summary, name_b, result_b)
        save_quality_excel(
            out / outputs["excel"],
            "Dual Image Compare",
            {
                key: value
                for key, value in manifest.items()
                if key not in {"quality_a", "quality_b"}
            },
            [
                {
                    "label": name_a,
                    "metrics": result_a["metrics"],
                },
                {
                    "label": name_b,
                    "metrics": result_b["metrics"],
                },
            ],
        )
        self.log.append(f"Saved image quality comparison: {out}")

    def run_viewer(self, save_output=True):
        self.set_image_save_ready("single", False)
        self.current_original = None
        self.current_result = None
        self.preview_result.set_array(None)
        self.clear_viewer_stats()
        raw_path = self.viewer_file.text()
        raw, self.raw_viewer_info = self.load_analysis_source(raw_path)
        display = raw_to_display_u8(raw, self.raw_viewer_info['value_max'])
        self.raw_viewer_quality = thermal_raw_quality_analysis(raw, self.raw_viewer_info['value_max'])
        self.update_single_quality_score_labels(
            self.raw_viewer_quality_labels,
            self.raw_viewer_quality,
        )
        self.current_original = raw; self.current_result = raw
        self.preview_result.set_array(display)
        self.update_image_preview_titles()
        stats = self.update_viewer_stats(raw)
        self.set_image_save_ready("single", True)
        self.log.append(f"\n[{self.raw_viewer_info['source_type'].title()}] File: {raw_path}")
        self.log.append(
            f"Resolution: {raw.shape[1]} x {raw.shape[0]}; Pixels: {raw.size:,}; "
            f"Source bit depth: {self.raw_viewer_info['bit_depth']}-bit"
        )
        self.log_stats(self.raw_viewer_info["source_type"].title(), raw, stats)
        self.log.append(f"P1={stats['p1']:.3f}  P99={stats['p99']:.3f}")
        self.log.append(
            "Total Score="
            f"{self.raw_viewer_quality['metrics']['overall_score']:.3f}"
        )
        self.log.append(f"Preview scale: fixed 0..{self.raw_viewer_info['value_max']} {self.raw_viewer_info['value_name']}")
        if not save_output:
            return

        out = self.make_output("raw_viewer_analysis")
        preview_path = out / "raw_preview.png"
        histogram_path = out / "raw_histogram.png"
        excel_path = out / "raw_analysis.xlsx"
        quality_excel_path = out / "quality_analysis.xlsx"
        base_path = out / "base.png"
        detail_path = out / "detail.png"
        manifest_path = out / "analysis.json"
        summary_path = out / "analysis_summary.txt"
        save_png(preview_path, display)
        raw_csv_path = None
        if self.raw_viewer_info["source_type"] == "raw":
            raw_csv_path = out / "original_raw_dn.csv"
            save_csv_matrix(raw_csv_path, raw, "%d")
            self.log.append(f"Saved original RAW DN CSV: {raw_csv_path}")
        save_rgb_png(histogram_path, self.raw_viewer_histogram_rgb)
        save_png(base_path, self.raw_viewer_quality["base_u8"])
        save_png(detail_path, self.raw_viewer_quality["detail_u8"])
        manifest = {
            "format_version": 1,
            "mode": "raw_viewer_analysis",
            "analyzed_at": datetime.now().astimezone().isoformat(
                timespec="seconds"
            ),
            "source_info": self.raw_viewer_info,
            "source_raw": str(Path(raw_path).resolve()),
            "resolution": {"width": raw.shape[1], "height": raw.shape[0]},
            "bit_depth": self.raw_viewer_info['bit_depth'],
            "raw_max": self.raw_viewer_info['value_max'],
            "statistics": dict(stats),
            "quality": quality_analysis_manifest(self.raw_viewer_quality),
            "preview_output": preview_path.name,
            "raw_csv_output": (
                raw_csv_path.name if raw_csv_path is not None else None
            ),
            "histogram_output": histogram_path.name,
            "histogram_display_bins": 256,
            "excel_output": excel_path.name,
            "quality_excel_output": quality_excel_path.name,
            "quality_layer_outputs": {
                "base": base_path.name,
                "detail": detail_path.name,
            },
        }
        with open(manifest_path, "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(summary_path, "w", encoding="utf-8") as summary:
            summary.write(f"Analyzed: {manifest['analyzed_at']}\n")
            summary.write(f"Source RAW: {manifest['source_raw']}\n")
            summary.write(
                f"Resolution: {raw.shape[1]}x{raw.shape[0]}\n"
            )
            summary.write(
                f"Bit depth / max: {self.raw_viewer_info['bit_depth']} / {self.raw_viewer_info['value_max']}\n"
            )
            self.write_raw_viewer_analysis_summary(summary, stats, self.raw_viewer_info["value_name"])
            write_quality_summary(
                summary,
                "Single Image Statistics",
                self.raw_viewer_quality,
            )
            summary.write(f"Preview: {preview_path.name}\n")
            summary.write(f"Histogram: {histogram_path.name}\n")
            summary.write(f"Excel analysis: {excel_path.name}\n")
            summary.write(f"Quality analysis: {quality_excel_path.name}\n")
        save_histogram_excel(
            excel_path,
            "Single Image Statistics Analysis",
            {
                key: value
                for key, value in manifest.items()
                if key not in {"statistics", "quality"}
            },
            [
                {
                    "label": "RAW",
                    "stats": stats,
                    "histogram": self.raw_viewer_histogram,
                }
            ],
            value_max=self.raw_viewer_info['value_max'],
            value_name=self.raw_viewer_info["value_name"],
        )
        save_quality_excel(
            quality_excel_path,
            "Single Image Statistics Quality",
            {
                key: value
                for key, value in manifest.items()
                if key not in {"statistics", "quality"}
            },
            [
                {
                    "label": "Single Image Statistics",
                    "metrics": self.raw_viewer_quality["metrics"],
                }
            ],
        )
        self.log.append(f"Saved RAW analysis: {manifest_path}")
        self.log.append(f"Saved RAW histogram: {histogram_path}")
        self.log.append(f"Saved RAW Excel analysis: {excel_path}")
        self.log.append(f"Saved image quality analysis: {quality_excel_path}")

    def run_average(self):
        avg, valid, skipped = average_raw_folder(self.avg_folder.text(), self.width, self.height)
        out = self.make_output("average_raw")
        suffix = self.avg_suffix_combo.currentText()
        avg_preview = raw_to_display_u8(avg, self.raw_max)
        raw_path = out / f"average_{suffix}.raw"; avg.astype(RAW_DTYPE).tofile(raw_path)
        save_png(out / f"average_preview_{suffix}.png", avg_preview)
        with open(out / "source_files.txt", "w", encoding="utf-8") as f:
            f.write("\n".join(valid))
        self.current_original = load_raw(valid[0], self.width, self.height); self.current_result = avg
        self.preview_result.set_array(avg_preview)
        self.log.append(f"Valid frames: {len(valid)}; Skipped: {len(skipped)}")
        self.log.append(f"Filename suffix: {suffix}")
        self.log_stats("First Frame", self.current_original); self.log_stats("Average RAW", avg)
        self.log.append(f"Preview scale: fixed 0..{self.raw_max} DN")
        self.log.append(f"Saved: {raw_path}")

    def _get_nuc_inputs(self):
        if self.nuc_mode.currentIndex() == 0:
            low, low_files, low_skip = average_raw_folder(self.low_folder.text(), self.width, self.height)
            high, high_files, high_skip = average_raw_folder(self.high_folder.text(), self.width, self.height)
            return low, high, low_files, high_files, low_skip, high_skip
        low = load_raw(self.low_single.text(), self.width, self.height)
        high = load_raw(self.high_single.text(), self.width, self.height)
        return low, high, [self.low_single.text()], [self.high_single.text()], [], []

    def run_nuc_calculate(self):
        low, high, low_files, high_files, low_skip, high_skip = self._get_nuc_inputs()
        gain, offset, valid, delta, low_avg, high_avg = calculate_nuc(low, high, self.min_delta.value())
        stored_k, gfx_b, table_valid, novatek_low_avg, novatek_high_avg = build_novatek_nuc_tables(
            low, high, self.min_delta.value()
        )
        valid = table_valid
        point_number = self.nuc_point_combo.currentData()
        out = self.make_output("nuc_calculate")
        low.astype(RAW_DTYPE).tofile(out / "low_input_used.raw"); high.astype(RAW_DTYPE).tofile(out / "high_input_used.raw")
        np.save(out / "gain.npy", gain); np.save(out / "offset.npy", offset)
        save_csv_matrix(out / "gain.csv", gain, "%.8f"); save_csv_matrix(out / "offset.csv", offset, "%.3f")
        k_path = out / f"point{point_number}_K.bin"
        b_path = out / f"point{point_number}_B.bin"
        stored_k.astype("<u2").tofile(k_path)
        gfx_b.astype("<i2").tofile(b_path)
        yy, xx = np.where(~valid)
        with open(out / "invalid_pixels.csv", "w", newline="") as f:
            wr = csv.writer(f); wr.writerow(["x", "y", "low", "high", "delta"])
            for y, x in zip(yy, xx): wr.writerow([x, y, int(low[y,x]), int(high[y,x]), float(delta[y,x])])
        with open(out / "nuc_summary.txt", "w", encoding="utf-8") as f:
            f.write(f"Novatek NUC point: {point_number}\nMode: {self.nuc_mode.currentText()}\nLow frames: {len(low_files)}\nHigh frames: {len(high_files)}\n")
            f.write(f"Low mean: {low_avg:.6f}\nHigh mean: {high_avg:.6f}\nTarget delta: {high_avg-low_avg:.6f}\n")
            f.write(f"Min delta: {self.min_delta.value()}\nValid pixels: {valid.sum()} / {valid.size}\n")
            if valid.any(): f.write(f"Gain mean/std: {gain[valid].mean():.8f} / {gain[valid].std():.8f}\nOffset mean/std: {offset[valid].mean():.4f} / {offset[valid].std():.4f}\n")
            f.write(f"Novatek integer LOW/HIGH mean: {novatek_low_avg} / {novatek_high_avg}\n")
            f.write(f"K format: uint16 Q8; B format: signed int16\n")
            f.write(f"Runtime formula: output = (K * (input - B)) >> {NOVATEK_NUC_GAIN_SHIFT}\n")
            f.write(f"K file: {k_path.name} ({k_path.stat().st_size} bytes)\n")
            f.write(f"B file: {b_path.name} ({b_path.stat().st_size} bytes)\n")
        low_preview = raw_to_display_u8(low, self.raw_max)
        high_preview = raw_to_display_u8(high, self.raw_max)
        save_png(out / "low_preview.png", low_preview); save_png(out / "high_preview.png", high_preview)
        # Preview the same quantized K/B formula used by the Novatek GFX path.
        corrected_high = apply_novatek_nuc(high, stored_k, gfx_b, self.raw_max)
        corrected_high_preview = raw_to_display_u8(corrected_high, self.raw_max)
        corrected_high.astype(RAW_DTYPE).tofile(out / "high_nuc_corrected.raw")
        save_png(out / "high_nuc_corrected_preview.png", corrected_high_preview)
        self.current_original = high; self.current_result = corrected_high
        self.preview_original.set_array(high_preview)
        self.preview_result.set_array(corrected_high_preview)
        self.log.append(f"Novatek NUC point={point_number}: {k_path.name}, {b_path.name}")
        self.log.append(f"Mode: {self.nuc_mode.currentText()}")
        self.log.append(f"LOW frames={len(low_files)}, HIGH frames={len(high_files)}, invalid/skipped={len(low_skip)+len(high_skip)}")
        self.log.append(f"Low mean={low_avg:.3f}, High mean={high_avg:.3f}, Target delta={high_avg-low_avg:.3f}")
        self.log.append(f"Valid NUC pixels={valid.sum():,}/{valid.size:,} ({valid.mean()*100:.3f}%)")
        if valid.any(): self.log.append(f"Gain mean={gain[valid].mean():.6f}, std={gain[valid].std():.6f}; Offset mean={offset[valid].mean():.3f}, std={offset[valid].std():.3f}")
        if valid.any():
            corrected_valid = corrected_high[valid].astype(np.float64)
            self.log.append(
                f"Corrected HIGH valid pixels: mean={corrected_valid.mean():.3f}, "
                f"std={corrected_valid.std():.3f}, min={corrected_valid.min():.0f}, "
                f"max={corrected_valid.max():.0f} DN"
            )
        if not valid.all():
            self.log.append(f"Invalid pixels kept uncorrected: {(~valid).sum():,}")
        self.log.append("Preview scale: fixed 0..Bit Depth Max (no percentile stretching)")

    def run_nuc_apply(self):
        raw = load_raw(self.apply_raw_file.text(), self.width, self.height)
        stored_k = load_novatek_nuc_bin(self.k_bin_file.text(), self.width, self.height, "<u2", "K")
        gfx_b = load_novatek_nuc_bin(self.b_bin_file.text(), self.width, self.height, "<i2", "B")
        corrected = apply_novatek_nuc(raw, stored_k, gfx_b, self.raw_max)
        out = self.make_output("nuc_apply")
        original_path = out / "original.raw"
        original_png_path = out / "original.png"
        corrected_path = out / "corrected_nuc.raw"
        corrected_png_path = out / "corrected_nuc.png"
        raw.astype(RAW_DTYPE).tofile(original_path)
        save_png(original_png_path, raw_to_display_u8(raw, self.raw_max))
        corrected.astype(RAW_DTYPE).tofile(corrected_path)
        save_png(corrected_png_path, raw_to_display_u8(corrected, self.raw_max))
        self.current_original = raw; self.current_result = corrected
        self.preview_original.set_array(raw_to_display_u8(raw, self.raw_max))
        self.preview_result.set_array(raw_to_display_u8(corrected, self.raw_max))
        self.log.append(f"K BIN: {self.k_bin_file.text()}")
        self.log.append(f"B BIN: {self.b_bin_file.text()}")
        self.log.append(f"Saved original RAW: {original_path}")
        self.log.append(f"Saved original PNG: {original_png_path}")
        self.log.append(f"Saved NUC RAW: {corrected_path}")
        self.log.append(f"Saved NUC PNG: {corrected_png_path}")

    def run_high_low(self):
        low = load_raw(self.an_low.text(), self.width, self.height); high = load_raw(self.an_high.text(), self.width, self.height)
        diff = high.astype(np.int32) - low.astype(np.int32); threshold = self.an_min_delta.value()
        valid = diff > threshold
        out = self.make_output("high_low_analysis")
        np.save(out / "delta.npy", diff)
        yy, xx = np.where(~valid)
        with open(out / "abnormal_pixels.csv", "w", newline="") as f:
            wr = csv.writer(f); wr.writerow(["x","y","low","high","delta","high_le_low","small_delta","saturated"])
            for y,x in zip(yy,xx):
                d=int(diff[y,x]); wr.writerow([x,y,int(low[y,x]),int(high[y,x]),d,int(d<=0),int(0<d<=threshold),int(low[y,x]>=self.raw_max or high[y,x]>=self.raw_max)])
        # Signed delta uses the fixed Bit Depth range; zero maps to middle gray.
        delta8 = signed_delta_to_display_u8(diff, self.raw_max)
        save_png(out / "difference_preview.png", delta8)
        self.current_original = low; self.current_result = diff
        self.preview_original.set_array(raw_to_display_u8(low, self.raw_max)); self.preview_result.set_array(delta8)
        self.log_stats("LOW", low); self.log_stats("HIGH", high)
        self.log.append(f"High > Low: {(diff>0).sum():,} ({(diff>0).mean()*100:.3f}%)")
        self.log.append(f"High = Low: {(diff==0).sum():,}; High < Low: {(diff<0).sum():,}")
        self.log.append(f"Delta Mean={diff.mean():.3f}, Std={diff.std():.3f}, Min={diff.min()}, Max={diff.max()}")
        self.log.append(f"Delta > {threshold} valid: {valid.sum():,}/{valid.size:,}")
        self.log.append(f"Difference preview scale: fixed {-self.raw_max}..{self.raw_max} DN")

    def run_drift(self):
        files = find_raw_files(self.drift_folder.text())
        if not files: raise FileNotFoundError("No RAW files found")
        frames=[]; valid_files=[]
        for p in files:
            try: frames.append(load_raw(p,self.width,self.height)); valid_files.append(p)
            except Exception as exc: self.log.append(f"Skip {Path(p).name}: {exc}")
        if not frames: raise ValueError("No valid frames")
        stack=np.stack(frames).astype(np.float64)
        pm=stack.mean(axis=0); ps=stack.std(axis=0); pmin=stack.min(axis=0); pmax=stack.max(axis=0); p2p=pmax-pmin
        frame_mean=stack.mean(axis=(1,2)); frame_std=stack.std(axis=(1,2))
        out=self.make_output("temporal_drift")
        pm.astype(RAW_DTYPE).tofile(out/"temporal_mean.raw")
        np.save(out/"pixel_std.npy",ps); np.save(out/"pixel_peak_to_peak.npy",p2p)
        mean_preview=raw_to_display_u8(pm,self.raw_max); std_preview=raw_to_display_u8(ps,self.raw_max); p2p_preview=raw_to_display_u8(p2p,self.raw_max)
        save_png(out/"temporal_mean_preview.png",mean_preview); save_png(out/"temporal_std_preview.png",std_preview); save_png(out/"peak_to_peak_preview.png",p2p_preview)
        with open(out/"frame_statistics.csv","w",newline="") as f:
            wr=csv.writer(f); wr.writerow(["frame_index","filename","frame_mean_dn","frame_spatial_std_dn"])
            for i,p in enumerate(valid_files): wr.writerow([i,Path(p).name,f"{frame_mean[i]:.6f}",f"{frame_std[i]:.6f}"])
        std_fail=ps>self.std_spec.value(); p2p_fail=p2p>self.p2p_spec.value()
        self.current_original=frames[0]; self.current_result=ps
        self.preview_result.set_array(std_preview)
        self.log.append(f"Frames={len(frames)}")
        self.log.append(f"Frame mean: mean={frame_mean.mean():.3f}, P-P={np.ptp(frame_mean):.3f}, std={frame_mean.std():.3f}")
        self.log.append(f"Pixel temporal std: mean={ps.mean():.3f}, 95%={np.percentile(ps,95):.3f}, max={ps.max():.3f}")
        self.log.append(f"Std > {self.std_spec.value():.1f}: {std_fail.sum():,} ({std_fail.mean()*100:.3f}%)")
        self.log.append(f"P-P > {self.p2p_spec.value():.1f}: {p2p_fail.sum():,} ({p2p_fail.mean()*100:.3f}%)")
        self.log.append(f"Preview scale: fixed 0..{self.raw_max} DN")

    def invalidate_bpr_preview(self, _value=None):
        self._bpr_preview = None
        self.bpr_coordinate_info.clear()
        if self.FUNCTIONS[self.current_function_index()] == "BPR Bad Pixel Marking":
            self.run_btn.setEnabled(False)
            self.preview_original.set_array(None)
            self.preview_result.set_array(None)
            self.current_original = self.current_result = None

    def auto_preview_bpr(self, _value=None):
        self.invalidate_bpr_preview()
        if self.FUNCTIONS[self.current_function_index()] != "BPR Bad Pixel Marking":
            return
        if not self.bpr_coordinate_file.text().strip():
            return
        if self.bpr_use_raw.isChecked() and not self.bpr_raw_file.text().strip():
            return
        try:
            coordinate_path = self.bpr_coordinate_file.text()
            coordinates, table_info = load_bpr_coordinates(coordinate_path)
            raw_path = ""
            if self.bpr_use_raw.isChecked():
                raw_path = self.bpr_raw_file.text()
                raw, info = self.load_analysis_source(raw_path)
                base_gray = raw_to_display_u8(raw, info["value_max"])
                background_mode = f"{info['value_name']} fixed scale 0..{info['value_max']}"
                base_filename = "image_fixed.png"
            else:
                base_gray = np.full((self.height, self.width), 255, dtype=np.uint8)
                background_mode = "White canvas"
                base_filename = "white_canvas.png"
            height, width = base_gray.shape
            valid_coordinates, outside_coordinates = split_coordinates_by_bounds(coordinates, width, height)
            coordinate_report = format_bpr_coordinate_report(table_info, coordinates, outside_coordinates, width, height)
            self.bpr_coordinate_info.setPlainText(coordinate_report)
            if not valid_coordinates:
                raise ValueError(f"No BPR coordinates are inside the {width}x{height} image.")
            radius = self.bpr_circle_radius.value()
            highlight_bright_points = (
                self.bpr_highlight_bright_points.isChecked()
            )
            highlight_radius = self.bpr_highlight_radius.value()
            base_rgb = np.repeat(base_gray[:, :, np.newaxis], 3, axis=2)
            marked_rgb = draw_bpr_markers(
                base_gray,
                valid_coordinates,
                radius,
                highlight_bright_points,
                highlight_radius,
            )
            self._bpr_preview = (coordinate_path, coordinates, table_info, outside_coordinates,
                                 coordinate_report, raw_path, background_mode, base_filename,
                                 width, height, radius, highlight_bright_points,
                                 highlight_radius, base_rgb, marked_rgb)
            self.current_original, self.current_result = base_rgb, marked_rgb
            self.preview_original.set_array(base_rgb)
            self.preview_result.set_array(marked_rgb)
            self.run_btn.setEnabled(True)
            self.bpr_note.setText(
                f"Coordinates: top-left (0,0), X=0..{width-1}, Y=0..{height-1}; image[y,x]. "
                "Preview updates automatically. Save writes the displayed result."
            )
            self.log.append(f"BPR preview: {width}x{height}, marked={len(valid_coordinates)}, outside={len(outside_coordinates)}; {background_mode}")
        except Exception as exc:
            self.log.append(f"BPR preview error: {exc}")

    def run_bpr_marking(self):
        cached = getattr(self, "_bpr_preview", None)
        if cached is None:
            return
        (coordinate_path, coordinates, table_info, outside_coordinates,
         coordinate_report, raw_path, background_mode, base_filename,
         width, height, radius, highlight_bright_points, highlight_radius,
         base_rgb, marked_rgb) = cached
        out = self.make_output("bpr_bad_pixel_marking")
        base_path = out / base_filename
        marked_path = out / "bad_pixels_marked.png"
        coordinate_output_path = out / "bad_pixel_coordinates.csv"
        summary_path = out / "bpr_marking_summary.txt"

        save_rgb_png(base_path, base_rgb)
        save_rgb_png(marked_path, marked_rgb)
        outside_set = set(outside_coordinates)
        with open(coordinate_output_path, "w", newline="", encoding="utf-8") as destination:
            writer = csv.writer(destination)
            writer.writerow(["x", "y", "in_bounds", "index"])
            for index, (x, y) in enumerate(coordinates, 1):
                writer.writerow([x, y, 0 if (x, y) in outside_set else 1, index])

        with open(summary_path, "w", encoding="utf-8") as summary:
            summary.write(f"BPR source: {coordinate_path}\n")
            summary.write(f"Source format: {table_info['format']}\n")
            summary.write(f"Resolution: {width}x{height}\n")
            summary.write(f"Background: {background_mode}\n")
            if raw_path:
                summary.write(f"Image source: {raw_path}\n")
            summary.write(f"Circle radius: {radius} px\n")
            summary.write(
                "Bright point highlight: "
                f"{'enabled' if highlight_bright_points else 'disabled'}\n"
            )
            if highlight_bright_points:
                summary.write(f"Bright point radius: {highlight_radius} px\n")
            summary.write("\n")
            summary.write(coordinate_report)
            summary.write("\n")

        self.log.append(f"Saved: {marked_path}")
        self.log.append(f"Coordinates: {coordinate_output_path}")

    def run_raw_algorithm(self):
        raw, self.raw_algorithm_info = self.load_analysis_source(self.raw_algorithm_file.text())
        raw8 = raw_to_display_u8(raw, self.raw_algorithm_info["value_max"])
        stages = self.raw_pipeline_editor.stages()
        pipeline_input, input_max, mapping_details = (
            self.prepare_raw_pipeline_input(raw)
        )
        total_started = time.perf_counter()
        processed, records = apply_algorithm_pipeline(
            pipeline_input,
            stages,
            input_max=input_max, allow_empty=True,
            pad_fieldscale_input=self.raw_algorithm_info["source_type"] == "image",
        )
        total_ms = (time.perf_counter() - total_started) * 1000.0

        # Image Algorithm Compare uses the complete RAW Full profile. The original image is
        # evaluated in native DN while the final pipeline image is evaluated in 8-bit;
        # both are internally normalized to 0..1 before scoring.
        quality_original = thermal_raw_quality_analysis(raw, self.raw_algorithm_info["value_max"])
        quality_final = thermal_raw_quality_analysis(processed, 255)
        self.raw_algorithm_quality_original = quality_original
        self.raw_algorithm_quality_final = quality_final
        self.update_quality_score_labels(
            self.raw_algorithm_quality_labels,
            quality_original,
            quality_final,
        )

        out = self.make_output("raw_algorithm_compare")
        original_path = out / "original_fixed.png"
        mapping_path = out / (
            "input_native_dn.png" if input_max > 255 else "input_mapped.png"
        )
        result_path = out / "final_pipeline.png"
        overview_path = out / "00_pipeline_overview.png"
        base_original_path = out / "base_original.png"
        base_final_path = out / "base_final.png"
        detail_original_path = out / "detail_original.png"
        detail_final_path = out / "detail_final.png"
        quality_excel_path = out / "quality_analysis.xlsx"
        save_png(original_path, raw8)
        raw_csv_path = None
        if self.raw_algorithm_info["source_type"] == "raw":
            raw_csv_path = out / "original_raw_dn.csv"
            save_csv_matrix(raw_csv_path, raw, "%d")
            self.log.append(f"Saved original RAW DN CSV: {raw_csv_path}")
        save_png(mapping_path, pipeline_input)
        for record in records:
            step_path = out / (
                f"step_{record['pipeline_index']:02d}_"
                f"{record['algorithm']}.png"
            )
            save_png(step_path, record["image"])
            record["output_file"] = step_path.name
        save_png(result_path, processed)
        save_png(base_original_path, quality_original["base_u8"])
        save_png(base_final_path, quality_final["base_u8"])
        save_png(detail_original_path, quality_original["detail_u8"])
        save_png(detail_final_path, quality_final["detail_u8"])

        mapping_label = self.raw_input_mapping_combo.currentText()
        overview_items = [
            ("00 Original Fixed", raw8, 255),
            (f"01 Pipeline Input - {mapping_label}", pipeline_input, input_max),
        ]
        for overview_index, record in enumerate(records, start=2):
            overview_items.append(
                (
                    f"{overview_index:02d} Step {record['pipeline_index']:02d} "
                    f"{record['name']}",
                    record["image"],
                    255,
                )
            )
        save_pipeline_overview_png(overview_path, overview_items)
        self.raw_algorithm_source = raw
        self.raw_pipeline_input = pipeline_input
        self.raw_pipeline_input_max = input_max
        self.raw_pipeline_final = processed
        self.raw_pipeline_records = records
        self.current_original = raw8
        self.update_image_preview_titles()
        self.preview_original.set_array(raw8)
        self.update_raw_stage_preview()
        comparison_path = out / "image_comparison.png"
        if self.algorithm_compare_mode.currentIndex() == 1:
            self.save_dual_swipe_preview(comparison_path)
        else:
            comparison = self.compose_dual_image_preview(
                [raw8, self.current_result], self.algorithm_image_names()
            )
            save_rgb_png(comparison_path, cv2.cvtColor(comparison, cv2.COLOR_BGR2RGB))
        self.log.append(f"Saved comparison preview: {comparison_path}")

        manifest = {
            "format_version": 2,
            "mode": "raw_algorithm_compare",
            "comparison_view": self.algorithm_compare_mode.currentText(),
            "swipe_position": self.dual_swipe_preview.position,
            "comparison_output": comparison_path.name,
            "comparison_image_names": list(self.algorithm_image_names()),
            "comparison_stage": self.raw_preview_stage_combo.currentText(),
            "comparison_stage_index": self.raw_preview_stage_combo.currentData(),
            "source": str(Path(self.raw_algorithm_file.text()).resolve()),
            "resolution": {"width": raw.shape[1], "height": raw.shape[0]},
            "source_info": self.raw_algorithm_info,
            "bit_depth": self.raw_algorithm_info["bit_depth"],
            "quality_profile": "RAW Full",
            "input_mapping": {
                "key": self.raw_input_mapping_combo.currentData(),
                "name": self.raw_input_mapping_combo.currentText(),
                "details": mapping_details,
                "input_max": input_max,
                "output_file": mapping_path.name,
            },
            "pipeline": self.raw_pipeline_editor.configuration(),
            "execution": pipeline_execution_json(records),
            "total_ms": round(total_ms, 3),
            "original_output": original_path.name,
            "raw_csv_output": (
                raw_csv_path.name if raw_csv_path is not None else None
            ),
            "final_output": result_path.name,
            "overview_output": overview_path.name,
            "quality_original": quality_analysis_manifest(quality_original),
            "quality_final": quality_analysis_manifest(quality_final),
            "quality_layer_outputs": {
                "base_original": base_original_path.name,
                "base_final": base_final_path.name,
                "detail_original": detail_original_path.name,
                "detail_final": detail_final_path.name,
            },
            "quality_excel_output": quality_excel_path.name,
        }
        for execution, record in zip(manifest["execution"], records):
            if "output_file" in record:
                execution["output_file"] = record["output_file"]
        with open(out / "pipeline.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(out / "pipeline_summary.txt", "w", encoding="utf-8") as summary:
            summary.write(f"Source: {manifest['source']}\n")
            summary.write(f"Input Mapping: {mapping_details}\n")
            summary.write(f"Pipeline: {self.raw_pipeline_editor.summary()}\n")
            summary.write(f"Overview: {overview_path.name}\n")
            summary.write(
                "Quality: Image Quality profile; flat regions are separated by gray level "
                "and connected region before robust noise/uniformity aggregation.\n"
            )
            write_quality_summary(summary, "Original RAW", quality_original)
            write_quality_summary(summary, "Final Pipeline", quality_final)
            summary.write(f"Quality Excel: {quality_excel_path.name}\n")
            for execution in manifest["execution"]:
                summary.write(
                    f"Step {execution['pipeline_index']}: {execution['name']} - "
                    f"{execution['elapsed_ms']:.3f} ms - "
                    f"{' | '.join(execution['details'])}\n"
                )
            summary.write(f"Pipeline Total: {total_ms:.3f} ms\n")

        save_quality_excel(
            quality_excel_path,
            "Image Algorithm Quality Compare",
            {
                key: value
                for key, value in manifest.items()
                if key not in {"quality_original", "quality_final"}
            },
            [
                {
                    "label": "Original RAW",
                    "metrics": quality_original["metrics"],
                },
                {
                    "label": "Final Pipeline",
                    "metrics": quality_final["metrics"],
                },
            ],
        )

        self.log.append(f"Input Mapping: {mapping_details}")
        self.log.append(f"Pipeline: {self.raw_pipeline_editor.summary()}")
        for record in records:
            self.log.append(
                f"Step {record['pipeline_index']} {record['name']}: "
                f"{record['elapsed_ms']:.3f} ms | "
                f"{' | '.join(record['details'])}"
            )
        self.log.append(f"Pipeline total: {total_ms:.3f} ms")
        self.log.append(
            "Image Quality: Original="
            f"{quality_original['metrics']['overall_score']:.2f}, Final="
            f"{quality_final['metrics']['overall_score']:.2f}"
        )
        self.log_stats("Original 8-bit", raw8)
        self.log_stats("Final 8-bit", processed)
        self.log.append(f"Saved: {original_path}")
        self.log.append(f"Saved: {result_path}")
        self.log.append(f"Saved overview: {overview_path}")
        self.log.append(f"Saved quality analysis: {quality_excel_path}")
        self.log.append(f"Pipeline settings: {out / 'pipeline.json'}")

    @staticmethod
    def camera_combo_device_id(combo):
        return combo.currentData(CAMERA_DEVICE_ID_ROLE)

    @staticmethod
    def camera_combo_friendly_name(combo):
        name = combo.currentData(CAMERA_FRIENDLY_NAME_ROLE)
        return str(name) if name else combo.currentText()

    @staticmethod
    def populate_camera_combo(
        combo,
        devices,
        default_index,
        preferred_device_id=None,
        preferred_name=None,
    ):
        combo.blockSignals(True)
        try:
            combo.clear()
            selected_item = -1
            for device in devices:
                camera_index = int(device["index"])
                friendly_name = str(device["name"])
                device_id = str(device["device_id"])
                combo.addItem(
                    f"{friendly_name} (Camera {camera_index})",
                    camera_index,
                )
                item_index = combo.count() - 1
                combo.setItemData(
                    item_index, device_id, CAMERA_DEVICE_ID_ROLE
                )
                combo.setItemData(
                    item_index, friendly_name, CAMERA_FRIENDLY_NAME_ROLE
                )
                if preferred_device_id and device_id == preferred_device_id:
                    selected_item = item_index

            if preferred_device_id and selected_item < 0:
                disconnected_name = preferred_name or "Previously selected camera"
                combo.addItem(f"Disconnected - {disconnected_name}", None)
                selected_item = combo.count() - 1
                combo.setItemData(
                    selected_item,
                    preferred_device_id,
                    CAMERA_DEVICE_ID_ROLE,
                )
                combo.setItemData(
                    selected_item,
                    disconnected_name,
                    CAMERA_FRIENDLY_NAME_ROLE,
                )

            if combo.count() == 0:
                combo.addItem("No camera detected", None)
                combo.setItemData(
                    0, "", CAMERA_DEVICE_ID_ROLE
                )
                combo.setItemData(
                    0, "No camera detected", CAMERA_FRIENDLY_NAME_ROLE
                )
                selected_item = 0
            elif selected_item < 0:
                selected_item = combo.findData(default_index)
                if selected_item < 0:
                    selected_item = 0
            combo.setCurrentIndex(selected_item)
        finally:
            combo.blockSignals(False)

    def selected_camera_info(self, combo):
        camera_index = combo.currentData()
        return {
            "index": None if camera_index is None else int(camera_index),
            "device_id": self.camera_combo_device_id(combo),
            "name": self.camera_combo_friendly_name(combo),
        }

    def refresh_camera_devices(self, _checked=False, log_changes=True):
        devices = enumerate_camera_devices()
        signature = tuple(
            (
                int(device["index"]),
                str(device["device_id"]),
                str(device["name"]),
            )
            for device in devices
        )
        if signature == self._camera_device_signature:
            return

        old_devices = {
            str(device["device_id"]): str(device["name"])
            for device in self.camera_devices
        }
        new_devices = {
            str(device["device_id"]): str(device["name"])
            for device in devices
        }
        selections = []
        for combo in (
            self.viewer_uvc_camera_combo,
            self.uvc_camera_combo,
            self.dual_uvc_camera_a,
            self.dual_uvc_camera_b,
        ):
            selections.append(
                (
                    self.camera_combo_device_id(combo),
                    self.camera_combo_friendly_name(combo),
                )
            )

        self.populate_camera_combo(
            self.viewer_uvc_camera_combo,
            devices,
            DEFAULT_CAMERA_INDEX,
            *selections[0],
        )
        self.populate_camera_combo(
            self.uvc_camera_combo,
            devices,
            DEFAULT_CAMERA_INDEX,
            *selections[1],
        )
        self.populate_camera_combo(
            self.dual_uvc_camera_a,
            devices,
            DEFAULT_CAMERA_INDEX,
            *selections[2],
        )
        default_b_index = 0 if DEFAULT_CAMERA_INDEX != 0 else 1
        self.populate_camera_combo(
            self.dual_uvc_camera_b,
            devices,
            default_b_index,
            *selections[3],
        )

        self.camera_devices = devices
        self._camera_device_signature = signature
        devices_by_id = {
            str(device["device_id"]): device for device in devices
        }
        for worker in (
            self.viewer_uvc_worker,
            self.uvc_worker,
            self.dual_uvc_worker_a,
            self.dual_uvc_worker_b,
        ):
            if worker is None or not worker.camera_device_id:
                continue
            device = devices_by_id.get(str(worker.camera_device_id))
            if device is None:
                worker.update_camera_device(None)
            else:
                worker.update_camera_device(
                    int(device["index"]), str(device["name"])
                )

        if log_changes and hasattr(self, "log"):
            added = [
                new_devices[device_id]
                for device_id in new_devices.keys() - old_devices.keys()
            ]
            removed = [
                old_devices[device_id]
                for device_id in old_devices.keys() - new_devices.keys()
            ]
            if added:
                self.log.append("Camera connected: " + ", ".join(added))
            if removed:
                self.log.append("Camera disconnected: " + ", ".join(removed))

    def uvc_viewer_camera_manifest(self):
        camera_info = self.selected_camera_info(
            self.viewer_uvc_camera_combo
        )
        resolution = None
        if self.viewer_uvc_frame is not None:
            resolution = {
                "width": int(self.viewer_uvc_frame.shape[1]),
                "height": int(self.viewer_uvc_frame.shape[0]),
            }
        return {
            "capture_mode": DEFAULT_UVC_CAPTURE_MODE,
            "selected_capture_mode": getattr(
                self.viewer_uvc_worker, "selected_capture_mode", None
            ),
            "camera_index": camera_info["index"],
            "camera_name": camera_info["name"],
            "camera_device_id": camera_info["device_id"],
            "requested_fps": int(DEFAULT_UVC_FPS),
            "reported_camera_fps": round(self.viewer_uvc_actual_fps, 3),
            "resolution": resolution,
        }

    @staticmethod
    def write_uvc_viewer_analysis_summary(summary, stats):
        summary.write("Gray statistics (8-bit):\n")
        summary.write(
            f"  Pixels / gray levels: {stats['count']:,} / "
            f"{stats['unique_levels']}\n"
        )
        summary.write(
            f"  Min / max: {stats['min']} / {stats['max']}\n"
        )
        summary.write(
            f"  Mean / median / std: {stats['mean']:.3f} / "
            f"{stats['median']:.3f} / {stats['std']:.3f}\n"
        )
        summary.write(
            f"  P1 / P5 / P95 / P99: {stats['p1']:.3f} / "
            f"{stats['p5']:.3f} / {stats['p95']:.3f} / "
            f"{stats['p99']:.3f}\n"
        )
        summary.write(f"  P99-P1: {stats['p99_p1']:.3f}\n")
        summary.write(
            f"  Entropy: {stats['entropy_bits']:.3f} bits\n"
        )
        summary.write(
            f"  Zero clip: {stats['zero_count']:,} "
            f"({stats['zero_percent']:.3f}%)\n"
        )
        summary.write(
            f"  255 clip: {stats['full_count']:,} "
            f"({stats['full_percent']:.3f}%)\n"
        )
        summary.write(
            f"  Y<=16: {stats['limited_low_count']:,} "
            f"({stats['limited_low_percent']:.3f}%)\n"
        )
        summary.write(
            f"  Y>=235: {stats['limited_high_count']:,} "
            f"({stats['limited_high_percent']:.3f}%)\n"
        )

    def update_uvc_viewer_button_states(self):
        running = self.viewer_uvc_worker is not None
        ready = running and self.viewer_uvc_frame is not None
        recording = self.viewer_uvc_recording_worker is not None
        self.viewer_uvc_camera_combo.setEnabled(not recording)
        self.viewer_uvc_start_btn.setEnabled(not running)
        self.viewer_uvc_stop_btn.setEnabled(running)
        self.viewer_uvc_capture_btn.setEnabled(ready)
        self.viewer_uvc_record_start_btn.setEnabled(ready and not recording)
        self.viewer_uvc_record_stop_btn.setEnabled(recording)
        self.viewer_uvc_quality_enabled.setEnabled(not recording)
        if self.current_function_index() == self.FUNCTIONS.index("Single UVC Statistic"):
            self.run_btn.setEnabled(ready)

    def update_uvc_viewer_quality_option(self, enabled):
        enabled = bool(enabled)
        self.viewer_uvc_quality_group.setEnabled(enabled)
        self.viewer_uvc_quality = None
        self.viewer_uvc_temporal_tracker.reset()
        self.viewer_uvc_recording_quality_timeline = []
        self.update_single_quality_score_labels(
            self.viewer_uvc_quality_labels, None
        )
        self.viewer_uvc_quality_chart.set_paused(not enabled)
        if enabled and self.viewer_uvc_frame is not None:
            self.update_uvc_viewer_analysis(self.viewer_uvc_frame)

    def update_dual_uvc_quality_option(self, enabled):
        self.dual_uvc_quality_group.setEnabled(bool(enabled))
        self.clear_dual_uvc_quality(reset_timeline=False)
        self.dual_uvc_quality_chart.set_paused(not enabled)

    def update_uvc_algorithm_quality_option(self, enabled):
        self.uvc_quality_group.setEnabled(bool(enabled))
        self.clear_uvc_algorithm_quality(reset_timeline=False)
        self.uvc_quality_chart.set_paused(not enabled)

    def capture_uvc_viewer(self, _checked=False):
        if self.viewer_uvc_worker is None or self.viewer_uvc_frame is None:
            QMessageBox.information(
                self,
                "No UVC Frame",
                "Start Single UVC Statistic and wait for an image first.",
            )
            return
        self.update_uvc_viewer_analysis(self.viewer_uvc_frame)
        out = self.make_output("uvc_viewer_capture")
        image_path = out / "uvc_viewer.png"
        histogram_path = out / "gray_histogram.png"
        excel_path = out / "gray_analysis.xlsx"
        quality_excel_path = out / "quality_analysis.xlsx"
        base_path = out / "base.png"
        detail_path = out / "detail.png"
        save_png(image_path, self.viewer_uvc_frame)
        save_rgb_png(histogram_path, self.viewer_uvc_histogram_rgb)
        quality_enabled = self.viewer_uvc_quality is not None
        if quality_enabled:
            save_png(base_path, self.viewer_uvc_quality["base_u8"])
            save_png(detail_path, self.viewer_uvc_quality["detail_u8"])
        manifest = {
            "format_version": 1,
            "mode": "uvc_viewer_capture",
            "captured_at": datetime.now().astimezone().isoformat(
                timespec="seconds"
            ),
            **self.uvc_viewer_camera_manifest(),
            "gray_analysis": dict(self.viewer_uvc_stats),
            "quality_enabled": quality_enabled,
            "image_output": image_path.name,
            "histogram_output": histogram_path.name,
            "excel_output": excel_path.name,
        }
        if quality_enabled:
            manifest["quality"] = quality_analysis_manifest(self.viewer_uvc_quality)
            manifest["quality_excel_output"] = quality_excel_path.name
            manifest["quality_layer_outputs"] = {
                "base": base_path.name,
                "detail": detail_path.name,
            }
        with open(out / "capture.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(
            out / "capture_summary.txt", "w", encoding="utf-8"
        ) as summary:
            summary.write(f"Captured: {manifest['captured_at']}\n")
            summary.write(
                f"Camera: {manifest['camera_name']} "
                f"(index {manifest['camera_index']})\n"
            )
            summary.write(
                f"Capture mode: {manifest['capture_mode']} -> "
                f"{manifest['selected_capture_mode']}\n"
            )
            summary.write(
                f"Resolution: {manifest['resolution']['width']}x"
                f"{manifest['resolution']['height']}\n"
            )
            self.write_uvc_viewer_analysis_summary(
                summary, manifest["gray_analysis"]
            )
            if quality_enabled:
                write_quality_summary(
                    summary, "Single UVC Statistic", self.viewer_uvc_quality
                )
            else:
                summary.write("Quality scoring: Disabled\n")
            summary.write(f"Image: {image_path.name}\n")
            summary.write(f"Histogram: {histogram_path.name}\n")
            summary.write(f"Excel analysis: {excel_path.name}\n")
        save_histogram_excel(
            excel_path,
            "Single UVC Statistic Capture Analysis",
            {
                key: value
                for key, value in manifest.items()
                if key != "gray_analysis"
            },
            [
                {
                    "label": "Capture",
                    "stats": self.viewer_uvc_stats,
                    "histogram": self.viewer_uvc_histogram,
                }
            ],
            value_max=255,
            value_name="Gray",
        )
        if quality_enabled:
            save_quality_excel(
                quality_excel_path,
                "Single UVC Statistic Capture Quality",
                {
                    key: value
                    for key, value in manifest.items()
                    if key not in {"gray_analysis", "quality"}
                },
                [{
                    "label": "Single UVC Statistic",
                    "metrics": self.viewer_uvc_quality["metrics"],
                }],
                temporal_timelines=(
                    self.viewer_uvc_quality_chart.excel_timelines()
                ),
                timeline_metric_definitions=(
                    self.viewer_uvc_quality_chart.excel_metric_definition()
                ),
            )
        self.log.append(f"Saved Single UVC Statistic capture: {image_path}")
        self.log.append(f"Saved Single UVC Statistic histogram: {histogram_path}")
        self.log.append(f"Saved Single UVC Statistic Excel analysis: {excel_path}")
        if quality_enabled:
            self.log.append(f"Saved Single UVC Statistic quality analysis: {quality_excel_path}")

    def start_uvc_viewer_recording(self, _checked=False):
        if self.viewer_uvc_recording_worker is not None:
            QMessageBox.information(
                self,
                "Recording Active",
                "Single UVC Statistic recording is already active.",
            )
            return
        if self.viewer_uvc_worker is None or self.viewer_uvc_frame is None:
            QMessageBox.information(
                self,
                "No UVC Frame",
                "Start Single UVC Statistic and wait for an image first.",
            )
            return
        self.update_uvc_viewer_analysis(self.viewer_uvc_frame)

        requested_fps = float(DEFAULT_UVC_FPS)
        recording_fps = (
            self.viewer_uvc_actual_fps
            if 1.0 <= self.viewer_uvc_actual_fps <= 240.0
            else requested_fps
        )
        out = self.make_output("uvc_viewer_recording")
        video_path = out / "uvc_viewer.avi"
        overview_path = out / "00_uvc_viewer_overview.png"
        histogram_path = out / "00_uvc_viewer_histogram.png"
        base_start_path = out / "base_start.png"
        detail_start_path = out / "detail_start.png"
        save_png(overview_path, self.viewer_uvc_frame)
        save_rgb_png(histogram_path, self.viewer_uvc_histogram_rgb)
        quality_enabled = self.viewer_uvc_quality is not None
        if quality_enabled:
            save_png(base_start_path, self.viewer_uvc_quality["base_u8"])
            save_png(detail_start_path, self.viewer_uvc_quality["detail_u8"])
        started_at = datetime.now().astimezone().isoformat(timespec="seconds")
        manifest = {
            "format_version": 1,
            "mode": "uvc_viewer_recording",
            "started_at": started_at,
            **self.uvc_viewer_camera_manifest(),
            "gray_analysis_at_start": dict(self.viewer_uvc_stats),
            "quality_enabled": quality_enabled,
            "recording_fps": round(recording_fps, 3),
            "recording_layout": "Single UVC Statistic image; no text overlay",
            "recording_output": video_path.name,
            "received_fps_output": "received_fps.csv",
            "overview_output": overview_path.name,
            "histogram_at_start_output": histogram_path.name,
        }
        if quality_enabled:
            manifest["quality_at_start"] = quality_analysis_manifest(
                self.viewer_uvc_quality
            )
            manifest["quality_layers_at_start"] = {
                "base": base_start_path.name,
                "detail": detail_start_path.name,
            }
        with open(out / "recording.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)

        worker = UvcViewerRecordingThread(video_path, recording_fps, self)
        worker.finished.connect(self.handle_uvc_viewer_recording_finished)
        worker.finished.connect(worker.deleteLater)
        try:
            fps_path = out / "received_fps.csv"
            logger = UvcFpsCsvLogger(fps_path, self.viewer_uvc_worker.status_changed.emit)
        except Exception as exc:
            worker.deleteLater()
            QMessageBox.critical(self, "FPS Logging", f"Cannot create FPS log: {exc}")
            return
        self.viewer_uvc_fps_logger = logger
        self.viewer_uvc_worker.fps_logger = logger
        self.log.append(f"Recording received FPS CSV: {fps_path}")
        self.viewer_uvc_recording_worker = worker
        self.viewer_uvc_recording_output_dir = out
        self.viewer_uvc_recording_manifest = manifest
        self.viewer_uvc_recording_start_histogram = (
            self.viewer_uvc_histogram.copy()
        )
        self.viewer_uvc_recording_quality_start = (
            {"metrics": dict(self.viewer_uvc_quality["metrics"])}
            if quality_enabled
            else None
        )
        self.viewer_uvc_recording_quality_timeline = []
        self.viewer_uvc_recording_stop_reason = None
        self.viewer_uvc_recording_started_perf = time.perf_counter()
        self.update_uvc_viewer_button_states()
        worker.start()
        worker.submit(
            self.viewer_uvc_frame,
            self.viewer_uvc_recording_started_perf,
        )
        self.log.append(
            f"Single UVC Statistic recording started: {video_path} "
            f"({recording_fps:.3f} FPS, {UVC_RECORDING_FOURCC})"
        )
        self.log.append(f"Saved recording overview: {overview_path}")
        self.log.append(f"Saved recording histogram: {histogram_path}")

    def write_uvc_viewer_recording_summary(self, worker, reason):
        out = self.viewer_uvc_recording_output_dir
        manifest = self.viewer_uvc_recording_manifest
        if out is None or manifest is None:
            return
        recording_stats = worker.result_summary or {}
        self.update_uvc_viewer_analysis(self.viewer_uvc_frame)
        final_histogram_path = out / "uvc_viewer_histogram_final.png"
        save_rgb_png(
            final_histogram_path, self.viewer_uvc_histogram_rgb
        )
        final_histogram = self.viewer_uvc_histogram.copy()
        excel_path = out / "recording_analysis.xlsx"
        quality_excel_path = out / "quality_analysis.xlsx"
        base_stop_path = out / "base_stop.png"
        detail_stop_path = out / "detail_stop.png"
        quality_enabled = bool(manifest.get("quality_enabled", True))
        if quality_enabled:
            save_png(base_stop_path, self.viewer_uvc_quality["base_u8"])
            save_png(detail_stop_path, self.viewer_uvc_quality["detail_u8"])
        manifest["stopped_at"] = datetime.now().astimezone().isoformat(
            timespec="seconds"
        )
        manifest["stop_reason"] = reason
        manifest["gray_analysis_at_stop"] = dict(self.viewer_uvc_stats)
        manifest["histogram_at_stop_output"] = final_histogram_path.name
        manifest["recording"] = recording_stats
        manifest["excel_output"] = excel_path.name
        if quality_enabled:
            manifest["quality_at_stop"] = quality_analysis_manifest(
                self.viewer_uvc_quality
            )
            manifest["quality_timeline"] = list(
                self.viewer_uvc_recording_quality_timeline
            )
            manifest["quality_excel_output"] = quality_excel_path.name
            manifest["quality_layers_at_stop"] = {
                "base": base_stop_path.name,
                "detail": detail_stop_path.name,
            }
        with open(out / "recording.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(
            out / "recording_summary.txt", "w", encoding="utf-8"
        ) as summary:
            summary.write(f"Started: {manifest['started_at']}\n")
            summary.write(f"Stopped: {manifest['stopped_at']}\n")
            summary.write(f"Reason: {reason}\n")
            summary.write(
                f"Camera: {manifest['camera_name']} "
                f"(index {manifest['camera_index']})\n"
            )
            summary.write(
                f"Capture mode: {manifest['capture_mode']} -> "
                f"{manifest['selected_capture_mode']}\n"
            )
            summary.write(
                f"Resolution: {manifest['resolution']['width']}x"
                f"{manifest['resolution']['height']}\n"
            )
            summary.write(
                f"Recording FPS: {manifest['recording_fps']:.3f}\n"
            )
            summary.write("Analysis at start:\n")
            self.write_uvc_viewer_analysis_summary(
                summary, manifest["gray_analysis_at_start"]
            )
            if quality_enabled:
                write_quality_summary(
                    summary,
                    "Single UVC Statistic Start",
                    self.viewer_uvc_recording_quality_start,
                )
            summary.write("Analysis at stop:\n")
            self.write_uvc_viewer_analysis_summary(
                summary, manifest["gray_analysis_at_stop"]
            )
            if quality_enabled:
                write_quality_summary(
                    summary, "Single UVC Statistic Stop", self.viewer_uvc_quality
                )
            else:
                summary.write("Quality scoring: Disabled\n")
            summary.write(
                f"Codec: {recording_stats.get('codec', 'N/A')}\n"
            )
            summary.write(
                f"Duration: "
                f"{recording_stats.get('duration_seconds', 0.0):.3f} s\n"
            )
            summary.write(
                f"Submitted / written frames: "
                f"{recording_stats.get('submitted_frames', 0)} / "
                f"{recording_stats.get('written_frames', 0)}\n"
            )
            summary.write(
                f"Duplicate / queue dropped / timestamp dropped: "
                f"{recording_stats.get('duplicate_frames', 0)} / "
                f"{recording_stats.get('queue_dropped_frames', 0)} / "
                f"{recording_stats.get('timestamp_dropped_frames', 0)}\n"
            )
            summary.write(
                f"File size: "
                f"{recording_stats.get('file_size_bytes', 0)} bytes\n"
            )
            summary.write(f"Video: {manifest['recording_output']}\n")
            summary.write(f"Overview: {manifest['overview_output']}\n")
            summary.write(
                f"Histogram at start: "
                f"{manifest['histogram_at_start_output']}\n"
            )
            summary.write(
                f"Histogram at stop: "
                f"{manifest['histogram_at_stop_output']}\n"
            )
            summary.write(f"Excel analysis: {excel_path.name}\n")
            if quality_enabled:
                summary.write(f"Quality Excel: {quality_excel_path.name}\n")
        start_histogram = self.viewer_uvc_recording_start_histogram
        if start_histogram is None:
            start_histogram = final_histogram
        save_histogram_excel(
            excel_path,
            "Single UVC Statistic Recording Analysis",
            {
                key: value
                for key, value in manifest.items()
                if key
                not in {
                    "gray_analysis_at_start",
                    "gray_analysis_at_stop",
                }
            },
            [
                {
                    "label": "Start",
                    "stats": manifest["gray_analysis_at_start"],
                    "histogram": start_histogram,
                },
                {
                    "label": "Stop",
                    "stats": manifest["gray_analysis_at_stop"],
                    "histogram": final_histogram,
                },
            ],
            value_max=255,
            value_name="Gray",
        )
        if quality_enabled:
            quality_start = self.viewer_uvc_recording_quality_start
            if quality_start is None:
                quality_start = self.viewer_uvc_quality
            save_quality_excel(
                quality_excel_path,
                "Single UVC Statistic Recording Quality",
                {
                    key: value
                    for key, value in manifest.items()
                    if key
                    not in {
                        "gray_analysis_at_start",
                        "gray_analysis_at_stop",
                        "quality_at_start",
                        "quality_at_stop",
                        "quality_timeline",
                    }
                },
                [
                    {"label": "Start", "metrics": quality_start["metrics"]},
                    {
                        "label": "Stop",
                        "metrics": self.viewer_uvc_quality["metrics"],
                    },
                ],
                temporal_timelines={
                    "Single UVC Statistic": list(
                        self.viewer_uvc_recording_quality_timeline
                    )
                },
            )
        self.log.append(f"Saved Single UVC Statistic recording Excel analysis: {excel_path}")
        if quality_enabled:
            self.log.append(f"Saved Single UVC Statistic recording quality: {quality_excel_path}")

    def close_uvc_viewer_fps_log(self):
        logger = getattr(self, "viewer_uvc_fps_logger", None)
        if logger is not None:
            logger.close()
            self.viewer_uvc_fps_logger = None

    def finalize_uvc_viewer_recording_worker(self, worker, reason):
        if worker is not self.viewer_uvc_recording_worker:
            return
        self.close_uvc_viewer_fps_log()
        try:
            self.write_uvc_viewer_recording_summary(worker, reason)
        except Exception as exc:
            self.log.append(f"Single UVC Statistic recording summary error: {exc}")
        out = self.viewer_uvc_recording_output_dir
        if worker.error_message:
            self.log.append(
                f"Single UVC Statistic recording error: {worker.error_message}"
            )
        elif out is not None:
            stats = worker.result_summary or {}
            self.log.append(
                f"Single UVC Statistic recording stopped: {out / 'uvc_viewer.avi'}; "
                f"{stats.get('written_frames', 0)} frames, "
                f"{stats.get('file_size_bytes', 0)} bytes"
            )
        self.viewer_uvc_recording_worker = None
        self.viewer_uvc_recording_output_dir = None
        self.viewer_uvc_recording_manifest = None
        self.viewer_uvc_recording_start_histogram = None
        self.viewer_uvc_recording_quality_start = None
        self.viewer_uvc_recording_quality_timeline = []
        self.viewer_uvc_recording_stop_reason = None
        self.viewer_uvc_recording_started_perf = 0.0
        self.update_uvc_viewer_button_states()

    def stop_uvc_viewer_recording(
        self,
        _checked=False,
        reason=None,
        log_message=True,
        wait_ms=3000,
    ):
        self.close_uvc_viewer_fps_log()
        worker = self.viewer_uvc_recording_worker
        if worker is None:
            return True
        stop_reason = reason or "Stopped by user"
        self.viewer_uvc_recording_stop_reason = stop_reason
        self.viewer_uvc_record_start_btn.setEnabled(False)
        self.viewer_uvc_record_stop_btn.setEnabled(False)
        worker.stop()
        if worker.isRunning() and not worker.wait(wait_ms):
            if log_message:
                self.log.append(
                    "Single UVC Statistic recording thread is still finishing the AVI file."
                )
            return False
        self.finalize_uvc_viewer_recording_worker(worker, stop_reason)
        return True

    def handle_uvc_viewer_recording_finished(self):
        worker = self.sender()
        if worker is not self.viewer_uvc_recording_worker:
            return
        reason = self.viewer_uvc_recording_stop_reason
        if reason is None:
            reason = (
                "Recording error"
                if worker.error_message
                else "Recording worker finished"
            )
        self.finalize_uvc_viewer_recording_worker(worker, reason)

    def start_uvc_viewer(self, _checked=False):
        if cv2 is None:
            message = "OpenCV not installed. pip install opencv-python"
            self.log.append(f"\nERROR: {message}")
            QMessageBox.critical(self, "Camera Error", message)
            return
        if not self.stop_uvc_viewer(log_message=False):
            self.viewer_uvc_restart_when_stopped = True
            self.log.append(
                "Previous Single UVC Statistic thread is still stopping; the camera "
                "will start automatically afterward."
            )
            return

        camera_info = self.selected_camera_info(
            self.viewer_uvc_camera_combo
        )
        camera_index = camera_info["index"]
        if camera_index is None:
            QMessageBox.warning(
                self,
                "Camera Disconnected",
                "Select a connected camera before starting Single UVC Statistic.",
            )
            return

        requested_fps = DEFAULT_UVC_FPS
        worker = UvcCaptureThread(
            camera_index,
            self.width,
            self.height,
            requested_fps,
            UVC_RECONNECT_SECONDS,
            self,
            frame_consumer=self.viewer_uvc_buffer.submit,
            capture_mode=DEFAULT_UVC_CAPTURE_MODE,
            camera_device_id=camera_info["device_id"],
            camera_name=camera_info["name"],
        )
        worker.camera_opened.connect(self.handle_uvc_viewer_camera_opened)
        worker.status_changed.connect(self.handle_uvc_viewer_status)
        worker.stream_interrupted.connect(
            self.handle_uvc_viewer_stream_interrupted
        )
        worker.finished.connect(self.handle_uvc_viewer_worker_finished)
        worker.finished.connect(worker.deleteLater)

        self.viewer_uvc_worker = worker
        self.viewer_uvc_frame = None
        self.clear_uvc_viewer_analysis()
        self.viewer_uvc_actual_fps = 0.0
        self.viewer_uvc_output_fps = 0.0
        self.viewer_uvc_frame_count = 0
        self.viewer_uvc_fps_started = time.perf_counter()
        self.viewer_uvc_buffer.clear()
        self.viewer_uvc_buffer_sequence = 0
        self.preview_result.set_array(None)
        self.current_result = None
        self.update_uvc_viewer_button_states()
        self.viewer_uvc_status.setText(
            f"Opening {camera_info['name']} (Camera {camera_index}, "
            f"{DEFAULT_UVC_CAPTURE_MODE})..."
        )
        self.viewer_uvc_display_timer.start()
        self.log.append(
            f"Starting Single UVC Statistic: {camera_info['name']} "
            f"(Camera {camera_index}, {requested_fps} FPS, "
            f"{DEFAULT_UVC_CAPTURE_MODE})"
        )
        worker.start()

    def stop_uvc_viewer(
        self, _checked=False, log_message=True, wait_ms=2000
    ):
        self.viewer_uvc_restart_when_stopped = False
        self.close_uvc_video_window()
        self.close_uvc_histogram_window()
        self.viewer_uvc_display_timer.stop()
        if not self.stop_uvc_viewer_recording(
            reason="Single UVC Statistic camera stopped",
            log_message=log_message,
            wait_ms=wait_ms,
        ):
            return False
        worker = self.viewer_uvc_worker
        was_running = worker is not None
        if worker is not None:
            worker.stop()
            if not worker.wait(wait_ms):
                self.viewer_uvc_start_btn.setEnabled(False)
                self.viewer_uvc_stop_btn.setEnabled(False)
                self.viewer_uvc_status.setText(
                    "Single UVC Statistic capture thread is stopping..."
                )
                return False
            if self.viewer_uvc_worker is worker:
                self.viewer_uvc_worker = None

        self.viewer_uvc_actual_fps = 0.0
        self.viewer_uvc_output_fps = 0.0
        self.viewer_uvc_temporal_tracker.reset()
        self.viewer_uvc_quality = None
        self.update_single_quality_score_labels(
            self.viewer_uvc_quality_labels, None
        )
        self.update_uvc_viewer_button_states()
        self.viewer_uvc_status.setText("Camera stopped")
        if was_running and log_message:
            self.log.append("Single UVC Statistic camera stopped")
        return True

    def switch_uvc_viewer_setting(self, _combo_index=None):
        if self.viewer_uvc_worker is not None:
            self.start_uvc_viewer()

    def handle_uvc_viewer_camera_opened(
        self, camera_index, actual_width, actual_height, actual_fps
    ):
        if self.sender() is not self.viewer_uvc_worker:
            return
        self.viewer_uvc_temporal_tracker.reset()
        self.viewer_uvc_quality = None
        self.update_single_quality_score_labels(
            self.viewer_uvc_quality_labels, None
        )
        self.viewer_uvc_actual_fps = float(actual_fps)
        camera_name = getattr(
            self.viewer_uvc_worker,
            "camera_name",
            f"Camera {camera_index}",
        )
        self.viewer_uvc_status.setText(
            f"{camera_name} (Camera {camera_index}) opened: "
            f"{actual_width}x{actual_height}, "
            f"reported {actual_fps:.1f} FPS; waiting for frame..."
        )

    def handle_uvc_viewer_status(self, message):
        if self.sender() is self.viewer_uvc_worker:
            self.viewer_uvc_status.setText(message)

    def handle_uvc_viewer_stream_interrupted(self, message):
        if self.sender() is not self.viewer_uvc_worker:
            return
        self.log.append(
            f"Single UVC Statistic automatically stopped: {message}"
        )
        if self.stop_uvc_viewer(log_message=False):
            self.viewer_uvc_status.setText(
                "Camera stopped — stream interrupted"
            )

    def consume_uvc_viewer_latest_frame(self):
        frame, sequence, frame_time = self.viewer_uvc_buffer.take_latest(
            self.viewer_uvc_buffer_sequence
        )
        if frame is None:
            return
        self.viewer_uvc_buffer_sequence = sequence
        gray = uvc_to_gray(frame)
        display = (
            gray.copy()
            if gray.dtype == np.uint8
            else raw_to_display_u8(gray, self.raw_max)
        )
        if self.app_settings["camera_mirror"]["camera_a"]:
            display = np.ascontiguousarray(display[:, ::-1])
        self.viewer_uvc_frame = display
        if self.viewer_uvc_quality_enabled.isChecked():
            self.viewer_uvc_temporal_tracker.add_frame(display, frame_time)
        if (
            self.viewer_uvc_last_analysis_time <= 0.0
            or frame_time - self.viewer_uvc_last_analysis_time
            >= UVC_VIEWER_ANALYSIS_INTERVAL_SECONDS
        ):
            self.update_uvc_viewer_analysis(display, frame_time)
        self.viewer_uvc_frame_count += 1
        elapsed = frame_time - self.viewer_uvc_fps_started
        if elapsed >= 0.5:
            self.viewer_uvc_output_fps = self.viewer_uvc_frame_count / elapsed
            self.viewer_uvc_frame_count = 0
            self.viewer_uvc_fps_started = frame_time

        if self.current_function_index() == self.FUNCTIONS.index("Single UVC Statistic"):
            self.current_result = display
            self.preview_result.set_array(display)
            self.update_uvc_video_window()
        recording_worker = self.viewer_uvc_recording_worker
        if recording_worker is not None and recording_worker.isRunning():
            recording_worker.submit(display, frame_time)
        self.update_uvc_viewer_button_states()
        camera_info = self.selected_camera_info(
            self.viewer_uvc_camera_combo
        )
        self.viewer_uvc_status.setText(
            f"{camera_info['name']} (Camera {camera_info['index']}) "
            f"{display.shape[1]}x{display.shape[0]} | "
            f"Output FPS {self.viewer_uvc_output_fps:.1f}"
        )

    def handle_uvc_viewer_worker_finished(self):
        if self.sender() is not self.viewer_uvc_worker:
            return
        restart = self.viewer_uvc_restart_when_stopped
        self.viewer_uvc_restart_when_stopped = False
        self.close_uvc_video_window()
        self.close_uvc_histogram_window()
        self.stop_uvc_viewer_recording(
            reason="Single UVC Statistic capture thread finished"
        )
        self.viewer_uvc_worker = None
        self.viewer_uvc_display_timer.stop()
        self.viewer_uvc_actual_fps = 0.0
        self.viewer_uvc_output_fps = 0.0
        self.viewer_uvc_temporal_tracker.reset()
        self.viewer_uvc_quality = None
        self.update_single_quality_score_labels(
            self.viewer_uvc_quality_labels, None
        )
        self.update_uvc_viewer_button_states()
        self.viewer_uvc_status.setText("Camera stopped")
        if (
            restart
            and self.current_function_index()
            == self.FUNCTIONS.index("Single UVC Statistic")
        ):
            QTimer.singleShot(0, self.start_uvc_viewer)

    def dual_uvc_image_names(self):
        return (
            self.dual_uvc_image_name_a.currentText(),
            self.dual_uvc_image_name_b.currentText(),
        )

    def dual_uvc_image_file_tags(self):
        return (
            self.dual_uvc_image_name_a.currentData(),
            self.dual_uvc_image_name_b.currentData(),
        )

    def dual_uvc_image_name_changed(self, _combo_index=None):
        name_a, name_b = self.dual_uvc_image_names()
        self.dual_uvc_group_a.setTitle(f"UVC Camera A — {name_a}")
        self.dual_uvc_group_b.setTitle(f"UVC Camera B — {name_b}")
        quality_headers = getattr(
            self.dual_uvc_quality_group,
            "comparison_header_labels",
            (),
        )
        if len(quality_headers) == 2:
            quality_headers[0].setText(name_a)
            quality_headers[1].setText(name_b)
        if hasattr(self, "dual_uvc_quality_chart"):
            self.dual_uvc_quality_chart.set_series_labels(
                camera_a=name_a,
                camera_b=name_b,
            )
        if self.dual_uvc_worker_a is None:
            self.dual_uvc_status_a.setText(f"{name_a} (Camera A) stopped")
        if self.dual_uvc_worker_b is None:
            self.dual_uvc_status_b.setText(f"{name_b} (Camera B) stopped")
        dual_page_index = len(self.FUNCTIONS) - 3
        if self.current_function_index() == dual_page_index:
            self.preview_original.set_title(name_a)
            self.preview_result.set_title(name_b)
            self.preview_original_caption.setText(name_a)
            self.preview_result_caption.setText(name_b)
        self.update_uvc_video_window()

    def dual_uvc_setting_changed(self, _combo_index=None):
        if (
            self.dual_uvc_worker_a is not None
            or self.dual_uvc_worker_b is not None
        ):
            self.start_dual_uvc_cameras()

    def dual_uvc_frames_ready(self):
        return (
            self.dual_uvc_worker_a is not None
            and self.dual_uvc_worker_b is not None
            and self.dual_uvc_frame_a is not None
            and self.dual_uvc_frame_b is not None
        )

    def clear_dual_uvc_quality(self, reset_timeline=True):
        self.dual_uvc_quality_a = None
        self.dual_uvc_quality_b = None
        self.dual_uvc_quality_last_update = 0.0
        self.dual_uvc_temporal_tracker_a.reset()
        self.dual_uvc_temporal_tracker_b.reset()
        if hasattr(self, "dual_uvc_quality_labels"):
            self.update_quality_score_labels(
                self.dual_uvc_quality_labels, None, None
            )
        if reset_timeline and hasattr(self, "dual_uvc_quality_chart"):
            self.dual_uvc_quality_chart.clear()

    def update_dual_uvc_quality(self, force=False):
        if not self.dual_uvc_quality_enabled.isChecked():
            return False
        if self.dual_uvc_frame_a is None or self.dual_uvc_frame_b is None:
            return False
        now = time.perf_counter()
        if (
            not force
            and now - self.dual_uvc_quality_last_update
            < DUAL_UVC_QUALITY_INTERVAL_SECONDS
        ):
            return False
        result_a = add_uvc_temporal_quality(
            thermal_uvc_quality_analysis(
                self.dual_uvc_frame_a,
                max_analysis_pixels=UVC_QUALITY_ANALYSIS_MAX_PIXELS,
            ),
            self.dual_uvc_temporal_tracker_a.result(),
        )
        result_b = add_uvc_temporal_quality(
            thermal_uvc_quality_analysis(
                self.dual_uvc_frame_b,
                max_analysis_pixels=UVC_QUALITY_ANALYSIS_MAX_PIXELS,
            ),
            self.dual_uvc_temporal_tracker_b.result(),
        )
        self.dual_uvc_quality_a = result_a
        self.dual_uvc_quality_b = result_b
        self.dual_uvc_quality_last_update = now
        self.update_quality_score_labels(
            self.dual_uvc_quality_labels, result_a, result_b
        )
        self.dual_uvc_quality_chart.append_metrics(
            now,
            camera_a=result_a["metrics"],
            camera_b=result_b["metrics"],
        )
        if self.dual_uvc_recording_worker is not None:
            elapsed = max(0.0, now - self.dual_uvc_recording_started_perf)
            self.dual_uvc_recording_quality_timeline.append(
                {
                    "elapsed_seconds": round(elapsed, 3),
                    "camera_a": dict(result_a["metrics"]),
                    "camera_b": dict(result_b["metrics"]),
                }
            )
        return True

    def update_dual_uvc_button_states(self):
        running = (
            self.dual_uvc_worker_a is not None
            or self.dual_uvc_worker_b is not None
        )
        ready = self.dual_uvc_frames_ready()
        recording = self.dual_uvc_recording_worker is not None
        self.dual_uvc_image_name_a.setEnabled(not recording)
        self.dual_uvc_image_name_b.setEnabled(not recording)
        self.dual_uvc_quality_enabled.setEnabled(not recording)
        self.dual_uvc_start_btn.setEnabled(not running)
        self.dual_uvc_stop_btn.setEnabled(running)
        self.dual_uvc_capture_btn.setEnabled(ready)
        self.dual_uvc_record_start_btn.setEnabled(ready and not recording)
        self.dual_uvc_record_stop_btn.setEnabled(recording)

    def start_dual_uvc_cameras(self, _checked=False):
        if cv2 is None:
            message = "OpenCV not installed. pip install opencv-python"
            self.log.append(f"\nERROR: {message}")
            QMessageBox.critical(self, "Camera Error", message)
            return
        camera_info_a = self.selected_camera_info(self.dual_uvc_camera_a)
        camera_info_b = self.selected_camera_info(self.dual_uvc_camera_b)
        if camera_info_a["index"] is None or camera_info_b["index"] is None:
            QMessageBox.warning(
                self,
                "Camera Disconnected",
                "Select two connected cameras before starting Dual UVC.",
            )
            return
        camera_a = camera_info_a["index"]
        camera_b = camera_info_b["index"]
        if camera_a == camera_b:
            if (
                self.dual_uvc_worker_a is not None
                or self.dual_uvc_worker_b is not None
            ):
                self.stop_dual_uvc_cameras(log_message=False)
            QMessageBox.warning(
                self,
                "Duplicate Camera Index",
                "Camera A and Camera B must use different CAMERA_INDEX values.",
            )
            return

        if not self.stop_dual_uvc_cameras(log_message=False):
            self.dual_uvc_restart_when_stopped = True
            self.log.append(
                "Previous dual UVC capture threads are still stopping; "
                "both cameras will restart afterward."
            )
            return
        self.dual_uvc_restart_when_stopped = False

        worker_a = UvcCaptureThread(
            camera_a,
            self.width,
            self.height,
            int(DEFAULT_UVC_FPS),
            UVC_RECONNECT_SECONDS,
            self,
            frame_consumer=self.dual_uvc_buffer_a.submit,
            capture_mode=DEFAULT_UVC_CAPTURE_MODE,
            camera_device_id=camera_info_a["device_id"],
            camera_name=camera_info_a["name"],
        )
        worker_b = UvcCaptureThread(
            camera_b,
            self.width,
            self.height,
            int(DEFAULT_UVC_FPS),
            UVC_RECONNECT_SECONDS,
            self,
            frame_consumer=self.dual_uvc_buffer_b.submit,
            capture_mode=DEFAULT_UVC_CAPTURE_MODE,
            camera_device_id=camera_info_b["device_id"],
            camera_name=camera_info_b["name"],
        )
        worker_a.camera_opened.connect(self.handle_dual_uvc_camera_opened_a)
        worker_b.camera_opened.connect(self.handle_dual_uvc_camera_opened_b)
        worker_a.status_changed.connect(self.handle_dual_uvc_status_a)
        worker_b.status_changed.connect(self.handle_dual_uvc_status_b)
        worker_a.stream_interrupted.connect(
            self.handle_dual_uvc_stream_interrupted
        )
        worker_b.stream_interrupted.connect(
            self.handle_dual_uvc_stream_interrupted
        )
        worker_a.finished.connect(self.handle_dual_uvc_worker_finished)
        worker_b.finished.connect(self.handle_dual_uvc_worker_finished)
        worker_a.finished.connect(worker_a.deleteLater)
        worker_b.finished.connect(worker_b.deleteLater)
        self.dual_uvc_worker_a = worker_a
        self.dual_uvc_worker_b = worker_b
        self.dual_uvc_frame_a = None
        self.dual_uvc_frame_b = None
        self.dual_uvc_buffer_a.clear()
        self.dual_uvc_buffer_b.clear()
        self.dual_uvc_buffer_sequence_a = 0
        self.dual_uvc_buffer_sequence_b = 0
        self.dual_uvc_frame_time_a = 0.0
        self.dual_uvc_frame_time_b = 0.0
        self.dual_uvc_actual_fps_a = 0.0
        self.dual_uvc_actual_fps_b = 0.0
        self.dual_uvc_output_fps_a = 0.0
        self.dual_uvc_output_fps_b = 0.0
        self.dual_uvc_frame_count_a = 0
        self.dual_uvc_frame_count_b = 0
        self.clear_dual_uvc_quality()
        now = time.perf_counter()
        self.dual_uvc_fps_started_a = now
        self.dual_uvc_fps_started_b = now
        name_a, name_b = self.dual_uvc_image_names()
        self.dual_uvc_status_a.setText(
            f"Opening {name_a} (Camera A, index {camera_a})..."
        )
        self.dual_uvc_status_b.setText(
            f"Opening {name_b} (Camera B, index {camera_b})..."
        )
        self.preview_original.set_array(None)
        self.preview_result.set_array(None)
        self.update_dual_uvc_button_states()
        self.dual_uvc_display_timer.start()
        self.log.append(
            f"Starting Dual UVC: A={name_a}, "
            f"{camera_info_a['name']}, index {camera_a} "
            f"({DEFAULT_UVC_CAPTURE_MODE}), B={name_b}, "
            f"{camera_info_b['name']}, index {camera_b} "
            f"({DEFAULT_UVC_CAPTURE_MODE})"
        )
        worker_a.start()
        worker_b.start()

    def stop_dual_uvc_cameras(
        self, _checked=False, log_message=True, wait_ms=2000
    ):
        self.dual_uvc_restart_when_stopped = False
        self.close_uvc_video_window()
        self.dual_uvc_display_timer.stop()
        self.stop_dual_uvc_recording(
            reason="Dual UVC cameras stopped", log_message=log_message
        )
        worker_a = self.dual_uvc_worker_a
        worker_b = self.dual_uvc_worker_b
        was_running = worker_a is not None or worker_b is not None
        for worker in (worker_a, worker_b):
            if worker is not None:
                worker.stop()

        all_stopped = True
        if worker_a is not None:
            if worker_a.wait(wait_ms):
                if self.dual_uvc_worker_a is worker_a:
                    self.dual_uvc_worker_a = None
            else:
                all_stopped = False
        if worker_b is not None:
            if worker_b.wait(wait_ms):
                if self.dual_uvc_worker_b is worker_b:
                    self.dual_uvc_worker_b = None
            else:
                all_stopped = False

        if not all_stopped:
            self.dual_uvc_start_btn.setEnabled(False)
            self.dual_uvc_stop_btn.setEnabled(False)
            self.dual_uvc_capture_btn.setEnabled(False)
            self.dual_uvc_record_start_btn.setEnabled(False)
            self.dual_uvc_record_stop_btn.setEnabled(False)
            self.log.append("Dual UVC capture threads are still stopping.")
            return False

        self.dual_uvc_actual_fps_a = 0.0
        self.dual_uvc_actual_fps_b = 0.0
        name_a, name_b = self.dual_uvc_image_names()
        self.dual_uvc_status_a.setText(f"{name_a} (Camera A) stopped")
        self.dual_uvc_status_b.setText(f"{name_b} (Camera B) stopped")
        self.clear_dual_uvc_quality()
        self.update_dual_uvc_button_states()
        if was_running and log_message:
            self.log.append("Dual UVC cameras stopped")
        return True

    def handle_dual_uvc_camera_opened_a(
        self, camera_index, actual_width, actual_height, actual_fps
    ):
        if self.sender() is not self.dual_uvc_worker_a:
            return
        self.clear_dual_uvc_quality()
        self.dual_uvc_actual_fps_a = float(actual_fps)
        name_a = self.dual_uvc_image_name_a.currentText()
        camera_name = getattr(
            self.dual_uvc_worker_a, "camera_name", f"Camera {camera_index}"
        )
        self.dual_uvc_status_a.setText(
            f"{name_a} (Camera A): {camera_name}, index {camera_index}, "
            f"{actual_width}x{actual_height}, "
            f"reported {actual_fps:.1f} FPS"
        )

    def handle_dual_uvc_camera_opened_b(
        self, camera_index, actual_width, actual_height, actual_fps
    ):
        if self.sender() is not self.dual_uvc_worker_b:
            return
        self.clear_dual_uvc_quality()
        self.dual_uvc_actual_fps_b = float(actual_fps)
        name_b = self.dual_uvc_image_name_b.currentText()
        camera_name = getattr(
            self.dual_uvc_worker_b, "camera_name", f"Camera {camera_index}"
        )
        self.dual_uvc_status_b.setText(
            f"{name_b} (Camera B): {camera_name}, index {camera_index}, "
            f"{actual_width}x{actual_height}, "
            f"reported {actual_fps:.1f} FPS"
        )

    def handle_dual_uvc_status_a(self, message):
        if self.sender() is self.dual_uvc_worker_a:
            self.dual_uvc_status_a.setText(message)

    def handle_dual_uvc_status_b(self, message):
        if self.sender() is self.dual_uvc_worker_b:
            self.dual_uvc_status_b.setText(message)

    def handle_dual_uvc_stream_interrupted(self, message):
        if self.sender() not in (
            self.dual_uvc_worker_a,
            self.dual_uvc_worker_b,
        ):
            return
        self.log.append(f"Dual UVC automatically stopped: {message}")
        if self.stop_dual_uvc_cameras(log_message=False):
            name_a, name_b = self.dual_uvc_image_names()
            self.dual_uvc_status_a.setText(
                f"{name_a} (Camera A) stopped — stream interrupted"
            )
            self.dual_uvc_status_b.setText(
                f"{name_b} (Camera B) stopped — stream interrupted"
            )

    def consume_dual_uvc_latest_frames(self):
        updated = False
        frame_a, sequence_a, frame_time_a = self.dual_uvc_buffer_a.take_latest(
            self.dual_uvc_buffer_sequence_a
        )
        if frame_a is not None:
            self.dual_uvc_buffer_sequence_a = sequence_a
            self.handle_dual_uvc_frame(
                "a", frame_a, frame_time_a, submit_recording=False
            )
            updated = True
        frame_b, sequence_b, frame_time_b = self.dual_uvc_buffer_b.take_latest(
            self.dual_uvc_buffer_sequence_b
        )
        if frame_b is not None:
            self.dual_uvc_buffer_sequence_b = sequence_b
            self.handle_dual_uvc_frame(
                "b", frame_b, frame_time_b, submit_recording=False
            )
            updated = True
        if updated:
            self.update_dual_uvc_quality()
            self.submit_dual_uvc_recording_pair()

    def handle_dual_uvc_frame_a(self, frame):
        if self.sender() is not self.dual_uvc_worker_a:
            return
        self.handle_dual_uvc_frame("a", frame)

    def handle_dual_uvc_frame_b(self, frame):
        if self.sender() is not self.dual_uvc_worker_b:
            return
        self.handle_dual_uvc_frame("b", frame)

    def handle_dual_uvc_frame(
        self, camera_key, frame, frame_time=None, submit_recording=True
    ):
        gray = uvc_to_gray(frame)
        display = (
            gray.copy()
            if gray.dtype == np.uint8
            else raw_to_display_u8(gray, self.raw_max)
        )
        mirror_key = "camera_a" if camera_key == "a" else "camera_b"
        if self.app_settings["camera_mirror"][mirror_key]:
            display = np.ascontiguousarray(display[:, ::-1])
        now = time.perf_counter() if frame_time is None else float(frame_time)
        if self.dual_uvc_quality_enabled.isChecked():
            if camera_key == "a":
                self.dual_uvc_temporal_tracker_a.add_frame(display, now)
            else:
                self.dual_uvc_temporal_tracker_b.add_frame(display, now)
        if camera_key == "a":
            self.dual_uvc_frame_a = display
            self.dual_uvc_frame_time_a = now
            self.dual_uvc_frame_count_a += 1
            elapsed = now - self.dual_uvc_fps_started_a
            if elapsed >= 0.5:
                self.dual_uvc_output_fps_a = (
                    self.dual_uvc_frame_count_a / elapsed
                )
                self.dual_uvc_frame_count_a = 0
                self.dual_uvc_fps_started_a = now
            self.dual_uvc_status_a.setText(
                f"{self.dual_uvc_image_name_a.currentText()} "
                f"(Camera A) {display.shape[1]}x{display.shape[0]} | "
                f"Output FPS {self.dual_uvc_output_fps_a:.1f}"
            )
        else:
            self.dual_uvc_frame_b = display
            self.dual_uvc_frame_time_b = now
            self.dual_uvc_frame_count_b += 1
            elapsed = now - self.dual_uvc_fps_started_b
            if elapsed >= 0.5:
                self.dual_uvc_output_fps_b = (
                    self.dual_uvc_frame_count_b / elapsed
                )
                self.dual_uvc_frame_count_b = 0
                self.dual_uvc_fps_started_b = now
            self.dual_uvc_status_b.setText(
                f"{self.dual_uvc_image_name_b.currentText()} "
                f"(Camera B) {display.shape[1]}x{display.shape[0]} | "
                f"Output FPS {self.dual_uvc_output_fps_b:.1f}"
            )

        dual_page_index = len(self.FUNCTIONS) - 3
        if self.current_function_index() == dual_page_index:
            if self.dual_uvc_frame_a is not None:
                self.current_original = self.dual_uvc_frame_a
                self.preview_original.set_array(self.dual_uvc_frame_a)
            if self.dual_uvc_frame_b is not None:
                self.current_result = self.dual_uvc_frame_b
                self.preview_result.set_array(self.dual_uvc_frame_b)

        recording_worker = self.dual_uvc_recording_worker
        if (
            submit_recording
            and recording_worker is not None
            and recording_worker.isRunning()
            and self.dual_uvc_frame_a is not None
            and self.dual_uvc_frame_b is not None
        ):
            recording_worker.submit(
                self.dual_uvc_frame_a,
                self.dual_uvc_frame_b,
                time.perf_counter(),
            )
        self.update_dual_uvc_button_states()
        self.update_uvc_video_window()

    def submit_dual_uvc_recording_pair(self):
        recording_worker = self.dual_uvc_recording_worker
        if (
            recording_worker is not None
            and recording_worker.isRunning()
            and self.dual_uvc_frame_a is not None
            and self.dual_uvc_frame_b is not None
        ):
            recording_worker.submit(
                self.dual_uvc_frame_a,
                self.dual_uvc_frame_b,
                time.perf_counter(),
            )

    def handle_dual_uvc_worker_finished(self):
        worker = self.sender()
        if worker is self.dual_uvc_worker_a:
            self.dual_uvc_worker_a = None
            self.dual_uvc_status_a.setText(
                f"{self.dual_uvc_image_name_a.currentText()} (Camera A) stopped"
            )
        elif worker is self.dual_uvc_worker_b:
            self.dual_uvc_worker_b = None
            self.dual_uvc_status_b.setText(
                f"{self.dual_uvc_image_name_b.currentText()} (Camera B) stopped"
            )
        else:
            return
        self.stop_dual_uvc_recording(
            reason="A Dual UVC capture thread finished"
        )
        self.update_dual_uvc_button_states()
        if self.dual_uvc_close_pending and (
            self.dual_uvc_worker_a is None
            and self.dual_uvc_worker_b is None
        ):
            self.dual_uvc_display_timer.stop()
            self.dual_uvc_close_pending = False
            QTimer.singleShot(0, self.close)
        elif self.dual_uvc_restart_when_stopped and (
            self.dual_uvc_worker_a is None
            and self.dual_uvc_worker_b is None
        ):
            self.dual_uvc_display_timer.stop()
            self.dual_uvc_restart_when_stopped = False
            QTimer.singleShot(0, self.start_dual_uvc_cameras)

    def dual_uvc_labels(self):
        return self.dual_uvc_image_names()

    def save_dual_uvc_overview(self, path):
        if self.dual_uvc_frame_a is None or self.dual_uvc_frame_b is None:
            raise ValueError("Both Dual UVC frames are required for overview.")
        label_a, label_b = self.dual_uvc_labels()
        comparison_bgr = UvcRecordingThread._compose_frame(
            self.dual_uvc_frame_a,
            self.dual_uvc_frame_b,
            label_a,
            label_b,
        )
        save_rgb_png(
            path, cv2.cvtColor(comparison_bgr, cv2.COLOR_BGR2RGB)
        )

    def dual_uvc_camera_manifest(self):
        camera_info_a = self.selected_camera_info(self.dual_uvc_camera_a)
        camera_info_b = self.selected_camera_info(self.dual_uvc_camera_b)
        return {
            "camera_a": {
                "image_name": self.dual_uvc_image_name_a.currentText(),
                "image_name_file_tag": self.dual_uvc_image_name_a.currentData(),
                "camera_index": camera_info_a["index"],
                "camera_name": camera_info_a["name"],
                "camera_device_id": camera_info_a["device_id"],
                "capture_mode": DEFAULT_UVC_CAPTURE_MODE,
                "selected_capture_mode": getattr(
                    self.dual_uvc_worker_a, "selected_capture_mode", None
                ),
                "requested_fps": int(DEFAULT_UVC_FPS),
                "reported_fps": round(self.dual_uvc_actual_fps_a, 3),
                "resolution": {
                    "width": int(self.dual_uvc_frame_a.shape[1]),
                    "height": int(self.dual_uvc_frame_a.shape[0]),
                },
            },
            "camera_b": {
                "image_name": self.dual_uvc_image_name_b.currentText(),
                "image_name_file_tag": self.dual_uvc_image_name_b.currentData(),
                "camera_index": camera_info_b["index"],
                "camera_name": camera_info_b["name"],
                "camera_device_id": camera_info_b["device_id"],
                "capture_mode": DEFAULT_UVC_CAPTURE_MODE,
                "selected_capture_mode": getattr(
                    self.dual_uvc_worker_b, "selected_capture_mode", None
                ),
                "requested_fps": int(DEFAULT_UVC_FPS),
                "reported_fps": round(self.dual_uvc_actual_fps_b, 3),
                "resolution": {
                    "width": int(self.dual_uvc_frame_b.shape[1]),
                    "height": int(self.dual_uvc_frame_b.shape[0]),
                },
            },
            "software_pair_delta_ms": round(
                abs(
                    self.dual_uvc_frame_time_a
                    - self.dual_uvc_frame_time_b
                )
                * 1000.0,
                3,
            ),
        }

    def capture_dual_uvc_pair(self):
        if not self.dual_uvc_quality_enabled.isChecked():
            QMessageBox.information(
                self,
                "Quality Scoring Disabled",
                "Enable Calculate Quality Score before exporting a Dual UVC quality comparison.",
            )
            return
        if not self.dual_uvc_frames_ready():
            QMessageBox.information(
                self,
                "No Dual UVC Frames",
                "Start both cameras and wait for Camera A and Camera B frames.",
            )
            return

        self.update_dual_uvc_quality(force=True)
        quality_a = self.dual_uvc_quality_a
        quality_b = self.dual_uvc_quality_b
        out = self.make_output("dual_uvc_compare")
        file_tag_a, file_tag_b = self.dual_uvc_image_file_tags()
        camera_a_path = out / f"camera_{file_tag_a}.png"
        camera_b_path = out / f"camera_{file_tag_b}.png"
        overview_path = out / "00_dual_uvc_overview.png"
        base_a_path = out / f"base_{file_tag_a}.png"
        base_b_path = out / f"base_{file_tag_b}.png"
        detail_a_path = out / f"detail_{file_tag_a}.png"
        detail_b_path = out / f"detail_{file_tag_b}.png"
        quality_excel_path = out / "quality_analysis.xlsx"
        save_png(camera_a_path, self.dual_uvc_frame_a)
        save_png(camera_b_path, self.dual_uvc_frame_b)
        self.save_dual_uvc_overview(overview_path)
        save_png(base_a_path, quality_a["base_u8"])
        save_png(base_b_path, quality_b["base_u8"])
        save_png(detail_a_path, quality_a["detail_u8"])
        save_png(detail_b_path, quality_b["detail_u8"])
        manifest = {
            "format_version": 1,
            "mode": "dual_uvc_compare",
            "captured_at": datetime.now().astimezone().isoformat(
                timespec="seconds"
            ),
            "pairing": "latest-frame software pairing; not hardware synchronized",
            **self.dual_uvc_camera_manifest(),
            "camera_a_output": camera_a_path.name,
            "camera_b_output": camera_b_path.name,
            "overview_output": overview_path.name,
            "quality_a": quality_analysis_manifest(quality_a),
            "quality_b": quality_analysis_manifest(quality_b),
            "quality_outputs": {
                "base_a": base_a_path.name,
                "base_b": base_b_path.name,
                "detail_a": detail_a_path.name,
                "detail_b": detail_b_path.name,
                "excel": quality_excel_path.name,
            },
        }
        with open(out / "capture.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(
            out / "capture_summary.txt", "w", encoding="utf-8"
        ) as summary:
            summary.write(f"Captured: {manifest['captured_at']}\n")
            summary.write(
                f"{manifest['camera_a']['image_name']}: "
                f"{manifest['camera_a']['camera_name']}, "
                f"index {manifest['camera_a']['camera_index']}, "
                f"{manifest['camera_a']['resolution']['width']}x"
                f"{manifest['camera_a']['resolution']['height']}, "
                f"requested {manifest['camera_a']['capture_mode']}, "
                f"selected {manifest['camera_a']['selected_capture_mode']}\n"
            )
            summary.write(
                f"{manifest['camera_b']['image_name']}: "
                f"{manifest['camera_b']['camera_name']}, "
                f"index {manifest['camera_b']['camera_index']}, "
                f"{manifest['camera_b']['resolution']['width']}x"
                f"{manifest['camera_b']['resolution']['height']}, "
                f"requested {manifest['camera_b']['capture_mode']}, "
                f"selected {manifest['camera_b']['selected_capture_mode']}\n"
            )
            summary.write(
                f"Software pair delta: "
                f"{manifest['software_pair_delta_ms']:.3f} ms\n"
            )
            summary.write(
                "Quality scores are no-reference and intended for same-scene "
                "relative comparison.\n"
            )
            write_quality_summary(
                summary, manifest["camera_a"]["image_name"], quality_a
            )
            write_quality_summary(
                summary, manifest["camera_b"]["image_name"], quality_b
            )
            summary.write(f"Overview: {overview_path.name}\n")
            summary.write(f"Quality Excel: {quality_excel_path.name}\n")
        save_quality_excel(
            quality_excel_path,
            "Dual UVC Quality Compare",
            {
                key: value
                for key, value in manifest.items()
                if key not in {"quality_a", "quality_b"}
            },
            [
                {
                    "label": manifest["camera_a"]["image_name"],
                    "metrics": quality_a["metrics"],
                },
                {
                    "label": manifest["camera_b"]["image_name"],
                    "metrics": quality_b["metrics"],
                },
            ],
            temporal_timelines=self.dual_uvc_quality_chart.excel_timelines(),
            timeline_metric_definitions=(
                self.dual_uvc_quality_chart.excel_metric_definition()
            ),
        )
        name_a, name_b = self.dual_uvc_image_names()
        self.log.append(f"Saved Dual UVC {name_a}: {camera_a_path}")
        self.log.append(f"Saved Dual UVC {name_b}: {camera_b_path}")
        self.log.append(f"Saved Dual UVC overview: {overview_path}")
        self.log.append(f"Saved Dual UVC quality analysis: {quality_excel_path}")

    def start_dual_uvc_recording(self, _checked=False):
        if self.dual_uvc_recording_worker is not None:
            QMessageBox.information(
                self,
                "Recording Active",
                "Dual UVC recording is already active.",
            )
            return
        if not self.dual_uvc_frames_ready():
            QMessageBox.information(
                self,
                "No Dual UVC Frames",
                "Start both cameras and wait for Camera A and Camera B frames.",
            )
            return

        quality_enabled = self.dual_uvc_quality_enabled.isChecked()
        quality_start_a = None
        quality_start_b = None
        if quality_enabled:
            self.update_dual_uvc_quality(force=True)
            quality_start_a = self.dual_uvc_quality_a
            quality_start_b = self.dual_uvc_quality_b

        fps_a = (
            self.dual_uvc_actual_fps_a
            if 1.0 <= self.dual_uvc_actual_fps_a <= 240.0
            else float(DEFAULT_UVC_FPS)
        )
        fps_b = (
            self.dual_uvc_actual_fps_b
            if 1.0 <= self.dual_uvc_actual_fps_b <= 240.0
            else float(DEFAULT_UVC_FPS)
        )
        recording_fps = min(fps_a, fps_b)
        out = self.make_output("dual_uvc_recording")
        file_tag_a, file_tag_b = self.dual_uvc_image_file_tags()
        video_path = out / "dual_uvc_comparison.avi"
        overview_path = out / "00_dual_uvc_overview.png"
        base_start_a_path = out / f"base_start_{file_tag_a}.png"
        base_start_b_path = out / f"base_start_{file_tag_b}.png"
        detail_start_a_path = out / f"detail_start_{file_tag_a}.png"
        detail_start_b_path = out / f"detail_start_{file_tag_b}.png"
        self.save_dual_uvc_overview(overview_path)
        if quality_enabled:
            save_png(base_start_a_path, quality_start_a["base_u8"])
            save_png(base_start_b_path, quality_start_b["base_u8"])
            save_png(detail_start_a_path, quality_start_a["detail_u8"])
            save_png(detail_start_b_path, quality_start_b["detail_u8"])
        started_at = datetime.now().astimezone().isoformat(timespec="seconds")
        manifest = {
            "format_version": 1,
            "mode": "dual_uvc_recording",
            "started_at": started_at,
            "pairing": "latest-frame software pairing; not hardware synchronized",
            **self.dual_uvc_camera_manifest(),
            "recording_fps": round(recording_fps, 3),
            "recording_output": video_path.name,
            "overview_output": overview_path.name,
            "quality_enabled": quality_enabled,
        }
        if quality_enabled:
            manifest["quality_at_start"] = {
                "camera_a": quality_analysis_manifest(quality_start_a),
                "camera_b": quality_analysis_manifest(quality_start_b),
            }
            manifest["quality_layers_at_start"] = {
                "base_a": base_start_a_path.name,
                "base_b": base_start_b_path.name,
                "detail_a": detail_start_a_path.name,
                "detail_b": detail_start_b_path.name,
            }
        with open(out / "recording.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)

        label_a, label_b = self.dual_uvc_labels()
        worker = UvcRecordingThread(
            video_path,
            recording_fps,
            self,
            left_label=label_a,
            right_label=label_b,
        )
        worker.finished.connect(self.handle_dual_uvc_recording_finished)
        worker.finished.connect(worker.deleteLater)
        self.dual_uvc_recording_worker = worker
        self.dual_uvc_recording_output_dir = out
        self.dual_uvc_recording_manifest = manifest
        self.dual_uvc_recording_quality_start_a = (
            {"metrics": dict(quality_start_a["metrics"])}
            if quality_enabled
            else None
        )
        self.dual_uvc_recording_quality_start_b = (
            {"metrics": dict(quality_start_b["metrics"])}
            if quality_enabled
            else None
        )
        self.dual_uvc_recording_quality_timeline = []
        self.dual_uvc_recording_stop_reason = None
        self.dual_uvc_recording_started_perf = time.perf_counter()
        self.update_dual_uvc_button_states()
        worker.start()
        worker.submit(
            self.dual_uvc_frame_a,
            self.dual_uvc_frame_b,
            self.dual_uvc_recording_started_perf,
        )
        self.log.append(
            f"Dual UVC recording started: {video_path} "
            f"({recording_fps:.3f} FPS, {UVC_RECORDING_FOURCC}, "
            f"quality {'enabled' if quality_enabled else 'disabled'})"
        )

    def write_dual_uvc_recording_summary(self, worker, reason):
        out = self.dual_uvc_recording_output_dir
        manifest = self.dual_uvc_recording_manifest
        if out is None or manifest is None:
            return
        stats = worker.result_summary or {}
        quality_enabled = bool(manifest.get("quality_enabled", True))
        quality_stop_a = None
        quality_stop_b = None
        quality_start_a = None
        quality_start_b = None
        name_a = manifest["camera_a"]["image_name"]
        name_b = manifest["camera_b"]["image_name"]
        file_tag_a = manifest["camera_a"]["image_name_file_tag"]
        file_tag_b = manifest["camera_b"]["image_name_file_tag"]
        quality_excel_path = out / "quality_analysis.xlsx"
        if quality_enabled:
            self.update_dual_uvc_quality(force=True)
            quality_stop_a = self.dual_uvc_quality_a
            quality_stop_b = self.dual_uvc_quality_b
            quality_start_a = self.dual_uvc_recording_quality_start_a
            quality_start_b = self.dual_uvc_recording_quality_start_b
            if quality_stop_a is None or quality_stop_b is None:
                raise RuntimeError(
                    "Dual UVC quality frames are unavailable at stop."
                )
            if quality_start_a is None or quality_start_b is None:
                quality_start_a = quality_stop_a
                quality_start_b = quality_stop_b
            base_stop_a_path = out / f"base_stop_{file_tag_a}.png"
            base_stop_b_path = out / f"base_stop_{file_tag_b}.png"
            detail_stop_a_path = out / f"detail_stop_{file_tag_a}.png"
            detail_stop_b_path = out / f"detail_stop_{file_tag_b}.png"
            save_png(base_stop_a_path, quality_stop_a["base_u8"])
            save_png(base_stop_b_path, quality_stop_b["base_u8"])
            save_png(detail_stop_a_path, quality_stop_a["detail_u8"])
            save_png(detail_stop_b_path, quality_stop_b["detail_u8"])
        manifest["ended_at"] = datetime.now().astimezone().isoformat(
            timespec="seconds"
        )
        manifest["stop_reason"] = reason
        manifest["success"] = worker.error_message is None and bool(
            stats.get("written_frames", 0)
        )
        manifest["error"] = worker.error_message
        manifest["recording"] = stats
        if quality_enabled:
            manifest["quality_at_stop"] = {
                "camera_a": quality_analysis_manifest(quality_stop_a),
                "camera_b": quality_analysis_manifest(quality_stop_b),
            }
            manifest["quality_timeline"] = list(
                self.dual_uvc_recording_quality_timeline
            )
            manifest["quality_layers_at_stop"] = {
                "base_a": base_stop_a_path.name,
                "base_b": base_stop_b_path.name,
                "detail_a": detail_stop_a_path.name,
                "detail_b": detail_stop_b_path.name,
            }
            manifest["quality_excel_output"] = quality_excel_path.name
        with open(out / "recording.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(
            out / "recording_summary.txt", "w", encoding="utf-8"
        ) as summary:
            summary.write(f"Started: {manifest['started_at']}\n")
            summary.write(f"Ended: {manifest['ended_at']}\n")
            summary.write(f"Stop reason: {reason}\n")
            summary.write(f"Success: {manifest['success']}\n")
            if worker.error_message:
                summary.write(f"Error: {worker.error_message}\n")
            summary.write(
                f"Images: {name_a} "
                f"(index {manifest['camera_a']['camera_index']}) / "
                f"{name_b} "
                f"(index {manifest['camera_b']['camera_index']})\n"
            )
            summary.write(
                f"Recording FPS: {manifest['recording_fps']:.3f}\n"
            )
            summary.write(
                f"Duration: {stats.get('duration_seconds', 0.0):.3f} s\n"
            )
            summary.write(
                f"Submitted / written / duplicate frames: "
                f"{stats.get('submitted_frames', 0)} / "
                f"{stats.get('written_frames', 0)} / "
                f"{stats.get('duplicate_frames', 0)}\n"
            )
            summary.write(
                f"Queue / timestamp dropped frames: "
                f"{stats.get('queue_dropped_frames', 0)} / "
                f"{stats.get('timestamp_dropped_frames', 0)}\n"
            )
            summary.write(
                f"File size: {stats.get('file_size_bytes', 0)} bytes\n"
            )
            if quality_enabled:
                summary.write(
                    "Quality scores are no-reference and intended for "
                    "same-scene relative comparison.\n"
                )
                summary.write("Quality at start:\n")
                write_quality_summary(
                    summary,
                    manifest["camera_a"]["image_name"],
                    quality_start_a,
                )
                write_quality_summary(
                    summary,
                    manifest["camera_b"]["image_name"],
                    quality_start_b,
                )
                summary.write("Quality at stop:\n")
                write_quality_summary(
                    summary,
                    manifest["camera_a"]["image_name"],
                    quality_stop_a,
                )
                write_quality_summary(
                    summary,
                    manifest["camera_b"]["image_name"],
                    quality_stop_b,
                )
            else:
                summary.write("Quality scoring: Disabled\n")
            summary.write(f"Overview: {manifest['overview_output']}\n")
            if quality_enabled:
                summary.write(f"Quality Excel: {quality_excel_path.name}\n")
        if quality_enabled:
            save_quality_excel(
                quality_excel_path,
                "Dual UVC Recording Quality",
                {
                    key: value
                    for key, value in manifest.items()
                    if key
                    not in {
                        "quality_at_start",
                        "quality_at_stop",
                        "quality_timeline",
                    }
                },
                [
                    {
                        "label": f"{name_a} Start",
                        "metrics": quality_start_a["metrics"],
                    },
                    {
                        "label": f"{name_b} Start",
                        "metrics": quality_start_b["metrics"],
                    },
                    {
                        "label": f"{name_a} Stop",
                        "metrics": quality_stop_a["metrics"],
                    },
                    {
                        "label": f"{name_b} Stop",
                        "metrics": quality_stop_b["metrics"],
                    },
                ],
                temporal_timelines={
                    name_a: [
                        {
                            "elapsed_seconds": item["elapsed_seconds"],
                            "metrics": item["camera_a"],
                        }
                        for item in self.dual_uvc_recording_quality_timeline
                    ],
                    name_b: [
                        {
                            "elapsed_seconds": item["elapsed_seconds"],
                            "metrics": item["camera_b"],
                        }
                        for item in self.dual_uvc_recording_quality_timeline
                    ],
                },
            )

    def finalize_dual_uvc_recording_worker(self, worker, reason):
        if worker is not self.dual_uvc_recording_worker:
            return
        try:
            self.write_dual_uvc_recording_summary(worker, reason)
        except Exception as exc:
            self.log.append(f"Dual UVC recording summary error: {exc}")
        out = self.dual_uvc_recording_output_dir
        video_path = None if out is None else out / "dual_uvc_comparison.avi"
        if worker.error_message:
            self.log.append(f"Dual UVC recording error: {worker.error_message}")
        else:
            stats = worker.result_summary or {}
            self.log.append(
                f"Dual UVC recording stopped: {video_path}; "
                f"written={stats.get('written_frames', 0)}, "
                f"duplicates={stats.get('duplicate_frames', 0)}, "
                f"dropped={stats.get('queue_dropped_frames', 0) + stats.get('timestamp_dropped_frames', 0)}"
            )
        self.dual_uvc_recording_worker = None
        self.dual_uvc_recording_output_dir = None
        self.dual_uvc_recording_manifest = None
        self.dual_uvc_recording_quality_start_a = None
        self.dual_uvc_recording_quality_start_b = None
        self.dual_uvc_recording_quality_timeline = []
        self.dual_uvc_recording_stop_reason = None
        self.dual_uvc_recording_started_perf = 0.0
        self.update_dual_uvc_button_states()

    def stop_dual_uvc_recording(
        self,
        _checked=False,
        reason=None,
        wait_ms=5000,
        log_message=True,
    ):
        worker = self.dual_uvc_recording_worker
        if worker is None:
            return True
        stop_reason = reason or "User requested"
        self.dual_uvc_recording_stop_reason = stop_reason
        worker.stop()
        self.dual_uvc_record_start_btn.setEnabled(False)
        self.dual_uvc_record_stop_btn.setEnabled(False)
        if not worker.wait(wait_ms):
            if log_message:
                self.log.append(
                    "Dual UVC recording thread is still finishing the AVI file."
                )
            return False
        self.finalize_dual_uvc_recording_worker(worker, stop_reason)
        return True

    def handle_dual_uvc_recording_finished(self):
        worker = self.sender()
        if worker is not self.dual_uvc_recording_worker:
            return
        reason = self.dual_uvc_recording_stop_reason
        if reason is None:
            reason = (
                "Recording error"
                if worker.error_message
                else "Recording worker finished"
            )
        self.finalize_dual_uvc_recording_worker(worker, reason)

    def uvc_result_title(self):
        return self.uvc_preview_stage_combo.currentText()

    def mark_uvc_processing_config_changed(self, _value=None):
        if self.uvc_recording_worker is not None:
            self.stop_uvc_recording(
                reason="Pipeline configuration changed"
            )
        self.uvc_processing_revision += 1
        if hasattr(self, "uvc_temporal_tracker_original"):
            self.clear_uvc_algorithm_quality()
        if hasattr(self, "uvc_capture_btn"):
            self.uvc_capture_btn.setEnabled(False)
        worker = self.uvc_processing_worker
        if worker is not None and worker.isRunning():
            try:
                config = self.build_uvc_processing_config()
            except Exception as exc:
                worker.set_config(None)
                self.uvc_status.setText(f"Pipeline configuration error: {exc}")
            else:
                worker.set_config(config)

    def update_uvc_pipeline_ui(self, _value=None):
        self.rebuild_pipeline_preview_combo(
            self.uvc_pipeline_editor, self.uvc_preview_stage_combo
        )
        self.mark_uvc_processing_config_changed()
        self.update_uvc_result_title()
        self.refresh_uvc_processed_preview()

    def update_uvc_result_title(self):
        if self.current_function_index() == len(self.FUNCTIONS) - 1:
            title = self.uvc_result_title()
            self.preview_result.set_title(title)
            self.preview_result_caption.setText(title)

    def build_uvc_processing_config(self):
        stages = self.uvc_pipeline_editor.stages()
        if not any(stage.enabled for stage in stages):
            raise ValueError("Add and enable at least one UVC algorithm step.")
        return {
            "revision": self.uvc_processing_revision,
            "capture_mode": DEFAULT_UVC_CAPTURE_MODE,
            "raw_max": self.raw_max,
            "stages": stages,
        }

    def refresh_uvc_processed_preview(self):
        worker = self.uvc_processing_worker
        if (
            self.uvc_original_source is None
            or worker is None
            or not worker.isRunning()
        ):
            return
        try:
            config = self.build_uvc_processing_config()
        except Exception as exc:
            message = str(exc)
            if message != self.uvc_last_processing_error:
                self.log.append(f"UVC frame processing error: {message}")
            self.uvc_last_processing_error = message
            self.uvc_processed_frame = None
            self.current_result = None
            self.preview_result.set_array(None)
            self.uvc_capture_btn.setEnabled(False)
            self.uvc_status.setText(f"Frame processing error: {message}")
            return
        worker.submit(self.uvc_original_source, config)

    def update_uvc_stage_preview(self, _combo_index=None):
        if self.uvc_processed_frame is None:
            self.update_uvc_result_title()
            return
        selected_pipeline_index = self.uvc_preview_stage_combo.currentData()
        preview = pipeline_record_image(
            self.uvc_pipeline_records,
            selected_pipeline_index,
            self.uvc_processed_frame,
        )
        self.current_result = preview
        if self.current_function_index() == len(self.FUNCTIONS) - 1:
            self.preview_result.set_array(preview)
        self.update_uvc_result_title()

    def process_uvc_frame(self, original_gray):
        return process_uvc_pipeline_frame(
            original_gray,
            self.build_uvc_processing_config(),
            clahe_cache={},
            fieldscale_states={},
        )

    def clear_uvc_algorithm_quality(self, reset_timeline=True):
        self.uvc_quality_original = None
        self.uvc_quality_processed = None
        self.uvc_quality_last_update = 0.0
        self.uvc_temporal_tracker_original.reset()
        self.uvc_temporal_tracker_processed.reset()
        if hasattr(self, "uvc_quality_labels"):
            self.update_quality_score_labels(
                self.uvc_quality_labels, None, None
            )
        if reset_timeline and hasattr(self, "uvc_quality_chart"):
            self.uvc_quality_chart.clear()

    def update_uvc_algorithm_quality(self, force=False):
        if not self.uvc_quality_enabled.isChecked():
            return False
        if self.uvc_original_frame is None or self.uvc_processed_frame is None:
            return False
        now = time.perf_counter()
        if (
            not force
            and now - self.uvc_quality_last_update
            < DUAL_UVC_QUALITY_INTERVAL_SECONDS
        ):
            return False
        original_result = add_uvc_temporal_quality(
            thermal_uvc_quality_analysis(
                self.uvc_original_frame,
                max_analysis_pixels=UVC_QUALITY_ANALYSIS_MAX_PIXELS,
            ),
            self.uvc_temporal_tracker_original.result(),
        )
        processed_result = add_uvc_temporal_quality(
            thermal_uvc_quality_analysis(
                self.uvc_processed_frame,
                max_analysis_pixels=UVC_QUALITY_ANALYSIS_MAX_PIXELS,
            ),
            self.uvc_temporal_tracker_processed.result(),
        )
        self.uvc_quality_original = original_result
        self.uvc_quality_processed = processed_result
        self.uvc_quality_last_update = now
        self.update_quality_score_labels(
            self.uvc_quality_labels, original_result, processed_result
        )
        self.uvc_quality_chart.append_metrics(
            now,
            original=original_result["metrics"],
            final=processed_result["metrics"],
        )
        if self.uvc_recording_worker is not None:
            elapsed = max(0.0, now - self.uvc_recording_started_perf)
            self.uvc_recording_quality_timeline.append(
                {
                    "elapsed_seconds": round(elapsed, 3),
                    "original": dict(original_result["metrics"]),
                    "final": dict(processed_result["metrics"]),
                }
            )
        return True

    def build_uvc_overview_items(self):
        input_label = (
            "Native DN" if self.uvc_pipeline_input_max > 255 else "UVC Gray"
        )
        items = [
            ("00 Original UVC", self.uvc_original_frame, 255),
            (
                f"01 Pipeline Input - {input_label}",
                self.uvc_pipeline_input,
                self.uvc_pipeline_input_max,
            ),
        ]
        for overview_index, record in enumerate(
            self.uvc_pipeline_records, start=2
        ):
            items.append(
                (
                    f"{overview_index:02d} Step {record['pipeline_index']:02d} "
                    f"{record['name']}",
                    record["image"],
                    255,
                )
            )
        return items

    def start_uvc_recording(self, _checked=False):
        if not self.uvc_quality_enabled.isChecked():
            QMessageBox.information(
                self,
                "Quality Scoring Disabled",
                "Enable Calculate Quality Score before recording a UVC Algorithm quality report.",
            )
            return
        if self.uvc_recording_worker is not None:
            QMessageBox.information(
                self, "Recording Active", "UVC recording is already active."
            )
            return
        if (
            self.uvc_worker is None
            or self.uvc_original_frame is None
            or self.uvc_processed_frame is None
            or self.uvc_pipeline_input is None
        ):
            QMessageBox.information(
                self,
                "No UVC Frame",
                "Start the camera and wait for both previews first.",
            )
            return

        self.update_uvc_algorithm_quality(force=True)
        quality_start_original = self.uvc_quality_original
        quality_start_processed = self.uvc_quality_processed

        requested_fps = float(DEFAULT_UVC_FPS)
        recording_fps = (
            self.uvc_actual_fps
            if 1.0 <= self.uvc_actual_fps <= 240.0
            else requested_fps
        )
        started_at = datetime.now().astimezone().isoformat(timespec="seconds")
        out = self.make_output("uvc_recording")
        video_path = out / "comparison_pipeline.avi"
        overview_path = out / "00_pipeline_overview.png"
        base_start_original_path = out / "base_start_original.png"
        base_start_final_path = out / "base_start_final.png"
        detail_start_original_path = out / "detail_start_original.png"
        detail_start_final_path = out / "detail_start_final.png"
        manifest_path = out / "pipeline.json"

        try:
            save_pipeline_overview_png(
                overview_path, self.build_uvc_overview_items()
            )
            save_png(base_start_original_path, quality_start_original["base_u8"])
            save_png(base_start_final_path, quality_start_processed["base_u8"])
            save_png(
                detail_start_original_path, quality_start_original["detail_u8"]
            )
            save_png(
                detail_start_final_path, quality_start_processed["detail_u8"]
            )
            manifest = {
                "format_version": 1,
                "mode": "uvc_recording",
                "started_at": started_at,
                "capture_mode": DEFAULT_UVC_CAPTURE_MODE,
                "selected_capture_mode": getattr(
                    self.uvc_worker, "selected_capture_mode", None
                ),
                "camera_index": self.selected_camera_info(
                    self.uvc_camera_combo
                )["index"],
                "camera_name": self.selected_camera_info(
                    self.uvc_camera_combo
                )["name"],
                "camera_device_id": self.selected_camera_info(
                    self.uvc_camera_combo
                )["device_id"],
                "requested_fps": requested_fps,
                "reported_camera_fps": round(self.uvc_actual_fps, 3),
                "recording_fps": round(recording_fps, 3),
                "source_resolution": {
                    "width": int(self.uvc_original_frame.shape[1]),
                    "height": int(self.uvc_original_frame.shape[0]),
                },
                "recording_layout": "Original UVC | Final Pipeline Output",
                "recording_output": video_path.name,
                "overview_output": overview_path.name,
                "bit_depth": self.bit_depth,
                "pipeline": self.uvc_pipeline_editor.configuration(),
                "execution_at_start": pipeline_execution_json(
                    self.uvc_pipeline_records
                ),
                "quality_at_start": {
                    "original": quality_analysis_manifest(
                        quality_start_original
                    ),
                    "final": quality_analysis_manifest(
                        quality_start_processed
                    ),
                },
                "quality_layers_at_start": {
                    "base_original": base_start_original_path.name,
                    "base_final": base_start_final_path.name,
                    "detail_original": detail_start_original_path.name,
                    "detail_final": detail_start_final_path.name,
                },
            }
            with open(manifest_path, "w", encoding="utf-8") as destination:
                json.dump(manifest, destination, ensure_ascii=False, indent=2)
        except Exception as exc:
            QMessageBox.critical(self, "Recording Error", str(exc))
            self.log.append(f"UVC recording setup error: {exc}")
            return

        worker = UvcRecordingThread(video_path, recording_fps, self)
        worker.finished.connect(self.handle_uvc_recording_finished)
        worker.finished.connect(worker.deleteLater)
        self.uvc_recording_worker = worker
        self.uvc_recording_output_dir = out
        self.uvc_recording_manifest = manifest
        self.uvc_recording_quality_start_original = {
            "metrics": dict(quality_start_original["metrics"]),
        }
        self.uvc_recording_quality_start_processed = {
            "metrics": dict(quality_start_processed["metrics"]),
        }
        self.uvc_recording_quality_timeline = []
        self.uvc_recording_stop_reason = None
        self.uvc_recording_started_perf = time.perf_counter()
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn.setEnabled(True)
        worker.start()
        worker.submit(
            self.uvc_original_frame,
            self.uvc_processed_frame,
            self.uvc_recording_started_perf,
        )
        self.log.append(
            f"UVC recording started: {video_path} "
            f"({recording_fps:.3f} FPS, {UVC_RECORDING_FOURCC})"
        )
        self.log.append(f"Saved recording overview: {overview_path}")

    def write_uvc_recording_summary(self, worker, reason):
        out = self.uvc_recording_output_dir
        manifest = self.uvc_recording_manifest
        if out is None or manifest is None:
            return
        stats = worker.result_summary or {}
        error_message = worker.error_message
        self.update_uvc_algorithm_quality(force=True)
        quality_stop_original = self.uvc_quality_original
        quality_stop_processed = self.uvc_quality_processed
        quality_start_original = self.uvc_recording_quality_start_original
        quality_start_processed = self.uvc_recording_quality_start_processed
        if quality_stop_original is None or quality_stop_processed is None:
            raise RuntimeError("UVC algorithm quality frames are unavailable at stop.")
        if quality_start_original is None or quality_start_processed is None:
            quality_start_original = quality_stop_original
            quality_start_processed = quality_stop_processed
        base_stop_original_path = out / "base_stop_original.png"
        base_stop_final_path = out / "base_stop_final.png"
        detail_stop_original_path = out / "detail_stop_original.png"
        detail_stop_final_path = out / "detail_stop_final.png"
        quality_excel_path = out / "quality_analysis.xlsx"
        save_png(base_stop_original_path, quality_stop_original["base_u8"])
        save_png(base_stop_final_path, quality_stop_processed["base_u8"])
        save_png(detail_stop_original_path, quality_stop_original["detail_u8"])
        save_png(detail_stop_final_path, quality_stop_processed["detail_u8"])
        manifest["ended_at"] = datetime.now().astimezone().isoformat(
            timespec="seconds"
        )
        manifest["stop_reason"] = reason
        manifest["success"] = error_message is None and bool(
            stats.get("written_frames", 0)
        )
        manifest["error"] = error_message
        manifest["recording"] = stats
        manifest["quality_at_stop"] = {
            "original": quality_analysis_manifest(quality_stop_original),
            "final": quality_analysis_manifest(quality_stop_processed),
        }
        manifest["quality_timeline"] = list(
            self.uvc_recording_quality_timeline
        )
        manifest["quality_layers_at_stop"] = {
            "base_original": base_stop_original_path.name,
            "base_final": base_stop_final_path.name,
            "detail_original": detail_stop_original_path.name,
            "detail_final": detail_stop_final_path.name,
        }
        manifest["quality_excel_output"] = quality_excel_path.name
        recorded_pipeline_text = " -> ".join(
            stage["name"]
            for stage in manifest["pipeline"]
            if stage.get("enabled")
        )
        with open(out / "pipeline.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)

        with open(
            out / "recording_summary.txt", "w", encoding="utf-8"
        ) as summary:
            summary.write(f"Started: {manifest['started_at']}\n")
            summary.write(f"Ended: {manifest['ended_at']}\n")
            summary.write(f"Stop reason: {reason}\n")
            summary.write(f"Success: {manifest['success']}\n")
            if error_message:
                summary.write(f"Error: {error_message}\n")
            summary.write(
                f"Camera: {manifest['camera_name']} "
                f"(index {manifest['camera_index']})\n"
            )
            summary.write(
                f"Camera device ID: {manifest['camera_device_id']}\n"
            )
            summary.write(f"Capture mode: {manifest['capture_mode']}\n")
            summary.write(
                f"Selected capture mode: {manifest['selected_capture_mode']}\n"
            )
            summary.write(
                f"Source resolution: "
                f"{manifest['source_resolution']['width']}x"
                f"{manifest['source_resolution']['height']}\n"
            )
            summary.write(
                f"Requested / reported / recording FPS: "
                f"{manifest['requested_fps']:.3f} / "
                f"{manifest['reported_camera_fps']:.3f} / "
                f"{manifest['recording_fps']:.3f}\n"
            )
            summary.write(f"Codec: {stats.get('codec', '--')}\n")
            summary.write(
                f"Duration: {stats.get('duration_seconds', 0.0):.3f} s\n"
            )
            summary.write(
                f"Submitted / written / duplicate frames: "
                f"{stats.get('submitted_frames', 0)} / "
                f"{stats.get('written_frames', 0)} / "
                f"{stats.get('duplicate_frames', 0)}\n"
            )
            summary.write(
                f"Queue / timestamp dropped frames: "
                f"{stats.get('queue_dropped_frames', 0)} / "
                f"{stats.get('timestamp_dropped_frames', 0)}\n"
            )
            summary.write(
                f"File size: {stats.get('file_size_bytes', 0)} bytes\n"
            )
            summary.write(f"Overview: {manifest['overview_output']}\n")
            summary.write(f"Pipeline: {recorded_pipeline_text}\n")
            summary.write(
                "Quality scores are no-reference and intended for same-scene "
                "relative comparison.\n"
            )
            summary.write("Quality at start:\n")
            write_quality_summary(
                summary, "Original UVC", quality_start_original
            )
            write_quality_summary(
                summary, "Final Pipeline", quality_start_processed
            )
            summary.write("Quality at stop:\n")
            write_quality_summary(
                summary, "Original UVC", quality_stop_original
            )
            write_quality_summary(
                summary, "Final Pipeline", quality_stop_processed
            )
            summary.write(f"Quality Excel: {quality_excel_path.name}\n")
        save_quality_excel(
            quality_excel_path,
            "UVC Algorithm Recording Quality",
            {
                key: value
                for key, value in manifest.items()
                if key
                not in {
                    "quality_at_start",
                    "quality_at_stop",
                    "quality_timeline",
                }
            },
            [
                {
                    "label": "Original Start",
                    "metrics": quality_start_original["metrics"],
                },
                {
                    "label": "Final Start",
                    "metrics": quality_start_processed["metrics"],
                },
                {
                    "label": "Original Stop",
                    "metrics": quality_stop_original["metrics"],
                },
                {
                    "label": "Final Stop",
                    "metrics": quality_stop_processed["metrics"],
                },
            ],
            temporal_timelines={
                "Original": [
                    {
                        "elapsed_seconds": item["elapsed_seconds"],
                        "metrics": item["original"],
                    }
                    for item in self.uvc_recording_quality_timeline
                ],
                "Final": [
                    {
                        "elapsed_seconds": item["elapsed_seconds"],
                        "metrics": item["final"],
                    }
                    for item in self.uvc_recording_quality_timeline
                ],
            },
        )

    def finalize_uvc_recording_worker(self, worker, reason):
        if worker is not self.uvc_recording_worker:
            return
        try:
            self.write_uvc_recording_summary(worker, reason)
        except Exception as exc:
            self.log.append(f"UVC recording summary error: {exc}")

        out = self.uvc_recording_output_dir
        video_path = None if out is None else out / "comparison_pipeline.avi"
        if worker.error_message:
            self.log.append(f"UVC recording error: {worker.error_message}")
            self.uvc_status.setText(
                f"Recording error: {worker.error_message}"
            )
        else:
            stats = worker.result_summary or {}
            self.log.append(
                f"UVC recording stopped: {video_path}; "
                f"written={stats.get('written_frames', 0)}, "
                f"duplicates={stats.get('duplicate_frames', 0)}, "
                f"dropped={stats.get('queue_dropped_frames', 0) + stats.get('timestamp_dropped_frames', 0)}"
            )
            self.uvc_status.setText(f"Recording saved: {video_path}")

        self.uvc_recording_worker = None
        self.uvc_recording_output_dir = None
        self.uvc_recording_manifest = None
        self.uvc_recording_quality_start_original = None
        self.uvc_recording_quality_start_processed = None
        self.uvc_recording_quality_timeline = []
        self.uvc_recording_stop_reason = None
        self.uvc_recording_started_perf = 0.0
        frame_ready = (
            self.uvc_worker is not None
            and self.uvc_original_frame is not None
            and self.uvc_processed_frame is not None
        )
        self.uvc_record_start_btn.setEnabled(frame_ready)
        self.uvc_record_stop_btn.setEnabled(False)

    def stop_uvc_recording(
        self,
        _checked=False,
        reason=None,
        wait_ms=5000,
        log_message=True,
    ):
        worker = self.uvc_recording_worker
        if worker is None:
            return True
        stop_reason = reason or "User requested"
        self.uvc_recording_stop_reason = stop_reason
        worker.stop()
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn.setEnabled(False)
        if not worker.wait(wait_ms):
            self.uvc_status.setText("Finishing UVC recording...")
            if log_message:
                self.log.append(
                    "UVC recording thread is still finishing the AVI file."
                )
            return False
        self.finalize_uvc_recording_worker(worker, stop_reason)
        return True

    def handle_uvc_recording_finished(self):
        worker = self.sender()
        if worker is not self.uvc_recording_worker:
            return
        reason = self.uvc_recording_stop_reason
        if reason is None:
            reason = (
                "Recording error"
                if worker.error_message
                else "Recording worker finished"
            )
        self.finalize_uvc_recording_worker(worker, reason)

    def start_uvc_camera(self, _checked=False):
        if cv2 is None:
            message = "OpenCV not installed. pip install opencv-python"
            self.log.append(f"\nERROR: {message}")
            QMessageBox.critical(self, "Camera Error", message)
            return
        try:
            processing_config = self.build_uvc_processing_config()
        except Exception as exc:
            message = str(exc)
            self.log.append(f"\nERROR: {message}")
            QMessageBox.critical(self, "UVC Pipeline Error", message)
            return

        if not self.stop_uvc_camera(log_message=False):
            self.uvc_restart_when_stopped = True
            self.log.append(
                "Previous capture thread is still stopping; camera will start "
                "automatically afterward."
            )
            return
        self.uvc_restart_when_stopped = False
        capture_mode = DEFAULT_UVC_CAPTURE_MODE
        camera_info = self.selected_camera_info(self.uvc_camera_combo)
        camera_index = camera_info["index"]
        if camera_index is None:
            QMessageBox.warning(
                self,
                "Camera Disconnected",
                "Select a connected camera before starting UVC.",
            )
            return
        requested_fps = DEFAULT_UVC_FPS

        processing_worker = UvcProcessingThread(self)
        processing_worker.set_config(processing_config)
        processing_worker.frame_processed.connect(self.handle_uvc_processed_frame)
        processing_worker.frame_failed.connect(self.handle_uvc_processing_failed)
        processing_worker.finished.connect(processing_worker.deleteLater)
        worker = UvcCaptureThread(
            camera_index,
            self.width,
            self.height,
            requested_fps,
            UVC_RECONNECT_SECONDS,
            self,
            frame_consumer=processing_worker.submit_frame,
            capture_mode=capture_mode,
            camera_device_id=camera_info["device_id"],
            camera_name=camera_info["name"],
        )
        worker.frame_seen.connect(self.handle_uvc_frame_seen)
        worker.camera_opened.connect(self.handle_uvc_camera_opened)
        worker.status_changed.connect(self.handle_uvc_worker_status)
        worker.stream_interrupted.connect(
            self.handle_uvc_stream_interrupted
        )
        worker.finished.connect(self.handle_uvc_worker_finished)
        worker.finished.connect(worker.deleteLater)
        self.uvc_worker = worker
        self.uvc_processing_worker = processing_worker
        self.uvc_original_source = None
        self.uvc_original_frame = None
        self.uvc_processed_frame = None
        self.clear_uvc_algorithm_quality()
        self.uvc_pipeline_input = None
        self.uvc_pipeline_input_max = 255
        self.uvc_pipeline_records = []
        self.uvc_pipeline_total_ms = 0.0
        self.uvc_last_processing_error = None
        self.uvc_frame_count = 0
        self.uvc_fps_display = 0.0
        now = time.perf_counter()
        self.uvc_last_time = now
        self.uvc_last_frame_time = now
        self.uvc_waiting_for_frame = False
        self.uvc_reconnect_requested = False
        self.uvc_actual_fps = 0.0
        self.preview_original.set_array(None)
        self.preview_result.set_array(None)
        self.uvc_watchdog_timer.start()
        self.uvc_start_btn.setEnabled(False)
        self.uvc_stop_btn.setEnabled(True)
        self.uvc_capture_btn.setEnabled(False)
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn.setEnabled(False)
        self.uvc_status.setText(
            f"Opening {camera_info['name']} (Camera {camera_index}, "
            f"{capture_mode})..."
        )
        self.log.append(
            f"Starting {camera_info['name']} (Camera {camera_index}) in "
            f"{capture_mode} mode "
            "using the background capture thread..."
        )
        processing_order = self.uvc_pipeline_editor.summary()
        self.log.append(f"UVC processing order: {processing_order}")
        self.log.append(
            "UVC processor uses latest-frame-wins; pending old frames are replaced."
        )
        processing_worker.start()
        worker.start()

    def stop_uvc_processing_worker(self, wait_ms=5000):
        worker = self.uvc_processing_worker
        if worker is None:
            return True
        worker.stop()
        if worker.isRunning() and not worker.wait(wait_ms):
            return False
        self.uvc_processing_worker = None
        return True

    def stop_uvc_camera(self, _checked=False, log_message=True):
        self.uvc_restart_when_stopped = False
        self.close_uvc_video_window()
        self.stop_uvc_recording(
            reason="Camera stopped", log_message=log_message
        )
        worker = self.uvc_worker
        was_running = (
            worker is not None or self.uvc_processing_worker is not None
        )
        self.uvc_watchdog_timer.stop()
        if worker is not None:
            worker.stop()
            if not worker.wait(2000):
                self.uvc_start_btn.setEnabled(False)
                self.uvc_stop_btn.setEnabled(False)
                self.uvc_capture_btn.setEnabled(False)
                self.uvc_record_start_btn.setEnabled(False)
                self.uvc_record_stop_btn.setEnabled(False)
                self.uvc_status.setText("Camera capture thread is stopping...")
                return False
            self.uvc_worker = None
        if not self.stop_uvc_processing_worker():
            self.uvc_start_btn.setEnabled(False)
            self.uvc_stop_btn.setEnabled(False)
            self.uvc_capture_btn.setEnabled(False)
            self.uvc_record_start_btn.setEnabled(False)
            self.uvc_record_stop_btn.setEnabled(False)
            self.uvc_status.setText("UVC processing thread is stopping...")
            return False
        self.uvc_start_btn.setEnabled(True)
        self.uvc_stop_btn.setEnabled(False)
        self.uvc_capture_btn.setEnabled(False)
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn.setEnabled(False)
        self.uvc_actual_fps = 0.0
        self.clear_uvc_algorithm_quality()
        self.uvc_status.setText("Camera stopped")
        if was_running and log_message:
            self.log.append("Camera stopped")
        return True

    def switch_uvc_camera(self, _combo_index):
        if self.uvc_worker is not None:
            self.start_uvc_camera()

    def handle_uvc_camera_opened(
        self, camera_index, actual_width, actual_height, actual_fps
    ):
        if self.sender() is not self.uvc_worker:
            return
        self.clear_uvc_algorithm_quality()
        self.uvc_actual_fps = float(actual_fps)
        self.uvc_last_frame_time = time.perf_counter()
        self.uvc_reconnect_requested = False
        camera_name = getattr(
            self.uvc_worker, "camera_name", f"Camera {camera_index}"
        )
        self.uvc_status.setText(
            f"{camera_name} (Camera {camera_index}) opened: "
            f"{actual_width}x{actual_height}, "
            f"reported {actual_fps:.1f} FPS; waiting for frame..."
        )
        self.log.append(
            f"{camera_name} (Camera {camera_index}) opened: "
            f"{actual_width}x{actual_height}, "
            f"requested {DEFAULT_UVC_FPS} FPS, "
            f"reported {actual_fps:.1f} FPS"
        )

    def handle_uvc_worker_status(self, message):
        if self.sender() is not self.uvc_worker:
            return
        self.uvc_status.setText(message)

    def handle_uvc_stream_interrupted(self, message):
        if self.sender() is not self.uvc_worker:
            return
        self.log.append(
            f"UVC Algorithm Compare automatically stopped: {message}"
        )
        if self.stop_uvc_camera(log_message=False):
            self.uvc_status.setText("Camera stopped — stream interrupted")

    def handle_uvc_worker_finished(self):
        if self.sender() is not self.uvc_worker:
            return
        self.stop_uvc_recording(reason="Camera capture thread finished")
        restart = self.uvc_restart_when_stopped
        close_pending = self.uvc_close_pending
        self.uvc_restart_when_stopped = False
        self.uvc_worker = None
        self.stop_uvc_processing_worker()
        self.uvc_watchdog_timer.stop()
        self.uvc_start_btn.setEnabled(True)
        self.uvc_stop_btn.setEnabled(False)
        self.uvc_capture_btn.setEnabled(False)
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn.setEnabled(False)
        self.uvc_actual_fps = 0.0
        self.uvc_status.setText("Camera stopped")
        if close_pending:
            self.uvc_close_pending = False
            QTimer.singleShot(0, self.close)
        elif (
            restart
            and self.current_function_index() == len(self.FUNCTIONS) - 1
        ):
            QTimer.singleShot(0, self.start_uvc_camera)

    def force_close_after_uvc_timeout(self):
        if not self.uvc_close_pending:
            return
        worker = self.uvc_worker
        if worker is not None and worker.isRunning():
            worker.terminate()
            if not worker.wait(500):
                QTimer.singleShot(500, self.force_close_after_uvc_timeout)
                return
        self.uvc_worker = None
        self.stop_uvc_processing_worker(wait_ms=500)
        self.uvc_close_pending = False
        QTimer.singleShot(0, self.close)

    def force_close_after_dual_uvc_timeout(self):
        if not self.dual_uvc_close_pending:
            return
        self.dual_uvc_display_timer.stop()
        still_running = False
        for worker in (self.dual_uvc_worker_a, self.dual_uvc_worker_b):
            if worker is not None and worker.isRunning():
                worker.terminate()
                if not worker.wait(500):
                    still_running = True
        if still_running:
            QTimer.singleShot(500, self.force_close_after_dual_uvc_timeout)
            return
        self.dual_uvc_worker_a = None
        self.dual_uvc_worker_b = None
        self.stop_dual_uvc_recording(
            reason="Application closing", wait_ms=500
        )
        self.dual_uvc_close_pending = False
        QTimer.singleShot(0, self.close)

    def check_uvc_stream(self):
        worker = self.uvc_worker
        if worker is None or not worker.isRunning():
            return
        elapsed = time.perf_counter() - self.uvc_last_frame_time
        if elapsed < UVC_FRAME_WAIT_SECONDS:
            return
        if not self.uvc_waiting_for_frame:
            self.uvc_waiting_for_frame = True
            self.log.append(
                "UVC frame paused; keeping the last image (possible FFC)."
            )
        self.uvc_status.setText(
            f"Waiting for UVC frame {elapsed:.1f}s (possible FFC); "
            "last image retained."
        )
        if (
            elapsed >= UVC_RECONNECT_SECONDS
            and not self.uvc_reconnect_requested
        ):
            self.uvc_reconnect_requested = True
            camera_info = self.selected_camera_info(self.uvc_camera_combo)
            self.log.append(
                f"No UVC frame for {elapsed:.1f}s; requesting automatic "
                f"reconnect for {camera_info['name']} "
                f"(Camera {camera_info['index']})."
            )
            worker.force_reconnect()

    def handle_uvc_frame_seen(self, frame_time):
        if self.sender() is not self.uvc_worker:
            return
        self.uvc_last_frame_time = frame_time
        self.uvc_reconnect_requested = False
        if self.uvc_waiting_for_frame:
            self.uvc_waiting_for_frame = False
            self.log.append("UVC stream resumed.")

    def handle_uvc_frame(self, frame):
        if self.sender() is not self.uvc_worker:
            return
        processing_worker = self.uvc_processing_worker
        if processing_worker is None or not processing_worker.isRunning():
            return
        if self.app_settings["camera_mirror"]["camera_a"]:
            frame = np.ascontiguousarray(frame[:, ::-1])
        try:
            config = self.build_uvc_processing_config()
        except Exception as exc:
            self.handle_uvc_processing_error(str(exc))
            return
        processing_worker.submit(frame, config)

    def handle_uvc_processed_frame(self, payload):
        if self.sender() is not self.uvc_processing_worker:
            return
        if payload["revision"] != self.uvc_processing_revision:
            return
        self.uvc_last_processing_error = None
        self.uvc_original_source = payload["original_source"]
        self.uvc_original_frame = payload["original_display"]
        self.uvc_pipeline_input = payload["pipeline_input"]
        self.uvc_pipeline_input_max = payload["pipeline_input_max"]
        self.uvc_processed_frame = payload["final_image"]
        self.uvc_pipeline_records = payload["records"]
        self.uvc_pipeline_total_ms = payload["total_ms"]
        now = time.perf_counter()
        if self.uvc_quality_enabled.isChecked():
            self.uvc_temporal_tracker_original.add_frame(
                self.uvc_original_frame, now
            )
            self.uvc_temporal_tracker_processed.add_frame(
                self.uvc_processed_frame, now
            )
        self.current_original = self.uvc_original_frame
        self.preview_original.set_array(self.uvc_original_frame)
        self.update_uvc_stage_preview()
        self.update_uvc_algorithm_quality()
        self.update_uvc_video_window()
        self.uvc_capture_btn.setEnabled(True)

        recording_worker = self.uvc_recording_worker
        if recording_worker is not None and recording_worker.isRunning():
            recording_worker.submit(
                self.uvc_original_frame,
                self.uvc_processed_frame,
                now,
            )
            self.uvc_record_start_btn.setEnabled(False)
            self.uvc_record_stop_btn.setEnabled(True)
        elif recording_worker is None:
            self.uvc_record_start_btn.setEnabled(True)
            self.uvc_record_stop_btn.setEnabled(False)
        else:
            self.uvc_record_start_btn.setEnabled(False)
            self.uvc_record_stop_btn.setEnabled(False)
        self.uvc_frame_count += 1
        elapsed = now - self.uvc_last_time
        if elapsed >= 0.5:
            self.uvc_fps_display = self.uvc_frame_count / elapsed
            self.uvc_frame_count = 0
            self.uvc_last_time = now
        requested_fps = DEFAULT_UVC_FPS
        budget_ms = 1000.0 / requested_fps
        budget_text = (
            "OVER BUDGET"
            if self.uvc_pipeline_total_ms > budget_ms
            else "within budget"
        )
        detail_parts = list(payload["preprocessing_details"])
        detail_parts.extend(
            f"S{record['pipeline_index']} {record['name']} "
            f"{record['elapsed_ms']:.1f}ms"
            for record in self.uvc_pipeline_records
        )
        detail_text = " | ".join(detail_parts)
        recording_text = ""
        if recording_worker is not None and recording_worker.isRunning():
            recording_text = (
                f"REC {max(0.0, now - self.uvc_recording_started_perf):.1f}s | "
            )
        selected_capture_mode = getattr(
            self.uvc_worker, "selected_capture_mode", None
        )
        selected_capture_name = UvcCaptureThread._capture_mode_name(
            selected_capture_mode or DEFAULT_UVC_CAPTURE_MODE
        )
        camera_info = self.selected_camera_info(self.uvc_camera_combo)
        self.uvc_status.setText(
            recording_text
            + f"{camera_info['name']} (Camera {camera_info['index']}) | "
            f"Capture {selected_capture_name} | "
            f"{self.uvc_original_frame.shape[1]}x"
            f"{self.uvc_original_frame.shape[0]} | "
            f"Output FPS {self.uvc_fps_display:.1f} | "
            f"Pipeline {self.uvc_pipeline_total_ms:.1f} ms / "
            f"{budget_ms:.1f} ms ({budget_text})"
            + (f" | {detail_text}" if detail_text else "")
        )

    def handle_uvc_processing_error(self, message):
        if self.uvc_recording_worker is not None:
            self.stop_uvc_recording(reason="UVC processing error")
        if message != self.uvc_last_processing_error:
            self.log.append(f"UVC frame processing error: {message}")
        self.uvc_last_processing_error = message
        self.uvc_processed_frame = None
        self.uvc_pipeline_records = []
        self.current_result = None
        self.preview_result.set_array(None)
        self.uvc_capture_btn.setEnabled(False)
        self.uvc_record_start_btn.setEnabled(False)
        self.uvc_record_stop_btn.setEnabled(False)
        self.uvc_status.setText(f"Frame processing error: {message}")

    def handle_uvc_processing_failed(self, payload):
        if self.sender() is not self.uvc_processing_worker:
            return
        if payload.get("revision") != self.uvc_processing_revision:
            return
        self.handle_uvc_processing_error(payload["message"])

    def capture_uvc_pair(self):
        if not self.uvc_quality_enabled.isChecked():
            QMessageBox.information(
                self,
                "Quality Scoring Disabled",
                "Enable Calculate Quality Score before exporting a UVC Algorithm quality comparison.",
            )
            return
        if self.uvc_original_frame is None or self.uvc_processed_frame is None:
            QMessageBox.information(
                self,
                "No UVC Frame",
                "Start the camera and wait for a frame first.",
            )
            return
        self.update_uvc_algorithm_quality(force=True)
        quality_original = self.uvc_quality_original
        quality_processed = self.uvc_quality_processed
        out = self.make_output("uvc_algorithm_compare")
        original_path = out / "original.png"
        pipeline_input_path = out / "pipeline_input.png"
        processed_path = out / "final_pipeline.png"
        comparison_path = out / "comparison_pipeline.png"
        overview_path = out / "00_pipeline_overview.png"
        base_original_path = out / "base_original.png"
        base_processed_path = out / "base_final.png"
        detail_original_path = out / "detail_original.png"
        detail_processed_path = out / "detail_final.png"
        quality_excel_path = out / "quality_analysis.xlsx"
        save_png(original_path, self.uvc_original_frame)
        save_png(pipeline_input_path, self.uvc_pipeline_input)
        step_output_files = {}
        for record in self.uvc_pipeline_records:
            step_path = out / (
                f"step_{record['pipeline_index']:02d}_"
                f"{record['algorithm']}.png"
            )
            save_png(step_path, record["image"])
            step_output_files[record["pipeline_index"]] = step_path.name
        save_png(processed_path, self.uvc_processed_frame)
        save_uvc_comparison_png(
            comparison_path,
            self.uvc_original_frame,
            self.uvc_processed_frame,
            "Final Pipeline Output",
        )
        input_label = (
            "Native DN" if self.uvc_pipeline_input_max > 255 else "UVC Gray"
        )
        overview_items = [
            ("00 Original UVC", self.uvc_original_frame, 255),
            (
                f"01 Pipeline Input - {input_label}",
                self.uvc_pipeline_input,
                self.uvc_pipeline_input_max,
            ),
        ]
        for overview_index, record in enumerate(
            self.uvc_pipeline_records, start=2
        ):
            overview_items.append(
                (
                    f"{overview_index:02d} Step {record['pipeline_index']:02d} "
                    f"{record['name']}",
                    record["image"],
                    255,
                )
            )
        save_pipeline_overview_png(overview_path, overview_items)
        save_png(base_original_path, quality_original["base_u8"])
        save_png(base_processed_path, quality_processed["base_u8"])
        save_png(detail_original_path, quality_original["detail_u8"])
        save_png(detail_processed_path, quality_processed["detail_u8"])
        manifest = {
            "format_version": 1,
            "mode": "uvc_algorithm_compare",
            "capture_mode": DEFAULT_UVC_CAPTURE_MODE,
            "selected_capture_mode": getattr(
                self.uvc_worker, "selected_capture_mode", None
            ),
            "camera_index": self.selected_camera_info(
                self.uvc_camera_combo
            )["index"],
            "camera_name": self.selected_camera_info(
                self.uvc_camera_combo
            )["name"],
            "camera_device_id": self.selected_camera_info(
                self.uvc_camera_combo
            )["device_id"],
            "requested_fps": DEFAULT_UVC_FPS,
            "resolution": {
                "width": int(self.uvc_original_frame.shape[1]),
                "height": int(self.uvc_original_frame.shape[0]),
            },
            "selected_thermal_resolution": {
                "width": self.width,
                "height": self.height,
            },
            "bit_depth": self.bit_depth,
            "pipeline": self.uvc_pipeline_editor.configuration(),
            "execution": pipeline_execution_json(self.uvc_pipeline_records),
            "total_ms": round(self.uvc_pipeline_total_ms, 3),
            "original_output": original_path.name,
            "pipeline_input_output": pipeline_input_path.name,
            "final_output": processed_path.name,
            "comparison_output": comparison_path.name,
            "overview_output": overview_path.name,
            "quality_original": quality_analysis_manifest(
                quality_original
            ),
            "quality_final": quality_analysis_manifest(quality_processed),
            "quality_layer_outputs": {
                "base_original": base_original_path.name,
                "base_final": base_processed_path.name,
                "detail_original": detail_original_path.name,
                "detail_final": detail_processed_path.name,
            },
            "quality_excel_output": quality_excel_path.name,
        }
        for execution in manifest["execution"]:
            output_file = step_output_files.get(execution["pipeline_index"])
            if output_file:
                execution["output_file"] = output_file
        with open(out / "pipeline.json", "w", encoding="utf-8") as destination:
            json.dump(manifest, destination, ensure_ascii=False, indent=2)
        with open(out / "pipeline_summary.txt", "w", encoding="utf-8") as summary:
            summary.write(
                f"Camera: {manifest['camera_name']} "
                f"(index {manifest['camera_index']})\n"
            )
            summary.write(
                f"Camera device ID: {manifest['camera_device_id']}\n"
            )
            summary.write(f"Capture mode: {manifest['capture_mode']}\n")
            summary.write(
                f"Selected capture mode: {manifest['selected_capture_mode']}\n"
            )
            summary.write(
                f"Actual resolution: {manifest['resolution']['width']}x"
                f"{manifest['resolution']['height']}\n"
            )
            summary.write(f"Overview: {overview_path.name}\n")
            summary.write(f"Pipeline: {self.uvc_pipeline_editor.summary()}\n")
            summary.write(
                "Quality scores are no-reference and intended for same-scene "
                "relative comparison.\n"
            )
            write_quality_summary(
                summary, "Original UVC", quality_original
            )
            write_quality_summary(
                summary, "Final Pipeline", quality_processed
            )
            summary.write(f"Quality Excel: {quality_excel_path.name}\n")
            for execution in manifest["execution"]:
                summary.write(
                    f"Step {execution['pipeline_index']}: {execution['name']} - "
                    f"{execution['elapsed_ms']:.3f} ms - "
                    f"{' | '.join(execution['details'])}\n"
                )
            summary.write(f"Total: {self.uvc_pipeline_total_ms:.3f} ms\n")
        save_quality_excel(
            quality_excel_path,
            "UVC Algorithm Quality Compare",
            {
                key: value
                for key, value in manifest.items()
                if key not in {"quality_original", "quality_final"}
            },
            [
                {
                    "label": "Original UVC",
                    "metrics": quality_original["metrics"],
                },
                {
                    "label": "Final Pipeline",
                    "metrics": quality_processed["metrics"],
                },
            ],
            temporal_timelines=self.uvc_quality_chart.excel_timelines(),
            timeline_metric_definitions=(
                self.uvc_quality_chart.excel_metric_definition()
            ),
        )
        self.log.append(f"Saved: {original_path}")
        self.log.append(f"Saved: {processed_path}")
        self.log.append(f"Saved comparison: {comparison_path}")
        self.log.append(f"Saved overview: {overview_path}")
        self.log.append(f"Saved quality analysis: {quality_excel_path}")
        self.log.append(f"Pipeline settings: {out / 'pipeline.json'}")

    def save_current_preview(self):
        if self.current_result is None:
            QMessageBox.information(
                self, "No Preview", "Run a function first."
            )
            return
        function_index = self.current_function_index()
        is_uvc_algorithm = (
            self.FUNCTIONS[function_index] == "UVC Algorithm Compare"
        )
        if is_uvc_algorithm:
            if self.uvc_original_frame is None:
                QMessageBox.information(
                    self,
                    "No UVC Frame",
                    "Start the camera and wait for both previews first.",
                )
                return
            out = self.make_output("manual_preview")
            original = self.uvc_original_frame.copy()
            processed = self.current_result.copy()
            algorithm_title = self.uvc_result_title()
            original_path = out / "original_uvc.png"
            processed_path = out / "processed_selected_pipeline.png"
            comparison_path = out / "comparison_selected_pipeline.png"
            save_png(original_path, original)
            save_png(processed_path, processed)
            save_uvc_comparison_png(
                comparison_path, original, processed, algorithm_title
            )
            self.log.append(f"Saved UVC preview: {algorithm_title}")
            self.log.append(f"Saved original: {original_path}")
            self.log.append(f"Saved processed: {processed_path}")
            self.log.append(f"Saved comparison: {comparison_path}")
            return

        out = self.make_output("manual_preview")
        path = out / "preview.png"
        if (self.FUNCTIONS[function_index] in ("Dual Image Compare", "Image Algorithm Compare")
                and self.image_comparison_mode().currentIndex() == 1):
            self.save_dual_swipe_preview(path)
            self.log.append(f"Saved swipe comparison preview: {path}")
            return
        is_bpr_marking = (
            self.FUNCTIONS[function_index] == "BPR Bad Pixel Marking"
        )
        if is_bpr_marking:
            preview_to_save = np.clip(
                self.current_result, 0, 255
            ).astype(np.uint8)
        elif self.FUNCTIONS[function_index] == "Single Image Statistics":
            preview_to_save = raw_to_display_u8(self.current_result, self.raw_viewer_info["value_max"])
        elif function_index == 4:
            preview_to_save = signed_delta_to_display_u8(
                self.current_result, self.raw_max
            )
        elif function_index in (0, 1, 2, 3, 5):
            preview_to_save = raw_to_display_u8(
                self.current_result, self.raw_max
            )
        else:
            preview_to_save = np.clip(
                self.current_result, 0, 255
            ).astype(np.uint8)
        if is_bpr_marking:
            save_rgb_png(path, preview_to_save)
        else:
            save_png(path, preview_to_save)
        self.log.append(f"Saved preview: {path}")

    def open_last_result(self):
        if not self.last_output_dir or not os.path.isdir(self.last_output_dir):
            QMessageBox.information(self,"No Result","No result folder yet."); return
        path=self.last_output_dir
        if sys.platform.startswith("win"): os.startfile(path)
        elif sys.platform=="darwin": os.system(f'open "{path}"')
        else: os.system(f'xdg-open "{path}" >/dev/null 2>&1 &')

    def closeEvent(self, event):
        self.uvc_restart_when_stopped = False
        self.dual_uvc_restart_when_stopped = False
        self.viewer_uvc_restart_when_stopped = False
        self.close_uvc_video_window()
        self.close_uvc_histogram_window()
        self.close_raw_histogram_window()
        self.camera_refresh_timer.stop()
        self.uvc_watchdog_timer.stop()
        self.viewer_uvc_display_timer.stop()
        self.dual_uvc_display_timer.stop()
        if not self.stop_uvc_viewer(
            log_message=False, wait_ms=3000
        ):
            QTimer.singleShot(500, self.close)
            event.ignore()
            return
        if not self.stop_dual_uvc_recording(
            reason="Application closing", wait_ms=3000
        ):
            QTimer.singleShot(500, self.close)
            event.ignore()
            return
        dual_workers = (
            self.dual_uvc_worker_a,
            self.dual_uvc_worker_b,
        )
        if any(
            worker is not None and worker.isRunning()
            for worker in dual_workers
        ):
            if not self.dual_uvc_close_pending:
                self.dual_uvc_close_pending = True
                self.dual_uvc_start_btn.setEnabled(False)
                self.dual_uvc_stop_btn.setEnabled(False)
                self.dual_uvc_capture_btn.setEnabled(False)
                self.dual_uvc_record_start_btn.setEnabled(False)
                self.dual_uvc_record_stop_btn.setEnabled(False)
                self.dual_uvc_status_a.setText(
                    f"Stopping {self.dual_uvc_image_name_a.currentText()} before exit..."
                )
                self.dual_uvc_status_b.setText(
                    f"Stopping {self.dual_uvc_image_name_b.currentText()} before exit..."
                )
                for dual_worker in dual_workers:
                    if dual_worker is not None:
                        dual_worker.stop()
                QTimer.singleShot(
                    UVC_CLOSE_TIMEOUT_MS,
                    self.force_close_after_dual_uvc_timeout,
                )
            event.ignore()
            return
        self.dual_uvc_worker_a = None
        self.dual_uvc_worker_b = None
        if not self.stop_uvc_recording(
            reason="Application closing", wait_ms=3000
        ):
            self.uvc_status.setText("Finishing recording before exit...")
            QTimer.singleShot(500, self.close)
            event.ignore()
            return
        worker = self.uvc_worker
        if worker is not None and worker.isRunning():
            if not self.uvc_close_pending:
                self.uvc_close_pending = True
                self.uvc_start_btn.setEnabled(False)
                self.uvc_stop_btn.setEnabled(False)
                self.uvc_capture_btn.setEnabled(False)
                self.uvc_record_start_btn.setEnabled(False)
                self.uvc_record_stop_btn.setEnabled(False)
                self.uvc_status.setText("Stopping camera before exit...")
                worker.stop()
                QTimer.singleShot(
                    UVC_CLOSE_TIMEOUT_MS,
                    self.force_close_after_uvc_timeout,
                )
            event.ignore()
            return
        self.uvc_worker = None
        self.stop_uvc_processing_worker()
        self.uvc_close_pending = False
        super().closeEvent(event)


def main():
    app=QApplication(sys.argv)
    app.setStyle("Fusion")
    w=ThermalRawToolkit(); w.showMaximized()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
