"""reports / quality."""

from pathlib import Path
import numpy as np
from thermal_toolkit.core.constants import ACTIVE_SPATIAL_QUALITY_WEIGHTS, ACTIVE_UVC_TEMPORAL_WEIGHTS
from thermal_toolkit.reports.histogram import _flatten_excel_fields


def write_quality_summary(summary, label, result):
    metrics = result["metrics"] if "metrics" in result else result
    if metrics.get("quality_profile") == "Fieldscale IQA":
        from thermal_toolkit.analysis.quality.fieldscale_iqa import FIELDS
        summary.write(f"{label} [Fieldscale IQA; no total score]:\n")
        for name, key in FIELDS:
            value = metrics.get(key)
            value_text = "N/A" if value is None else f"{value:.8f}"
            reason = result.get("unavailable_reasons", {}).get(key, "")
            summary.write(f"  {name}: {value_text}" + (f" ({reason})" if reason else "") + "\n")
        return

    def score_text(key):
        value = metrics.get(key)
        return "Collecting..." if value is None else f"{float(value):.3f}"

    profile = metrics.get("quality_profile", "Quality")
    summary.write(f"{label} quality [{profile}] (0-100, higher is better):\n")
    summary.write(
        f"  Total Score / detail clarity / contrast: "
        f"{score_text('overall_score')} / "
        f"{score_text('detail_clarity_score')} / "
        f"{score_text('contrast_quality_score')}\n"
    )
    summary.write(
        f"  Flat cleanliness / thermal uniformity / clipping control / "
        f"artifact control: {score_text('flat_cleanliness_score')} / "
        f"{score_text('thermal_uniformity_score')} / "
        f"{score_text('clipping_control_score')} / "
        f"{score_text('artifact_control_score')}\n"
    )
    if metrics.get("quality_confidence_score") is not None:
        summary.write(
            f"  Quality confidence / flat groups / gray bands: "
            f"{score_text('quality_confidence_score')} / "
            f"{int(metrics.get('flat_region_group_count', 0))} / "
            f"{int(metrics.get('flat_gray_band_count', 0))}\n"
        )
    if "flicker_control_score" in metrics:
        summary.write(
            f"  Flicker control / temporal stability: "
            f"{score_text('flicker_control_score')} / "
            f"{score_text('temporal_stability_score')}\n"
        )
        stable_text = score_text("stable_region_percent")
        summary.write(
            f"  Temporal window / frames / stable region: "
            f"{float(metrics.get('temporal_window_seconds', 0.0)):.3f} s / "
            f"{int(metrics.get('temporal_frame_count', 0))} / "
            f"{stable_text}%\n"
        )
    summary.write(
        f"  Flat noise / base residual noise: "
        f"{float(metrics.get('flat_noise_dn', 0.0)):.3f} / "
        f"{float(metrics.get('base_residual_noise_dn', 0.0)):.3f} DN\n"
    )
    if metrics.get("flat_noise_worst_dn") is not None:
        summary.write(
            f"  Flat noise typical / worst: "
            f"{float(metrics.get('flat_noise_typical_dn', 0.0)):.3f} / "
            f"{float(metrics.get('flat_noise_worst_dn', 0.0)):.3f} DN\n"
        )
        summary.write(
            f"  Uniformity typical / worst / row-column banding: "
            f"{float(metrics.get('uniformity_typical_dn', 0.0)):.3f} / "
            f"{float(metrics.get('uniformity_worst_dn', 0.0)):.3f} / "
            f"{float(metrics.get('row_column_banding_dn', 0.0)):.3f} DN\n"
        )
    if metrics.get("impulse_outlier_control_score") is not None:
        summary.write(
            f"  Robust noise / impulse control / impulse pixels: "
            f"{score_text('robust_noise_score')} / "
            f"{score_text('impulse_outlier_control_score')} / "
            f"{float(metrics.get('impulse_outlier_percent', 0.0)):.4f}%\n"
        )
    summary.write(
        f"  Detail SNR / halo ratio / seam periodicity: "
        f"{float(metrics.get('detail_snr', 0.0)):.3f} / "
        f"{float(metrics.get('halo_ratio', 0.0)):.3f} / "
        f"{float(metrics.get('seam_periodicity_score', 0.0)):.3f}\n"
    )


def save_quality_excel(
    path,
    report_title,
    metadata,
    analyses,
    temporal_timelines=None,
    timeline_metric_definitions=None,
):
    if analyses and all(item["metrics"].get("quality_profile") == "Fieldscale IQA" for item in analyses):
        from thermal_toolkit.reports.image_quality import save_image_quality_excel
        return save_image_quality_excel(path, report_title, metadata, analyses)
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel output requires openpyxl. Install dependencies with: "
            "pip install -r requirements_thermal_raw_toolkit.txt"
        ) from exc

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append([report_title])
    summary_sheet.merge_cells("A1:C1")
    summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor="2F75B5")
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    summary_sheet.append(["Section", "Field", "Value"])
    for field_name, value in _flatten_excel_fields(metadata):
        summary_sheet.append(["Report", field_name, value])
    for metric_name, value in ACTIVE_SPATIAL_QUALITY_WEIGHTS.items():
        summary_sheet.append(
            ["Spatial Weight", metric_name, f"{value}%"]
        )
    for metric_name, value in ACTIVE_UVC_TEMPORAL_WEIGHTS.items():
        summary_sheet.append(
            ["UVC Temporal Weight", metric_name, f"{value}%"]
        )
    score_metric_names = [
        "overall_score",
        "detail_clarity_score",
        "contrast_quality_score",
        "flat_cleanliness_score",
        "thermal_uniformity_score",
        "clipping_control_score",
        "artifact_control_score",
    ]
    if analyses and all(
        "quality_confidence_score" in analysis["metrics"] for analysis in analyses
    ):
        score_metric_names.append("quality_confidence_score")
    if analyses and all(
        "flicker_control_score" in analysis["metrics"] for analysis in analyses
    ):
        score_metric_names.extend(
            ["flicker_control_score", "temporal_stability_score"]
        )

    def score_display_name(metric_name):
        if metric_name == "overall_score":
            return "Total Score"
        return metric_name

    def excel_value(value):
        if value is None:
            return None
        if isinstance(value, (bool, np.bool_)):
            return bool(value)
        if isinstance(value, (int, float, np.integer, np.floating)):
            return float(value)
        return str(value)

    for analysis in analyses:
        for metric_name in score_metric_names:
            summary_sheet.append(
                [
                    analysis["label"],
                    score_display_name(metric_name),
                    excel_value(analysis["metrics"][metric_name]),
                ]
            )
    for cell in summary_sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    summary_sheet.freeze_panes = "A3"
    summary_sheet.auto_filter.ref = f"A2:C{summary_sheet.max_row}"
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 36
    summary_sheet.column_dimensions["C"].width = 48

    metrics_sheet = workbook.create_sheet("Score Comparison")
    metrics_sheet.append(["Score"] + [analysis["label"] for analysis in analyses])
    for metric_name in score_metric_names:
        metrics_sheet.append(
            [score_display_name(metric_name)]
            + [excel_value(analysis["metrics"][metric_name]) for analysis in analyses]
        )
    diagnostic_metric_names = [
        metric_name
        for metric_name in analyses[0]["metrics"]
        if metric_name not in score_metric_names
    ]
    diagnostics_sheet = workbook.create_sheet("Measurements")
    diagnostics_sheet.append(
        ["Measurement"] + [analysis["label"] for analysis in analyses]
    )
    for metric_name in diagnostic_metric_names:
        diagnostics_sheet.append(
            [metric_name]
            + [excel_value(analysis["metrics"][metric_name]) for analysis in analyses]
        )
    for worksheet in (metrics_sheet, diagnostics_sheet):
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            worksheet.column_dimensions[column_cells[0].column_letter].width = 20

    temporal_sheet = None
    timeline_chart_columns = []
    if temporal_timelines:
        temporal_sheet = workbook.create_sheet("Temporal Timeline")
        timeline_metrics = timeline_metric_definitions or [
            ("Total Score", "overall_score"),
            ("Spatial", "spatial_quality_score"),
            ("Flicker Control", "flicker_control_score"),
            ("Temporal Stability", "temporal_stability_score"),
            ("Flicker DN", "flicker_dn"),
            ("Temporal Noise DN", "temporal_noise_dn"),
            ("Stable Region %", "stable_region_percent"),
        ]
        headers = ["Elapsed Seconds"]
        for stream_label in temporal_timelines:
            for display_name, _metric_name in timeline_metrics:
                headers.append(f"{stream_label} {display_name}")
        temporal_sheet.append(headers)
        max_samples = max(len(items) for items in temporal_timelines.values())
        first_timeline = next(iter(temporal_timelines.values()))
        for sample_index in range(max_samples):
            elapsed = (
                first_timeline[sample_index]["elapsed_seconds"]
                if sample_index < len(first_timeline)
                else None
            )
            row = [elapsed]
            for items in temporal_timelines.values():
                metrics = (
                    items[sample_index]["metrics"]
                    if sample_index < len(items)
                    else {}
                )
                row.extend(
                    excel_value(metrics.get(metric_name))
                    for _display_name, metric_name in timeline_metrics
                )
            temporal_sheet.append(row)
        for cell in temporal_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        temporal_sheet.freeze_panes = "A2"
        temporal_sheet.auto_filter.ref = temporal_sheet.dimensions
        temporal_sheet.column_dimensions["A"].width = 18
        for column_index in range(2, temporal_sheet.max_column + 1):
            temporal_sheet.column_dimensions[
                temporal_sheet.cell(1, column_index).column_letter
            ].width = 24
        metrics_per_stream = len(timeline_metrics)
        for stream_index in range(len(temporal_timelines)):
            first_column = 2 + stream_index * metrics_per_stream
            timeline_chart_columns.extend(
                first_column + metric_index
                for metric_index, (_display_name, metric_name) in enumerate(
                    timeline_metrics
                )
                if metric_name.endswith("_score")
            )
        temporal_summary_sheet = workbook.create_sheet("Temporal Summary")
        temporal_summary_sheet.append(
            ["Stream", "Score", "Average", "P10", "Worst"]
        )
        for stream_label, items in temporal_timelines.items():
            for display_name, metric_name in timeline_metrics:
                values = np.asarray(
                    [
                        item["metrics"].get(metric_name)
                        for item in items
                        if item["metrics"].get(metric_name) is not None
                    ],
                    dtype=np.float64,
                )
                temporal_summary_sheet.append(
                    [
                        stream_label,
                        display_name,
                        None if values.size == 0 else float(np.mean(values)),
                        None
                        if values.size == 0
                        else float(np.percentile(values, 10.0)),
                        None if values.size == 0 else float(np.min(values)),
                    ]
                )
        for cell in temporal_summary_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        temporal_summary_sheet.freeze_panes = "A2"
        temporal_summary_sheet.auto_filter.ref = temporal_summary_sheet.dimensions
        for column_cells in temporal_summary_sheet.columns:
            temporal_summary_sheet.column_dimensions[
                column_cells[0].column_letter
            ].width = 22

    charts_sheet = workbook.create_sheet("Charts")
    score_chart = BarChart()
    score_chart.type = "col"
    score_chart.style = 10
    score_chart.title = "Quality Score Comparison"
    score_chart.y_axis.title = "Score (0-100)"
    score_chart.height = 10
    score_chart.width = 24
    score_chart.add_data(
        Reference(
            metrics_sheet,
            min_col=2,
            max_col=1 + len(analyses),
            min_row=1,
            max_row=1 + len(score_metric_names),
        ),
        titles_from_data=True,
    )
    score_chart.set_categories(
        Reference(
            metrics_sheet,
            min_col=1,
            min_row=2,
            max_row=1 + len(score_metric_names),
        )
    )
    charts_sheet.add_chart(score_chart, "A1")
    if temporal_sheet is not None and temporal_sheet.max_row >= 2:
        temporal_chart = LineChart()
        temporal_chart.style = 13
        temporal_chart.title = "UVC Temporal Quality Timeline"
        temporal_chart.y_axis.title = "Score (0-100)"
        temporal_chart.x_axis.title = "Elapsed Seconds"
        temporal_chart.height = 11
        temporal_chart.width = 24
        for column_index in timeline_chart_columns:
            temporal_chart.add_data(
                Reference(
                    temporal_sheet,
                    min_col=column_index,
                    min_row=1,
                    max_row=temporal_sheet.max_row,
                ),
                titles_from_data=True,
            )
        temporal_chart.set_categories(
            Reference(
                temporal_sheet,
                min_col=1,
                min_row=2,
                max_row=temporal_sheet.max_row,
            )
        )
        charts_sheet.add_chart(temporal_chart, "A22")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
