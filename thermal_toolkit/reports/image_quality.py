"""Fieldscale-style metrics, without legacy weighted total scores."""

from pathlib import Path
from thermal_toolkit.analysis.quality.fieldscale_iqa import FIELDS, HELP
from thermal_toolkit.reports.histogram import _flatten_excel_fields


def save_image_quality_excel(path, report_title, metadata, analyses):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([report_title])
    summary.append(["Field", "Value"])
    for key, value in _flatten_excel_fields(metadata):
        summary.append([key, value])
    summary.append(["Profile", "Fieldscale IQA; no aggregate score"])
    summary.append(["Reference", "Paired metrics require the corresponding original; N/A = unavailable"])
    comparison = workbook.create_sheet("Metric Comparison")
    comparison.append(["Metric"] + [analysis["label"] for analysis in analyses])
    for name, key in FIELDS:
        comparison.append([name] + [analysis["metrics"].get(key) for analysis in analyses])
    definitions = workbook.create_sheet("Definitions")
    definitions.append(["Metric", "Definition", "Unavailable"])
    for name, key in FIELDS:
        definitions.append([name, HELP[key], "Blank cell means N/A, not zero"])
    for analysis in analyses:
        for key, reason in analysis.get("unavailable_reasons", {}).items():
            definitions.append([analysis["label"], key, reason])
    charts = workbook.create_sheet("Charts")
    for index, (name, _) in enumerate(FIELDS):
        chart = BarChart()
        chart.title = name
        chart.y_axis.title = "Metric value"
        chart.add_data(Reference(comparison, min_col=2, max_col=len(analyses) + 1,
                                 min_row=index + 2, max_row=index + 2), from_rows=True)
        chart.set_categories(Reference(comparison, min_col=2, max_col=len(analyses) + 1,
                                        min_row=1, max_row=1))
        chart.legend = None
        chart.width, chart.height = 20, 9
        charts.add_chart(chart, f"A{1 + index * 19}")
    for sheet in (summary, comparison, definitions):
        sheet.freeze_panes = "B2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 35
        for column in ("B", "C", "D"):
            sheet.column_dimensions[column].width = 35
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
