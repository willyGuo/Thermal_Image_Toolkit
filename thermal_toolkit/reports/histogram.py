"""reports / histogram."""

from pathlib import Path
import json
import numpy as np
import re
from thermal_toolkit.core.opencv import cv2


def render_gray_histogram_rgb(
    histogram,
    stats,
    width=800,
    height=360,
    value_max=255,
    title="Gray Histogram (256 bins)",
    show_cdf=True,
):
    canvas = np.full((height, width, 3), (23, 25, 29), dtype=np.uint8)
    left, right, top, bottom = 58, width - 22, 48, height - 48
    plot_width = right - left
    plot_height = bottom - top

    cv2.rectangle(canvas, (left, top), (right, bottom), (67, 73, 82), 1)
    for fraction in (0.25, 0.5, 0.75):
        y = int(round(bottom - fraction * plot_height))
        cv2.line(canvas, (left, y), (right, y), (45, 49, 56), 1)

    peak_count = max(1, int(np.max(histogram)))
    for bin_index, count in enumerate(histogram):
        x0 = left + int(bin_index * plot_width / 256.0)
        x1 = left + int((bin_index + 1) * plot_width / 256.0)
        y = bottom - int(round(float(count) * plot_height / peak_count))
        cv2.rectangle(
            canvas,
            (x0, max(top, y)),
            (max(x0, x1), bottom),
            (184, 126, 67),
            -1,
        )

    cumulative = np.cumsum(histogram).astype(np.float64)
    if show_cdf and cumulative[-1] > 0:
        cumulative /= cumulative[-1]
        cdf_points = []
        for bin_index, value in enumerate(cumulative):
            x = left + int(round(bin_index * plot_width / 255.0))
            y = bottom - int(round(value * plot_height))
            cdf_points.append((x, y))
        cv2.polylines(
            canvas,
            [np.asarray(cdf_points, dtype=np.int32)],
            False,
            (72, 205, 235),
            2,
            cv2.LINE_AA,
        )

    marker_specs = (
        (stats["p1"], "P1", (105, 190, 105)),
        (stats["median"], "P50", (225, 225, 225)),
        (stats["p99"], "P99", (90, 145, 245)),
    )
    for marker_index, (value, label, color) in enumerate(marker_specs):
        x = max(
            left,
            min(
                left + int(round(float(value) * plot_width / value_max)),
                right,
            ),
        )
        cv2.line(canvas, (x, top), (x, bottom), color, 1, cv2.LINE_AA)
        cv2.putText(
            canvas,
            f"{label} {value:.1f}",
            (
                max(left, min(x - 28, right - 72)),
                top + 16 + marker_index * 16,
            ),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.42,
            color,
            1,
            cv2.LINE_AA,
        )

    tick_values = [
        int(round(value_max * fraction))
        for fraction in (0.0, 0.25, 0.5, 0.75, 1.0)
    ]
    for value in tick_values:
        x = left + int(round(value * plot_width / value_max))
        cv2.line(canvas, (x, bottom), (x, bottom + 5), (150, 157, 166), 1)
        cv2.putText(
            canvas,
            str(value),
            (x - max(7, len(str(value)) * 4), bottom + 22),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.43,
            (190, 196, 204),
            1,
            cv2.LINE_AA,
        )

    cv2.putText(
        canvas,
        title,
        (left, 24),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.62,
        (235, 238, 242),
        1,
        cv2.LINE_AA,
    )
    cv2.putText(
        canvas,
        f"Peak count {peak_count:,}",
        (left, 42),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.42,
        (170, 177, 186),
        1,
        cv2.LINE_AA,
    )
    if show_cdf:
        cv2.line(canvas, (right - 145, 22), (right - 115, 22), (72, 205, 235), 2)
        cv2.putText(
            canvas,
            "CDF",
            (right - 107, 27),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.45,
            (205, 211, 218),
            1,
            cv2.LINE_AA,
        )
    return cv2.cvtColor(canvas, cv2.COLOR_BGR2RGB)


def _excel_cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Path):
        return str(value)
    return json.dumps(value, ensure_ascii=False)


def _flatten_excel_fields(values, prefix=""):
    for key, value in values.items():
        field_name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            yield from _flatten_excel_fields(value, field_name)
        else:
            yield field_name, _excel_cell_value(value)


def save_histogram_excel(
    path,
    report_title,
    metadata,
    analyses,
    value_max=255,
    value_name="Gray",
):
    """Save statistics, Histogram/CDF data, and native Excel charts."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel output requires openpyxl. Install dependencies with: "
            "pip install -r requirements_thermal_raw_toolkit.txt"
        ) from exc

    normalized_analyses = []
    used_sheet_names = set()
    for analysis_index, analysis in enumerate(analyses, start=1):
        label = str(analysis["label"])
        histogram = np.asarray(analysis["histogram"], dtype=np.int64)
        if histogram.ndim != 1 or histogram.size == 0:
            raise ValueError(
                f"Excel Histogram must be a non-empty 1D array, got {histogram.shape}."
            )
        base_name = re.sub(r"[\\/*?:\[\]]", "_", f"Histogram {label}")[:31]
        sheet_name = base_name or f"Histogram {analysis_index}"
        suffix = 2
        while sheet_name in used_sheet_names:
            suffix_text = f" {suffix}"
            sheet_name = f"{base_name[:31 - len(suffix_text)]}{suffix_text}"
            suffix += 1
        used_sheet_names.add(sheet_name)
        normalized_analyses.append(
            {
                "label": label,
                "stats": dict(analysis["stats"]),
                "histogram": histogram,
                "sheet_name": sheet_name,
            }
        )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append([report_title])
    summary_sheet.merge_cells("A1:C1")
    summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor="2F75B5")
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    summary_sheet.append(["Section", "Field", "Value"])
    for cell in summary_sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    for field_name, value in _flatten_excel_fields(metadata):
        summary_sheet.append(["Report", field_name, value])
    for analysis in normalized_analyses:
        for field_name, value in _flatten_excel_fields(analysis["stats"]):
            summary_sheet.append([analysis["label"], field_name, value])
    summary_sheet.freeze_panes = "A3"
    summary_sheet.auto_filter.ref = f"A2:C{summary_sheet.max_row}"
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 32
    summary_sheet.column_dimensions["C"].width = 48

    histogram_sheets = []
    for analysis in normalized_analyses:
        worksheet = workbook.create_sheet(analysis["sheet_name"])
        histogram_sheets.append((analysis, worksheet))
        worksheet.append(
            [
                "Bin",
                f"{value_name} Start",
                f"{value_name} End",
                f"{value_name} Center",
                f"{analysis['label']} Count",
                f"{analysis['label']} Percentage",
                f"{analysis['label']} Cumulative Count",
                f"{analysis['label']} CDF Percentage",
            ]
        )
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
        histogram = analysis["histogram"]
        bin_width = (int(value_max) + 1) / histogram.size
        total_count = int(histogram.sum())
        cumulative_count = 0
        for bin_index, count_value in enumerate(histogram):
            count = int(count_value)
            cumulative_count += count
            range_start = int(np.ceil(bin_index * bin_width))
            range_end = int(np.ceil((bin_index + 1) * bin_width) - 1)
            range_start = max(0, min(range_start, int(value_max)))
            range_end = max(range_start, min(range_end, int(value_max)))
            center = (range_start + range_end) / 2.0
            percentage = count / total_count if total_count else 0.0
            cdf_percentage = (
                cumulative_count / total_count if total_count else 0.0
            )
            worksheet.append(
                [
                    bin_index,
                    range_start,
                    range_end,
                    center,
                    count,
                    percentage,
                    cumulative_count,
                    cdf_percentage,
                ]
            )
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.column_dimensions["A"].width = 10
        worksheet.column_dimensions["B"].width = 16
        worksheet.column_dimensions["C"].width = 16
        worksheet.column_dimensions["D"].width = 17
        worksheet.column_dimensions["E"].width = 16
        worksheet.column_dimensions["F"].width = 16
        worksheet.column_dimensions["G"].width = 20
        worksheet.column_dimensions["H"].width = 18
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, 6).number_format = "0.000000%"
            worksheet.cell(row_index, 8).number_format = "0.000000%"

    charts_sheet = workbook.create_sheet("Charts")
    histogram_chart = BarChart()
    histogram_chart.type = "col"
    histogram_chart.style = 10
    histogram_chart.title = "Histogram Comparison"
    histogram_chart.y_axis.title = "Pixel Count"
    histogram_chart.x_axis.title = value_name
    histogram_chart.gapWidth = 0
    histogram_chart.height = 10
    histogram_chart.width = 24
    cdf_chart = LineChart()
    cdf_chart.style = 13
    cdf_chart.title = "CDF Comparison"
    cdf_chart.y_axis.title = "Cumulative Percentage"
    cdf_chart.x_axis.title = value_name
    cdf_chart.height = 10
    cdf_chart.width = 24
    for _analysis, worksheet in histogram_sheets:
        histogram_last_row = len(_analysis["histogram"]) + 1
        categories = Reference(
            worksheet, min_col=4, min_row=2, max_row=histogram_last_row
        )
        histogram_chart.add_data(
            Reference(worksheet, min_col=5, min_row=1, max_row=histogram_last_row),
            titles_from_data=True,
        )
        histogram_chart.set_categories(categories)
        cdf_chart.add_data(
            Reference(worksheet, min_col=8, min_row=1, max_row=histogram_last_row),
            titles_from_data=True,
        )
        cdf_chart.set_categories(categories)
    cdf_chart.y_axis.numFmt = "0%"
    charts_sheet.add_chart(histogram_chart, "A1")
    charts_sheet.add_chart(cdf_chart, "A21")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
