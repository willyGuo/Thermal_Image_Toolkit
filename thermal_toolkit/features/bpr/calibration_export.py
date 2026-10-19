"""Runtime Excel export for bad-pixel calibration (existing openpyxl dependency)."""
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from thermal_toolkit.features.bpr.calibration import pixel_threshold_details, OVERLAP_ROWS, overlap_statistics


def save_bad_pixel_workbook(path, result, params, a, b, metadata):
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append(['Bad Pixel Calibration Summary'])
    summary.append(['Total pixels', int(a.size)])
    summary.append(['Resolution', f'{a.shape[1]} x {a.shape[0]}'])
    summary.append([])
    summary.append(['Method', 'Status', 'Candidate pixels', 'Ratio'])
    for key, title in (('dc', 'DC Level'), ('response', 'Responsivity'), ('netd', 'NETD')):
        enabled = params[key]
        count = metadata['method_counts'][title]
        summary.append([title, 'Enabled' if enabled else 'Disabled',
                        count if enabled else 'N/A', count / a.size if enabled else 'N/A'])
    summary.append(['Union (unique)', 'Completed', metadata['union_count'],
                    metadata['union_count'] / a.size])
    summary.append(['Ratios use all image pixels. Union counts overlapping points once.'])
    summary.append([])
    summary.append(['DPC table', 'Generated' if metadata['dpc_table_generated'] else 'Not generated'])
    summary.append(['Skip reason', metadata['dpc_table_skip_reason'] or '--'])
    summary.append(['NETD frames', metadata['temporal_count']])
    summary.append(['Source A', metadata['source_a']])
    summary.append(['Source B', metadata['source_b'] or 'N/A'])
    summary.append([])
    summary.append(['Method overlap', 'Candidate pixels', 'Ratio'])
    overlaps = overlap_statistics(result['masks'])
    for label, keys in OVERLAP_ROWS:
        enabled = all(params[key] for key in keys)
        summary.append([label, overlaps[label] if enabled else 'N/A',
                        overlaps[label] / a.size if enabled else 'N/A'])
        summary.cell(summary.max_row, 3).number_format = '0.000%'
    summary.append(['First 7 rows partition Union. Last 3 pair counts include All three; do not sum.'])
    summary.cell(summary.max_row, 1).alignment = Alignment(wrap_text=True)
    summary.row_dimensions[summary.max_row].height = 60
    for row in range(6, 10):
        summary.cell(row, 4).number_format = '0.000%'
    for row in (5, 18):
        for cell in summary[row]:
            cell.fill = PatternFill('solid', fgColor='24465A')
            cell.font = Font(color='FFFFFF', bold=True)
    for column, width in (('A', 42), ('B', 48), ('C', 22), ('D', 18)):
        summary.column_dimensions[column].width = width
    for row in (10, 13, 15, 16):
        for cell in summary[row]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        summary.row_dimensions[row].height = 48
    summary.freeze_panes = 'C6'
    summary.sheet_view.showGridLines = False
    definitions = (
        ('dc', 'DC Level', ['x', 'y', 'RAW A (DN)', 'Local median (DN)', 'Local std (DN)',
                           'Absolute difference (DN)', 'Threshold (DN)', 'Excess (DN)']),
        ('response', 'Responsivity', ['x', 'y', 'RAW A (DN)', 'RAW B (DN)', 'B-A (DN)',
                                    'Deviation (%)', 'Threshold (%)', 'Excess (percentage points)', 'Local 9x9 response median (DN)',
                                    'Absolute response difference (DN)', 'Threshold (DN)', 'Excess (DN)']),
        ('netd', 'NETD', ['x', 'y', 'NETD (mK)', 'Global mean (mK)', 'Global std (mK)',
                         'Absolute deviation (mK)', 'Threshold (mK)', 'Excess (mK)', 'Reason']))
    for key, title, headers in definitions:
        sheet = book.create_sheet(title)
        count = int(result['masks'][title].sum())
        sheet.append([title, 'Enabled' if params[key] else 'Disabled', 'Pixels', count,
                      'Ratio', count / a.size if params[key] else 'N/A'])
        sheet['F1'].number_format = '0.000%'
        sheet.append(['Coordinates: original RAW, top-left (0,0). Each sheet lists only its method hits.'])
        sheet.append(['DPC table', 'Generated' if metadata['dpc_table_generated'] else metadata['dpc_table_skip_reason']])
        setting = params[{'dc':'dc_sigma','response':'response_percent','netd':'netd_sigma'}[key]]
        sheet.append(['Threshold setting', setting, '%' if key == 'response' else 'sigma',
                      'DC window' if key == 'dc' else ('Response window' if key == 'response' else ''),
                      params['window'] if key == 'dc' else (9 if key == 'response' else '')])
        headers = headers + ['Matched methods']
        sheet.append(headers)
        yy, xx = np.where(result['masks'][title])
        for y, x in zip(yy, xx):
            d = pixel_threshold_details(result, params, x, y)
            if key == 'dc':
                row = [int(x), int(y), int(a[y,x]), float(result['local_median'][y,x]),
                       float(result['local_std'][y,x]), d['dc_compared_value_dn'],
                       d['dc_threshold_dn'], d['dc_excess_dn']]
            elif key == 'response':
                row = [int(x), int(y), int(a[y,x]), int(b[y,x]), float(result['response_dn'][y,x]),
                       d['response_compared_value_percentage_points'], d['response_threshold_percentage_points'],
                       d['response_excess_percentage_points'], float(result['response_local_reference_dn'][y,x]),
                       abs(float(result['response_dn'][y,x] - result['response_local_reference_dn'][y,x])),
                       params['response_percent'] / 100 * abs(float(result['response_local_reference_dn'][y,x])),
                       max(0., abs(float(result['response_dn'][y,x] - result['response_local_reference_dn'][y,x])) -
                           params['response_percent'] / 100 * abs(float(result['response_local_reference_dn'][y,x])))]
            else:
                value = result['netd_mk'][y,x]
                row = [int(x), int(y), float(value) if np.isfinite(value) else 'N/A',
                       result['netd_mean_mk'], result['netd_std_mk'], d['netd_compared_value_mk'],
                       d['netd_threshold_mk'], d['netd_excess_mk'], d['netd_reason']]
            row.append('; '.join(name for name, mask in result['masks'].items() if mask[y,x]))
            sheet.append(row)
        sheet.freeze_panes = 'C6'
        sheet.auto_filter.ref = f'A5:{get_column_letter(len(headers))}{max(5, sheet.max_row)}'
        for cell in sheet[5]:
            cell.fill = PatternFill('solid', fgColor='24465A')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[5].height = 32
        for col, title_text in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(col)].width = 9 if col <= 2 else max(19, min(30, len(title_text)+2))
        for row in sheet.iter_rows(min_row=6):
            for cell in row[2:]:
                if isinstance(cell.value, float):
                    cell.number_format = '0.000'
        sheet.sheet_view.showGridLines = False
    try:
        book.save(path)
    finally:
        book.close()
