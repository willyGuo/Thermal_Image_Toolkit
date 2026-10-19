"""features / nuc / export."""

import numpy as np
from thermal_toolkit.core.constants import RAW_DTYPE
from thermal_toolkit.features.nuc.definitions import (
    NUC_ANALYSIS_FIELDS,
    NUC_ANALYSIS_HELP,
    NUC_CHANGE_HELP,
    NUC_COMPARISON_ROIS,
    NUC_ROI_METRICS,
)


NUC_STATISTICS_XLSX_HEADERS = ('\u9ed1\u9ad4\u6eab\u5ea6_\u651d\u6c0f', '\u8655\u7406\u968e\u6bb5', 'RAW\u6700\u5c0f\u503c_DN', 'RAW\u6700\u5927\u503c_DN', 'RAW\u5168\u5e45\u5e73\u5747\u4eae\u5ea6_DN', 'RAW\u5168\u5e45\u7a7a\u9593\u6a19\u6e96\u5dee_DN', 'RAW\u7b2c1\u767e\u5206\u4f4d\u503c_DN', 'RAW\u7b2c99\u767e\u5206\u4f4d\u503c_DN', 'RAW\u6709\u6548\u7070\u968e\u6578', 'RAW\u96f6\u503c\u50cf\u7d20\u6578', 'RAW\u6eff\u523b\u5ea6\u50cf\u7d20\u6578_16383', '\u4e2d\u592e\u5e73\u5747\u4eae\u5ea6_DN', '\u4e2d\u592e\u7a7a\u9593\u6a19\u6e96\u5dee_DN', '\u56db\u5468\u908a\u7de3\u5e73\u5747\u4eae\u5ea6_DN', '\u5404\u6a6b\u5217\u5e73\u5747\u4eae\u5ea6\u7684\u6a19\u6e96\u5dee_DN', '\u5404\u76f4\u6b04\u5e73\u5747\u4eae\u5ea6\u7684\u6a19\u6e96\u5dee_DN', '\u540464\u4e5864\u5340\u584a\u5e73\u5747\u4eae\u5ea6\u7684\u6a19\u6e96\u5dee_DN', '\u540464\u4e5864\u5340\u584a\u5e73\u5747\u4eae\u5ea6\u7684\u6700\u5927\u6700\u5c0f\u5dee_DN', '\u5c40\u90e8\u4e2d\u4f4d\u6578\u6b98\u5dee\u7684\u7a69\u5065\u6a19\u6e96\u5dee\u4f30\u8a08_DN', '\u5c40\u90e8\u96e2\u7fa4\u50cf\u7d20\u5224\u5b9a\u9580\u6abb_DN', '\u5c40\u90e8\u5019\u9078\u96e2\u7fa4\u50cf\u7d20\u6578_\u975e\u78ba\u5b9a\u58de\u9ede', 'JPG\u5e73\u5747\u7070\u968e_0\u81f3255', 'JPG\u7a7a\u9593\u6a19\u6e96\u5dee_\u7070\u968e', 'RAW\u8207JPG\u76f8\u95dc\u4fc2\u6578_\u539f\u65b9\u5411', 'RAW\u6c34\u5e73\u93e1\u50cf\u5f8c\u8207JPG\u76f8\u95dc\u4fc2\u6578', 'RAW\u6a94\u6848\u96dc\u6e4a_SHA256', '\u91cf\u6e2c\u5ea7\u6a19\u7cfb\u7d71', '\u65b9\u584a\u6982\u7565\u7bc4\u570d_X\u8d77\u9ede', '\u65b9\u584a\u6982\u7565\u7bc4\u570d_X\u7d42\u9ede_\u4e0d\u542b', '\u65b9\u584a\u6982\u7565\u7bc4\u570d_Y\u8d77\u9ede', '\u65b9\u584a\u6982\u7565\u7bc4\u570d_Y\u7d42\u9ede_\u4e0d\u542b', '\u65b9\u584a\u91cf\u6e2c\u5340_X\u8d77\u9ede', '\u65b9\u584a\u91cf\u6e2c\u5340_X\u7d42\u9ede_\u4e0d\u542b', '\u65b9\u584a\u91cf\u6e2c\u5340_Y\u8d77\u9ede', '\u65b9\u584a\u91cf\u6e2c\u5340_Y\u7d42\u9ede_\u4e0d\u542b', '\u5de6\u5074\u53c3\u8003\u5340_X\u8d77\u9ede', '\u5de6\u5074\u53c3\u8003\u5340_X\u7d42\u9ede_\u4e0d\u542b', '\u53f3\u5074\u53c3\u8003\u5340_X\u8d77\u9ede', '\u53f3\u5074\u53c3\u8003\u5340_X\u7d42\u9ede_\u4e0d\u542b', '\u5de6\u53f3\u53c3\u8003\u5340_Y\u8d77\u9ede', '\u5de6\u53f3\u53c3\u8003\u5340_Y\u7d42\u9ede_\u4e0d\u542b', '\u65b9\u584a\u5e73\u5747\u4eae\u5ea6_DN', '\u5de6\u5074\u53c3\u8003\u5340\u5e73\u5747\u4eae\u5ea6_DN', '\u53f3\u5074\u53c3\u8003\u5340\u5e73\u5747\u4eae\u5ea6_DN', '\u5de6\u53f3\u5468\u570d\u5e73\u5747\u4eae\u5ea6_DN', '\u65b9\u584a\u6e1b\u5468\u570d\u4eae\u5ea6\u5dee_DN_\u6b63\u4eae\u8ca0\u6697')


def save_nuc_statistics_xlsx(path, raw, corrected, report):
    """Wide, one-row-per-stage report following statistics.xlsx column order."""
    import hashlib
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "statistics"
    headers = list(NUC_STATISTICS_XLSX_HEADERS)
    headers[10] = f"RAW滿刻度像素數_{(1 << report['bit_depth']) - 1}"
    keys = ("min", "max", "mean", "std", "p1_DN", "p99_DN", "unique_levels", "zero_pixels",
            "fullscale_pixels", "center_mean_DN", "center_std_DN", "rim_mean_DN", "row_mean_std_DN",
            "column_mean_std_DN", "block64_mean_std_DN", "block64_mean_range_DN", "local_residual_sigma_DN",
            "local_candidate_threshold_DN", "local_candidate_pixels")
    extra_fields = []
    for name in report["roi_analysis"]["regions"]:
        for key in NUC_ROI_METRICS:
            title = NUC_ANALYSIS_HELP[key].split("：", 1)[0]
            for mode, suffix in (("value", "數值"), ("amount", "下降量"), ("percent", "下降率%")):
                headers.append(f"{name}_{title}_{suffix}")
                extra_fields.append((name, key, mode))
    for title in ("RAW標準差", "RAW峰對峰", "左右平均", "方塊與左右亮度差絕對值"):
        headers.extend((title + "_下降量_DN", title + "_下降率%"))
    headers.extend(("來源RAW檔案", "N", "位深"))
    sheet.append(headers)
    regions = report["roi_analysis"]["regions"]
    comparison = report["roi_analysis"]["comparison"]
    block = NUC_COMPARISON_ROIS["方塊"]
    left = NUC_COMPARISON_ROIS["左側參考"]
    right = NUC_COMPARISON_ROIS["右側參考"]
    for stage, image, title in (("original", raw, "原始RAW"), ("corrected", corrected, "NUC後")):
        stats = report[stage]
        values = ["N/A", title, *[stats[key] for key in keys], *(["N/A"] * 4),
                  hashlib.sha256(image.astype(RAW_DTYPE).tobytes(order="C")).hexdigest(),
                  "RAW_top_left_zero_based_end_exclusive",
                  block[0], block[2], block[1], block[3], block[0], block[2], block[1], block[3],
                  left[0], left[2], right[0], right[2], left[1], left[3],
                  regions["方塊"].get(stage, {}).get("mean"), regions["左側參考"].get(stage, {}).get("mean"),
                  regions["右側參考"].get(stage, {}).get("mean"), comparison.get(stage, {}).get("reference_mean_DN"),
                  comparison.get(stage, {}).get("block_minus_reference_DN")]
        for name, key, mode in extra_fields:
            values.append(regions[name].get(stage, {}).get(key) if mode == "value" else
                          (regions[name]["reduction"][key][mode] if stage == "corrected" else None))
        for key in ("std", "peak_to_peak"):
            values.extend((report["reduction_dn"][key], report["reduction_percent"][key]) if stage == "corrected" else (None, None))
        for key in ("reference_mean_DN", "absolute_block_difference_DN"):
            reduction = comparison.get("reduction", {}).get(key, {})
            values.extend((reduction.get("amount"), reduction.get("percent")) if stage == "corrected" else (None, None))
        values.extend((report["raw_file"], report["N"], report["bit_depth"]))
        sheet.append(["N/A" if value is None else value for value in values])
        sheet.cell(sheet.max_row, len(headers) - 2).data_type = "s"
    for column, cell in enumerate(sheet[1], 1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305870")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.row_dimensions[1].height = 64
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
    for column, key in enumerate(keys, 3):
        sheet.cell(1, column).comment = Comment(NUC_ANALYSIS_HELP[key], "NUC_Analyze")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    notes = workbook.create_sheet("說明")
    for row in (
        ("項目", "說明"),
        ("資料來源", "本次載入的原始 RAW，以及套用本次 K/B/N 後的 RAW；每個階段一列。"),
        ("缺少的輸入", "未提供黑體溫度及 JPG，因此相關欄位填 N/A；不由 DN 推算溫度。"),
        ("方塊範圍", "本次量測使用 X=[310,370)、Y=[205,265)，60×60。與參考檔原本 44×36 量測範圍不同。"),
        ("左右合併", report["roi_analysis"]["comparison_definition"]),
        ("下降量", "原始 − NUC後；下降率 = 下降量 / 原始 × 100%。NUC後列記錄相對原始的下降量；原始列不適用。原始值為 0 時下降率為 N/A。"),
        ("亮度差改善", "使用方塊減左右平均之絕對值；正下降量表示差距縮小，負值表示差距變大。"),
        ("SHA256", "各列對應此次保存的 original.raw / corrected_nuc.raw 檔案內容。"),
        ("N", report["N"]), ("位深", report["bit_depth"]),
        ("原始RAW", report["raw_file"]), ("K表", report["k_file"]), ("B表", report["b_file"]),
    ):
        notes.append(row)
        if isinstance(row[1], str):
            notes.cell(notes.max_row, 2).data_type = "s"
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 100
    for row in notes:
        row[1].alignment = Alignment(wrap_text=True)
        notes.row_dimensions[row[0].row].height = 40
    workbook.save(path)
    workbook.close()


def save_nuc_analysis_xlsx(path, raw, corrected, stored_k, gfx_b, report):
    """Save the center ROI as coordinate-aligned matrices and Chinese statistics."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    x0, y0, x1, y1 = NUC_COMPARISON_ROIS["方塊"]
    roi = np.s_[y0:y1, x0:x1]
    workbook = Workbook()
    info = workbook.active
    info.title = "分析設定"
    info.append(["項目", "內容"])
    for row in (
        ("原始 RAW 檔案", report["raw_file"]), ("K 表檔案", report["k_file"]),
        ("B 表檔案", report["b_file"]), ("影像寬度", report["width"]),
        ("影像高度", report["height"]), ("位深", report["bit_depth"]), ("N", report["N"]),
        ("校正公式", "輸出 = K × (輸入 − B) ÷ 2^N；向下取整後裁切至位深範圍"),
        ("K 格式", "uint16，儲存的整數 K 值；實際增益為 K / 2^N"),
        ("B 格式", "int16，保留正負號"),
        ("區域定義", "方塊 60×60：X=[310,370)、Y=[205,265)。超出影像範圍時標示 N/A，不縮放或裁切區域。"),
        ("左側參考區", "X=[280,304)、Y=[216,252)，24×36"),
        ("右側參考區", "X=[376,400)、Y=[216,252)，24×36"),
        ("左右合併與改善定義", report["roi_analysis"]["comparison_definition"]),
        ("座標定義", "原始 RAW 左上角為 (0, 0)，X 向右、Y 向下；終點不包含"),
        ("X 起點", x0), ("X 終點（不含）", x1), ("Y 起點", y0), ("Y 終點（不含）", y1),
        ("區域像素數", (x1 - x0) * (y1 - y0)),
        ("矩陣讀法", "每區各有 RAW、NUC 後、K、B 四張表；首列為原圖 X 座標、首欄為原圖 Y 座標"),
        ("下降量與下降率", "原始 − 校正後；下降率 = 下降量 / 原始 × 100%。負值代表增加；原始為 0 時為 N/A"),
        ("統計注意事項", "標準差使用 ddof=0。DN 不是溫度；均勻度應在均勻場景下比較。大量裁切也可能降低標準差。"),
    ):
        info.append(row)
        if isinstance(row[1], str):
            info.cell(info.max_row, 2).data_type = "s"
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 95
    for row in info.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        info.row_dimensions[row[0].row].height = 32

    summary = workbook.create_sheet("統計結果")
    summary.append(["範圍", "指標", "單位", "原始 RAW", "NUC 後", "下降量", "下降率 (%)", "中文說明"])

    def add_comparison(scope, title, unit, before, after, help_text=""):
        decrease = before - after if before is not None and after is not None else None
        percent = decrease / before * 100 if decrease is not None and before != 0 else None
        summary.append([scope, title, unit, *["N/A" if v is None else v for v in (before, after, decrease, percent)], help_text])

    for title, key, unit in NUC_ANALYSIS_FIELDS:
        explanation = NUC_ANALYSIS_HELP[key]
        add_comparison("全圖分析", explanation.split("：", 1)[0], unit,
                       report["original"][key], report["corrected"][key], explanation)
    for name, region in report["roi_analysis"]["regions"].items():
        for key in NUC_ROI_METRICS:
            explanation = NUC_ANALYSIS_HELP[key]
            unit = next(unit for _, metric, unit in NUC_ANALYSIS_FIELDS if metric == key)
            add_comparison(name, explanation.split("：", 1)[0], unit,
                           region.get("original", {}).get(key), region.get("corrected", {}).get(key),
                           (f"區域 {region['xyxy']}，終點不含。" if region["xyxy"] else report["roi_analysis"]["comparison_definition"]) if region["valid"] else "區域超出影像範圍，無有效統計。")
    comparison = report["roi_analysis"]["comparison"]
    for key, title in (("reference_mean_DN", "左右參考平均"), ("block_minus_reference_DN", "方塊減左右參考平均"), ("absolute_block_difference_DN", "方塊與左右亮度差絕對值")):
        add_comparison("區域比較", title, "DN", comparison.get("original", {}).get(key),
                       comparison.get("corrected", {}).get(key), report["roi_analysis"]["comparison_definition"])
    clipping = [("全圖", report["clipped_low_count"], report["clipped_high_count"])]
    clipping.extend((name, region.get("clipped_low_count", "N/A"), region.get("clipped_high_count", "N/A"))
                    for name, region in report["roi_analysis"]["regions"].items())
    for scope, low, high in clipping:
        summary.append([scope, "下限裁切像素數", "count", "N/A", low, "N/A", "N/A", NUC_CHANGE_HELP["clipped"]])
        summary.append([scope, "上限裁切像素數", "count", "N/A", high, "N/A", "N/A", NUC_CHANGE_HELP["clipped"]])
    for col, width in zip("ABCDEFGH", (20, 32, 12, 18, 18, 18, 18, 85)):
        summary.column_dimensions[col].width = width
    for row in summary.iter_rows(min_row=2):
        for cell in row[3:7]:
            cell.number_format = "0.000000"
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row[0].row].height = 44
    summary.auto_filter.ref = summary.dimensions

    for name, region in report["roi_analysis"]["regions"].items():
        if region["xyxy"] is None:
            continue  # Combined reference uses the two existing coordinate-preserving sheets.
        x0, y0, x1, y1 = region["xyxy"]
        for title, data in (("原始RAW", raw), ("NUC後", corrected), ("K值", stored_k), ("B值", gfx_b)):
            sheet = workbook.create_sheet(name + title)
            if not region["valid"]:
                sheet.append(["N/A：區域超出影像範圍"])
                continue
            sheet.append(["Y / X", *range(x0, x1)])
            for y in range(y0, y1):
                sheet.append([y, *[int(value) for value in data[y, x0:x1]]])
            sheet.freeze_panes = "B2"
            for col in range(1, x1 - x0 + 2):
                sheet.column_dimensions[get_column_letter(col)].width = 11

    for sheet in workbook:
        if not sheet.freeze_panes:
            sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305870")
    workbook.save(path)
    workbook.close()
