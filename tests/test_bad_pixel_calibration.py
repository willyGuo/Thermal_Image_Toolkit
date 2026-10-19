import os
from openpyxl import load_workbook
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np

from thermal_toolkit.features.bpr.calibration import detect_dead_pixels, encode_dpc_table, calibrate_bad_pixels, temporal_statistics, pixel_threshold_details, overlap_statistics, OVERLAP_ROWS
from thermal_toolkit.features.bpr.coordinates import load_bpr_sie_coordinates


class BadPixelCalibrationTests(unittest.TestCase):
    def test_exact_equality_and_sie_bytes_round_trip(self):
        a = np.array([[100, 0, 65535], [300, 400, 500]], dtype=np.uint16)
        b = np.array([[101, 0, 65535], [299, 400, 501]], dtype=np.uint16)
        points = detect_dead_pixels(a, b)
        self.assertEqual(points, [(1, 0), (2, 0), (1, 1)])
        payload = encode_dpc_table(points)
        self.assertEqual(len(payload), 16384)
        self.assertEqual(payload[:14], bytes.fromhex("00 80 01 00 02 00 01 80 01 00 ff ff 00 00"))
        self.assertEqual(payload[14:], bytes(16384 - 14))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "dpc_table.bin"
            path.write_bytes(payload)
            decoded, _ = load_bpr_sie_coordinates(path)
            self.assertEqual(decoded, points)

    def test_empty_identical_and_mismatched_inputs(self):
        a = np.ones((3, 4), dtype=np.uint16)
        self.assertEqual(detect_dead_pixels(a, a + 1), [])
        self.assertEqual(encode_dpc_table([])[:4], bytes.fromhex("ff ff 00 00"))
        with self.assertRaisesRegex(ValueError, "identical"):
            detect_dead_pixels(a, a.copy())
        with self.assertRaises(ValueError):
            detect_dead_pixels(a, a[:1])

    def test_capacity_sorting_and_coordinate_validation(self):
        self.assertEqual(encode_dpc_table([(2, 1), (1, 0), (2, 1)]),
                         encode_dpc_table([(1, 0), (2, 1)]))
        self.assertEqual(len(encode_dpc_table([(0, y) for y in range(4095)])), 16384)
        with self.assertRaisesRegex(ValueError, "at most"):
            encode_dpc_table([(x, 0) for x in range(4096)])
        for point in ((-1, 0), (32768, 0), (0, 32767)):
            with self.assertRaises(ValueError):
                encode_dpc_table([point])


class ThreeMethodTests(unittest.TestCase):
    def test_dc_response_and_netd_union(self):
        a = np.full((21, 21), 7500.)
        b = a + 1000
        a[4,4] += 100
        b[4,4] += 100  # DC-only offset, normal response
        b[10,10] = a[10,10] + 500  # response outlier
        noise = np.full_like(a, 3.)
        noise[16,16] = 100  # temporal-noise outlier
        r = calibrate_bad_pixels(a,b,temporal_std=noise,netd=True)
        self.assertEqual(set(r['coordinates']), {(4,4),(10,10),(16,16)})
        self.assertTrue(r['masks']['DC Level'][4,4])
        self.assertFalse(r['masks']['Responsivity'][4,4])
        self.assertTrue(r['masks']['Responsivity'][10,10])
        self.assertTrue(r['masks']['NETD'][16,16])
        self.assertAlmostEqual(r['netd_mk'][0,0], 30)

    def test_threshold_evidence(self):
        a = np.full((11,11), 1000.)
        b = a + 1000
        a[3,3] += 100
        b[3,3] += 100
        b[5,5] = a[5,5]
        params = dict(dc=True, response=True, netd=True, dc_sigma=5.,
                      response_percent=5., netd_sigma=5.)
        r = calibrate_bad_pixels(a,b,temporal_std=np.ones_like(a),**params)
        row = pixel_threshold_details(r,params,3,3)
        self.assertEqual(row['dc_compared_value_dn'], 100)
        self.assertEqual(row['dc_threshold_dn'], 0)
        self.assertEqual(row['dc_excess_dn'], 100)
        self.assertEqual(row['response_excess_percentage_points'], 0)
        row = pixel_threshold_details(r,params,5,5)
        self.assertEqual(row['response_threshold_percentage_points'], 5)
        self.assertEqual(row['response_excess_percentage_points'], 95)
        self.assertEqual(row['netd_reason'], 'zero_or_opposite_response')
        self.assertEqual(row['netd_excess_mk'], 'N/A')
        params['netd'] = False
        self.assertEqual(pixel_threshold_details(r,params,5,5)['netd_threshold_mk'], 'N/A')

    def test_response_uses_local_reference_and_strict_boundary(self):
        a = np.full((41,81), 7000.)
        b = a + 1000
        b[:,41:] = a[:,41:] + 2000
        # Two normal regions must both pass despite different global responses.
        r = calibrate_bad_pixels(a,b,dc=False)
        self.assertFalse(r['masks']['Responsivity'][20,20])
        self.assertFalse(r['masks']['Responsivity'][20,60])
        b[20,20] = a[20,20] + 1050
        b[20,60] = a[20,60] + 2101
        r = calibrate_bad_pixels(a,b,dc=False,window=3)
        self.assertFalse(r['masks']['Responsivity'][20,20])
        self.assertTrue(r['masks']['Responsivity'][20,60])
        self.assertEqual(r['response_local_reference_dn'][20,60],2000)
        b[20,20] += 1
        self.assertTrue(calibrate_bad_pixels(a,b,dc=False)['masks']['Responsivity'][20,20])

    def test_overlap_partitions_and_inclusive_pairs(self):
        bits = np.arange(1,8)[None,:]
        masks = {name: (bits & bit) != 0 for name, bit in
                 (('DC Level',1), ('Responsivity',2), ('NETD',4))}
        counts = overlap_statistics(masks)
        self.assertEqual([counts[label] for label, _ in OVERLAP_ROWS[:7]], [1]*7)
        self.assertEqual([counts[label] for label, _ in OVERLAP_ROWS[7:]], [2]*3)
        self.assertEqual(sum(counts[label] for label, _ in OVERLAP_ROWS[:7]), 7)

    def test_slide_temporal_example(self):
        frames = [np.full((3,3), v) for v in (7996,7998,8000,8002,8004)]
        mean, std = temporal_statistics(list(range(5)), lambda i: frames[i])
        self.assertEqual(mean[0,0],8000)
        self.assertAlmostEqual(std[0,0], np.sqrt(8))
        r = calibrate_bad_pixels(np.full((3,3),7500),np.full((3,3),8500),
                                 temporal_std=std,netd=True)
        self.assertAlmostEqual(r['netd_mk'][0,0],28.284271247)

    def test_validation_and_zero_response(self):
        a = np.full((11,11),1000.)
        b = a+1000
        b[0,0] = a[0,0]
        r = calibrate_bad_pixels(a,b,dc=False,response=False,netd=True,
                                 temporal_std=np.ones_like(a))
        self.assertIn((0,0),r['coordinates'])
        for kwargs in ({'temperature_b':25, 'netd':True}, {'window':4}, {'netd':True},
                       {'dc':False,'response':False}):
            with self.assertRaises(ValueError):
                calibrate_bad_pixels(a,b,**kwargs)
        with self.assertRaises(ValueError):
            temporal_statistics([0],lambda _:a)


class BadPixelCalibrationGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from thermal_toolkit.core.qt import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def test_run_exports_and_invalid_input_creates_no_folder(self):
        from thermal_toolkit.ui.main_window import ThermalRawToolkit
        with patch.object(ThermalRawToolkit, "refresh_camera_devices"), tempfile.TemporaryDirectory() as folder:
            window = ThermalRawToolkit()
            try:
                index = window.FUNCTIONS.index("Bad pixel calibration")
                window.change_page(index)
                self.assertTrue(window.bad_pixel_dc.isChecked())
                self.assertTrue(window.bad_pixel_response.isChecked())
                self.assertTrue(window.bad_pixel_netd.isChecked())
                self.assertEqual(window.bad_pixel_temperature_b.value(), 60)
                window.bad_pixel_netd.setChecked(False)
                self.assertEqual(window.stack.currentIndex(), index)
                self.assertEqual(window.stack.count(), len(window.FUNCTIONS))
                self.assertEqual(window.category_badge.text(), "Calibration")
                self.assertEqual(window.run_btn.text(), "Run / Generate dpc_table.bin")
                resolution_index = next(i for i in range(window.resolution_combo.count())
                                        if tuple(window.resolution_combo.itemData(i)) == (256, 192))
                window.resolution_combo.setCurrentIndex(resolution_index)
                a = np.full((192, 256), 1000, dtype="<u2")
                b = np.full_like(a, 2000)
                b[0, 0] = a[0, 0]
                b[108, 231] = a[108, 231]
                b[191, 255] = a[191, 255]
                source_a, source_b = Path(folder) / "a.raw", Path(folder) / "b.raw"
                a.tofile(source_a)
                b.tofile(source_b)
                window.bad_pixel_raw_a.setText(str(source_a))
                window.bad_pixel_raw_b.setText(str(source_b))
                output_root = Path(folder) / "output"
                window.output_root.setText(str(output_root))
                self.assertFalse(output_root.exists())
                window.run_btn.click()
                output = Path(window.last_output_dir)
                self.assertEqual((output / "temperature_a.raw").read_bytes(), source_a.read_bytes())
                self.assertEqual((output / "temperature_b.raw").read_bytes(), source_b.read_bytes())
                decoded, _ = load_bpr_sie_coordinates(output / "dpc_table.bin")
                self.assertEqual(decoded, [(0, 0), (231, 108), (255, 191)])
                self.assertEqual((output / "dpc_table.bin").stat().st_size, 16384)
                book = load_workbook(output / 'bad_pixel_coordinates.xlsx', data_only=True)
                self.assertEqual(book.sheetnames, ['Summary', 'DC Level', 'Responsivity', 'NETD'])
                self.assertEqual(book['Summary']['C9'].value, 3)
                self.assertAlmostEqual(book['Summary']['D9'].value, 3 / a.size)
                self.assertEqual(book['Summary']['B12'].value, 'Generated')
                self.assertEqual(book['Responsivity'].max_row, 8)
                self.assertEqual(book['Responsivity']['G6'].value, 5)
                self.assertEqual(book['Responsivity']['H6'].value, 95)
                self.assertEqual(book['NETD']['B1'].value, 'Disabled')
                book.close()
                self.assertIn("3 / 49,152", window.bad_pixel_calibration_status.text())
                temporal_folder = Path(folder) / 'sequence'
                temporal_folder.mkdir()
                for n, shift in enumerate((-2, 0, 2)):
                    frame = a.copy()
                    frame = (frame.astype(np.int32) + shift).astype('<u2')
                    frame[80, 80] = 1000 + shift * 50
                    frame.tofile(temporal_folder / f'{n}.raw')
                window.bad_pixel_temporal.setText(str(temporal_folder))
                window.bad_pixel_netd.setChecked(True)
                window.run_bad_pixel_calibration()
                netd_output = Path(window.last_output_dir)
                import json
                metadata = json.loads((netd_output / 'calibration_summary.json').read_text())
                self.assertEqual(metadata['temporal_count'], 3)
                self.assertEqual(metadata['union_count'], 4)
                self.assertEqual(metadata['method_counts']['NETD'], 4)
                for source in temporal_folder.iterdir():
                    copied = next((netd_output / 'temporal_inputs').glob('*_' + source.name))
                    self.assertEqual(copied.read_bytes(), source.read_bytes())
                window.bad_pixel_netd.setChecked(False)
                before = list(output_root.iterdir())
                window.bad_pixel_raw_b.setText(str(source_a))
                self.assertIsNone(window.current_result)
                with self.assertRaisesRegex(ValueError, "identical"):
                    window.run_bad_pixel_calibration()
                self.assertEqual(list(output_root.iterdir()), before)
                source_b.write_bytes(b"short")
                window.bad_pixel_raw_b.setText(str(source_b))
                with self.assertRaisesRegex(ValueError, "RAW size"):
                    window.run_bad_pixel_calibration()
                self.assertEqual(list(output_root.iterdir()), before)
                b.fill(1000)
                b[0, 0] = 2000
                b.tofile(source_b)
                with self.assertRaisesRegex(ValueError, "Median response"):
                    window.run_bad_pixel_calibration()
                self.assertEqual(list(output_root.iterdir()), before)
                window.bad_pixel_response.setChecked(False)
                self.assertFalse(window.bad_pixel_raw_b.isEnabled())
                self.assertFalse(window.bad_pixel_temperature_a.isEnabled())
                window.bad_pixel_raw_b.setText('does-not-exist.raw')
                a[30,30] = 1200
                a.tofile(source_a)
                window.run_bad_pixel_calibration()
                dc_output = Path(window.last_output_dir)
                self.assertFalse((dc_output / 'temperature_b.raw').exists())
                dc_stats = json.loads((dc_output / 'calibration_summary.json').read_text())
                self.assertIsNone(dc_stats['source_b'])
                self.assertIsNone(dc_stats['response_reference_dn'])
                self.assertEqual(dc_stats['union_count'], 1)
                book = load_workbook(dc_output / 'bad_pixel_coordinates.xlsx')
                self.assertEqual(book['DC Level'].max_row, 6)
                self.assertEqual(book['Responsivity']['B1'].value, 'Disabled')
                book.close()
                window.bad_pixel_response.setChecked(True)
                self.assertTrue(window.bad_pixel_raw_b.isEnabled())
                window.bad_pixel_raw_b.setText('')
                with self.assertRaisesRegex(ValueError, 'requires temperature B'):
                    window.run_bad_pixel_calibration()
                window.bad_pixel_raw_b.setText(str(source_b))
                a.fill(1000)
                a.tofile(source_a)
                b.fill(2000)
                b.flat[:45000:9] = 1500
                b.tofile(source_b)
                window.run_bad_pixel_calibration()
                overflow_output = Path(window.last_output_dir)
                self.assertFalse((overflow_output / 'dpc_table.bin').exists())
                self.assertIsNotNone(window.current_result)
                self.assertEqual((overflow_output / 'temperature_b.raw').read_bytes(), source_b.read_bytes())
                book = load_workbook(overflow_output / 'bad_pixel_coordinates.xlsx', read_only=True)
                self.assertEqual(book['Responsivity'].max_row, 5005)
                self.assertIn('4,095', book['Responsivity']['B3'].value)
                book.close()
                stats = json.loads((overflow_output / 'calibration_summary.json').read_text())
                self.assertFalse(stats['dpc_table_generated'])
                self.assertEqual(stats['union_count'], 5000)
                self.assertEqual(stats['method_counts']['Responsivity'], 5000)
                self.assertEqual(sum(stats['method_combination_counts'].values()), 5000)
                self.assertAlmostEqual(stats['union_percent'], 5000 / a.size * 100)
                self.assertIn('NOT generated', window.bad_pixel_calibration_status.text())
                self.assertIn('4,095', stats['dpc_table_skip_reason'])
            finally:
                window.close()


if __name__ == "__main__":
    unittest.main()
